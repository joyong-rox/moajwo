"""반영 전 격리 기동 점검 — 후보 backend·frontend 이미지(ID)를 운영과 분리해 띄우고 기동·DB 연결·CA·TLS·핵심 취합·내려받기를
확인한다 (docs/DEPLOY.md 11장). 운영 컨테이너·데이터 볼륨·포트를 쓰지 않는다: 새 내부 네트워크(--internal, 바깥 연결 없음)와 새 임시
볼륨에 띄우고 호스트 포트를 열지 않으며, 끝나면 이 실행이 만든 컨테이너·볼륨·네트워크만 지우고 남지 않았는지 확인한다. 자료는 점검
안에서 만든 합성 자료뿐이고, 비밀번호는 점검 DB에만 쓰는 임의 값이라 기록하지 않는다.

사용법: python scripts/image_smoke.py <backend 이미지> <frontend 이미지> <결과.json>
종료 코드: 0 통과 · 4 점검 실패(측정한 점검이 실패) · 3 도구 오류(빠진·중복·형식 오류 결과, docker 명령 실패, 정리 실패·확인 불가, 출력 실패)
· 2 사용법 · 1 예상치 못한 중단(처리하지 못한 예외). 통과는 필수 점검 22개(CHECKS_IN 17 + 호스트 5)가 각각 정확히 한 번, ok가 참거짓 값 True로 모두 있을 때만이다.
측정한 실패와 도구 오류가 함께면: 정리 실패·확인 불가·출력 실패가 있으면 3(남은 자원·기록부터 봐야 함), 그 밖에는 4가 먼저다 —
실패한 id와 도구 오류는 어느 쪽이든 결과 JSON과 출력에 모두 남는다. 상태 조회가 docker 쪽에서 실패해 끝난 점검은 대상의 실패로
적지 않고 도구 오류로 남긴다(미측정). 컨테이너 안 한 단계가 예외로 끝나 하지 못한 점검은 사유(not_run)와 함께 실패로 남는다 —
대부분 후보 쪽 원인(응답 없음·모듈 없음·파일 손상)이기 때문이다.

요청은 backend 컨테이너 안의 Python이 보낸다 (호스트 포트를 열지 않기 위해):
  1. 기동: backend HEALTHCHECK healthy, 새 SQLite DB의 alembic current = heads, frontend 응답
  2. DB 연결: GET /health → status ok · database ok
  3. CA·TLS: 기본 신뢰 저장소가 OS CA 묶음을 읽음. 점검 안에서 만든 시험 CA로 TLS 서버를 띄워
     Python ssl — 시험 CA + 맞는 이름 연결 / 기본 신뢰 저장소는 모르는 CA 거절 / 이름 다름 거절,
     psycopg(libpq, sslmode=verify-full) — 시험 CA + 맞는 이름은 TLS 성립 / 이름 다름 거절 / sslrootcert=system은 모르는 CA 거절
  4. frontend: 첫 화면 200·CSP·nosniff, 프록시 /api/health
  5. 핵심 취합·내려받기 (frontend 프록시를 거쳐): 최초 소유자 설정 → CSV·XLSX 업로드 → 공통 key 연결 확인·만들기 → 결과 내려받아
     내용 확인 → 묶기 확인·만들기 → 결과 내려받기
  6. 로그: docker logs가 종료 0이고 기동 표지가 있어야 '읽음' — backend Traceback·줄 머리 ERROR/CRITICAL·5xx 응답 없음,
     frontend [emerg]·[alert]·[crit]·[error] 없음
  7. 정리: 만든 자원마다 삭제 결과 + 목록 조회(종료 0)로 남지 않았음을 확인 — 조회 실패는 '남은 것 없음'이 아니다. 컨테이너는
     익명 볼륨과 함께 지우고(rm -f -v), 격리 확인 때 본 익명 볼륨도 남지 않았는지 확인한다. 삭제 오류 원문은 결과 JSON에 남긴다
"""

import json
import os
import re
import secrets
import subprocess
import sys
import time

os.environ["NoDefaultCurrentDirectoryInExePath"] = "1"  # Windows: docker를 현재 폴더(저장소)에서 먼저 찾지 않게 — 단독 실행에도
TOOL = "image_smoke 2"
PREFIX = "moajwo-smoke-"  # 이 도구가 만드는 자원 이름의 머리 — 실행마다 시각·pid·무작위 값을 붙인다
PROD_VOLUME = "moajwo_moajwo-data"
RESULT_MARK = "@@RESULT "

# 점검 id는 바뀌지 않는 식별자다 — 통과 판단은 이 목록이 모두, 각각 한 번, ok=True로 있는지로 한다 (이름 문구·개수로 하지 않음)
CHECKS_IN = [  # backend 컨테이너 안에서 (CHECK_PY)
    ("boot.alembic_head", "기동: 새 SQLite DB의 alembic current = heads"),
    ("db.health", "DB 연결: GET /health → status ok · database ok"),
    ("ca.default_store", "CA: 기본 신뢰 저장소가 OS CA 묶음을 읽음"),
    ("tls.py.trusted", "TLS(Python ssl): 시험 CA + 맞는 이름 → 연결"),
    ("tls.py.unknown_ca", "TLS(Python ssl): 기본 신뢰 저장소 → 모르는 CA 거절"),
    ("tls.py.wrong_name", "TLS(Python ssl): 이름 다름 → 거절"),
    ("tls.pg.trusted", "TLS(psycopg·libpq verify-full): 시험 CA + 맞는 이름 → TLS 성립"),
    ("tls.pg.wrong_name", "TLS(psycopg·libpq verify-full): 이름 다름 → 거절"),
    ("tls.pg.system_unknown_ca", "TLS(psycopg·libpq verify-full): sslrootcert=system → 모르는 CA 거절"),
    ("front.index", "frontend: 첫 화면 200 · CSP · nosniff"),
    ("front.proxy_health", "frontend 프록시: /api/health → backend ok"),
    ("flow.setup", "핵심 취합 준비: 설치 초기화 명령(app.adminctl init — 점검 DB)으로 admin 생성 · 로그인 200 · 세션 쿠키"),
    ("flow.upload", "업로드(프록시): CSV·XLSX 2개를 읽음"),
    ("flow.keyjoin_plan", "공통 key 연결 확인(plan): 준비됨"),
    ("flow.keyjoin_run", "공통 key 연결 만들기(run): 201"),
    ("flow.keyjoin_download", "공통 key 연결 결과 내려받기(프록시): 기준 key 3개 · K001 행에 두 파일 값"),
    ("flow.bundle", "묶기 확인·만들기·내려받기(프록시): 시트 2개 이상"),
]
CHECKS_HOST_BEFORE = [
    ("host.backend_healthy", "기동: backend healthy (HEALTHCHECK)"),
    ("host.frontend_up", "기동: frontend 응답 (컨테이너 안 8080)"),
]
CHECKS_HOST_AFTER = [
    ("host.backend_log", "로그: backend 로그를 읽음 · Traceback·ERROR/CRITICAL 줄·5xx 응답 없음"),
    ("host.frontend_log", "로그: frontend 로그를 읽음 · [emerg]·[alert]·[crit]·[error] 없음"),
    ("host.cleanup", "정리: 이 실행의 컨테이너(익명 볼륨 포함)·볼륨·네트워크를 지웠고 남지 않음을 확인"),
]
REQUIRED = CHECKS_HOST_BEFORE + CHECKS_IN + CHECKS_HOST_AFTER  # 22개, 이 순서로 기록
NAMES = dict(REQUIRED)
LOG_MARK = {"backend": "Application startup complete", "frontend": "Configuration complete; ready for start up"}
# docker 쪽 연결 실패 — 명령이 컨테이너에 닿지 못했으므로 점검 대상의 실패가 아니라 조회 실패(도구 오류)로 본다
DOCKER_DOWN = re.compile(r"error during connect|Cannot connect to the Docker daemon|is the docker daemon running", re.IGNORECASE)
ANON_VOLUME = re.compile(r"[0-9a-f]{64}")  # 이름 없는(익명) 볼륨 — 이미지의 VOLUME 선언으로 docker run이 만든다

CHECK_PY = r"""
import io, json, os, secrets, socket, ssl, subprocess, sys, tempfile, threading, time, urllib.error, urllib.request

ITEMS = dict(json.loads(__CHECKS_IN__))
R = []
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def rec(cid, ok, detail=None):
    if cid not in ITEMS or any(x["id"] == cid for x in R):  # 모르는 id·두 번째 기록은 남기지 않는다 (호스트가 빠짐으로 잡음)
        return
    R.append({"id": cid, "item": ITEMS[cid], "ok": ok is True, "detail": detail})


def step(ids, fn):
    err = None
    try:
        fn()
    except Exception as e:  # 한 단계의 예외가 나머지 단계를 막지 않게
        err = type(e).__name__ + ": " + str(e)[:300]
    for cid in ids:  # 이 단계에서 기록하지 못한 점검은 사유와 함께 실패로 남긴다 (빠진 채 넘어가지 않게)
        if not any(x["id"] == cid for x in R):
            rec(cid, False, {"not_run": True, "reason": err or "앞 점검이 실패해 하지 않음"})


def sh(*a):
    r = subprocess.run(list(a), capture_output=True, text=True, cwd="/app", timeout=120)
    return r.returncode, r.stdout + "\n" + r.stderr


OPENER = urllib.request.build_opener(urllib.request.ProxyHandler({}))
COOKIE = {}


def http(method, url, body=None, files=None, timeout=180):
    headers, data = {}, None
    if COOKIE.get("v"):
        headers["Cookie"] = COOKIE["v"]
    if files is not None:
        bd = "smoke" + secrets.token_hex(8)
        data = b"".join(
            f'--{bd}\r\nContent-Disposition: form-data; name="files"; filename="{n}"\r\nContent-Type: {t}\r\n\r\n'.encode() + c + b"\r\n"
            for n, c, t in files
        ) + f"--{bd}--\r\n".encode()
        headers["Content-Type"] = "multipart/form-data; boundary=" + bd
    elif body is not None:
        data, headers["Content-Type"] = json.dumps(body).encode(), "application/json"
    req = urllib.request.Request(url, data=data, method=method, headers=headers)
    try:
        with OPENER.open(req, timeout=timeout) as r:
            return r.status, {k.lower(): v for k, v in r.headers.items()}, r.read()
    except urllib.error.HTTPError as e:
        return e.code, {k.lower(): v for k, v in e.headers.items()}, e.read()


def js(b):
    try:
        return json.loads(b)
    except ValueError:
        return {}


def boot():
    rc1, cur = sh("alembic", "current")
    rc2, heads = sh("alembic", "heads")
    cur_l = [x.strip() for x in cur.splitlines() if "(head)" in x]
    head_l = [x.strip() for x in heads.splitlines() if "(head)" in x]
    rec("boot.alembic_head", rc1 == 0 and rc2 == 0 and len(head_l) == 1 and cur_l[:1] == head_l, {"current": cur_l, "heads": head_l})


def db():
    s, _, b = http("GET", "http://127.0.0.1:8000/health")
    d = js(b)
    rec("db.health", s == 200 and d.get("status") == "ok" and d.get("database") == "ok", {"status": s, "body": d})


TMP = tempfile.mkdtemp(prefix="smoke-tls-")
HS = []  # 시험 서버 쪽에서 본 TLS 성립 결과: (psycopg 서버인지, 결과)
TLS = {}


def ossl(*a):
    r = subprocess.run(["openssl", *a], capture_output=True, text=True, cwd=TMP, timeout=60)
    if r.returncode:
        raise RuntimeError("openssl " + a[0] + ": " + r.stderr[-200:])


def serve(pg):
    ls = socket.socket()
    ls.bind(("127.0.0.1", 0))
    ls.listen(8)

    def loop():
        while True:
            c, _ = ls.accept()
            try:
                c.settimeout(10)
                if pg:  # PostgreSQL 첫 요청(SSLRequest)에 'S'로 답한 뒤 TLS — 그 뒤 시작 메시지를 받으면 끊는다
                    if c.recv(8)[4:8] != (80877103).to_bytes(4, "big"):
                        c.close()
                        continue
                    c.sendall(b"S")
                t = TLS["server"].wrap_socket(c, server_side=True)
                HS.append((pg, "ok"))
                if not pg:
                    t.sendall(b"ok")
                try:
                    t.recv(1024)
                except Exception:
                    pass
                t.close()
            except Exception as e:
                HS.append((pg, type(e).__name__))
                c.close()

    threading.Thread(target=loop, daemon=True).start()
    return ls.getsockname()[1]


def py_client(port, cafile, name):
    ctx = ssl.create_default_context(cafile=cafile) if cafile else ssl.create_default_context()
    try:
        with socket.create_connection(("127.0.0.1", port), timeout=10) as s, ctx.wrap_socket(s, server_hostname=name) as t:
            return ["ok", t.recv(8).decode()]
    except ssl.SSLCertVerificationError as e:
        return ["verify-failed", e.verify_message]
    except Exception as e:
        return ["error", type(e).__name__ + ": " + str(e)[:160]]


def tls():
    ossl("req", "-x509", "-newkey", "rsa:2048", "-nodes", "-keyout", "ca.key", "-out", "ca.pem", "-days", "2",
         "-subj", "/CN=moajwo-smoke-ca")
    ossl("req", "-newkey", "rsa:2048", "-nodes", "-keyout", "srv.key", "-out", "srv.csr", "-subj", "/CN=smoke.test")
    with open(os.path.join(TMP, "ext.cnf"), "w") as f:
        f.write("subjectAltName=DNS:smoke.test\nextendedKeyUsage=serverAuth\n")
    ossl("x509", "-req", "-in", "srv.csr", "-CA", "ca.pem", "-CAkey", "ca.key", "-CAcreateserial", "-out", "srv.pem",
         "-days", "2", "-extfile", "ext.cnf")
    TLS["ca"] = os.path.join(TMP, "ca.pem")
    TLS["server"] = ssl.SSLContext(ssl.PROTOCOL_TLS_SERVER)
    TLS["server"].load_cert_chain(os.path.join(TMP, "srv.pem"), os.path.join(TMP, "srv.key"))
    paths = ssl.get_default_verify_paths()
    stats = ssl.create_default_context().cert_store_stats()
    real = os.path.realpath(paths.openssl_cafile)
    rec("ca.default_store", real == "/etc/ssl/certs/ca-certificates.crt" and stats.get("x509_ca", 0) >= 100,
        {"cafile": paths.openssl_cafile, "realpath": real, "x509_ca": stats.get("x509_ca"), "openssl": ssl.OPENSSL_VERSION})
    port = serve(False)
    a, b, c = py_client(port, TLS["ca"], "smoke.test"), py_client(port, None, "smoke.test"), py_client(port, TLS["ca"], "other.test")
    rec("tls.py.trusted", a == ["ok", "ok"], {"result": a})
    rec("tls.py.unknown_ca", b[0] == "verify-failed", {"result": b})
    rec("tls.py.wrong_name", c[0] == "verify-failed" and "mismatch" in (c[1] or "").lower(), {"result": c})


def pg_client(port, rootcert, host):
    import psycopg

    HS.clear()
    info = (f"host={host} hostaddr=127.0.0.1 port={port} dbname=smoke user=smoke sslmode=verify-full sslrootcert={rootcert}"
            " gssencmode=disable connect_timeout=10")
    try:
        psycopg.connect(info).close()
        msg = "connected"
    except Exception as e:
        msg = " ".join(str(e).split())[:300]
    for _ in range(30):  # 서버 쪽 결과가 기록될 때까지 잠깐 기다림
        if HS:
            break
        time.sleep(0.1)
    return msg, [h for p, h in HS if p]


def pg_tls():
    import psycopg

    if "server" not in TLS:
        raise RuntimeError("시험 인증서가 없음 (앞 CA·TLS 단계 실패)")
    port = serve(True)
    a, b, c = pg_client(port, TLS["ca"], "smoke.test"), pg_client(port, TLS["ca"], "other.test"), pg_client(port, "system", "smoke.test")
    lib = {"libpq": psycopg.pq.version(), "impl": psycopg.pq.__impl__}
    rec("tls.pg.trusted", "ok" in a[1] and "certificate verify failed" not in a[0] and "does not match" not in a[0],
        {**lib, "server": a[1], "client": a[0]})
    rec("tls.pg.wrong_name", "does not match host name" in b[0], {"server": b[1], "client": b[0]})
    rec("tls.pg.system_unknown_ca", "certificate verify failed" in c[0], {"server": c[1], "client": c[0]})


def front():
    s, h, _ = http("GET", "http://frontend:8080/")
    rec("front.index", s == 200 and bool(h.get("content-security-policy")) and h.get("x-content-type-options") == "nosniff",
        {"status": s, "csp": bool(h.get("content-security-policy")), "nosniff": h.get("x-content-type-options")})
    s, _, b = http("GET", "http://frontend:8080/api/health")
    d = js(b)
    rec("front.proxy_health", s == 200 and d.get("database") == "ok", {"status": s, "body": d})


API = "http://frontend:8080/api"


def flow():
    import openpyxl

    # 공개 최초 설정은 종료됐다 — 설치 담당자의 초기화 명령(app.adminctl init, 점검 DB)으로 admin을 만들고 로그인한다.
    # 비밀번호는 점검 컨테이너 안의 본인만 읽는 임시 파일로만 넘기고 바로 지운다 (출력·기록 없음)
    pw = "Sm-" + secrets.token_hex(12)
    pwfile = "/tmp/smoke-admin-pw"
    with open(pwfile, "w", encoding="utf-8", opener=lambda p, f: os.open(p, f, 0o600)) as fh:
        fh.write(pw + "\n")
    rc, out = sh(sys.executable, "-m", "app.adminctl", "init", "--yes", "--password-file", pwfile)
    os.remove(pwfile)
    s, h, _ = http("POST", API + "/auth/login", {"username": "admin", "password": pw})
    cookie = (h.get("set-cookie") or "").split(";")[0]
    COOKIE["v"] = cookie
    rec("flow.setup", rc == 0 and s == 200 and "=" in cookie,
        {"init_rc": rc, "init_tail": out.strip()[-160:], "login": s, "cookie": bool(cookie)})  # fmt: skip
    csv_a = "코드,인구\nK001,100\nK002,200\nK003,300\n".encode("utf-8")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "예산"
    for row in (["코드", "예산"], ["K001", 10], ["K002", 20], ["K004", 40]):
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    s, _, b = http("POST", API + "/files", files=[("smoke_a.csv", csv_a, "text/csv"), ("smoke_b.xlsx", buf.getvalue(), XLSX)])
    up = js(b) if s == 200 else []
    up = up if isinstance(up, list) else []
    ids = [x.get("id") for x in up if isinstance(x, dict) and x.get("id") and not x.get("error")]
    rec("flow.upload", s == 200 and len(ids) == 2,
        {"status": s, "rows": [x.get("data_rows") for x in up if isinstance(x, dict)], "errors": [x.get("error") for x in up if isinstance(x, dict)]})
    if len(ids) != 2:
        return
    fa, fb = ids
    body = {"inputs": [{"file_id": f, "sheet_id": "s0", "key_columns": ["c0"], "columns": [{"column_id": "c1"}]} for f in (fa, fb)],
            "scope": "base", "base": {"file_id": fa, "sheet_id": "s0"}}
    s, _, b = http("POST", API + "/keyjoin/plan", body)
    plan = js(b)
    rec("flow.keyjoin_plan", s == 200 and plan.get("ready") is True,
        {"status": s, "result": plan.get("result"), "problems": [p.get("text") for p in plan.get("problems", []) if isinstance(p, dict)][:3]})
    s, _, b = http("POST", API + "/keyjoin/run", {**body, "expected_basis": plan.get("basis") or "x" * 16, "request_id": secrets.token_hex(16)})
    run = js(b)
    rec("flow.keyjoin_run", s == 201 and bool(run.get("id")), {"status": s, "id": run.get("id"), "state": run.get("status")})
    s, h, b = http("GET", API + f"/runs/{run.get('id')}/download")
    ok = s == 200 and h.get("content-type", "").startswith(XLSX)
    detail = {"status": s, "type": h.get("content-type"), "bytes": len(b)}
    if ok:
        book = openpyxl.load_workbook(io.BytesIO(b), read_only=True)

        def head(w):
            h = list(next(w.iter_rows(values_only=True), []))
            while h and h[-1] is None:
                h.pop()
            return h

        def norm(v):
            return str(int(v)) if isinstance(v, float) and v.is_integer() else str(v)

        sheet = next((w for w in book.worksheets if head(w) == plan.get("header")), None)
        rows = [list(r) for r in sheet.iter_rows(values_only=True)][1:] if sheet is not None else []
        keys = sorted({v for r in rows for v in r if isinstance(v, str) and v.startswith("K00")})
        k1 = [norm(v) for v in next((r for r in rows if "K001" in r), [])]
        detail.update({"sheets": book.sheetnames, "header": plan.get("header"), "first_row": [head(w) for w in book.worksheets][:1],
                       "data_sheet": sheet.title if sheet is not None else None, "keys": keys, "K001": k1})
        ok = sheet is not None and keys == ["K001", "K002", "K003"] and "100" in k1 and "10" in k1
    rec("flow.keyjoin_download", ok, detail)
    s, _, b = http("POST", API + "/bundle/plan", {"file_ids": [fa, fb]})
    bp = js(b)
    s2, _, b2 = http("POST", API + "/bundle/run", {"file_ids": [fa, fb], "expected_basis": bp.get("basis") or "x" * 16})
    br = js(b2)
    s3, h3, b3 = http("GET", API + f"/runs/{br.get('id')}/download")
    ok = s == 200 and s2 == 201 and s3 == 200 and h3.get("content-type", "").startswith(XLSX)
    detail = {"plan": s, "run": s2, "download": s3}
    if ok:
        names = openpyxl.load_workbook(io.BytesIO(b3), read_only=True).sheetnames
        detail["sheets"] = len(names)
        ok = len(names) >= 2
    rec("flow.bundle", ok, detail)


step(["boot.alembic_head"], boot)
step(["db.health"], db)
step(["ca.default_store", "tls.py.trusted", "tls.py.unknown_ca", "tls.py.wrong_name"], tls)
step(["tls.pg.trusted", "tls.pg.wrong_name", "tls.pg.system_unknown_ca"], pg_tls)
step(["front.index", "front.proxy_health"], front)
step(["flow.setup", "flow.upload", "flow.keyjoin_plan", "flow.keyjoin_run", "flow.keyjoin_download", "flow.bundle"], flow)
print("@@RESULT " + json.dumps(R, ensure_ascii=True))  # 한 줄 ASCII — 줄 구분 문자가 값에 있어도 결과 줄이 끊기지 않게
"""


def run(*args: str, stdin: bytes | None = None, timeout: int = 600) -> tuple[int, str, str]:
    """docker 명령 하나 — (종료 코드, 표준출력, 표준오류). 시간 초과·실행 실패는 종료 코드 -1 (조회 실패로 다룬다)."""
    env = dict(os.environ, MSYS_NO_PATHCONV="1")  # Windows Git Bash가 /app/... 인자를 경로로 바꾸지 않게
    try:
        r = subprocess.run(["docker", *args], input=stdin, capture_output=True, env=env, timeout=timeout, check=False)
    except (OSError, subprocess.TimeoutExpired) as e:
        return -1, "", f"{type(e).__name__}: {e}"
    return r.returncode, (r.stdout or b"").decode("utf-8", "replace"), (r.stderr or b"").decode("utf-8", "replace")


def must(*args: str, stdin: bytes | None = None, timeout: int = 600) -> str:
    """실패하면 안 되는 docker 명령 — 실패는 예외(도구 오류)."""
    rc, out, err = run(*args, stdin=stdin, timeout=timeout)
    if rc != 0:
        raise RuntimeError(f"docker {' '.join(args[:2])} 실패 (코드 {rc}): {err.strip()[:400]}")
    return out.strip()


def wait_for(probe, seconds: int) -> tuple[bool, bool, str]:
    """probe() → (조회 성공, 조건 만족, 보이는 값, 더 기다릴 필요 없음). 돌려줌: (만족, 마지막 조회 성공, 마지막 값·오류).
    만족하지 못하고 끝났는데 마지막 조회가 실패였으면 판정할 수 없다 — 부르는 쪽은 대상의 실패가 아니라 도구 오류로 남긴다."""
    end, last = time.monotonic() + seconds, ""
    while True:
        queried, ok, last, final = probe()
        if ok:
            return True, True, last
        if final or time.monotonic() >= end:
            return False, queried, last
        time.sleep(2)


def validate_child(items: object) -> tuple[dict[str, dict], list[str]]:
    """backend 컨테이너가 돌려준 결과 검사 — 목록이어야 하고, CHECKS_IN의 id가 각각 정확히 한 번, ok는 참거짓 값."""
    want = {cid for cid, _ in CHECKS_IN}
    if not isinstance(items, list):
        return {}, [f"컨테이너 결과가 목록이 아님 ({type(items).__name__})"]
    got: dict[str, dict] = {}
    errors: list[str] = []
    for x in items:
        cid = x.get("id") if isinstance(x, dict) else None
        if not isinstance(x, dict):
            errors.append(f"컨테이너 결과 항목이 객체가 아님 ({type(x).__name__})")
        elif not isinstance(cid, str) or cid not in want:
            errors.append(f"모르는 점검 id: {cid!r}")
        elif cid in got:
            errors.append(f"중복 점검 id: {cid}")
        elif type(x.get("ok")) is not bool:
            errors.append(f"{cid}: ok가 참거짓 값이 아님 ({x.get('ok')!r})")
        else:  # 컨테이너 안 단계의 예외로 하지 못한 점검은 detail.not_run과 사유를 달고 실패(측정한 결과)로 온다
            got[cid] = {"id": cid, "item": NAMES[cid], "ok": x["ok"], "detail": x.get("detail"), "measured": True}
    errors += [f"빠진 점검: {cid}" for cid, _ in CHECKS_IN if cid not in got]
    return got, errors


def read_log(name: str, kind: str) -> tuple[bool, str, dict]:
    """docker logs — 돌려줌: (명령 성공, 로그 글, 설명). 명령 실패는 도구 오류이고, 성공했는데 기동 표지가 없으면 로그를 믿을 수 없어
    점검 실패로 본다 (빈 결과를 '오류 없음'으로 읽지 않는다)."""
    rc, out, err = run("logs", name, timeout=120)
    # 컨테이너의 표준오류는 docker logs의 표준오류로 나온다 — 오류 표시가 아니라 로그의 일부. 줄 머리 검사가 맞게 줄을 나눠 잇는다
    text = out + ("" if not out or out.endswith("\n") else "\n") + err
    return rc == 0, text, {"rc": rc, "lines": text.count("\n"), "start_mark": LOG_MARK[kind] in text,
                           "stderr_tail": err.strip()[-200:] if rc != 0 else ""}  # fmt: skip


def cleanup(made: dict[str, list[str]], tag: str, anon: list[str]) -> tuple[bool, dict, list[str]]:
    """이 실행이 만든 자원만 지우고(이름이 이 실행의 tag로 시작), 종류별 목록을 다시 조회해 남지 않았음을 확인한다.
    목록 조회 실패는 '남은 것 없음'이 아니라 확인 불가(도구 오류)다. 돌려줌: (모두 지워짐 확인, 자원별 결과, 오류)."""
    remove = {"container": ("rm", "-f", "-v"), "volume": ("volume", "rm"), "network": ("network", "rm")}  # -v: 컨테이너의 익명 볼륨도
    listing = LISTING
    detail: dict = {k: {} for k in remove}
    errors: list[str] = []
    for kind in ("container", "volume", "network"):  # 컨테이너를 먼저 지워야 볼륨·네트워크가 풀린다
        for name in made[kind]:
            assert name.startswith(tag) and name != PROD_VOLUME, name
            rc, _, err = run(*remove[kind], name, timeout=120)
            detail[kind][name] = {"remove_rc": rc, "remove_err": err.strip()[-200:] if rc else ""}
    ok = True
    for kind, cmd in listing.items():
        rc, out, err = run(*cmd, timeout=60)
        if rc != 0:
            ok = False
            errors.append(f"정리 확인 실패: {kind} 목록 조회 종료 {rc} {err.strip()[:200]}")
            for name in made[kind]:
                detail[kind][name]["left"] = None
            continue
        present = {x.strip() for x in out.splitlines() if x.strip()}
        for name in made[kind]:
            detail[kind][name]["left"] = name in present
            if name in present:
                ok = False
                d = detail[kind][name]
                errors.append(f"정리 실패: {kind} {name}가 남음 (삭제 종료 {d['remove_rc']}: {d['remove_err'] or '오류 문구 없음'})")
        if kind == "volume":  # 격리 확인 때 본 익명 볼륨 — rm -v로 지웠어야 한다 (이름이 tag로 시작하지 않아 따로 지우지 않음)
            for name in anon:
                detail["volume"].setdefault("anonymous", {})[name] = {"left": name in present}
                if name in present:
                    ok = False
                    errors.append(f"정리 실패: 점검 컨테이너의 익명 볼륨 {name}가 남음")
        stray = sorted(x for x in present if x.startswith(tag) and x not in made[kind])
        if stray:  # 기록하기 전에 만들어진 것 — 지우지 않고(이 실행이 만든 것인지 확인하지 못함) 알린다
            ok = False
            detail[kind]["unrecorded"] = stray
            errors.append(f"정리 확인: 기록하지 않은 {kind} {stray}")
    return ok, detail, errors


LISTING = {"container": ("ps", "-a", "--format", "{{.Names}}"), "volume": ("volume", "ls", "--format", "{{.Name}}"),
           "network": ("network", "ls", "--format", "{{.Name}}")}  # fmt: skip


def previous_leftovers() -> dict[str, list[str]]:
    """앞 실행이 남긴 점검 자원(이름이 PREFIX로 시작). 목록 조회 실패는 '없음'이 아니라 예외(도구 오류)다."""
    found: dict[str, list[str]] = {}
    for kind, cmd in LISTING.items():
        rc, out, err = run(*cmd, timeout=60)
        if rc != 0:
            raise RuntimeError(f"시작 전 {kind} 목록 조회 실패 (종료 {rc}: {err.strip()[:160]})")
        names = sorted(x.strip() for x in out.splitlines() if x.strip().startswith(PREFIX))
        if names:
            found[kind] = names
    return found


def verify_isolation(name: str, net: str, vol: str | None, anon: list[str]) -> dict:
    """컨테이너가 이 실행의 내부 네트워크 하나에만 붙고, 호스트 포트가 없고, (backend는) 이 실행의 임시 볼륨만 쓰는지 조회해 확인한다.
    다르면 예외 — 운영과 분리됐다고 확인하지 못한 채 점검하지 않는다."""
    rc, internal, err = run("network", "inspect", net, "--format", "{{.Internal}}", timeout=60)
    fmt = "{{json .HostConfig.PortBindings}}|{{range .Mounts}}{{.Type}}:{{.Name}}:{{.Destination}};{{end}}|{{range $k, $v := .NetworkSettings.Networks}}{{$k}};{{end}}"
    rc2, out, err2 = run("inspect", name, "--format", fmt, timeout=60)
    if rc != 0 or rc2 != 0:
        raise RuntimeError(f"격리 상태 조회 실패 (network inspect {rc}: {err.strip()[:120]} · inspect {rc2}: {err2.strip()[:120]})")
    ports, mounts, nets = (out.strip().split("|") + ["", "", ""])[:3]
    for m in mounts.split(";"):  # 익명 볼륨은 판정 전에 기록 — 아래에서 멈춰도 정리 때 남지 않았는지 확인한다
        parts = m.split(":")
        if len(parts) >= 2 and parts[0] == "volume" and ANON_VOLUME.fullmatch(parts[1]) and parts[1] not in anon:
            anon.append(parts[1])
    want_mounts = f"volume:{vol}:/app/data;" if vol else ""
    seen = {"internal_network": internal.strip(), "port_bindings": ports or "null", "mounts": mounts, "networks": nets}
    problems = []
    if internal.strip() != "true":
        problems.append("네트워크가 --internal이 아님")
    if ports not in ("{}", "null", ""):
        problems.append(f"호스트 포트 연결 있음 {ports}")
    if mounts != want_mounts:
        problems.append(f"볼륨 연결이 다름 {mounts or '(없음)'} (기대 {want_mounts or '(없음)'})")
    if nets != f"{net};":
        problems.append(f"붙은 네트워크가 다름 {nets}")
    if PROD_VOLUME in mounts:
        problems.append("운영 데이터 볼륨이 연결됨")
    if problems:
        raise RuntimeError(f"{name} 격리 확인 실패: {'; '.join(problems)}")
    return seen


def write_atomic(path: str, data: dict) -> None:
    tmp = f"{path}.tmp-{os.getpid()}"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=1)
    os.replace(tmp, path)


def main() -> int:
    sys.stdout.reconfigure(errors="replace")  # Windows 콘솔(cp949)에서 못 쓰는 글자로 멈추지 않게
    if len(sys.argv) != 4:
        print(__doc__)
        return 2
    b_img, f_img, out = sys.argv[1:4]
    if os.path.exists(out):  # 앞 실행의 결과가 이번 결과로 읽히지 않게
        os.remove(out)
    tag = f"moajwo-smoke-{time.strftime('%Y%m%d%H%M%S')}-{os.getpid()}-{secrets.token_hex(3)}"
    net, vol, bname, fname = tag + "-net", tag + "-data", tag + "-backend", tag + "-frontend"
    assert vol != PROD_VOLUME
    checks: dict[str, dict] = {}
    errors: list[str] = []
    made: dict[str, list[str]] = {"container": [], "volume": [], "network": []}
    anon: list[str] = []  # 점검 컨테이너에 붙은 익명 볼륨 (격리 확인에서 봄)

    def rec(cid: str, ok: bool, detail: object = None) -> None:
        checks[cid] = {"id": cid, "item": NAMES[cid], "ok": ok is True, "detail": detail, "measured": True}

    started = time.strftime("%Y-%m-%dT%H:%M:%S%z")
    ids = {"backend": b_img, "frontend": f_img}
    cleanup_detail: dict = {}
    isolation: dict = {"tag": tag, "verified": False}
    try:
        ids = {k: must("image", "inspect", img, "--format", "{{.Id}}") for k, img in ids.items()}
        leftovers = previous_leftovers()  # 앞 실행이 남긴 점검 자원 — 지우지 않고(이 실행이 만든 것이 아님) 알리고 멈춘다
        if leftovers:
            raise RuntimeError(f"앞 실행이 남긴 점검 자원이 있음 — 확인 뒤 직접 지울 것: {leftovers}")
        made["network"].append(net)  # 만들기 전에 기록 — 만들다 실패해도 정리 대상이 되게
        must("network", "create", "--internal", net)
        made["volume"].append(vol)
        must("volume", "create", vol)
        made["container"].append(bname)
        must("run", "-d", "--name", bname, "--network", net, "--network-alias", "backend", "-v", f"{vol}:/app/data",
             "-e", "OLLAMA_BASE_URL=http://127.0.0.1:9", "--memory", "1g", "--memory-swap", "1g", "--cpus", "2", ids["backend"])  # fmt: skip
        isolation["backend"] = verify_isolation(bname, net, vol, anon)  # 고정 문구가 아니라 조회한 값 — 다르면 점검하지 않고 멈춘다

        def health() -> tuple[bool, bool, str, bool]:
            rc, o, e = run("inspect", bname, "--format", "{{.State.Status}} {{.State.Health.Status}} {{.RestartCount}}", timeout=60)
            if rc != 0:
                return False, False, f"조회 실패 {rc}: {e.strip()[:160]}", False
            words = o.split()
            return True, words[1:2] == ["healthy"], o.strip(), words[:1] in (["exited"], ["dead"])

        healthy, queried, last = wait_for(health, 150)
        if healthy or queried:
            rec("host.backend_healthy", healthy, {"state": last})
        else:  # 마지막 상태 조회가 실패 — backend의 실패로 적지 않는다
            errors.append(f"backend 상태를 조회하지 못해 판정할 수 없음 (마지막 조회 실패): {last}")
        if healthy:
            made["container"].append(fname)
            must("run", "-d", "--name", fname, "--network", net, "--network-alias", "frontend", "--memory", "256m", "--memory-swap", "256m",
                 "--cpus", "1", ids["frontend"])  # fmt: skip
            isolation["frontend"] = verify_isolation(fname, net, None, anon)
            isolation["verified"] = True

            def front_up() -> tuple[bool, bool, str, bool]:
                # wget 자체 읽기 시간 제한(-T 10)을 docker 명령 시간 제한(30초)보다 짧게 — 멈춘 frontend는 wget 실패(측정한 실패)로,
                # docker 명령 자체의 시간 초과·데몬 연결 실패만 조회 실패로 남는다
                rc, _, e = run("exec", fname, "wget", "-q", "-T", "10", "-O", "/dev/null", "http://127.0.0.1:8080/", timeout=30)
                return rc != -1 and not DOCKER_DOWN.search(e), rc == 0, f"wget 종료 {rc} {e.strip()[:160]}", False

            up, queried, last = wait_for(front_up, 60)
            if up or queried:
                rec("host.frontend_up", up, {"last": last})
            else:  # docker 쪽에서 실패(시간 초과·데몬 연결 실패) — frontend의 실패로 적지 않는다
                errors.append(f"frontend 응답을 조회하지 못해 판정할 수 없음 (docker 쪽 실패): {last}")
            if up:
                rc, text, err = run("exec", "-i", bname, "python", "-I", "-", stdin=CHECK_PY.replace("__CHECKS_IN__", repr(json.dumps(CHECKS_IN))).encode("utf-8"), timeout=900)
                lines = [x.rstrip("\r") for x in text.split("\n") if x.startswith(RESULT_MARK)]  # splitlines는 U+2028 등에서도 끊는다
                if rc != 0:
                    errors.append(f"컨테이너 점검 실행 종료 {rc}: {err.strip()[-300:]}")
                if len(lines) != 1:
                    errors.append(f"컨테이너 결과 표시 줄이 {len(lines)}개 (1개여야 함) · 출력 끝: {text[-300:]!r}")
                else:
                    try:
                        got, errs = validate_child(json.loads(lines[0][len(RESULT_MARK) :]))
                    except ValueError as e:
                        got, errs = {}, [f"컨테이너 결과 JSON 해석 실패: {e}"]
                    checks.update(got)
                    errors += errs
        for name, kind, cid in ((bname, "backend", "host.backend_log"), (fname, "frontend", "host.frontend_log")):
            if name not in made["container"]:
                continue
            done, text, info = read_log(name, kind)
            if not done:
                errors.append(f"{kind} 로그 조회 실패 (docker logs 종료 {info['rc']}: {info['stderr_tail']}) — '오류 없음'으로 보지 않음")
                continue
            if kind == "backend":  # uvicorn 'ERROR:    …'·alembic 'ERROR [..]'는 줄 머리 — 그리고 Traceback, 5xx 응답
                info.update({"traceback": text.count("Traceback"), "error": len(re.findall(r"(?m)^(?:ERROR|CRITICAL)\b", text)),
                             "http_5xx": len(re.findall(r'" 5\d\d ', text))})  # fmt: skip
                rec(cid, info["start_mark"] and info["traceback"] == 0 and info["error"] == 0 and info["http_5xx"] == 0, info)
            else:
                info.update({lv: text.count(f"[{lv}]") for lv in ("emerg", "alert", "crit", "error")})
                rec(cid, info["start_mark"] and all(info[lv] == 0 for lv in ("emerg", "alert", "crit", "error")), info)
    except Exception as e:  # noqa: BLE001 — 기록하고 정리로 간다
        errors.append(f"점검 진행 중 오류: {type(e).__name__}: {str(e)[:400]}")
    finally:
        try:
            cleaned, cleanup_detail, cerrs = cleanup(made, tag, anon)
        except Exception as e:  # noqa: BLE001 — 정리 자체의 오류도 원래 오류와 함께 남긴다
            cleaned, cleanup_detail, cerrs = False, {"made": made, "anonymous": anon}, [f"정리 중 오류: {type(e).__name__}: {str(e)[:300]}"]
        errors += cerrs
        if not cerrs and cleaned:  # 정리 실패·확인 불가는 도구 오류로 남고 이 점검은 미측정이 된다
            rec("host.cleanup", True, cleanup_detail)
    rows = []
    for cid, name in REQUIRED:  # 필수 22개를 정한 순서로 — 측정하지 못한 것은 미측정으로 남긴다
        rows.append(checks.get(cid) or {"id": cid, "item": name, "ok": False, "detail": None, "measured": False})
    failed = [r["id"] for r in rows if r["measured"] and r["ok"] is not True]
    unmeasured = [r["id"] for r in rows if not r["measured"]]
    # 정리 실패·확인 불가는 측정한 실패보다 먼저 도구 오류(3) — 남은 자원부터 봐야 한다. 실패한 id는 그대로 failed에 남는다
    status = "error" if cerrs or not cleaned else ("fail" if failed else ("error" if errors or unmeasured else "pass"))
    result = {"tool": TOOL, "started": started, "ended": time.strftime("%Y-%m-%dT%H:%M:%S%z"), "images": ids, "status": status,
              "ok": status == "pass", "required": len(REQUIRED), "passed": sum(1 for r in rows if r["measured"] and r["ok"] is True),
              "failed": failed, "unmeasured": unmeasured, "errors": errors,
              "isolation": isolation, "cleanup": cleanup_detail,
              "checks": rows}  # fmt: skip
    try:
        for r in rows:
            mark = "통과" if r["measured"] and r["ok"] is True else ("실패" if r["measured"] else "미측정")
            print(f"{mark} | {r['id']} | {r['item']} | {json.dumps(r['detail'], ensure_ascii=False)[:300]}")
        for e in errors:
            print(f"도구 오류 | {e}")
        print(f"요약 | 격리 기동 점검 {status} {result['passed']}/{len(REQUIRED)} · 실패 {len(failed)} · 미측정 {len(unmeasured)} · 도구 오류 {len(errors)}"
              f" | backend {ids['backend']} | frontend {ids['frontend']}")  # fmt: skip
        sys.stdout.flush()
    except Exception as e:  # noqa: BLE001 — 출력에 실패하면 통과로 남기지 않는다 (결과 파일 쓰기 실패와 같게 도구 오류 3)
        result["errors"].append(f"출력 실패: {type(e).__name__}: {e}")
        result["status"], result["ok"] = "error", False
    try:
        write_atomic(out, result)
    except OSError as e:
        sys.stderr.write(f"결과 파일을 쓰지 못함: {e}\n")
        return 3
    return {"pass": 0, "fail": 4}.get(result["status"], 3)


if __name__ == "__main__":
    sys.exit(main())
