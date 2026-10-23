"""9장 — 시연 5개 파일의 독립 기대값을 새 계정(B)으로 확인 (지시문 9장 표). 측정과 기대값을 함께 적는다 (판정은 비교만).
기대값: 5×50행×15열(250행·3,750값), 묶기 원본 영역 5×51×15, 발신Mail건수 합계 85,740·평균 342.96·최솟값 41·최댓값 2,150·개수 250,
부처별 합계 37,550·14,510·11,355·4,735·17,590, 요약 끔 5시트. 통합 결과의 열 수는 지금 항목 정의로 다시 잰다 (27열은 예전 기준).
사용법: python s9_demo_expected.py <출력 json> <입력 파일 폴더>"""

import io
import json
import secrets
import sys
from pathlib import Path

import openpyxl

from diag_common import XLSX, Api

OUT, SRC = Path(sys.argv[1]), Path(sys.argv[2])
FILES = [f"{m} 사용량.xlsx" for m in ("행안부", "기획예산처", "재정경제부", "외교부", "통일부")]
EXPECT_SUMS = {"행안부": 37550, "기획예산처": 14510, "재정경제부": 11355, "외교부": 4735, "통일부": 17590}
DEMO = {"날짜": "base_date", "기준일자": "base_date", "통계일자": "base_date", "기관코드": "org_code",
        "조직명": "department", "조직 명": "department", "부서명": "department", "발신Mail건수": "count"}  # fmt: skip
owner = Api().login("owner")
name, pw = "s9b-" + secrets.token_hex(3), "Sb-" + secrets.token_urlsafe(14) + "7z"
assert owner.call("POST", "/users", {"username": name, "display_name": "기대값 확인 B", "status": "active", "password": pw})[0] == 201
u = Api()
assert u.call("POST", "/auth/login", {"username": name, "password": pw})[0] == 200
rec: dict = {"files": []}
fids = []
for fname in FILES:
    s, body, _ = u.upload(fname, (SRC / fname).read_bytes(), XLSX)
    fid = body[0]["id"]
    pv = u.call("GET", f"/files/{fid}/preview?sheet_id=s0")[1]
    cols = pv["columns"]
    u.call("PUT", f"/files/{fid}/mapping/decisions", {"decisions": [{"sheet_id": "s0", "column_id": c["column_id"], "field_key": DEMO.get(c["name"])} for c in cols]})
    fids.append(fid)
    rec["files"].append({"file": fid, "name": fname, "columns": len(cols), "data_rows": pv.get("data_rows"), "total_rows": pv.get("total_rows")})
rec["input"] = {"rows": sum(f["data_rows"] for f in rec["files"]), "values": sum(f["data_rows"] * f["columns"] for f in rec["files"])}

# 부처별 합계 (파일 하나씩 분석). 통일부는 같은 조직 반복 행을 '서로 다른 기록'으로 확인해 모두 더한다 (시연과 같음)
sums = {}
for fname, fid in zip(FILES, fids, strict=True):
    # 표에 단위가 없어 '건'으로 확인, 같은 조직 반복 행은 '서로 다른 기록'으로 확인 (시연과 같은 확인)
    req = {"field_key": "count", "targets": [{"file_id": fid, "unit": "건", "rows_are_records": True}]}
    s, pv, _ = u.call("POST", "/analysis/preview", req)
    s, run, _ = u.call("POST", "/analysis/run", {**req, "expected_fingerprint": pv.get("fingerprint")})
    if s != 201:
        sums[fname] = {"status": s, "detail": str(run)[:200]}
        continue
    groups = u.call("GET", f"/analysis/runs/{run['id']}/groups?limit=10")[1]["items"]
    g = groups[0]
    sums[fname] = {"sum": g["sum"], "count": g["count"], "expected": EXPECT_SUMS[fname.split()[0]]}
rec["per_ministry"] = sums

# 묶기 요약 끔 → 5시트, 원본 값 대조
body = {"file_ids": fids}
s, plan, _ = u.call("POST", "/bundle/plan", body)
s, run, _ = u.call("POST", "/bundle/run", {**body, "expected_basis": plan["basis"]})
book = u.call("GET", f"/runs/{run['id']}/download", raw=True)[1]
wb = openpyxl.load_workbook(io.BytesIO(book), data_only=True)
mism, compared = 0, 0
for fname, ws in zip(FILES, wb.worksheets, strict=False):
    src = openpyxl.load_workbook(SRC / fname, data_only=True).worksheets[0]
    src_rows = [list(r) for r in src.iter_rows(values_only=True)]
    out_rows = [list(r) for r in ws.iter_rows(values_only=True)]
    lead = len(out_rows[0]) - len(src_rows[0])  # 앱이 앞에 붙인 열 (원본 행·행 종류)
    for i, row in enumerate(src_rows):
        for j, v in enumerate(row):
            compared += 1
            got = out_rows[i][lead + j] if i < len(out_rows) and lead + j < len(out_rows[i]) else "(없음)"
            same = got == v or (isinstance(v, (int, float)) and isinstance(got, (int, float)) and float(got) == float(v))
            if not same and not (hasattr(v, "isoformat") and str(got).startswith(v.isoformat()[:10])):
                mism += 1
rec["bundle_summary_off"] = {"sheets": wb.sheetnames, "sheet_count": len(wb.sheetnames), "cells_compared": compared, "mismatches": mism,
                             "sheet_dims": [[ws.title, ws.max_row, ws.max_column] for ws in wb.worksheets]}  # fmt: skip

# 통합표 (5개 파일, 검토 → 완료 → 통합) — 행·열·빈칸을 지금 기준으로 잰다
s, review, _ = u.call("POST", "/reviews", {"file_ids": fids})
for item in review["items"]:
    if item["category"] == "pending":
        cands = (item.get("payload") or {}).get("candidates") or []
        choice = cands[0].get("code") if cands and item["kind"] == "entity" else None
        u.call("POST", f"/reviews/{review['id']}/items/{item['id']}/decision", {"decision": "approved" if (choice or item["kind"] == "pii") else "rejected", "choice": choice})
s, done, _ = u.call("POST", f"/reviews/{review['id']}/complete")
rec["integration"] = {"complete": s, "summary": (done or {}).get("summary") if isinstance(done, dict) else str(done)[:200]}
if s == 200:
    s, run, _ = u.call("POST", f"/reviews/{review['id']}/integrate", {"include_unmapped": True})
    if s == 201:
        wb = openpyxl.load_workbook(io.BytesIO(u.call("GET", f"/runs/{run['id']}/download", raw=True)[1]), data_only=True)
        ws = wb["통합 데이터"]
        rows = list(ws.iter_rows(values_only=True))
        blanks = sum(1 for r in rows[1:] for c in r if c is None)
        rec["integration"] |= {"status": s, "rows": len(rows) - 1, "columns": len(rows[0]), "blank_cells": blanks, "header": list(rows[0])}
    else:
        rec["integration"] |= {"status": s, "detail": str(run)[:200]}
print(json.dumps(rec, ensure_ascii=False, indent=1)[:4000])
OUT.write_text(json.dumps(rec, ensure_ascii=False, indent=1, default=str), encoding="utf-8")
