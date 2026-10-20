"""기존 업무 흐름이 접근성 수정 뒤에도 그대로 되는지 (진단 격리 스택, 합성 자료, 화면 조작 위주).

A. 가입·최초 설정 화면(도움말 추가)의 320 CSS px 재배치
B. 업로드(화면, 키보드) → 열 매핑 확정(API — 이번 수정과 무관한 화면) → 홈에서 파일 고르기·취합 방식 라디오 →
   사전 검토 열기 → 승인 → 검토 완료 → 통합 결과 만들기 → 내려받기
C. 파일별 묶기: 요약 끔 → 엑셀 만들기 → 내려받기, 요약 켬 → 엑셀 만들기 → 내려받기
사용법: python flow_check.py <출력 폴더>
"""

import io
import json
import os
import re
import sys
from pathlib import Path

import openpyxl
from playwright.sync_api import Page, sync_playwright

from diag_common import EV, URL, USERS, Api, creds
from diag_scope import browser_proxy, route_guard

OUT = Path(sys.argv[1])
OUT.mkdir(parents=True, exist_ok=True)
IDS = json.loads((EV / "seed-ids.json").read_text(encoding="utf-8"))
RES: list[dict] = []
BLOCKED: list[str] = []
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
ONLY = [x for x in os.environ.get("ONLY_FLOWS", "").split(",") if x]  # 예: ONLY_FLOWS=bundle (비우면 전부)


def rec(step: str, ok: bool, detail="") -> None:
    RES.append({"step": step, "ok": bool(ok), "detail": str(detail)[:300]})
    print(f"{'통과' if ok else '실패'} | {step} | {str(detail)[:200]}", flush=True)


def login(pg: Page, who: str) -> None:
    pg.goto(URL + "/")
    pg.get_by_label("아이디").fill(USERS[who])
    pg.get_by_label("비밀번호", exact=True).fill(creds()[who])
    pg.locator(".login-card button[type=submit]").click()
    pg.wait_for_selector(".shell", timeout=60000)
    pg.wait_for_load_state("networkidle")


def sheets_of(path: Path) -> dict:
    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        return {ws.title: ws.max_row for ws in wb.worksheets}
    finally:
        wb.close()  # 읽기 전용 모드는 파일을 열어 둔다 (닫지 않으면 Windows에서 삭제 실패)


def flow_signup_reflow(b) -> None:
    cx = b.new_context(viewport={"width": 320, "height": 640}, locale="ko-KR")
    cx.route("**/*", route_guard(URL, BLOCKED))
    pg = cx.new_page()
    pg.goto(URL + "/")
    pg.get_by_role("button", name=re.compile("회원가입")).first.click()
    pg.wait_for_timeout(500)
    over = pg.evaluate("document.scrollingElement.scrollWidth - document.documentElement.clientWidth")
    pg.screenshot(path=str(OUT / "signup-320.png"), full_page=True)
    rec("A. 회원가입 화면 320px 페이지 가로 넘침 없음", over <= 0, f"{over}px, 도움말 {pg.locator('.login-help').count()}개")
    cx.close()


def flow_integrate(b) -> None:
    cx = b.new_context(viewport={"width": 1280, "height": 900}, locale="ko-KR", accept_downloads=True)
    cx.route("**/*", route_guard(URL, BLOCKED))
    pg = cx.new_page()
    login(pg, "user1")
    wb = openpyxl.Workbook()
    for r in [["주민등록 인구 현황 2026년 9월 (단위: 명)"], ["시도", "인구수"], ["강원특별자치도", 501], ["충청북도", 502]]:
        wb.active.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    # 키보드 조작 요소: 수정 전은 영역(role=button), 접근성 보완 뒤는 영역 안의 '파일 선택' 버튼
    drop = pg.locator(".wk-home .upload-drop[role=button], .wk-home .upload-drop .wk-drop-button").first
    drop.focus()
    with pg.expect_file_chooser() as fc:
        pg.keyboard.press("Enter")
    fc.value.set_files(files=[{"name": "흐름확인_인구.xlsx", "mimeType": XLSX, "buffer": buf.getvalue()}])
    pg.wait_for_function("() => !document.querySelector('.upload-drop.is-busy')", timeout=120000)
    pg.wait_for_timeout(800)
    status = pg.locator(".wk-upload [role=status]").first.inner_text()
    card = pg.locator(".upload-file-name", has_text="흐름확인_인구.xlsx").first
    fid = int(re.search(r"#(\d+)", card.inner_text()).group(1))
    rec("B1. 화면에서 키보드로 업로드", fid > 0, f"#{fid} · 상태 안내: {status}")
    api = Api().login("user1")
    s, pv, _ = api.call("GET", f"/files/{fid}/preview?sheet_id=s0")
    cols = {x["name"]: x["column_id"] for x in pv["columns"]}
    s, _, _ = api.call("PUT", f"/files/{fid}/mapping/decisions", {"decisions": [{"sheet_id": "s0", "column_id": cols["시도"], "field_key": "sido"},
                                                                                  {"sheet_id": "s0", "column_id": cols["인구수"], "field_key": "population"}]})  # fmt: skip
    rec("B2. 열 매핑 확정 (API)", s == 200, s)
    pg.goto(URL + "/")
    pg.wait_for_load_state("networkidle")
    pg.get_by_label("검사할 파일 검색").fill("흐름확인")
    pg.wait_for_timeout(600)
    pg.get_by_label(re.compile(rf"#{fid} .*검증 대상")).check()
    table = pg.locator("input[name=wk-result-mode][value=table]")
    pg.locator(".wk-mode-card", has=table).locator(".wk-mode-text strong").click()  # 라벨을 눌러 '하나의 통합표' 선택
    rec("B3. 파일 고르기·취합 방식 라디오(라벨 클릭)", table.is_checked(), pg.evaluate("[...document.querySelectorAll('input[name=wk-result-mode]')].map(e => e.value + ':' + e.checked)"))
    pg.get_by_role("button", name=re.compile("사전 검토 열기")).click()
    panel = pg.locator("section[aria-label='통합 전 사전 검토']")
    panel.locator(".review-score-value").wait_for(timeout=120000)
    pg.wait_for_timeout(800)
    approved = 0
    for _ in range(10):
        btn = panel.get_by_role("button", name="승인", exact=True)
        if not btn.count() or not btn.first.is_enabled():
            break
        btn.first.click()
        pg.wait_for_timeout(800)
        approved += 1
    rec("B4. 사전 검토 열기·확인 항목 승인", True, f"승인 {approved}건 · 점수 영역: {panel.locator('.review-score-value').first.aria_snapshot()!r}")
    done = panel.get_by_role("button", name=re.compile("^검토 완료"))
    if done.count() and done.first.is_enabled():
        done.first.click()
        pg.wait_for_timeout(1500)
    make = pg.get_by_role("button", name=re.compile("통합 결과 만들기"))
    make.first.wait_for(timeout=60000)
    make.first.click()
    link = pg.locator("a", has_text=re.compile(r"통합결과_.*\.xlsx")).first
    link.wait_for(timeout=120000)
    with pg.expect_download(timeout=60000) as d:
        link.click()
    path = OUT / "flow-integration.xlsx"
    d.value.save_as(path)
    sh = sheets_of(path)
    rec("B5. 검토 완료 → 통합 결과 만들기 → 내려받기", "통합 데이터" in sh, sh)
    path.unlink()
    cx.close()


def flow_bundle(b) -> None:
    cx = b.new_context(viewport={"width": 1280, "height": 900}, locale="ko-KR", accept_downloads=True)
    cx.route("**/*", route_guard(URL, BLOCKED))
    pg = cx.new_page()
    login(pg, "user1")
    pg.goto(f"{URL}/bundle?files={IDS['user1']['basic']},{IDS['user1']['many']}")
    pg.wait_for_load_state("networkidle")
    pg.wait_for_timeout(800)
    sw = pg.locator("label:has(input.peer.sr-only[type=checkbox]) input[type=checkbox]").first
    for want in (False, True):
        if sw.is_checked() != want:
            sw.focus()
            pg.keyboard.press("Space")
            pg.wait_for_timeout(1200)
        if want:  # 요약을 켜면 항목이 하나 이상 있어야 만들 수 있다 (기존 규칙) — 첫 후보 항목을 기본 계산(합계)으로 추가
            pg.get_by_role("button", name="항목 추가", exact=True).click()
            pg.wait_for_timeout(1500)
        make = pg.get_by_role("button", name=re.compile("엑셀 만들기"))
        enabled = make.first.is_enabled()
        if not enabled:
            rec(f"C. 요약 {'켬' if want else '끔'} — 엑셀 만들기", False, "버튼 비활성: " + pg.locator(".wk-why").first.inner_text()[:150] if pg.locator(".wk-why").count() else "버튼 비활성")
            continue
        make.first.click()
        link = pg.get_by_role("link", name="내려받기").first
        link.wait_for(timeout=120000)
        with pg.expect_download(timeout=60000) as d:
            link.click()
        path = OUT / f"flow-bundle-{'on' if want else 'off'}.xlsx"
        d.value.save_as(path)
        sh = sheets_of(path)
        rec(f"C. 요약 {'켬' if want else '끔'} — 엑셀 만들기 → 내려받기", True, f"스위치 {sw.is_checked()} · 시트 {sh}")
        path.unlink()
        pg.goto(f"{URL}/bundle?files={IDS['user1']['basic']},{IDS['user1']['many']}")
        pg.wait_for_load_state("networkidle")
        pg.wait_for_timeout(800)
    cx.close()


def main() -> None:
    with sync_playwright() as p:
        b = p.chromium.launch(channel="msedge", headless=True, proxy=browser_proxy(URL))
        for name, fn in (("signup", flow_signup_reflow), ("integrate", flow_integrate), ("bundle", flow_bundle)):
            if ONLY and name not in ONLY:
                continue
            try:
                fn(b)
            except Exception as e:  # noqa: BLE001 — 실패로 기록 (통과로 세지 않음)
                rec(f"오류: {name}", False, f"{type(e).__name__}: {str(e)[:250]}")
        b.close()
    (OUT / "flow-check.json").write_text(json.dumps({"results": RES, "blocked_out_of_scope": BLOCKED}, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"RESULT 통과 {sum(r['ok'] for r in RES)} · 실패 {sum(not r['ok'] for r in RES)} · 막힌 범위 밖 요청 {len(BLOCKED)}")


if __name__ == "__main__":
    main()
