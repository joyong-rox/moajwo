"""진단 스크립트 공용 도구 — 진단 격리 스택(127.0.0.1:8086)만 상대한다. 운영 주소는 쓰지 않는다.

요청 허용 범위는 diag_scope.py가 정한다 (DIAG_URL은 http://127.0.0.1:8086만 허용, 리다이렉트도 그 안에서만).
계정 비밀번호는 실행 때 무작위로 만들어 Git 밖 증거 폴더(DIAG_EVIDENCE, 기본 <저장소>/backups/diag-20261020)의
creds.json에만 둔다.
"""

import http.cookiejar
import io
import json
import os
import secrets
import subprocess
import urllib.error
import urllib.request
import uuid
from pathlib import Path

import openpyxl

from diag_scope import ScopeError, check_url, diag_base, opener

URL = diag_base(os.environ.get("DIAG_URL"))  # 허용 밖이면 ScopeError로 멈춘다
API = URL + "/api"
ROOT = Path(__file__).resolve().parents[2]
EV = Path(os.environ.get("DIAG_EVIDENCE", ROOT / "backups" / "diag-20261020"))
CREDS = EV / "creds.json"
BACKEND = "moajwo-diag-backend-1"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
USERS = {"owner": "diag-owner", "user1": "diag-user1", "user2": "diag-user2", "user3": "diag-user3"}


def creds() -> dict[str, str]:
    if CREDS.exists():
        return json.loads(CREDS.read_text(encoding="utf-8"))
    c = {k: "Dg-" + secrets.token_urlsafe(14) + "7z" for k in (*USERS, "pending")}
    EV.mkdir(parents=True, exist_ok=True)
    CREDS.write_text(json.dumps(c), encoding="utf-8")
    return c


class Api:
    def __init__(self) -> None:
        self.jar = http.cookiejar.CookieJar()
        self.op = opener(URL, urllib.request.HTTPCookieProcessor(self.jar))

    def call(self, method: str, path: str, body=None, data=None, headers=None, raw=False):
        if body is not None:
            data, headers = json.dumps(body).encode(), {"Content-Type": "application/json", **(headers or {})}
        if not path.startswith("/"):
            raise ScopeError(f"API 경로는 /로 시작해야 함: {path!r}")
        req = urllib.request.Request(check_url(API + path, URL), data=data, method=method, headers=headers or {})
        try:
            with self.op.open(req, timeout=300) as r:
                c = r.read()
                return r.status, c if raw else _parse(c), dict(r.headers)
        except urllib.error.HTTPError as e:
            c = e.read()
            return e.code, c if raw else _parse(c), dict(e.headers)

    def login(self, who: str) -> "Api":
        s, b, _ = self.call("POST", "/auth/login", {"username": USERS[who], "password": creds()[who]})
        if s != 200:
            raise RuntimeError(f"{who} 로그인 실패: {s}")
        return self

    def upload(self, name: str, content: bytes, ctype: str = XLSX):
        b = uuid.uuid4().hex
        body = (f'--{b}\r\nContent-Disposition: form-data; name="files"; filename="{name}"\r\n'
                f"Content-Type: {ctype}\r\n\r\n").encode() + content + f"\r\n--{b}--\r\n".encode()  # fmt: skip
        return self.call("POST", "/files", data=body, headers={"Content-Type": f"multipart/form-data; boundary={b}"})


def _parse(c: bytes):
    if c and c[:1] in b"[{":
        try:
            return json.loads(c)
        except ValueError:
            return c
    return c


def xlsx(sheets: dict[str, list[list]]) -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith(("=", "+", "-", "@")):
                    cell.data_type = "s"  # 시험 자료: 수식처럼 보이는 글자를 글자로 둔다
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def db(sql: str) -> list:
    """진단 스택 DB 읽기 전용 조회 (운영 DB 아님)."""
    code = f"import sqlite3,json; print(json.dumps(sqlite3.connect('file:/app/data/moajwo.db?mode=ro', uri=True).execute({sql!r}).fetchall(), default=str))"
    p = subprocess.run(["docker", "exec", BACKEND, "python", "-c", code], capture_output=True, text=True, encoding="utf-8")
    if p.returncode != 0:
        raise RuntimeError(f"진단 DB 조회 실패: {p.stderr[-400:]}")
    return json.loads(p.stdout)


# 측정 결과의 출처 (a11y_regress.py가 출처 표·재사용 판단에 쓴다). 빌드 자산 파일 이름(내용 해시)으로 어떤 빌드를 쟀는지 구분한다
ASSET_JS = r"() => { const s = [...document.scripts].map(x => x.src).find(x => /\/assets\/index-[^/]+\.js$/.test(x)); return s ? s.split('/').pop() : null }"


def run_meta(script: str, label: str, asset: str | None) -> dict:
    from datetime import datetime

    return {"script": script, "label": label, "started": datetime.now().astimezone().isoformat(timespec="seconds"), "url": URL,
            "frontend_asset": asset, "frontend_image": os.environ.get("DIAG_FE_IMAGE") or "미기재 (DIAG_FE_IMAGE 없음)"}  # fmt: skip
