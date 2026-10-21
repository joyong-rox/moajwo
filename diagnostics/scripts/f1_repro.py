"""F-1 재현·확인 — 가린 뒤 같아지는 구분 값이 분석·묶기 판정에 미치는 영향 (격리 진단 스택, 합성 자료·user2만, 측정만 저장).

수정 전 이미지와 수정 후 이미지에서 같은 요청을 보내 그룹·필터·반복·중복·구분(열 이름)·묶기 요약·저장값을 기록한다.
원문 이메일·전화가 응답·저장 결과·Excel에 새로 나오는지도 찾는다 (찾은 개수만 남기고 원문은 남기지 않음).
사용법: python f1_repro.py <출력 폴더> <라벨>
"""

import io
import json
import sys
from pathlib import Path

import openpyxl
from playwright.sync_api import sync_playwright

from diag_common import ASSET_JS, URL, USERS, Api, creds, db, run_meta, xlsx
from diag_scope import browser_proxy, route_guard

OUT, LABEL = Path(sys.argv[1]), sys.argv[2]
OUT.mkdir(parents=True, exist_ok=True)
TITLE = "주민등록 인구 현황 2026년 9월 (단위: 명)"
TAG = f"F1_{LABEL}_"
RAW = ["kim1@example.com", "kim2@example.com", "010-1234-5678", "010-9999-5678"]  # 합성 원문 — 결과에 나오면 안 됨
REC: dict[str, list] = {}
BLOCKED: list[str] = []


def rec(group: str, item: dict) -> None:
    REC.setdefault(group, []).append(item)
    print(group, json.dumps(item, ensure_ascii=False)[:260], flush=True)


def raw_hits(text: str) -> int:
    return sum(text.count(r) for r in RAW)


def up(u: Api, name: str, rows: list[list], sheet: str = "Sheet1") -> int:
    s, b, _ = u.upload(f"{TAG}{name}", xlsx({sheet: rows}))
    return b[0]["id"]


def confirm(u: Api, fid: int, mapping: dict[str, str | None]) -> None:
    pv = u.call("GET", f"/files/{fid}/preview?sheet_id=s0")[1]
    cols = {x["name"]: x["column_id"] for x in pv["columns"]}
    s, b, _ = u.call("PUT", f"/files/{fid}/mapping/decisions", {"decisions": [{"sheet_id": "s0", "column_id": cols[n], "field_key": k} for n, k in mapping.items()]})
    if s != 200:
        raise RuntimeError(f"확정 {fid}: {s} {b}")
    u.call("POST", f"/files/{fid}/pii/acknowledge", {})


def analyze(u: Api, what: str, body: dict) -> dict:
    s, pv, _ = u.call("POST", "/analysis/preview", body)
    out = {"what": what, "preview": s, "problems": pv.get("problems") if isinstance(pv, dict) else str(pv)[:300],
           "targets": [(t["status"], t["reasons"]) for t in pv.get("targets", [])] if isinstance(pv, dict) else None,
           "filter_values": pv.get("filter_values") if isinstance(pv, dict) else None, "raw_in_preview": raw_hits(json.dumps(pv, ensure_ascii=False))}  # fmt: skip
    if s == 200 and not pv["problems"] and all(t["status"] != "needs_input" for t in pv["targets"]):
        sr, run, _ = u.call("POST", "/analysis/run", body | {"expected_fingerprint": pv["fingerprint"]})
        out["run"] = sr
        if sr == 201:
            rid = run["id"]
            groups = u.call("GET", f"/analysis/runs/{rid}/groups?offset=0&limit=200")[1]["items"]
            detail = u.call("GET", f"/analysis/runs/{rid}/detail?offset=0&limit=200")[1]["items"]
            meta = u.call("GET", f"/analysis/runs/{rid}")[1]
            s2, content, _ = u.call("GET", f"/analysis/runs/{rid}/download", raw=True)
            wb = openpyxl.load_workbook(io.BytesIO(content))
            xl_groups = [[c.value for c in r] for r in wb["결과"].iter_rows(min_row=2)]
            stored = db(f"select summary, conditions, inputs, result from analysis_runs where id = {int(rid)}")[0]
            out |= {"id": rid, "groups": [(g["keys"], g["count"], g["sum"]) for g in groups], "detail_keys": [(d["row"], d["keys"]) for d in detail],
                    "excel_groups": xl_groups, "policy": meta["conditions"].get("policy"),
                    "stored_group_keys_masked": ["*" in json.dumps(g, ensure_ascii=False) for g in json.loads(stored[3])["groups"]],
                    "raw_in_stored": raw_hits(" ".join(stored)), "raw_in_api": raw_hits(json.dumps([groups, detail, meta], ensure_ascii=False)),
                    "raw_in_excel": raw_hits(json.dumps([[c.value for c in r] for ws in wb.worksheets for r in ws.iter_rows()], ensure_ascii=False, default=str))}  # fmt: skip
        else:
            out["run_body"] = str(run)[:300]
    rec("analysis", out)
    return out


def main() -> None:
    u = Api().login("user2")
    hdr = ["시도", "시군구", "인구수"]
    # 1) 그룹: 원문이 다르고 가림값이 같은 이메일 (시도가 달라 반복 아님)
    f_mail = up(u, "이메일충돌.xlsx", [[TITLE], hdr, ["서울특별시", "가팀 kim1@example.com", 10], ["부산광역시", "가팀 kim2@example.com", 20]])
    confirm(u, f_mail, {"시도": "sido", "시군구": "sigungu", "인구수": "population"})
    a1 = analyze(u, "그룹: 이메일 충돌 (시군구 기준)", {"field_key": "population", "targets": [{"file_id": f_mail}], "group_by": ["sigungu"]})
    # 2) 그룹: 전화번호 충돌
    f_tel = up(u, "전화충돌.xlsx", [[TITLE], hdr, ["서울특별시", "나팀 010-1234-5678", 10], ["부산광역시", "나팀 010-9999-5678", 20]])
    confirm(u, f_tel, {"시도": "sido", "시군구": "sigungu", "인구수": "population"})
    analyze(u, "그룹: 전화 충돌", {"field_key": "population", "targets": [{"file_id": f_tel}], "group_by": ["sigungu"]})
    # 3) 그룹 기준 여럿 (시도·시군구)
    analyze(u, "그룹: 시도·시군구 함께", {"field_key": "population", "targets": [{"file_id": f_mail}], "group_by": ["sido", "sigungu"]})
    # 4) 같은 원문 두 행 (서로 다른 기록으로 확인) → 같은 그룹이어야 함
    f_same = up(u, "같은원문.xlsx", [[TITLE], hdr, ["서울특별시", "가팀 kim1@example.com", 10], ["서울특별시", "가팀 kim1@example.com", 20]])
    confirm(u, f_same, {"시도": "sido", "시군구": "sigungu", "인구수": "population"})
    analyze(u, "같은 원문 반복 (서로 다른 기록 확인)", {"field_key": "population", "targets": [{"file_id": f_same, "rows_are_records": True}], "group_by": ["sigungu"]})
    # 5) 반복 판정: 같은 시도, 원문이 다른 시군구 → 반복이 아님
    f_rep = up(u, "반복판정.xlsx", [[TITLE], hdr, ["서울특별시", "가팀 kim1@example.com", 10], ["서울특별시", "가팀 kim2@example.com", 20]])
    confirm(u, f_rep, {"시도": "sido", "시군구": "sigungu", "인구수": "population"})
    analyze(u, "반복 판정: 같은 시도·원문이 다른 시군구", {"field_key": "population", "targets": [{"file_id": f_rep}], "group_by": ["sigungu"]})
    # 6) 파일 사이 중복 판정
    f_b = up(u, "중복판정B.xlsx", [[TITLE], hdr, ["부산광역시", "가팀 kim1@example.com", 30]])
    confirm(u, f_b, {"시도": "sido", "시군구": "sigungu", "인구수": "population"})
    f_a = up(u, "중복판정A.xlsx", [[TITLE], hdr, ["부산광역시", "가팀 kim2@example.com", 20]])
    confirm(u, f_a, {"시도": "sido", "시군구": "sigungu", "인구수": "population"})
    analyze(u, "파일 사이 중복 판정: 같은 시도·원문이 다른 시군구", {"field_key": "population", "targets": [{"file_id": f_a}, {"file_id": f_b}], "group_by": ["sigungu"]})
    # 6b) 반복 판정: 지역 코드 칸만 있는 자료, 원문이 다르고 가림값이 같은 코드 값 → 반복이 아님
    f_code = up(u, "지역코드판정.xlsx", [[TITLE], ["지역코드", "인구수"], ["kim1@example.com", 10], ["kim2@example.com", 20]])
    confirm(u, f_code, {"지역코드": "region_code", "인구수": "population"})
    analyze(u, "반복 판정: 지역 코드 칸만·원문이 다른 값", {"field_key": "population", "targets": [{"file_id": f_code}]})
    # 7) 구분 필터: 한 원문 값만 고르려는 경우
    fv = (a1.get("filter_values") or {}).get("sigungu") or []
    for v in fv:
        analyze(u, f"필터: 시군구 = {v}", {"field_key": "population", "targets": [{"file_id": f_mail}], "group_by": ["sigungu"], "filters": {"dims": {"sigungu": [v]}}})
    analyze(u, "필터: 가린 값 그대로(구분 번호 없이)", {"field_key": "population", "targets": [{"file_id": f_mail}], "group_by": ["sigungu"],
                                                    "filters": {"dims": {"sigungu": ["가팀 k***@example.com"]}}})  # fmt: skip
    # 8) 열 이름의 구분(variant): 원문이 다르고 가림값이 같은 두 열 이름 (두 파일)
    v1 = up(u, "구분열1.xlsx", [[TITLE], ["시도", "인구_kim1@example.com"], ["대전광역시", 5]])
    v2 = up(u, "구분열2.xlsx", [[TITLE], ["시도", "인구_kim2@example.com"], ["울산광역시", 6]])
    for f in (v1, v2):  # 미리보기의 열 이름은 가려져 오므로 열 번호(c0·c1)로 확정한다
        s, body, _ = u.call("PUT", f"/files/{f}/mapping/decisions", {"decisions": [{"sheet_id": "s0", "column_id": "c0", "field_key": "sido"},
                                                                               {"sheet_id": "s0", "column_id": "c1", "field_key": "population"}]})  # fmt: skip
        if s != 200:
            raise RuntimeError(f"확정 {f}: {s} {body}")
    analyze(u, "열 이름 구분: 가림값이 같은 두 열", {"field_key": "population", "targets": [{"file_id": v1}, {"file_id": v2}]})
    # 9) 묶기: 요약 끔·켬 (같은 시도·원문이 다른 시군구 파일)
    for name, summary in (("요약 끔", None), ("요약 켬", {"items": [{"field_key": "population", "aggregates": ["sum", "count"]}]})):
        body = {"file_ids": [f_rep]} | ({"summary": summary} if summary else {})
        before = db("select (select count(*) from runs)")[0][0]
        s, plan, _ = u.call("POST", "/bundle/plan", body)
        item = {"what": f"묶기 {name}", "plan": s, "problems": plan.get("problems") if isinstance(plan, dict) else str(plan)[:300],
                "raw_in_plan": raw_hits(json.dumps(plan, ensure_ascii=False))}  # fmt: skip
        if s == 200 and not plan.get("problems") and plan.get("ready", True):
            sr, run, _ = u.call("POST", "/bundle/run", body | {"expected_basis": plan["basis"]})
            item["run"] = sr
            if sr in (200, 201):
                s2, content, _ = u.call("GET", f"/runs/{run.get('id') or run.get('run_id')}/download", raw=True)
                wb = openpyxl.load_workbook(io.BytesIO(content))
                item["sheets"] = wb.sheetnames
                if "통계 요약" in wb.sheetnames:
                    item["summary_rows"] = [[c.value for c in r][:4] for r in wb["통계 요약"].iter_rows() if any(c.value is not None for c in r)][:30]
                item["raw_in_summary_sheet"] = raw_hits(json.dumps([[c.value for c in r] for r in wb["통계 요약"].iter_rows()], ensure_ascii=False, default=str)) if "통계 요약" in wb.sheetnames else None
            else:
                item["run_body"] = str(run)[:300]
        item["runs_delta"] = db("select (select count(*) from runs)")[0][0] - before
        rec("bundle", item)
    # 10) 화면: 그룹 결과 표 (실제 브라우저)
    rid = a1.get("id")
    if rid:
        with sync_playwright() as p:
            b = p.chromium.launch(channel="msedge", headless=True, proxy=browser_proxy(URL))
            cx = b.new_context(viewport={"width": 1280, "height": 900}, locale="ko-KR")
            cx.route("**/*", route_guard(URL, BLOCKED))
            pg = cx.new_page()
            pg.goto(URL + "/")
            pg.get_by_label("아이디").fill(USERS["user2"])
            pg.get_by_label("비밀번호", exact=True).fill(creds()["user2"])
            pg.locator(".login-card button[type=submit]").click()
            pg.wait_for_selector(".shell", timeout=60000)
            asset = pg.evaluate(ASSET_JS)
            pg.goto(f"{URL}/analysis?id={rid}")
            pg.locator("table.an-result tbody tr").first.wait_for(timeout=60000)
            pg.wait_for_timeout(800)
            rows = pg.locator("table.an-result tbody tr").evaluate_all("rs => rs.map(r => [...r.cells].map(c => c.innerText.trim()))")
            body_text = pg.locator("main").inner_text()
            rec("browser", {"what": "분석 결과 화면 그룹 표", "id": rid, "rows": rows, "raw_in_page": raw_hits(body_text)})
            pg.screenshot(path=str(OUT / f"{LABEL}-analysis-{rid}.png"), full_page=True)
            REC["meta"] = [run_meta("f1_repro.py", LABEL, asset)]
            b.close()
    REC["blocked_out_of_scope"] = BLOCKED
    (OUT / f"f1-repro-{LABEL}.json").write_text(json.dumps(REC, ensure_ascii=False, indent=1, default=str), encoding="utf-8")


if __name__ == "__main__":
    main()
