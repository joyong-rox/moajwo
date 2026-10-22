"""U-03 확인: 통계 요약·데이터 통계 분석 결과에 표준 항목('건수')과 함께 계산한 원본 열('발신Mail건수')이 화면·Excel·실행 기록에
나오는지 (격리 서버, 새 계정, 측정만 저장 — 판정하지 않음).

시연 5개 파일을 시연과 같은 매핑(날짜·기준일자·통계일자 → 기준일자, 기관코드, 조직명·조직 명·부서명 → 담당부서, 발신Mail건수 → 건수,
나머지 → 해당 없음)으로 확정한 뒤, 묶기 '통계 요약'(건수, 합계·평균·최솟값·최댓값·개수)과 데이터 통계 분석을 화면으로 실행한다.
같은 대상이 반복되는 시트는 시연처럼 '서로 다른 기록 — 반복 행 모두 더함'을 고른다. 독립 기대값과 비교할 수 있게 값도 적는다.
사용법: python summary_source_check.py <출력 폴더> <라벨> <새 계정 이름> <입력 파일 폴더>
"""

import io
import json
import secrets
import sys
from pathlib import Path

import openpyxl
from playwright.sync_api import sync_playwright

from diag_common import ASSET_JS, URL, XLSX, Api, run_meta
from diag_scope import browser_proxy, route_guard

OUT, LABEL, NAME, SRC = Path(sys.argv[1]), sys.argv[2], sys.argv[3], Path(sys.argv[4])
FILES = [f"{m} 사용량.xlsx" for m in ("행안부", "기획예산처", "재정경제부", "외교부", "통일부")]
DEMO = {"날짜": "base_date", "기준일자": "base_date", "통계일자": "base_date", "기관코드": "org_code",
        "조직명": "department", "조직 명": "department", "부서명": "department", "발신Mail건수": "count"}  # fmt: skip
OUT.mkdir(parents=True, exist_ok=True)
REC: dict[str, list] = {}
BLOCKED: list[str] = []


def rec(group: str, item: dict) -> None:
    REC.setdefault(group, []).append(item)
    print(group, json.dumps(item, ensure_ascii=False)[:400], flush=True)


def text(loc) -> str:  # noqa: ANN001 — Playwright Locator
    return " ".join(loc.inner_text().split())


def main() -> None:
    owner = Api().login("owner")
    password = "Sc-" + secrets.token_urlsafe(14) + "7z"
    s, b, _ = owner.call("POST", "/users", {"username": NAME, "display_name": f"요약 확인 {LABEL}", "status": "active", "password": password})
    if s != 201:
        raise RuntimeError(f"계정 생성 실패: {s} {b}")
    (OUT / f"creds-{LABEL}.json").write_text(json.dumps({NAME: password}), encoding="utf-8")
    u = Api()
    if u.call("POST", "/auth/login", {"username": NAME, "password": password})[0] != 200:
        raise RuntimeError("새 계정 로그인 실패")
    fids = []
    for name in FILES:
        s, body, _ = u.upload(name, (SRC / name).read_bytes(), XLSX)
        if s != 200:
            raise RuntimeError(f"업로드 실패 {name}: {s}")
        fid = body[0]["id"]
        cols = u.call("GET", f"/files/{fid}/preview?sheet_id=s0")[1]["columns"]
        decisions = [{"sheet_id": "s0", "column_id": c["column_id"], "field_key": DEMO.get(c["name"])} for c in cols]
        s, body, _ = u.call("PUT", f"/files/{fid}/mapping/decisions", {"decisions": decisions})
        if s != 200:
            raise RuntimeError(f"확정 실패 {name}: {s} {body}")
        fids.append(fid)
        rec("setup", {"file": fid, "name": name, "confirmed": len(decisions), "count_column": [c["name"] for c in cols if DEMO.get(c["name"]) == "count"]})

    with sync_playwright() as p:
        br = p.chromium.launch(channel="msedge", headless=True, proxy=browser_proxy(URL))
        cx = br.new_context(viewport={"width": 1280, "height": 900}, locale="ko-KR")
        cx.route("**/*", route_guard(URL, BLOCKED))
        pg = cx.new_page()
        pg.goto(URL + "/")
        pg.get_by_label("아이디").fill(NAME)
        pg.get_by_label("비밀번호", exact=True).fill(password)
        pg.locator(".login-card button[type=submit]").click()
        pg.wait_for_selector(".shell", timeout=60000)
        asset = pg.evaluate(ASSET_JS)

        # 1) 파일별 시트로 묶기 + 통계 요약 (화면)
        pg.goto(f"{URL}/bundle?files={','.join(map(str, fids))}")
        pg.get_by_text("통계 요약 시트 추가").first.click()
        pg.locator(".bd-add select").first.wait_for(timeout=60000)
        pg.locator(".bd-add select").first.select_option("count")
        pg.get_by_role("button", name="항목 추가").click()
        head = pg.locator(".bd-item-head").first
        head.wait_for(timeout=60000)
        pg.wait_for_function("() => [...document.querySelectorAll('.bd-item-head')].some(h => h.innerText.includes('원본 열') || h.innerText.includes('포함'))", timeout=60000)
        pg.wait_for_timeout(800)
        rec("bundle-plan", {"item_head": text(head), "warns": [text(x) for x in pg.locator(".bd-item .mapping-warn").all()]})
        for box in pg.get_by_label("반복 행 모두 더함").all():  # 시연처럼 같은 대상이 반복되는 시트만 확인 (화면에 보이는 것만)
            if not box.is_checked():
                box.check()
        pg.wait_for_timeout(1500)
        make = pg.get_by_role("button", name="엑셀 만들기")
        pg.wait_for_function("() => [...document.querySelectorAll('button')].some(b => b.innerText.startsWith('엑셀 만들기') && !b.disabled)", timeout=60000)
        compose = pg.locator(".bd-step li.is-summary")
        rec("bundle-compose", {"summary_line": text(compose) if compose.count() else None})
        pg.screenshot(path=str(OUT / f"{LABEL}-1-bundle-plan.png"), full_page=True)
        make.click()
        pg.locator(".bd-made").wait_for(timeout=120000)
        made = text(pg.locator(".bd-made"))
        run_id = int(made.split("#", 1)[1].split()[0])
        rec("bundle-made", {"text": made, "run": run_id})
        s, content, _ = u.call("GET", f"/runs/{run_id}/download", raw=True)
        wb = openpyxl.load_workbook(io.BytesIO(content))
        rows = [[c for c in r] for r in wb["통계 요약"].iter_rows(values_only=True)]
        at = next(i for i, r in enumerate(rows) if isinstance(r[0], str) and r[0].startswith("항목: "))
        table_at = next(i for i in range(at, len(rows)) if rows[i][0] in ("범위", "기간"))
        rec("bundle-excel", {"sheets": wb.sheetnames, "item": rows[at][0], "source_row": rows[at + 1][:2],
                             "table": [r[:6] for r in rows[table_at : table_at + 3]]})  # fmt: skip

        # 2) 실행 기록 화면
        pg.goto(URL + "/runs")
        line = pg.locator("table.runs-table tbody tr", has_text=f"#{run_id}").first
        line.wait_for(timeout=60000)
        rec("runs", {"row": text(line)})
        pg.screenshot(path=str(OUT / f"{LABEL}-2-runs.png"), full_page=True)

        # 3) 데이터 통계 분석 (화면): 항목 → 파일 → 대상 확인 → (반복 확인) → 실행 → 결과
        pg.goto(URL + "/analysis")
        pg.get_by_label("분석할 항목").select_option("count")
        for fid in fids:
            pg.locator("label.an-file", has_text=f"#{fid} ").locator("input").check()
        pg.get_by_role("button", name="대상 확인 (계산 안 함)").click()
        pg.locator("fieldset.an-step legend", has_text="2. 대상").wait_for(timeout=60000)
        summary = pg.locator("fieldset.an-step", has_text="2. 대상").locator("p.mw-muted").first
        rec("analysis-preview", {"summary": text(summary)})
        for box in pg.get_by_label("반복 행은 서로 다른 기록").all():  # 같은 대상이 반복되는 시트에만 보인다
            if not box.is_checked():
                box.check()
        if pg.get_by_role("button", name="다시 확인 (계산 안 함)").is_enabled():
            pg.get_by_role("button", name="다시 확인 (계산 안 함)").click()
            pg.wait_for_timeout(1500)
        run_btn = pg.get_by_role("button", name="통계 분석 실행")
        rec("analysis-run-button", {"enabled": run_btn.is_enabled()})
        pg.screenshot(path=str(OUT / f"{LABEL}-3-analysis-preview.png"), full_page=True)
        if run_btn.is_enabled():
            run_btn.click()
            pg.locator("section.an[aria-label='통계 분석 결과'] h2").wait_for(timeout=120000)
            pg.locator("table.an-result tbody tr").first.wait_for(timeout=60000)
            dl = pg.locator("section.an dl.an-conditions")
            terms = [text(x) for x in dl.locator("dt").all()]
            defs = [text(x) for x in dl.locator("dd").all()]
            cells = pg.locator("table.an-result tbody tr").first.evaluate("r => [...r.cells].map(c => c.innerText.trim())")
            rec("analysis-result", {"title": pg.title(), "h2": text(pg.locator("section.an h2").first), "conditions": dict(zip(terms, defs, strict=False)),
                                    "first_row": cells, "url": pg.evaluate("location.pathname + location.search")})  # fmt: skip
            pg.screenshot(path=str(OUT / f"{LABEL}-4-analysis-result.png"), full_page=True)
        REC["meta"] = [run_meta("summary_source_check.py", LABEL, asset)]
        br.close()
    REC["blocked_out_of_scope"] = BLOCKED
    (OUT / f"summary-source-{LABEL}.json").write_text(json.dumps(REC, ensure_ascii=False, indent=1, default=str), encoding="utf-8")


if __name__ == "__main__":
    main()
