"""C-02 확인: 데이터 통계 분석의 합계·평균·최솟값·최댓값이 API·화면·Excel에서 어떤 수로 나오는지
(격리 서버, 새 계정, 합성 자료, 측정만 저장 — 판정하지 않음).

사례 (시도·인구 표, 인구를 합계·평균·최솟값·최댓값·개수로):
- decimal: 0.1, 0.2 / decimal2: 10.4, 20.4 — 소수 합
- bins: 0.1, 0.2를 '0.3 이하(포함)'·'초과' 구간으로
- big: 2^53 + 1, 1 (CSV — xlsx 숫자 칸은 2^53을 넘는 정수를 담지 못한다)
- sixteen: 1,234,567,890,123,456, 4 — 16자리 (Excel 숫자 칸은 15자리까지만 정확)
- won: 1,000,000,000,001, 999,999,999,999, 500 — 1원·500원 차이
- mixed: -5, 0, '-', 'x', 10 — 음수·0·결측·비공개
API는 응답 글자 그대로(JSON 글자)도 적는다 — 브라우저가 읽기 전 값이다. 화면은 결과 표 첫 행과 결과 조건의 주의 문구,
Excel은 '결과' 첫 행과 '근거 행' 값 칸의 값·칸 종류(n 숫자, s 글자)를 적는다.
사용법: python precision_check.py <출력 폴더> <라벨> <새 계정 이름>
"""

import io
import json
import re
import secrets
import sys
from pathlib import Path

import openpyxl
from playwright.sync_api import sync_playwright

from diag_common import ASSET_JS, URL, XLSX, Api, run_meta, xlsx
from diag_scope import browser_proxy, route_guard

OUT, LABEL, NAME = Path(sys.argv[1]), sys.argv[2], sys.argv[3]
OUT.mkdir(parents=True, exist_ok=True)
TITLE = "주민등록 인구 현황 2026년 9월 (단위: 명)"
HEAD = ["시도", "인구"]
REGIONS = ["서울특별시", "부산광역시", "대구광역시", "인천광역시", "광주광역시", "대전광역시"]
MAP = {"시도": "sido", "인구": "population"}
AGGS = ["sum", "mean", "min", "max", "count"]
BINS = {"on": "sum", "ranges": [{"label": "0.3 이하", "max": 0.3, "max_inclusive": True},
                                {"label": "초과", "min": 0.3, "min_inclusive": False}]}  # fmt: skip
CASES = {
    "decimal": ([0.1, 0.2], {}),
    "decimal2": ([10.4, 20.4], {}),
    "bins": ([0.1, 0.2], {"bins": BINS}),
    "big": ([2**53 + 1, 1], {}),
    "sixteen": ([1234567890123456, 4], {}),
    "won": ([1_000_000_000_001, 999_999_999_999, 500], {}),
    "mixed": ([-5, 0, "-", "x", 10], {}),
}
CSV = {"big"}
BLOCKED: list[str] = []


def table(values: list) -> list[list]:
    return [[TITLE], HEAD, *[[REGIONS[i], v] for i, v in enumerate(values)]]


def content(case: str, values: list) -> tuple[str, bytes, str]:
    if case in CSV:
        text = "\r\n".join(",".join(str(x) for x in row) for row in table(values))
        return f"정밀도_{case}.csv", text.encode("utf-8"), "text/csv"
    return f"정밀도_{case}.xlsx", xlsx({"현황": table(values)}), XLSX


def raw_numbers(raw: bytes) -> dict[str, list[str]]:
    """응답 JSON 글자에 적힌 수 그대로 (파싱 전)."""
    text = raw.decode("utf-8")
    return {k: re.findall(rf'"{k}":\s*(-?[0-9][0-9.eE+-]*|null)', text) for k in ("sum", "mean", "min", "max")}


def cells(ws, row: int) -> list[list]:  # noqa: ANN001 — openpyxl Worksheet
    return [[None if c.value is None else str(c.value), c.data_type] for c in ws[row]]


def main() -> None:
    owner = Api().login("owner")
    password = "Pc-" + secrets.token_urlsafe(14) + "7z"
    s, b, _ = owner.call("POST", "/users", {"username": NAME, "display_name": f"정밀도 확인 {LABEL}", "status": "active", "password": password})
    if s != 201:
        raise RuntimeError(f"계정 생성 실패: {s} {b}")
    u = Api()
    if u.call("POST", "/auth/login", {"username": NAME, "password": password})[0] != 200:
        raise RuntimeError("새 계정 로그인 실패")
    rec: dict[str, dict] = {}
    for case, (values, extra) in CASES.items():
        name, data, ctype = content(case, values)
        s, body, _ = u.upload(name, data, ctype)
        if s != 200:
            raise RuntimeError(f"업로드 실패 {name}: {s} {body}")
        fid = body[0]["id"]
        cols = u.call("GET", f"/files/{fid}/preview?sheet_id=s0")[1]["columns"]
        decisions = [{"sheet_id": "s0", "column_id": c["column_id"], "field_key": MAP.get(c["name"])} for c in cols]
        s, body, _ = u.call("PUT", f"/files/{fid}/mapping/decisions", {"decisions": decisions})
        if s != 200:
            raise RuntimeError(f"확정 실패 {name}: {s} {body}")
        req = {"field_key": "population", "targets": [{"file_id": fid}], "aggregates": AGGS, **extra}
        s, pv, _ = u.call("POST", "/analysis/preview", req)
        if s != 200:
            raise RuntimeError(f"대상 확인 실패 {case}: {s} {pv}")
        s, run, _ = u.call("POST", "/analysis/run", {**req, "expected_fingerprint": pv["fingerprint"]})
        if s != 201:
            raise RuntimeError(f"실행 실패 {case}: {s} {run}")
        rid = run["id"]
        detail = u.call("GET", f"/analysis/runs/{rid}")[1]
        s, raw, _ = u.call("GET", f"/analysis/runs/{rid}/groups?limit=200", raw=True)
        [g] = json.loads(raw)["items"]
        s, book, _ = u.call("GET", f"/analysis/runs/{rid}/download", raw=True)
        wb = openpyxl.load_workbook(io.BytesIO(book))
        result, rows = wb["결과"], wb["근거 행"]
        rec[case] = {
            "file": fid, "run": rid, "values": [str(v) for v in values], "upload": ctype,
            "api": {k: g.get(k) for k in (*AGGS, "partial", "problem_rows", "bin")},
            "api_text": raw_numbers(raw),
            "warnings": detail["conditions"].get("warnings", []),
            "policy": detail["conditions"].get("policy"),
            "excel": {"결과_머리": cells(result, 1), "결과_1행": cells(result, 2),
                      "근거_값": [cells(rows, i)[-1] for i in range(2, rows.max_row + 1)]},
        }  # fmt: skip
        print(case, json.dumps(rec[case], ensure_ascii=False, default=str)[:700], flush=True)

    with sync_playwright() as p:
        br = p.chromium.launch(channel="msedge", headless=True, proxy=browser_proxy(URL))
        cx = br.new_context(viewport={"width": 1280, "height": 900}, locale="ko-KR")
        cx.route("**/*", route_guard(URL, BLOCKED))
        pg = cx.new_page()
        pg.goto(URL + "/")
        pg.get_by_label("아이디").fill(NAME)
        pg.get_by_label("비밀번호", exact=True).fill(password)
        pg.locator(".login-card button[type=submit]").click()
        pg.wait_for_selector(".shell", timeout=60000)
        asset = pg.evaluate(ASSET_JS)
        for case, item in rec.items():
            pg.goto(f"{URL}/analysis?id={item['run']}")
            row = pg.locator(".an-result tbody tr").first
            row.wait_for(timeout=60000)
            heads = [" ".join(t.split()) for t in pg.locator(".an-result thead th").all_inner_texts()]
            shown = [" ".join(t.split()) for t in row.locator("td").all_inner_texts()]
            item["screen"] = {"columns": heads, "first_row": shown,
                              "warnings": [" ".join(t.split()) for t in pg.locator(".mapping-warn").all_inner_texts()]}  # fmt: skip
            # 브라우저가 API 응답을 읽은 수 (JavaScript Number) — 화면 표시 전 값
            item["screen"]["js_numbers"] = pg.evaluate(
                """async (rid) => { const r = await fetch(`/api/analysis/runs/${rid}/groups?limit=200`, {credentials: 'same-origin'});
                   const g = (await r.json()).items[0]; return Object.fromEntries(['sum','mean','min','max'].map(k => [k, String(g[k])])) }""",
                item["run"],
            )
            if case in ("big", "decimal", "sixteen"):
                pg.screenshot(path=str(OUT / f"{LABEL}-{case}.png"), full_page=True)
            print("screen", case, json.dumps(item["screen"], ensure_ascii=False)[:500], flush=True)
        br.close()
    meta = run_meta("precision_check.py", LABEL, asset)
    (OUT / f"precision-{LABEL}.json").write_text(json.dumps({"meta": meta, "blocked": BLOCKED, "cases": rec}, ensure_ascii=False, indent=1, default=str), encoding="utf-8")


if __name__ == "__main__":
    main()
