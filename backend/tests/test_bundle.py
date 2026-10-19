"""파일별 시트로 묶기·통계 요약: 실제 업로드 → (확정) → 결과 구성 확인 → 만들기 → 결과 엑셀 내용으로 확인한다.

합성 자료만 쓴다."""

import hashlib
import io
import re
import sqlite3
import zipfile
from typing import Any

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app.config import DATA_DIR
from tests.test_accounts import owner  # noqa: F401
from tests.test_auth import _client, app_real_login  # noqa: F401

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
POP = "주민등록 인구 현황 2026년 9월 (단위: 명)"
HEAD = ["시도", "시군구", "인구수"]
MAP = {"시도": "sido", "시군구": "sigungu", "인구수": "population"}
# 같은 항목·단위·기준 기간, 지역은 모두 다름: a 10·20 / b 30 / c 40·50·60 / d 70·80·90·100
DATASET = {
    "a": [["서울특별시", "종로구", 10], ["서울특별시", "용산구", 20]],
    "b": [["서울특별시", "성동구", 30]],
    "c": [["서울특별시", "광진구", 40], ["서울특별시", "동대문구", 50], ["서울특별시", "중랑구", 60]],
    "d": [["서울특별시", "성북구", 70], ["서울특별시", "강북구", 80], ["서울특별시", "도봉구", 90],
          ["서울특별시", "노원구", 100]],
}  # fmt: skip


def _xlsx(sheets: dict[str, list[list]], title: str | None = POP) -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        if title:
            ws.append([title])
        for r in rows:
            ws.append(r)
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("=HYPERLINK"):
                    cell.data_type = "s"  # 시험 자료: '=HYPERLINK…'는 글자 칸 (그 밖의 '=…'는 진짜 수식)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _upload(c: TestClient, name: str, data: bytes) -> int:
    res = c.post("/files", files=[("files", (name, data, XLSX))])
    assert res.status_code == 200 and res.json()[0]["id"], res.text
    return res.json()[0]["id"]


def _confirm(c: TestClient, fid: int, mapping: dict[str, str | None], sheet: str = "s0") -> None:
    cols = {col["name"]: col["column_id"] for col in c.get(f"/files/{fid}/preview?sheet={sheet}").json()["columns"]}
    items = [{"sheet_id": sheet, "column_id": cols[n], "field_key": k} for n, k in mapping.items()]
    res = c.put(f"/files/{fid}/mapping/decisions", json={"decisions": items})
    assert res.status_code == 200, res.text


def _four(c: TestClient, tag: str, head: list[str] = HEAD, confirm: bool = True) -> list[int]:
    fids = []
    for name, rows in DATASET.items():
        fid = _upload(c, f"{name}.xlsx", _xlsx({f"{tag}{name}": [head, *rows]}))
        if confirm:
            _confirm(c, fid, MAP)
        fids.append(fid)
    return fids


def _plan(c: TestClient, body: dict[str, Any], code: int = 200) -> dict[str, Any]:
    res = c.post("/bundle/plan", json=body)
    assert res.status_code == code, res.text
    return res.json()


def _make(c: TestClient, body: dict[str, Any]) -> tuple[dict[str, Any], openpyxl.Workbook, dict[str, Any]]:
    plan = _plan(c, body)
    assert plan["ready"], plan["problems"]
    res = c.post("/bundle/run", json={**body, "expected_basis": plan["basis"]})
    assert res.status_code == 201, res.text
    run = res.json()
    dl = c.get(f"/runs/{run['id']}/download")
    assert dl.status_code == 200
    return run, openpyxl.load_workbook(io.BytesIO(dl.content)), plan


def _db(sql: str, *args: Any) -> list[tuple]:
    with sqlite3.connect(DATA_DIR / "moajwo.db") as con:
        return con.execute(sql, args).fetchall()


def _summary_rows(wb: openpyxl.Workbook) -> list[list[Any]]:
    return [list(r) for r in wb["통계 요약"].iter_rows(values_only=True)]


def _table(rows: list[list[Any]], item: str) -> tuple[list[Any], list[list[Any]]]:
    """'항목: …' 블록의 결과 표 (머리글, 값 행들)."""
    start = next(i for i, r in enumerate(rows) if r[0] == f"항목: {item}")
    aggs = ("합계", "평균", "최솟값", "최댓값", "유효 숫자 개수")
    head = next(i for i in range(start, len(rows)) if rows[i][0] in ("범위", "기간") and rows[i][1] in aggs)
    end = next(i for i in range(head, len(rows)) if rows[i][0] == "자료")
    trim = [[v for v in r if v is not None] for r in rows[head:end]]
    return trim[0], trim[1:]


@pytest.fixture(scope="module")
def bd(app_real_login, owner) -> dict[str, TestClient]:  # noqa: F811
    made = {"owner": owner}
    for name in ("bd-user1", "bd-user2"):
        res = owner.post("/users", json={"username": name, "display_name": name.upper(), "status": "active",
                                         "password": "pass-" + name})  # fmt: skip
        assert res.status_code == 201, res.text
        made[name] = _client(app_real_login, name, "pass-" + name)
    return made


# ---------------------------------------------------------------- 4개·5개 시트, 계산값


def test_four_files_without_summary_make_four_sheets_and_never_call_the_engine(bd, monkeypatch) -> None:
    from app import analysis, mapping

    def boom(*a: Any, **k: Any) -> Any:
        raise AssertionError("요약을 끄면 분석 엔진·모델을 부르지 않아야 함")

    u = bd["bd-user1"]
    fids = _four(u, "off", confirm=False)  # 매핑을 확정하지 않아도 묶을 수 있다
    monkeypatch.setattr(analysis, "evaluate", boom)
    monkeypatch.setattr(analysis, "compute", boom)
    monkeypatch.setattr(mapping, "get_engine", boom)
    run, wb, plan = _make(u, {"file_ids": fids})
    assert wb.sheetnames == ["a", "b", "c", "d"] and plan["sheet_count"] == 4 and plan["summary"] is None
    assert run["summary"]["mode"] == "sheets" and run["summary"]["stats_summary"] is False
    assert run["summary"]["verified"] is False and run["review_session_id"] is None
    ws = wb["c"]
    assert [c.value for c in ws[1]] == ["원본 행", "행 종류", *HEAD]
    assert [[c.value for c in r] for r in ws.iter_rows(min_row=2)] == [
        [3, "데이터", "서울특별시", "광진구", 40],
        [4, "데이터", "서울특별시", "동대문구", 50],
        [5, "데이터", "서울특별시", "중랑구", 60],
    ]
    assert run["row_count"] == 10 and run["file_name"].startswith("시트묶음_")


def test_four_files_with_summary_make_five_sheets_with_engine_values(bd) -> None:
    u = bd["bd-user1"]
    fids = _four(u, "on")
    item = {"field_key": "population", "aggregates": ["sum", "mean", "min", "max", "count"]}
    run, wb, plan = _make(u, {"file_ids": fids, "summary": {"items": [item]}})
    assert wb.sheetnames == ["a", "b", "c", "d", "통계 요약"] and plan["sheet_count"] == 5
    (it,) = plan["summary"]["items"]
    assert it["coverage"]["text"] == "4개 파일 중 4개 포함" and it["unit"] == "명" and it["periods"] == ["2026-09"]
    assert {lk["column"] for lk in it["links"]} == {"인구수"} and all(lk["status"] == "ok" for lk in it["links"])
    rows = _summary_rows(wb)
    head, values = _table(rows, "인구수")
    assert head == ["범위", "합계", "평균", "최솟값", "최댓값", "유효 숫자 개수"]
    assert values == [["전체", 550, 55, 10, 100, 10]]
    text = " ".join(str(v) for r in rows for v in r if v is not None)
    assert "자동으로 다시 계산되지 않습니다" in text and "파일별 평균을 다시 평균하지 않음" in text
    assert "4개 파일 중 4개 포함" in text and "전체 집계" in text
    saved = run["summary"]["summary_items"][0]
    assert saved["groups"][0]["sum"] == 550 and saved["groups"][0]["mean"] == 55 and run["summary"]["stats_summary"]
    # 같은 조건의 데이터 분석과 같은 값 (같은 엔진)
    body = {"field_key": "population", "targets": [{"file_id": f, "sheet_id": "s0"} for f in fids]}
    pv = u.post("/analysis/preview", json=body).json()
    res = u.post("/analysis/run", json={**body, "expected_fingerprint": pv["fingerprint"]})
    (g,) = u.get(f"/analysis/runs/{res.json()['id']}/groups").json()["items"]
    assert (g["sum"], g["mean"], g["min"], g["max"], g["count"]) == (550, 55, 10, 100, 10)


def test_only_selected_items_and_methods_appear(bd) -> None:
    u = bd["bd-user1"]
    head = ["시도", "시군구", "인구수(명)", "집행액(원)"]  # 단위는 열 이름에 (제목에 단위가 둘이면 판단하지 않음)
    fids = []
    for i, (name, rows) in enumerate(DATASET.items()):
        data = [[*r, (i + 1) * 1000] for r in rows]
        fid = _upload(u, f"{name}.xlsx", _xlsx({"현황": [head, *data]}, title="2026년 9월 현황"))
        _confirm(u, fid, {"시도": "sido", "시군구": "sigungu", "인구수(명)": "population", "집행액(원)": "expenditure"})
        fids.append(fid)
    items = [
        {"field_key": "population", "aggregates": ["sum", "mean"]},
        {"field_key": "expenditure", "aggregates": ["sum"]},
    ]
    _, wb, plan = _make(u, {"file_ids": fids, "summary": {"items": items}})
    rows = _summary_rows(wb)
    assert _table(rows, "인구수") == (["범위", "합계", "평균"], [["전체", 550, 55]])
    # 집행액: 행마다 같은 값이 붙어 a 2행×1000, b 1×2000, c 3×3000, d 4×4000
    assert _table(rows, "집행액") == (["범위", "합계"], [["전체", 2000 + 2000 + 9000 + 16000]])
    labels = [r[0] for r in rows if isinstance(r[0], str) and r[0].startswith("항목: ")]
    assert labels == ["항목: 인구수", "항목: 집행액"]
    assert not any(v in ("최솟값", "최댓값", "유효 숫자 개수") for r in rows for v in r)


# ---------------------------------------------------------------- 연결: 확정한 뜻으로, 빠진 연결·중복·단위·기간


def test_links_follow_confirmed_meaning_not_column_order_and_missing_links_must_be_excluded(bd) -> None:
    u = bd["bd-user1"]
    fids = []
    for i, (name, rows) in enumerate(DATASET.items()):
        if i == 1:  # b는 열 순서가 다르다
            fid = _upload(u, f"{name}.xlsx", _xlsx({"현황": [["인구수", "시군구", "시도"], *[r[::-1] for r in rows]]}))
            _confirm(u, fid, MAP)
        elif i == 3:  # d는 인구수 열을 확정하지 않았다
            fid = _upload(u, f"{name}.xlsx", _xlsx({"현황": [HEAD, *rows]}))
            _confirm(u, fid, {"시도": "sido", "시군구": "sigungu"})
        else:
            fid = _upload(u, f"{name}.xlsx", _xlsx({"현황": [HEAD, *rows]}))
            _confirm(u, fid, MAP)
        fids.append(fid)
    item = {"field_key": "population", "aggregates": ["sum", "count"]}
    plan = _plan(u, {"file_ids": fids, "summary": {"items": [item]}})
    assert not plan["ready"] and any("확정된 열이 없음" in p and "빼기" in p for p in plan["problems"])
    res = u.post("/bundle/run", json={"file_ids": fids, "summary": {"items": [item]}, "expected_basis": plan["basis"]})
    assert res.status_code == 409 and _db("SELECT count(*) FROM runs WHERE output_path LIKE '%시트묶음%' AND "
                                          "summary LIKE '%' || ? || '%'", plan["basis"])[0][0] == 0  # fmt: skip
    item["links"] = [{"file_id": fids[3], "sheet_id": "s0", "exclude": True, "exclude_reason": "인구수 없음"}]
    _, wb, plan = _make(u, {"file_ids": fids, "summary": {"items": [item]}})
    (it,) = plan["summary"]["items"]
    assert it["coverage"]["text"] == "4개 파일 중 3개 포함"
    assert _table(_summary_rows(wb), "인구수") == (
        ["범위", "합계", "유효 숫자 개수"],
        [["전체", 10 + 20 + 30 + 40 + 50 + 60, 6]],
    )
    text = " ".join(str(v) for r in _summary_rows(wb) for v in r if v is not None)
    assert "뺌 (사용자)" in text and "인구수 없음" in text
    assert [c.value for c in wb["b"][1]] == [
        "원본 행",
        "행 종류",
        "인구수",
        "시군구",
        "시도",
    ]  # 데이터 시트는 원래 열 순서


def test_unit_and_period_conflicts_follow_analysis_policy(bd) -> None:
    u = bd["bd-user1"]
    a = _upload(u, "단위a.xlsx", _xlsx({"현황": [HEAD, *DATASET["a"]]}))
    b = _upload(
        u, "단위b.xlsx", _xlsx({"현황": [HEAD, *DATASET["b"]]}, title="주민등록 인구 현황 2026년 9월 (단위: 천명)")
    )
    for f in (a, b):
        _confirm(u, f, MAP)
    item = {"field_key": "population", "aggregates": ["sum"]}
    plan = _plan(u, {"file_ids": [a, b], "summary": {"items": [item]}})
    assert not plan["ready"] and any("단위가 다른 자료" in p and "자동 환산 없음" in p for p in plan["problems"])
    # 기간이 다르면 나눠 계산한다 (시점 값을 더해 하나로 만들지 않음)
    c = _upload(
        u, "기간c.xlsx", _xlsx({"현황": [HEAD, *DATASET["b"]]}, title="주민등록 인구 현황 2026년 8월 (단위: 명)")
    )
    _confirm(u, c, MAP)
    _, wb, _ = _make(u, {"file_ids": [a, c], "summary": {"items": [item]}})
    assert _table(_summary_rows(wb), "인구수") == (["기간", "합계"], [["2026-08", 30], ["2026-09", 30]])


def test_duplicate_file_content_and_item_twice_are_blocked_for_summary(bd) -> None:
    u = bd["bd-user1"]
    data = _xlsx({"현황": [HEAD, *DATASET["a"]]})
    a1, a2 = _upload(u, "같은a.xlsx", data), _upload(u, "같은a.xlsx", data)
    for f in (a1, a2):
        _confirm(u, f, MAP)
    plan = _plan(u, {"file_ids": [a1, a2]})
    assert plan["ready"] and [s["name"] for s in plan["sheets"]] == ["같은a", "같은a (2)"]
    assert any("같은 내용의 파일" in w for w in plan["warnings"])
    assert plan["sheets"][1]["renamed"] and "번호를 붙임" in plan["sheets"][1]["renamed"][0]
    item = {"field_key": "population", "aggregates": ["sum"]}
    plan = _plan(u, {"file_ids": [a1, a2], "summary": {"items": [item]}})
    assert not plan["ready"] and any("같은 내용의 파일" in p for p in plan["problems"])
    plan = _plan(u, {"file_ids": [a1], "summary": {"items": [item, {**item, "aggregates": ["mean"]}]}})
    assert not plan["ready"] and any("두 번 골랐습니다" in p for p in plan["problems"])


# ---------------------------------------------------------------- 여러 시트·시트 이름


def test_multi_sheet_file_needs_explicit_sheets_and_names_follow_excel_rules(bd) -> None:
    u = bd["bd-user1"]
    multi = _upload(u, "여러시트.xlsx", _xlsx({"1월": [HEAD, *DATASET["a"]], "2월": [HEAD, *DATASET["b"]],
                                               "메타정보": [["설명"], ["통계표 설명"]]}))  # fmt: skip
    long = _upload(
        u,
        "아주긴파일이름이라서엑셀시트이름서른한자를넘어가는자료_2026년9월.xlsx",
        _xlsx({"현황": [HEAD, *DATASET["c"]]}),
    )
    odd = _upload(u, "괄호[1]:별*표?.xlsx", _xlsx({"현황": [HEAD, *DATASET["d"]]}))
    summary = _upload(u, "통계 요약.xlsx", _xlsx({"현황": [HEAD, *DATASET["b"]]}))
    plan = _plan(u, {"file_ids": [multi, long]})
    assert not plan["ready"] and any("시트가 2개" in p and "직접 고르세요" in p for p in plan["problems"])
    assert plan["files"][0]["skipped_sheets"] == ["메타정보"]
    body = {"file_ids": [multi, long, odd, summary],
            "sheets": [{"file_id": odd, "sheet_id": "s0"}, {"file_id": multi, "sheet_id": "s1"},
                       {"file_id": multi, "sheet_id": "s0"}]}  # fmt: skip
    run, wb, plan = _make(u, body)
    names = wb.sheetnames
    assert names[:3] == ["괄호_1__별_표_", "여러시트-2월", "여러시트-1월"] and len(names) == 5
    assert names[3] == "아주긴파일이름이라서엑셀시트이름서른한자를넘어가는자료_2026년9월"[:31] and len(names[3]) == 31
    assert names[4] == "통계 요약 (2)"  # 요약을 꺼도 예약 이름은 피한다 (켜고 끌 때 이름이 바뀌지 않게)
    shown = {s["name"]: s for s in plan["sheets"]}
    assert any("쓸 수 없는 글자" in r for r in shown["괄호_1__별_표_"]["renamed"])
    assert any("31자" in r for r in shown[names[3]]["renamed"])
    assert [c.value for c in wb["여러시트-2월"]["E"][1:]] == [30]
    assert run["summary"]["sheet_order"] == names


# ---------------------------------------------------------------- 값 보존
# 앞자리 0·긴 식별자·소수·음수·0·결측·수식·글자 수식


def _with_cached(data: bytes, cell: str, value: str) -> bytes:
    """openpyxl은 수식 결과값을 저장하지 않으므로, 결과값이 저장된 수식 칸을 시험용으로 만든다."""
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for info in src.infolist():
            body = src.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                body = re.sub(rf'(<c r="{cell}"[^>]*>)<f>([^<]*)</f><v\s*/?>(</v>)?'.encode(),
                              rf"\1<f>\2</f><v>{value}</v>".encode(), body)  # fmt: skip
                assert f"<v>{value}</v>".encode() in body
            z.writestr(info, body)
    return out.getvalue()


def test_values_are_preserved_and_formula_cells_without_cached_value_stay_blank(bd) -> None:
    u = bd["bd-user1"]
    head = ["시도", "코드", "식별부호", "인구수", "비율"]
    rows = [
        ["서울특별시", "0011", "12345678901234567890", 10, 0.125],
        ["부산광역시", "0022", "98765432109876543210", -5, 0],
        ["대구광역시", "0033", "11112222333344445555", "=D2*2", "-"],
        ["인천광역시", "0044", "55554444333322221111", "=D3+1", None],
        ["합계", None, None, "=SUM(D2:D5)", None],
        ["※ 자료: 시험", None, None, None, None],
        ["광주광역시", "0055", '=HYPERLINK("http://x")', 7, 1.5],
    ]
    # 제목 1행·머리글 2행 → 자료는 3행부터. D5(=D2*2)만 결과값 4가 저장되고 D6(=D3+1)·D7(=SUM)은 결과값 없음
    data = _with_cached(_xlsx({"현황": [head, *rows]}), "D5", "4")
    fid = _upload(u, "값보존.xlsx", data)
    stored = _db("SELECT stored_path, sha256 FROM source_files WHERE id = ?", fid)[0]
    before = hashlib.sha256((DATA_DIR / stored[0]).read_bytes()).hexdigest()
    plan = _plan(u, {"file_ids": [fid]})
    (s,) = plan["sheets"]
    assert s["formula_missing"] == 2 and s["formula_cached"] == 1 and s["formula_checked"]
    assert any("0으로 채우지 않음" in w for w in plan["warnings"])
    run, wb, _ = _make(u, {"file_ids": [fid]})
    ws = wb["값보존"]
    got = {r[0].value: r for r in ws.iter_rows(min_row=2)}
    assert got[3][3].value == "0011" and got[3][3].data_type == "s"
    assert got[3][4].value == "12345678901234567890" and got[3][4].data_type == "s"
    assert got[3][5].value == 10 and got[3][6].value == 0.125
    assert got[4][5].value == -5 and got[4][6].value == 0
    assert got[5][5].value == 4 and got[5][5].comment is None  # 결과값이 저장된 수식 → 그 값
    assert got[5][6].value == "-"  # 결측 표기는 그대로 (0으로 바꾸지 않음)
    assert got[6][5].value is None and "0으로 채우지 않음" in got[6][5].comment.text  # 결과값 없는 수식
    assert got[7][1].value == "소계·합계" and got[7][5].value is None and got[7][5].comment is not None
    assert got[8][1].value == "주석" and got[8][2].value == "※ 자료: 시험"
    assert got[9][4].value == '=HYPERLINK("http://x")' and got[9][4].data_type == "s"  # 수식으로 저장하지 않음
    assert len(got) == 7  # 소계·주석 행도 지우지 않음
    sheet = run["summary"]["sheets"][0]
    assert sheet["formula_missing"] == 2 and sheet["formula_cached"] == 1 and sheet["written_rows"] == 7
    assert hashlib.sha256((DATA_DIR / stored[0]).read_bytes()).hexdigest() == before == stored[1]


# ---------------------------------------------------------------- 바뀐 뒤 거절·요약 실패·권한·개인정보·한도


def test_changes_after_plan_are_rejected_and_nothing_is_made(bd) -> None:
    u = bd["bd-user1"]
    fids = _four(u, "stale")
    item = {"field_key": "population", "aggregates": ["sum"]}
    body = {"file_ids": fids, "summary": {"items": [item]}}
    plan = _plan(u, body)
    runs_before = _db("SELECT count(*) FROM runs")[0][0]
    _confirm(u, fids[0], {"인구수": None})  # 확인 뒤 매핑이 바뀜
    res = u.post("/bundle/run", json={**body, "expected_basis": plan["basis"]})
    assert res.status_code == 409 and "다시 확인" in res.json()["detail"]["message"]
    _confirm(u, fids[0], MAP)
    plan = _plan(u, body)
    res = u.post("/bundle/run", json={**body, "summary": {"items": [{**item, "aggregates": ["mean"]}]},
                                      "expected_basis": plan["basis"]})  # fmt: skip
    assert res.status_code == 409  # 조건이 바뀐 요청
    res = u.post("/bundle/run", json={**body, "sheets": [{"file_id": f, "sheet_id": "s0"} for f in fids[::-1]],
                                      "expected_basis": plan["basis"]})  # fmt: skip
    assert res.status_code == 409  # 순서가 바뀐 요청
    assert _db("SELECT count(*) FROM runs")[0][0] == runs_before


def test_summary_failure_makes_no_file_and_asks_to_turn_it_off(bd, monkeypatch) -> None:
    from fastapi import HTTPException

    from app import analysis

    u = bd["bd-user1"]
    fids = _four(u, "fail")
    body = {"file_ids": fids, "summary": {"items": [{"field_key": "population", "aggregates": ["sum"]}]}}
    plan = _plan(u, body)

    def fail(ev: Any) -> Any:
        raise HTTPException(422, "그룹이 너무 많습니다")

    monkeypatch.setattr(analysis, "compute", fail)
    runs_before = _db("SELECT count(*) FROM runs")[0][0]
    outputs = set((DATA_DIR / "outputs").glob("*.xlsx"))
    res = u.post("/bundle/run", json={**body, "expected_basis": plan["basis"]})
    assert res.status_code == 409 and "통계 요약 시트 추가" in res.json()["detail"]["message"]
    assert (
        _db("SELECT count(*) FROM runs")[0][0] == runs_before and set((DATA_DIR / "outputs").glob("*.xlsx")) == outputs
    )


def test_other_users_files_and_pii_rules(bd) -> None:
    u1, u2, own = bd["bd-user1"], bd["bd-user2"], bd["owner"]
    mine = _upload(u1, "내자료.xlsx", _xlsx({"현황": [HEAD, *DATASET["a"]]}))
    assert u2.post("/bundle/plan", json={"file_ids": [mine]}).status_code == 404  # 남의 파일은 없는 것처럼
    plan = _plan(own, {"file_ids": [mine]})  # 최고 관리자는 볼 수 있지만 묶을 수는 없다
    assert not plan["ready"] and any("다른 사용자의 파일" in p for p in plan["problems"])
    assert own.post("/bundle/run", json={"file_ids": [mine], "expected_basis": plan["basis"]}).status_code == 409
    contact = [["기관", "담당", "연락처", "인원"], ["행정안전부", "홍길동", "010-1234-5678", 3]]
    pii = _upload(u1, "연락처.xlsx", _xlsx({"현황": contact}, title=None))
    plan = _plan(u1, {"file_ids": [pii]})
    assert (
        not plan["ready"] and plan["files"][0]["pii"] == "needs_ack" and any("포함 확인" in p for p in plan["problems"])
    )
    assert u1.post(f"/files/{pii}/pii/acknowledge").status_code == 200
    run, wb, plan = _make(u1, {"file_ids": [pii]})
    assert plan["files"][0]["pii"] == "acknowledged" and any("가리지 않고" in w for w in plan["warnings"])
    assert wb["연락처"]["E2"].value == "010-1234-5678"  # 확인한 파일은 원본 값 (통합표와 같은 정책)
    assert u2.get(f"/runs/{run['id']}/download").status_code == 404


def test_limits_warn_and_never_truncate(bd, monkeypatch) -> None:
    from app import bundle

    u = bd["bd-user1"]
    fids = _four(u, "limit", confirm=False)
    monkeypatch.setattr(bundle, "MAX_ROWS", 5)
    plan = _plan(u, {"file_ids": fids})
    assert not plan["ready"] and any("10행" in p and "나눠 만드세요" in p for p in plan["problems"])
    assert plan["total_rows"] == 10 and plan["limits"]["rows"] == 5
    res = u.post("/bundle/run", json={"file_ids": fids, "expected_basis": plan["basis"]})
    assert res.status_code == 409


def test_runs_list_shows_bundle_mode_and_existing_requests_keep_old_behaviour(bd) -> None:
    u = bd["bd-user1"]
    fids = _four(u, "list", confirm=False)
    run, _, _ = _make(u, {"file_ids": fids})
    listed = {r["id"]: r for r in u.get("/runs").json()}
    assert listed[run["id"]]["summary"]["mode"] == "sheets" and listed[run["id"]]["status"] == "done"
    assert u.get(f"/runs/{run['id']}").json()["summary"]["sheet_order"] == ["a", "b", "c", "d"]
    # 새 선택값을 빼면(요약 없음) 예전처럼 요약 시트를 만들지 않는다
    assert u.post("/bundle/plan", json={"file_ids": fids}).json()["summary"] is None
    assert u.post("/bundle/plan", json={"file_ids": fids, "unknown": 1}).status_code == 422


def test_summary_keeps_blank_not_zero_and_upper_region_dedup(bd) -> None:
    u = bd["bd-user1"]
    blank = _upload(u, "빈칸.xlsx", _xlsx({"현황": [HEAD, ["서울특별시", "종로구", 10], ["서울특별시", "용산구", "-"],
                                                   ["서울특별시", "성동구", 30]]}))  # fmt: skip
    upper = _upload(u, "상하위.xlsx", _xlsx({"현황": [HEAD, ["부산광역시", None, 300], ["부산광역시", "중구", 100],
                                                     ["부산광역시", "서구", 200]]}))  # fmt: skip
    for f in (blank, upper):
        _confirm(u, f, MAP)
    item = {"field_key": "population", "aggregates": ["sum", "mean", "count"]}
    run, wb, plan = _make(u, {"file_ids": [blank, upper], "summary": {"items": [item]}})
    # '-'는 0이 아니라 빠지고(평균 분모에도 없음), 하위 합계와 같은 부산광역시 300은 중복으로 빠진다
    assert _table(_summary_rows(wb), "인구수") == (["범위", "합계", "평균", "유효 숫자 개수"], [["전체", 340, 85, 4]])
    saved = run["summary"]["summary_items"][0]
    assert saved["partial"] is True and saved["rows_excluded"] == 2
    text = " ".join(str(v) for r in _summary_rows(wb) for v in r if v is not None)
    assert "부분 집계" in text and "제외 사유" in text
    assert wb["빈칸"]["E3"].value == "-" and wb["빈칸"]["A3"].value == 4  # 데이터 시트는 원래 값 그대로 (원본 4행)


def test_one_sheet_picked_from_multi_sheet_file_keeps_sheet_name(bd) -> None:
    u = bd["bd-user1"]
    multi = _upload(u, "두시트.xlsx", _xlsx({"1월": [HEAD, *DATASET["a"]], "2월": [HEAD, *DATASET["b"]]}))
    plan = _plan(u, {"file_ids": [multi], "sheets": [{"file_id": multi, "sheet_id": "s1"}]})
    assert plan["ready"] and [s["name"] for s in plan["sheets"]] == ["두시트-2월"]


# ---------------------------------------------------------------- 계산 엔진이 요구하는 선택
# 화면에서 고른 값이 결과에 남는지


def _sheet_text(wb: openpyxl.Workbook) -> str:
    return " ".join(str(v) for r in _summary_rows(wb) for v in r if v is not None)


def test_region_level_choice_upper_100_lower_90_is_recorded(bd) -> None:
    u = bd["bd-user1"]
    fid = _upload(u, "상하위차이.xlsx", _xlsx({"현황": [HEAD, ["서울특별시", None, 100], ["서울특별시", "종로구", 40],
                                                       ["서울특별시", "용산구", 50]]}))  # fmt: skip
    _confirm(u, fid, MAP)
    item = {"field_key": "population", "aggregates": ["sum"]}
    plan = _plan(u, {"file_ids": [fid], "summary": {"items": [item]}})
    (it,) = plan["summary"]["items"]
    assert not plan["ready"] and any("집계 수준을 고르세요" in p for p in plan["problems"])
    assert it["region_levels"]["mismatch"] == 1 and it["region_levels"]["conflicts"][0]["difference"] == 10
    for level, total, text in (("upper", 100, "상위 지역만"), ("lower", 90, "하위 지역만")):
        run, wb, plan = _make(u, {"file_ids": [fid], "summary": {"items": [{**item, "region_level": level}]}})
        assert plan["summary"]["items"][0]["choices"]["region_level"] == level
        assert _table(_summary_rows(wb), "인구수") == (["범위", "합계"], [["전체", total]])
        sheet = _sheet_text(wb)
        assert "집계 수준" in sheet and text in sheet and "차이 10" in sheet and "선택한 지역 수준" not in sheet
        saved = run["summary"]["summary_items"][0]
        assert saved["region_level"] == level and text in saved["region_level_text"]
        assert saved["level_conflicts"] and "차이 10" in saved["level_conflicts"][0]


def test_variant_unit_and_repeated_row_choices_are_required_and_recorded(bd) -> None:
    u = bd["bd-user1"]
    item = {"field_key": "population", "aggregates": ["sum"]}
    head = ["시도", "시군구", "인구수(남)", "인구수(여)"]
    v = _upload(u, "남여.xlsx", _xlsx({"현황": [head, ["광주광역시", "동구", 11, 12], ["광주광역시", "서구", 21, 22]]}))
    _confirm(u, v, {"시도": "sido", "시군구": "sigungu", "인구수(남)": "population", "인구수(여)": "population"})
    plan = _plan(u, {"file_ids": [v], "summary": {"items": [item]}})
    assert not plan["ready"] and plan["summary"]["items"][0]["variants"] == ["남", "여"]
    run, wb, _ = _make(u, {"file_ids": [v], "summary": {"items": [{**item, "variant": "남"}]}})
    assert _table(_summary_rows(wb), "인구수")[1] == [["전체", 32]]
    assert "구분" in _sheet_text(wb) and "남 (사용자 선택)" in _sheet_text(wb)
    assert (
        run["summary"]["summary_items"][0]["variant"] == "남" and run["summary"]["summary_items"][0]["variant_chosen"]
    )
    x1 = _upload(u, "명.xlsx", _xlsx({"현황": [HEAD, ["울산광역시", "중구", 7]]}))
    x2 = _upload(
        u, "천명.xlsx", _xlsx({"현황": [HEAD, ["울산광역시", "남구", 3]]}, title="인구 현황 2026년 9월 (단위: 천명)")
    )
    for f in (x1, x2):
        _confirm(u, f, MAP)
    plan = _plan(u, {"file_ids": [x1, x2], "summary": {"items": [{**item, "unit": "명"}]}})
    assert not plan["ready"] and any(
        "단위가 다름" in p and "빼기" in p for p in plan["problems"]
    )  # 다른 단위는 확인 필요
    drop = [{"file_id": x2, "sheet_id": "s0", "exclude": True}]
    body = {"file_ids": [x1, x2], "summary": {"items": [{**item, "unit": "명", "links": drop}]}}
    run, wb, _ = _make(u, body)
    assert _table(_summary_rows(wb), "인구수")[1] == [["전체", 7]]
    assert "계산 단위" in _sheet_text(wb) and "자동 환산 없음" in _sheet_text(wb)
    assert run["summary"]["summary_items"][0]["unit_choice"] == "명"
    d1 = _upload(u, "반복행.xlsx", _xlsx({"현황": [HEAD, ["인천광역시", "중구", 10], ["인천광역시", "중구", 20],
                                                  ["인천광역시", "동구", 5]]}))  # fmt: skip
    _confirm(u, d1, MAP)
    plan = _plan(u, {"file_ids": [d1], "summary": {"items": [item]}})
    assert not plan["ready"] and plan["summary"]["items"][0]["links"][0]["repeated"]["keys"] == 1
    link = {"file_id": d1, "sheet_id": "s0", "rows_are_records": True}
    run, wb, _ = _make(u, {"file_ids": [d1], "summary": {"items": [{**item, "links": [link]}]}})
    assert _table(_summary_rows(wb), "인구수")[1] == [["전체", 35]]
    assert "서로 다른 기록" in _sheet_text(wb) and run["summary"]["summary_items"][0]["records"]
    # 반복 행 확인으로 파일 간 중복 차단을 우회하지 않는다
    d2 = _upload(u, "반복행상대.xlsx", _xlsx({"현황": [HEAD, ["인천광역시", "중구", 7]]}))
    _confirm(u, d2, MAP)
    plan = _plan(u, {"file_ids": [d1, d2], "summary": {"items": [{**item, "links": [link]}]}})
    assert not plan["ready"] and any("같은 대상" in p for p in plan["problems"])


def test_base_date_conflict_is_not_bypassed_by_region_level_choice(bd) -> None:
    u = bd["bd-user1"]
    head = ["시도", "시군구", "기준일", "인구수"]
    mapping = {"시도": "sido", "시군구": "sigungu", "기준일": "base_date", "인구수": "population"}
    up = _upload(u, "기준일상위.xlsx", _xlsx({"현황": [head, ["서울특별시", None, "2026-09-15", 100]]}))
    low = _upload(u, "기준일하위.xlsx", _xlsx({"현황": [head, ["서울특별시", "종로구", "2026-09-30", 40]]}))
    for f in (up, low):
        _confirm(u, f, mapping)
    for level in (None, "upper", "lower"):
        item = {"field_key": "population", "aggregates": ["sum"], "region_level": level}
        plan = _plan(u, {"file_ids": [up, low], "summary": {"items": [item]}})
        assert not plan["ready"] and any("기준일" in p for p in plan["problems"]), (level, plan["problems"])


def test_server_still_rejects_links_to_sheets_not_in_the_result(bd) -> None:
    u = bd["bd-user1"]
    a = _upload(u, "연결a.xlsx", _xlsx({"현황": [HEAD, *DATASET["a"]]}))
    b = _upload(u, "연결b.xlsx", _xlsx({"현황": [HEAD, *DATASET["b"]]}))
    item = {
        "field_key": "population",
        "aggregates": ["sum"],
        "links": [{"file_id": a, "sheet_id": "s0", "exclude": True}],
    }
    res = u.post("/bundle/plan", json={"file_ids": [b], "summary": {"items": [item]}})
    assert res.status_code == 422 and "결과에 넣지 않은 시트의 연결" in res.json()["detail"]
