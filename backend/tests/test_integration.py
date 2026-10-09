import hashlib
import io
import sqlite3
from datetime import datetime
from typing import Any

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app.config import DATA_DIR

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
DATA_SHEET, CHANGE_SHEET, SUMMARY_SHEET = "통합 데이터", "변경 내역", "처리 요약"


def _xlsx(rows: list[list], merges: tuple[str, ...] = ()) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "현황"
    for r in rows:
        ws.append(r)
    for m in merges:
        ws.merge_cells(m)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


ROWS = [
    ["기관명", "부서", "집행액", "시도", "시군구", "기준일자", "비고"],
    ["행안부", "재정과", "1,000", "서울", "강서구", datetime(2026, 9, 30), "  확인  "],
    [" 기상청 ", None, "2,500", "서울특별시", "강서구", datetime(2026, 9, 30), None],
    ["기상청", "총무과", 3000, "부산광역시", "강서구", datetime(2026, 9, 30), "정상"],
]


@pytest.fixture(scope="module")
def client():
    from app import files, integration, review, validation_api
    from app.db import Base, engine
    from app.main import app
    from app.mapping import MappingEngine, load_dictionary

    def no_llm(system: str, user: str, schema: dict) -> dict[str, Any]:
        return {"key": "none", "confidence": 0, "reason": "해당 없음"}

    mapping_engine = MappingEngine(load_dictionary(), None, no_llm)
    patch = pytest.MonkeyPatch()
    for module in (files, validation_api, review, integration):
        patch.setattr(module, "get_engine", lambda: mapping_engine)
    Base.metadata.create_all(engine)
    with TestClient(app) as c:
        glossary_csv = (
            "코드,정식명칭,유형,시도코드,시군구코드,별칭\n"
            "1741000,행정안전부,중앙행정기관,,,행안부\n"
            "1360000,기상청,중앙행정기관,,,\n"
            "6110000,서울특별시,광역자치단체,11,,\n"
            "3150000,서울특별시 강서구,기초자치단체,11,500,\n"
            "6260000,부산광역시,광역자치단체,26,,\n"
            "3360000,부산광역시 강서구,기초자치단체,26,440,\n"
        )
        res = c.post("/glossary/import", files=[("file", ("g.csv", glossary_csv.encode("utf-8"), "text/csv"))])
        assert res.json()["applied"], res.text
        yield c
    patch.undo()


def _ready_review(client: TestClient, data: bytes, name: str) -> tuple[int, int]:
    fid = client.post("/files", files=[("files", (name, data, XLSX))]).json()[0]["id"]
    review = client.post("/reviews", json={"file_ids": [fid], "expected_period": "2026-09"}).json()
    for item in [i for i in review["items"] if i["category"] == "pending"]:
        candidates = item["payload"].get("candidates") or []
        choice = (
            (item["payload"].get("field_key") or (candidates[0].get("key") if candidates else None))
            if item["kind"] == "mapping"
            else (candidates[0].get("code") if candidates else None)
        )
        decision = "approved" if (choice or item["kind"] == "pii") else "rejected"
        res = client.post(
            f"/reviews/{review['id']}/items/{item['id']}/decision", json={"decision": decision, "choice": choice}
        )
        assert res.status_code == 200, res.text
    done = client.post(f"/reviews/{review['id']}/complete")
    assert done.status_code == 200, done.text
    return review["id"], fid


def _db() -> sqlite3.Connection:
    return sqlite3.connect(DATA_DIR / "moajwo.db")


def test_integration_creates_three_sheets_with_painted_changes(client: TestClient) -> None:
    rid, fid = _ready_review(client, _xlsx(ROWS, ("B2:B3",)), "통합대상.xlsx")
    with _db() as con:
        stored, original_hash = con.execute(
            "select stored_path, sha256 from source_files where id = ?", (fid,)
        ).fetchone()

    res = client.post(f"/reviews/{rid}/integrate")
    assert res.status_code == 201, res.text
    run = res.json()
    assert run["status"] == "done" and run["row_count"] == 3

    # 결과 파일: 새 파일, outputs 폴더, 원본은 그대로
    content = client.get(f"/runs/{run['id']}/download").content
    wb = openpyxl.load_workbook(io.BytesIO(content))
    assert wb.sheetnames == [DATA_SHEET, CHANGE_SHEET, SUMMARY_SHEET]
    assert hashlib.sha256((DATA_DIR / stored).read_bytes()).hexdigest() == original_hash

    data = wb[DATA_SHEET]
    header = [c.value for c in data[1]]
    assert header[:3] == ["원본 파일", "원본 시트", "원본 행"]
    col = {name: i + 1 for i, name in enumerate(header)}
    assert {"기관명", "기관명 코드", "담당부서", "집행액", "시도", "시군구", "기준일자", "비고"} <= set(col)

    with _db() as con:
        changes = con.execute(
            "select output_row, column, kind, before, after, reason, confidence from run_changes where run_id = ?",
            (run["id"],),
        ).fetchall()
    cell_changes = {(r, c): (k, b, a) for r, c, k, b, a, _, _ in changes if k != "column"}
    assert len(changes) == run["change_count"] and all(reason and conf > 0 for *_, reason, conf in changes)

    # 색칠한 셀 = run_changes (1:1), 칠한 셀에는 메모(이전 값·근거)
    painted = {
        (cell.row, header[cell.column - 1])
        for row in data.iter_rows(min_row=2)
        for cell in row
        if cell.fill is not None and cell.fill.fgColor.rgb not in (None, "00000000")
    }
    assert painted == set(cell_changes)
    for r, c in cell_changes:
        cell = data.cell(row=r, column=col[c])
        assert cell.comment is not None and "이전 값" in cell.comment.text

    kinds = {k for k, *_ in cell_changes.values()}
    assert {"number", "ffill", "entity", "trim"} <= kinds
    assert cell_changes[(2, "집행액")] == ("number", "1,000", "1000")
    assert cell_changes[(2, "기관명")][0] == "entity" and cell_changes[(2, "기관명")][2] == "행정안전부"
    assert cell_changes[(2, "시도")] == ("entity", "서울", "서울특별시")
    assert cell_changes[(3, "담당부서")] == ("ffill", None, "재정과")
    assert cell_changes[(3, "기관명")][0] == "trim"
    assert (2, "시군구") not in cell_changes  # '강서구'는 그대로 (약칭 시도 '서울'이어도 '특별시 강서구'가 되지 않음)
    assert (2, "기준일자") not in cell_changes  # 날짜는 원본 형식 그대로
    assert data.cell(row=2, column=col["기관명 코드"]).value == "1741000"
    assert data.cell(row=4, column=col["시군구 코드"]).value == "3360000"  # 부산 강서구 (동명 시군구 구분)

    # 변경 내역 시트 = run_changes 전체 (열 이름 매핑 포함)
    sheet_rows = list(wb[CHANGE_SHEET].iter_rows(min_row=2, values_only=True))
    assert len(sheet_rows) == len(changes)
    assert any(r[9] == "열 이름 매핑" and r[6] == "부서" and r[5] == "담당부서" for r in sheet_rows)

    summary = [list(r) for r in wb[SUMMARY_SHEET].iter_rows(values_only=True)]
    flat = [str(v) for r in summary for v in r if v is not None]
    assert "통합대상.xlsx" in flat and "처리 시각" in flat and "경고 건수" in flat
    assert any(r[0] == "통합 행 수" and r[1] == 3 for r in summary if r)

    # 같은 검토로 다시 실행하면 항상 새 파일
    again = client.post(f"/reviews/{rid}/integrate").json()
    assert again["file_name"] != run["file_name"]
    outputs = sorted(p.name for p in (DATA_DIR / "outputs").glob("*.xlsx"))
    assert run["file_name"] in outputs and again["file_name"] in outputs
    assert not list((DATA_DIR / "outputs").glob("*.tmp"))

    page = client.get(f"/runs/{run['id']}/changes", params={"kind": "number", "limit": 2}).json()
    assert page["total"] >= 2 and len(page["items"]) == 2


def test_integration_requires_completed_review(client: TestClient) -> None:
    fid = client.post("/files", files=[("files", ("n.xlsx", _xlsx(ROWS), XLSX))]).json()[0]["id"]
    review = client.post("/reviews", json={"file_ids": [fid]}).json()
    res = client.post(f"/reviews/{review['id']}/integrate")
    assert res.status_code == 409 and "검토 완료" in res.json()["detail"]
    assert client.post("/reviews/999999/integrate").status_code == 404


def test_missing_change_record_fails_without_output(client: TestClient, monkeypatch: pytest.MonkeyPatch) -> None:
    from app import integration

    rid, _ = _ready_review(client, _xlsx(ROWS, ("B2:B3",)), "누락검사.xlsx")
    real = integration.write_workbook
    before = set((DATA_DIR / "outputs").glob("*.xlsx"))

    def lossy(path, built, summary):  # 기록 하나를 빼먹고 칠하는 상황을 흉내 낸다
        painted = real(path, built, summary)
        return painted - 1

    monkeypatch.setattr(integration, "write_workbook", lossy)
    res = client.post(f"/reviews/{rid}/integrate")
    assert res.status_code == 409 and "변경 기록 불일치" in res.json()["detail"]
    assert set((DATA_DIR / "outputs").glob("*.xlsx")) == before  # 결과 파일을 남기지 않음
    failed = client.get("/runs", params={"review_id": rid}).json()[0]
    assert failed["status"] == "failed" and failed["change_count"] == 0
    with _db() as con:
        assert con.execute("select count(*) from run_changes where run_id = ?", (failed["id"],)).fetchone()[0] == 0


def test_review_preview_matches_integration(client: TestClient) -> None:
    """검토 패널에 예고한 자동 보정(숫자·병합 채움·공백)이 통합 결과의 run_changes와 같은 칸 수인지."""
    rid, _ = _ready_review(client, _xlsx(ROWS, ("B2:B3",)), "일치검사.xlsx")
    review = client.get(f"/reviews/{rid}").json()
    preview: dict[str, int] = {}
    for item in review["items"]:
        if item["category"] == "auto" and item["kind"] in ("number", "ffill", "trim"):
            preview[item["kind"]] = preview.get(item["kind"], 0) + item["payload"]["count"]
    # 시군구는 열 값('강서구')이 바뀌지 않으므로 보정으로 예고하지 않는다
    assert not [i for i in review["items"] if i["kind"] == "entity" and i["column"] == "시군구"]

    run = client.post(f"/reviews/{rid}/integrate").json()
    changes = client.get(f"/runs/{run['id']}/changes", params={"limit": 1000}).json()["items"]
    actual: dict[str, int] = {}
    for ch in changes:
        if ch["kind"] in ("number", "ffill", "trim"):
            actual[ch["kind"]] = actual.get(ch["kind"], 0) + 1
    assert preview == actual and {"number", "ffill", "trim"} <= set(actual)


def _csv_review(client: TestClient, name: str) -> int:
    """사전에 없는 열(설치장소상세위치·운영메모)이 있는 CSV. 매핑 확인 항목은 모두 거부('해당 없음')한다."""
    lines = ["기관명,집행액,설치장소상세위치,운영메모", "행안부,1000,본관 1층,", "기상청,2500,,야간 운영"]
    data = "\r\n".join(lines).encode("cp949")
    fid = client.post("/files", files=[("files", (name, data, "text/csv"))]).json()[0]["id"]
    review = client.post("/reviews", json={"file_ids": [fid]}).json()
    for item in [i for i in review["items"] if i["category"] == "pending"]:
        if item["kind"] == "mapping":
            body = {"decision": "rejected"}
        else:
            candidates = item["payload"].get("candidates") or []
            body = {"decision": "approved", "choice": candidates[0]["code"]} if candidates else {"decision": "rejected"}
        assert client.post(f"/reviews/{review['id']}/items/{item['id']}/decision", json=body).status_code == 200
    assert client.post(f"/reviews/{review['id']}/complete").status_code == 200
    return review["id"]


def test_unmapped_columns_kept_with_original_names_when_asked(client: TestClient) -> None:
    rid = _csv_review(client, "설치정보.csv")
    plain = client.post(f"/reviews/{rid}/integrate").json()  # 기본: 넣지 않음 (기존 동작)
    assert "설치장소상세위치" not in plain["summary"]["columns"]
    assert {d["column"] for d in plain["summary"]["dropped_columns"]} >= {"설치장소상세위치", "운영메모"}

    res = client.post(f"/reviews/{rid}/integrate", json={"include_unmapped": True})
    assert res.status_code == 201, res.text
    run = res.json()
    cols = run["summary"]["columns"]
    assert {"설치장소상세위치", "운영메모"} <= set(cols) and run["summary"]["include_unmapped"] is True
    assert not [d for d in run["summary"]["dropped_columns"] if d["column"] in ("설치장소상세위치", "운영메모")]
    wb = openpyxl.load_workbook(io.BytesIO(client.get(f"/runs/{run['id']}/download").content))
    data = wb[DATA_SHEET]
    header = [c.value for c in data[1]]
    j = header.index("설치장소상세위치") + 1
    assert [data.cell(row=r, column=j).value for r in (2, 3)] == ["본관 1층", None]
    summary = {r[0]: r[1] for r in wb[SUMMARY_SHEET].iter_rows(values_only=True) if r and r[0]}
    assert summary["표준 항목이 없는 열"] == "원래 이름으로 포함 (2개)"


def test_blank_csv_cells_are_not_changes(client: TestClient) -> None:
    """CSV의 빈 칸('')을 '바뀐 셀'로 기록·색칠하지 않는다 (실제 자료에서 빈 칸 수천 개가 변경으로 잡히던 문제)."""
    rid = _csv_review(client, "빈칸.csv")
    run = client.post(f"/reviews/{rid}/integrate", json={"include_unmapped": True}).json()
    changes = client.get(f"/runs/{run['id']}/changes", params={"limit": 1000}).json()["items"]
    blanks = [c for c in changes if c["column"] in ("설치장소상세위치", "운영메모")]
    assert blanks == [] and "other" not in run["summary"]["change_counts"]


def test_sigungu_full_name_kept_without_sido_column(client: TestClient) -> None:
    """시도 열이 없으면 '서울특별시 강서구'를 '강서구'로 줄이지 않는다.

    실제 자료(발급제증명 시군구명)에서 6만여 칸이 바뀌던 문제.
    """
    lines = ["시군구명,건수", "서울특별시 강서구,10", "부산광역시 강서구,20"]
    fid = client.post(
        "/files", files=[("files", ("시군구만.csv", "\r\n".join(lines).encode("cp949"), "text/csv"))]
    ).json()[0]["id"]
    review = client.post("/reviews", json={"file_ids": [fid]}).json()
    for item in [i for i in review["items"] if i["category"] == "pending"]:
        client.post(f"/reviews/{review['id']}/items/{item['id']}/decision", json={"decision": "rejected"})
    assert client.post(f"/reviews/{review['id']}/complete").status_code == 200
    run = client.post(f"/reviews/{review['id']}/integrate").json()
    wb = openpyxl.load_workbook(io.BytesIO(client.get(f"/runs/{run['id']}/download").content))
    data = wb[DATA_SHEET]
    header = [c.value for c in data[1]]
    j = header.index("시군구") + 1
    assert [data.cell(row=r, column=j).value for r in (2, 3)] == ["서울특별시 강서구", "부산광역시 강서구"]
    assert [data.cell(row=r, column=header.index("시군구 코드") + 1).value for r in (2, 3)] == ["3150000", "3360000"]
    assert "entity" not in run["summary"]["change_counts"]
