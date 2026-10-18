"""검토 결과와 열 매핑 확정의 일치 — 검사한 뒤 확정이 바뀌면 이전 검사 결과로 완료·통합하지 않는다 (실제 쿠키 로그인).

재현 (cfef11a): 대상 열이 미확정인 채로 검토를 만들고 검토의 확인 필요 항목에서 expenditure로 승인하면, 합계를 다시
검사하지 않고 완료·통합되어 합계 오류(2,000,000,000,000,000 ≠ 하위 합계 1,000,000,000,000,300) 자료가 통합됐다.
완료한 뒤 열 매핑 화면에서 확정을 바꿔도 이전 검토로 통합됐다.

시험마다 새 전용 계정을 쓴다 — 열 매핑 화면의 확정은 그 사용자의 학습 별칭이 되어 같은 이름의 열이 다음 검토에서
자동 적용되므로, 시험끼리 섞이지 않게 한다.
"""

import csv
import io
import itertools
import json
from typing import Any

import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import func, select

from tests.test_accounts import owner  # noqa: F401
from tests.test_auth import _client, app_real_login  # noqa: F401

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
TITLE = "2026년 9월 사업별 집행 현황 (단위: 원)"
BAD = [["합계", "2000000000000000"], ["사업A", "1000000000000000"], ["사업B", "100"], ["사업C", "200"]]
OK = [["합계", "1000000000000300"], ["사업A", "1000000000000000"], ["사업B", "100"], ["사업C", "200"]]
COLUMNS = ["비식별처리비용", "식별부호"]
_users = itertools.count(1)


@pytest.fixture
def user(app_real_login, owner) -> TestClient:  # noqa: F811
    name = f"rc-user{next(_users)}"
    res = owner.post("/users", json={"username": name, "display_name": name.upper(), "status": "active",
                                     "password": f"pass-{name}"})  # fmt: skip
    assert res.status_code == 201, res.text
    return _client(app_real_login, name, f"pass-{name}")


def _upload(c: TestClient, name: str, header: list[str], rows: list[list], kind: str = "csv") -> int:
    data = [[TITLE], header, *rows]
    if kind == "csv":
        buf = io.StringIO()
        csv.writer(buf, lineterminator="\n").writerows(data)
        content, ctype = buf.getvalue().encode("utf-8"), "text/csv"
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        for r in data:
            ws.append(r)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.data_type = "s"  # 모두 글자 셀
        out = io.BytesIO()
        wb.save(out)
        content, ctype = out.getvalue(), XLSX
    res = c.post("/files", files=[("files", (f"{name}.{kind}", content, ctype))])
    assert res.status_code == 200 and res.json()[0]["id"], res.text
    return res.json()[0]["id"]


def _confirm(c: TestClient, fid: int, mapping: dict[str, str | None]) -> None:
    """열 매핑 화면의 확정 (검토 밖)."""
    pv = c.get(f"/files/{fid}/preview").json()
    ids = {col["name"]: col["column_id"] for col in pv["columns"]}
    items = [{"sheet_id": pv["sheet_id"], "column_id": ids[n], "field_key": k} for n, k in mapping.items()]
    assert c.put(f"/files/{fid}/mapping/decisions", json={"decisions": items}).status_code == 200


def _preview_column(c: TestClient, fid: int, column: str) -> list:
    pv = c.get(f"/files/{fid}/preview").json()
    j = [col["name"] for col in pv["columns"]].index(column)
    return [r["cells"][j] for r in pv["rows"]]


def _review(c: TestClient, fids: list[int]) -> dict:
    res = c.post("/reviews", json={"file_ids": fids})
    assert res.status_code == 201, res.text
    return res.json()


def _pending(review: dict, column: str) -> dict:
    return next(
        i for i in review["items"] if i["category"] == "pending" and i["kind"] == "mapping" and i["column"] == column
    )


def _decide(c: TestClient, review: dict, item: dict, body: dict) -> dict:
    res = c.post(f"/reviews/{review['id']}/items/{item['id']}/decision", json=body)
    assert res.status_code == 200, res.text
    return res.json()


def _state() -> dict[str, Any]:
    """통합 실행 기록·변경 내역·결과 파일 수 (거절된 요청이 아무것도 만들지 않았는지)."""
    from app.db import SessionLocal
    from app.integration import OUTPUT_DIR
    from app.models import IntegrationRun, RunChange

    with SessionLocal() as db:
        runs = db.scalar(select(func.count(IntegrationRun.id)))
        changes = db.scalar(select(func.count(RunChange.id)))
    return {"runs": runs, "changes": changes, "files": sorted(p.name for p in OUTPUT_DIR.glob("*"))}


def _excel(c: TestClient, run_id: int, label: str) -> list[tuple[Any, str]]:
    grid = list(openpyxl.load_workbook(io.BytesIO(c.get(f"/runs/{run_id}/download").content)).worksheets[0].iter_rows())
    k = next(i for i, cell in enumerate(grid[0]) if cell.value == label)
    return [(row[k].value, row[k].data_type) for row in grid[1:] if row[k].value is not None]


def _sum_mismatch(review: dict) -> list[str]:
    return [i["reason"] for i in review["items"] if i["category"] == "block" and i["kind"] == "sum_mismatch"]


# ---------------------------------------------------------------- 검토 중 승인


@pytest.mark.parametrize("column", COLUMNS)
@pytest.mark.parametrize("kind", ["csv", "xlsx"])
def test_quantity_approved_in_review_needs_recheck_before_complete_and_integrate(user, column: str, kind: str) -> None:
    fid = _upload(user, f"검토중승인_{column}", ["구분", column], BAD, kind)
    _confirm(user, fid, {"구분": "department"})  # 대상 열은 미확정
    review = _review(user, [fid])
    assert review["counts"]["block"] == 0  # 글자로 남은 합계는 검사하지 못했다 (미확정)
    assert review["needs_recheck"] is False

    review = _decide(user, review, _pending(review, column), {"decision": "approved", "choice": "expenditure"})
    reason = f"검토중승인_{column}.{kind} '{column}' 열: 검사에 쓰지 않음 → 확정 집행액"
    assert (review["needs_recheck"], review["can_complete"], review["recheck_reasons"]) == (True, False, [reason])
    assert review["summary"].startswith("다시 검사 필요 — 이 검사 결과로 완료·통합할 수 없습니다")
    assert user.get(f"/reviews/{review['id']}").json()["recheck_reasons"] == [reason]

    before = _state()
    done = user.post(f"/reviews/{review['id']}/complete")
    assert done.status_code == 409 and done.json()["detail"].startswith("다시 검사 필요"), done.text
    after = user.get(f"/reviews/{review['id']}").json()
    assert (after["status"], after["completed_at"]) == ("in_review", None)
    assert user.post(f"/reviews/{review['id']}/integrate", json={"include_unmapped": True}).status_code == 409
    assert _state() == before  # 완료·실행 기록·결과 파일·변경 내역 모두 생기지 않음

    rechecked = user.post(f"/reviews/{review['id']}/recheck", json={"replace": {}})
    assert rechecked.status_code == 201, rechecked.text
    new = rechecked.json()
    assert new["needs_recheck"] is False and new["can_complete"] is False
    (message,) = _sum_mismatch(new)  # 최신 확정(집행액)으로 계산한 합계 검사
    assert "2,000,000,000,000,000" in message and "하위 3행 합계 1,000,000,000,000,300" in message
    assert "차이 999,999,999,999,700" in message
    assert user.get(f"/reviews/{review['id']}").json()["status"] == "superseded"  # 이전 검토 기록은 남는다
    assert user.post(f"/reviews/{new['id']}/complete").status_code == 409
    assert user.post(f"/reviews/{new['id']}/integrate", json={"include_unmapped": True}).status_code == 409
    assert _state() == before
    assert _preview_column(user, fid, column) == ["2000000000000000", "1000000000000000", "100", "200"]  # 원문 그대로


@pytest.mark.parametrize("column", COLUMNS)
def test_normal_total_approved_in_review_integrates_after_recheck(user, column: str) -> None:
    fid = _upload(user, f"정상승인_{column}", ["구분", column], OK)
    _confirm(user, fid, {"구분": "department"})
    review = _review(user, [fid])
    review = _decide(user, review, _pending(review, column), {"decision": "approved", "choice": "expenditure"})
    assert review["needs_recheck"] is True
    assert user.post(f"/reviews/{review['id']}/complete").status_code == 409

    new = user.post(f"/reviews/{review['id']}/recheck", json={"replace": {}}).json()
    assert (new["counts"]["block"], new["needs_recheck"], new["can_complete"]) == (0, False, True)
    assert user.post(f"/reviews/{new['id']}/complete").status_code == 200
    run = user.post(f"/reviews/{new['id']}/integrate", json={"include_unmapped": True})
    assert run.status_code == 201, run.text
    assert _excel(user, run.json()["id"], "집행액") == [(1000000000000000, "n"), (100, "n"), (200, "n")]
    assert _preview_column(user, fid, column) == ["1000000000000300", "1000000000000000", "100", "200"]


def test_review_approvals_reuse_the_check_only_when_its_input_is_unchanged(user) -> None:
    """검사가 읽지 않는 항목(비고)으로 빈 칸·숫자 없는 글자 열을 연결하거나 '해당 없음'이면 기존 검사 결과를 이어 쓰고,
    숫자 열을 수량 항목으로 연결하면 합계·단위 비교 입력이 바뀌므로 다시 검사해야 한다."""
    rows = [["합계", "1000000000000300", "6", "가"], ["사업A", "1000000000000000", "1", "나"],
            ["사업B", "100", "2", "다"], ["사업C", "200", "3", "라"]]  # fmt: skip
    fid = _upload(user, "영향없는승인", ["구분", "식별부호", "처리량값", "참고문구"], rows)
    _confirm(user, fid, {"구분": "department"})
    review = _review(user, [fid])
    review = _decide(user, review, _pending(review, "식별부호"), {"decision": "approved", "choice": "note"})
    assert review["needs_recheck"] is False  # 검사가 읽지 않는 항목 + 빈 칸·숫자 없는 글자 열
    review = _decide(user, review, _pending(review, "참고문구"), {"decision": "rejected"})
    assert review["needs_recheck"] is False  # '해당 없음' (아직 '처리량값'이 확인 필요로 남음)
    review = _decide(user, review, _pending(review, "처리량값"), {"decision": "approved", "choice": "count"})
    assert (review["needs_recheck"], review["can_complete"]) == (True, False)  # 숫자 열 → 수량 항목
    assert review["recheck_reasons"] == ["영향없는승인.csv '처리량값' 열: 검사에 쓰지 않음 → 확정 건수"]
    assert user.post(f"/reviews/{review['id']}/complete").status_code == 409
    review = user.post(f"/reviews/{review['id']}/recheck", json={"replace": {}}).json()
    assert (review["needs_recheck"], review["can_complete"]) == (False, True)
    assert user.post(f"/reviews/{review['id']}/complete").status_code == 200
    run = user.post(f"/reviews/{review['id']}/integrate", json={"include_unmapped": True})
    assert run.status_code == 201, run.text
    assert _excel(user, run.json()["id"], "비고") == [("1000000000000000", "s"), ("100", "s"), ("200", "s")]
    assert _excel(user, run.json()["id"], "건수") == [(1, "n"), (2, "n"), (3, "n")]


def test_changing_an_approval_back_to_a_neutral_choice_clears_recheck(user) -> None:
    column = COLUMNS[0]
    fid = _upload(user, "승인변경", ["구분", column], BAD)
    _confirm(user, fid, {"구분": "department"})
    review = _review(user, [fid])
    item = _pending(review, column)
    assert _decide(user, review, item, {"decision": "approved", "choice": "expenditure"})["needs_recheck"] is True
    review = _decide(user, review, item, {"decision": "rejected"})  # 다시 '해당 없음' — 검사할 때와 같은 계산
    assert (review["needs_recheck"], review["can_complete"]) == (False, True)


# ---------------------------------------------------------------- 완료 뒤 다른 화면에서 확정 변경


@pytest.mark.parametrize("column", COLUMNS)
def test_completed_review_rejects_integration_after_quantity_confirmed_elsewhere(user, column: str) -> None:
    fid = _upload(user, f"완료뒤수량_{column}", ["구분", column], BAD)
    _confirm(user, fid, {"구분": "department"})
    review = _review(user, [fid])
    review = _decide(user, review, _pending(review, column), {"decision": "rejected"})  # '해당 없음'으로 완료
    assert user.post(f"/reviews/{review['id']}/complete").status_code == 200

    _confirm(user, fid, {column: "expenditure"})  # 열 매핑 화면에서 수량으로 확정
    got = user.get(f"/reviews/{review['id']}").json()
    reason = f"완료뒤수량_{column}.csv '{column}' 열: 확정 해당 없음 → 확정 집행액"
    assert (got["status"], got["needs_recheck"], got["recheck_reasons"]) == ("ready", True, [reason])
    assert got["summary"].startswith("다시 검사 필요 — 이 검사 결과로 통합할 수 없습니다")

    before = _state()
    res = user.post(f"/reviews/{review['id']}/integrate", json={"include_unmapped": True})
    assert res.status_code == 409, res.text
    assert res.json()["detail"].startswith("검사한 뒤 열 매핑이 바뀌어 이 검토 결과로 통합할 수 없습니다")
    assert reason in res.json()["detail"]
    assert _state() == before

    new = user.post(f"/reviews/{review['id']}/recheck", json={"replace": {}}).json()
    (message,) = _sum_mismatch(new)
    assert "1,000,000,000,000,300" in message and "999,999,999,999,700" in message


@pytest.mark.parametrize("after", ["note", None], ids=["비고", "해당 없음"])
def test_completed_review_rejects_integration_after_quantity_changed_to_text(user, after: str | None) -> None:
    column = COLUMNS[0]
    fid = _upload(user, "완료뒤글자", ["구분", column], OK, "xlsx")
    _confirm(user, fid, {"구분": "department", column: "expenditure"})
    review = _review(user, [fid])
    assert review["can_complete"] is True
    assert user.post(f"/reviews/{review['id']}/complete").status_code == 200
    first = user.post(f"/reviews/{review['id']}/integrate", json={"include_unmapped": True})
    assert first.status_code == 201, first.text
    assert _excel(user, first.json()["id"], "집행액") == [(1000000000000000, "n"), (100, "n"), (200, "n")]

    _confirm(user, fid, {column: after})
    shown = "비고" if after else "해당 없음"
    got = user.get(f"/reviews/{review['id']}").json()
    assert got["recheck_reasons"] == [f"완료뒤글자.xlsx '{column}' 열: 확정 집행액 → 확정 {shown}"]
    before = _state()
    assert user.post(f"/reviews/{review['id']}/integrate", json={"include_unmapped": True}).status_code == 409
    assert _state() == before
    assert user.get(f"/runs/{first.json()['id']}").status_code == 200  # 이전 결과는 그대로
    assert user.get(f"/runs/{first.json()['id']}/download").status_code == 200

    new = user.post(f"/reviews/{review['id']}/recheck", json={"replace": {}}).json()
    assert user.post(f"/reviews/{new['id']}/complete").status_code == 200
    run = user.post(f"/reviews/{new['id']}/integrate", json={"include_unmapped": True})
    assert run.status_code == 201, run.text
    label = "비고" if after else column
    assert _excel(user, run.json()["id"], label) == [("1000000000000000", "s"), ("100", "s"), ("200", "s")]
    assert _preview_column(user, fid, column) == ["1000000000000300", "1000000000000000", "100", "200"]


def test_integration_uses_checked_state_when_decision_changes_during_build(user, monkeypatch) -> None:
    """통합이 확정을 읽은 뒤 다른 요청이 확정을 바꿔도, 이 실행은 검사한 확정으로 만들고 다음 요청은 거절한다."""
    from app import integration
    from app.db import SessionLocal
    from app.models import MappingDecision

    column = COLUMNS[0]
    fid = _upload(user, "경합", ["구분", column], OK)
    _confirm(user, fid, {"구분": "department", column: "expenditure"})
    review = _review(user, [fid])
    assert user.post(f"/reviews/{review['id']}/complete").status_code == 200
    original, changed = integration.load_parsed, []

    def change_then_load(db, file_id):  # 확정 비교(스냅숏) 다음, 파일을 읽는 순간에 끼어든 다른 요청
        if not changed:
            with SessionLocal() as other:
                d = other.scalar(
                    select(MappingDecision).where(
                        MappingDecision.source_file_id == file_id, MappingDecision.column_name == column
                    )
                )
                d.field_key = "note"
                other.commit()
            changed.append(file_id)
        return original(db, file_id)

    monkeypatch.setattr(integration, "load_parsed", change_then_load)
    run = user.post(f"/reviews/{review['id']}/integrate", json={"include_unmapped": True})
    assert run.status_code == 201, run.text
    assert changed == [fid]
    assert _excel(user, run.json()["id"], "집행액") == [(1000000000000000, "n"), (100, "n"), (200, "n")]

    got = user.get(f"/reviews/{review['id']}").json()
    assert got["recheck_reasons"] == [f"경합.csv '{column}' 열: 확정 집행액 → 확정 비고"]
    before = _state()
    res = user.post(f"/reviews/{review['id']}/integrate", json={"include_unmapped": True})
    assert res.status_code == 409 and "확정 집행액 → 확정 비고" in res.json()["detail"]
    assert _state() == before


# ---------------------------------------------------------------- 제외 파일·개인정보·이전 버전 검토


def test_excluded_file_changes_do_not_block_but_included_ones_do(user) -> None:
    column = COLUMNS[1]
    fid = _upload(user, "본자료", ["구분", column], OK)
    _confirm(user, fid, {"구분": "department", column: "expenditure"})
    pii_rows = [r + ["rc-contact@example.com"] for r in OK]
    pid = _upload(user, "연락처자료", ["구분", column, "담당자 메일"], pii_rows)
    _confirm(user, pid, {"구분": "department", column: "expenditure"})

    review = _review(user, [fid, pid])
    pii = next(i for i in review["items"] if i["kind"] == "pii" and i["category"] == "pending")
    review = _decide(user, review, pii, {"decision": "rejected"})  # 개인정보 파일 제외
    assert [f["id"] for f in review["files"] if f["excluded"]] == [pid]
    _confirm(user, pid, {column: None})  # 제외한 파일의 확정 변경은 이 통합과 관계없다
    got = user.get(f"/reviews/{review['id']}").json()
    assert (got["needs_recheck"], got["can_complete"]) == (False, True)
    assert user.post(f"/reviews/{review['id']}/complete").status_code == 200
    run = user.post(f"/reviews/{review['id']}/integrate", json={"include_unmapped": True})
    assert run.status_code == 201, run.text
    assert [f["id"] for f in run.json()["summary"]["files"] if f["excluded"]] == [pid]
    assert _excel(user, run.json()["id"], "집행액") == [(1000000000000000, "n"), (100, "n"), (200, "n")]

    included = _review(user, [fid, pid])  # 개인정보 포함을 승인한 파일은 통합 대상이므로 확정 변경을 본다
    for item in [i for i in included["items"] if i["category"] == "pending"]:  # 포함 승인 + 메일 열은 '해당 없음'
        included = _decide(user, included, item, {"decision": "approved" if item["kind"] == "pii" else "rejected"})
    assert (included["needs_recheck"], included["can_complete"]) == (False, True)
    _confirm(user, pid, {column: "expenditure"})
    got = user.get(f"/reviews/{included['id']}").json()
    assert got["recheck_reasons"] == [f"연락처자료.csv '{column}' 열: 확정 해당 없음 → 확정 집행액"]
    assert user.post(f"/reviews/{included['id']}/complete").status_code == 409


def test_review_params_do_not_expose_the_checked_state(user) -> None:
    fid = _upload(user, "조건표시", ["구분", COLUMNS[0]], OK)
    _confirm(user, fid, {"구분": "department", COLUMNS[0]: "expenditure"})
    review = _review(user, [fid])
    assert "check_basis" not in review["params"] and "verified_decisions" not in review["params"]
    assert review["params"]["file_ids"] == [fid]


@pytest.mark.parametrize("record", [None, "decisions_only"], ids=["기록 없음", "확정만 기록(이전 보완)"])
def test_review_without_a_check_basis_is_not_treated_as_unchanged(user, record: str | None) -> None:
    """검사 기준이 없는 검토(이전 버전·이전 보완의 확정만 담은 기록)는 같은지 확인할 수 없으므로 다시 검사해야 한다.
    이미 만든 결과의 조회·내려받기는 그대로다. (실제 이전 이미지로 만든 검토는 격리 환경에서 따로 확인)"""
    from app.check_basis import LEGACY_REASON
    from app.db import SessionLocal
    from app.models import ReviewSession

    column = COLUMNS[0]
    fid = _upload(user, f"이전검토_{record}", ["구분", column], OK)
    _confirm(user, fid, {"구분": "department", column: "expenditure"})
    review = _review(user, [fid])
    assert user.post(f"/reviews/{review['id']}/complete").status_code == 200
    first = user.post(f"/reviews/{review['id']}/integrate", json={"include_unmapped": True})
    assert first.status_code == 201, first.text
    with SessionLocal() as db:
        session = db.get(ReviewSession, review["id"])
        params = json.loads(session.params)
        params.pop("check_basis")
        if record == "decisions_only":
            params["verified_decisions"] = {str(fid): {"자료": {"구분": "department", column: "expenditure"}}}
        session.params = json.dumps(params, ensure_ascii=False)
        db.commit()

    got = user.get(f"/reviews/{review['id']}").json()
    reason = LEGACY_REASON
    assert (got["needs_recheck"], got["recheck_reasons"], got["can_complete"]) == (True, [reason], False)
    assert "verified_decisions" not in got["params"]
    before = _state()
    res = user.post(f"/reviews/{review['id']}/integrate", json={"include_unmapped": True})
    assert res.status_code == 409 and res.json()["detail"].startswith(reason), res.text
    assert _state() == before
    assert user.get(f"/runs/{first.json()['id']}").status_code == 200  # 이미 만든 결과는 그대로
    assert user.get(f"/runs/{first.json()['id']}/download").status_code == 200
    new = user.post(f"/reviews/{review['id']}/recheck", json={"replace": {}}).json()  # 다시 검사하면 기준이 생긴다
    assert (new["needs_recheck"], new["can_complete"]) == (False, True)
    assert user.post(f"/reviews/{new['id']}/complete").status_code == 200
    assert user.post(f"/reviews/{new['id']}/integrate", json={"include_unmapped": True}).status_code == 201


# ---------------------------------------------------------------- 합계 검사 문구: 실제 불일치와 일부를 읽지 못한 경우


def test_sum_mismatch_message_counts_children_left_as_text(user) -> None:
    rows = [["합계", "1000"], ["사업A", "500"], ["사업B", "약 300"], ["사업C", "100"]]
    fid = _upload(user, "일부글자", ["구분", "집행액"], rows)
    _confirm(user, fid, {"구분": "department", "집행액": "expenditure"})
    result = user.post("/validation", json={"file_ids": [fid]}).json()
    (block,) = [b for b in result["blocks"] if b["code"] == "sum_mismatch"]
    assert block["message"].startswith("'합계' 행의 '집행액' 값 1,000이(가) 하위 2행 합계 600와(과) 다릅니다")
    assert "(차이 400)." in block["message"]
    assert block["message"].endswith("글자로 남아 합계에 넣지 않은 하위 값 1개가 있습니다.")
    assert [w["code"] for w in result["warnings"]].count("quantity_not_number") == 1  # 읽을 수 없는 값은 경고로도

    fid = _upload(user, "실제불일치", ["구분", "집행액"], [["합계", "1000"], ["사업A", "500"], ["사업B", "100"]])
    _confirm(user, fid, {"구분": "department", "집행액": "expenditure"})
    (block,) = [b for b in user.post("/validation", json={"file_ids": [fid]}).json()["blocks"]]
    assert block["message"].endswith("(차이 400).")  # 모두 읽은 값의 실제 불일치에는 덧붙이지 않는다
