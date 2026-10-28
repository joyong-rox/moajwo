"""파일 읽기 한도 (환경별 메모리 예산 — docs/MEMORY.md): 칸 수·한 칸 글자 수·동시에 읽는 파일 수.

한도를 넘는 입력은 메모리를 크게 쓰기 전에 거절하고, 읽기 차례를 얻지 못한 업로드는 기록·원본을 남기지 않는다
(입력 오류가 아니므로 '읽지 못한 파일'로 두지 않음). 시험은 한도만 낮춰 작은 파일로 본다.
"""

import csv
import io
import threading
import time

import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import func, select

from app import parsing
from app.config import DATA_DIR
from app.db import Base, SessionLocal, engine
from app.main import app
from app.models import SourceFile

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture(scope="module")
def client():
    Base.metadata.create_all(engine)
    with TestClient(app) as c:
        yield c


def _csv(rows: list[list]) -> bytes:
    buf = io.StringIO()
    csv.writer(buf, lineterminator="\n").writerows(rows)
    return buf.getvalue().encode("utf-8")


def _xlsx(rows: list[list]) -> bytes:
    wb = openpyxl.Workbook()
    for r in rows:
        wb.active.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _state() -> tuple[int, list[str]]:
    with SessionLocal() as db:
        count = db.scalar(select(func.count(SourceFile.id)))
    return count, sorted(p.name for p in (DATA_DIR / "uploads").glob("*"))


ROWS = [["시도", "인구"], *[["서울특별시", i] for i in range(50)]]  # 102칸


def test_limits_are_reported_for_the_screen(client: TestClient) -> None:
    body = client.get("/files/limits").json()
    assert body["max_file_cells"] == parsing.MAX_FILE_CELLS and body["max_cell_chars"] == 32_767
    assert body["max_upload_mib"] >= 1 and body["max_xlsx_uncompressed_mib"] >= 1 and body["parse_slots"] >= 1


def test_too_many_cells_are_not_read(client: TestClient, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(parsing, "MAX_FILE_CELLS", 100)
    analyzed: list[str] = []
    real = parsing.analyze_sheet
    monkeypatch.setattr(parsing, "analyze_sheet", lambda *a, **k: analyzed.append(a[0]) or real(*a, **k))
    for name, data, ctype in (("많은칸.csv", _csv(ROWS), "text/csv"), ("많은칸.xlsx", _xlsx(ROWS), XLSX)):
        body = client.post("/files", files=[("files", (name, data, ctype))]).json()[0]
        assert "칸이 너무 많아 읽지 않았습니다" in body["error"] and "한도 100칸" in body["error"], body
    assert analyzed == []  # 행을 정리·분석하기 전에 멈춘다
    monkeypatch.setattr(parsing, "MAX_FILE_CELLS", 102)  # 경계: 한도와 같으면 읽는다
    body = client.post("/files", files=[("files", ("경계.csv", _csv(ROWS), "text/csv"))]).json()[0]
    assert body["error"] is None and body["data_rows"] == 50


def test_a_csv_cell_longer_than_excel_allows_is_not_read(client: TestClient) -> None:
    rows = [["시도", "설명"], ["서울특별시", "가" * 32_767], ["부산광역시", "나" * 32_768]]
    body = client.post("/files", files=[("files", ("긴칸.csv", _csv(rows), "text/csv"))]).json()[0]
    assert "3행에 32,767자가 넘는 칸" in body["error"], body


def test_a_busy_upload_leaves_no_record_or_file(client: TestClient, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(parsing, "PARSE_WAIT_S", 0.2)
    before = _state()
    assert parsing._PARSE_SLOTS.acquire(timeout=5)  # 다른 요청이 읽는 중인 것처럼 차례를 모두 잡아 둔다
    try:
        slots = parsing.PARSE_SLOTS
        for _ in range(slots - 1):
            assert parsing._PARSE_SLOTS.acquire(timeout=5)
        body = client.post("/files", files=[("files", ("대기.csv", _csv(ROWS), "text/csv"))]).json()[0]
    finally:
        for _ in range(slots):
            parsing._PARSE_SLOTS.release()
    assert body["id"] is None and "다른 파일을 읽는 중" in body["error"], body
    assert _state() == before  # 기록·원본 모두 남기지 않는다 — 다시 올리면 된다
    body = client.post("/files", files=[("files", ("대기.csv", _csv(ROWS), "text/csv"))]).json()[0]
    assert body["error"] is None  # 차례가 비면 그대로 읽는다


def test_reading_an_uploaded_file_while_busy_is_503(client: TestClient, monkeypatch: pytest.MonkeyPatch) -> None:
    fid = client.post("/files", files=[("files", ("나중.csv", _csv(ROWS), "text/csv"))]).json()[0]["id"]
    monkeypatch.setattr(parsing, "PARSE_WAIT_S", 0.2)
    slots = parsing.PARSE_SLOTS
    for _ in range(slots):
        assert parsing._PARSE_SLOTS.acquire(timeout=5)
    try:
        res = client.post(f"/files/{fid}/mapping?use_llm=false")
    finally:
        for _ in range(slots):
            parsing._PARSE_SLOTS.release()
    assert res.status_code == 503 and res.headers["retry-after"] == "10" and "다른 파일을 읽는 중" in res.text
    assert client.post(f"/files/{fid}/mapping?use_llm=false").status_code == 200


def test_different_files_are_not_parsed_at_the_same_time(tmp_path, monkeypatch: pytest.MonkeyPatch) -> None:
    """같은 파일 중복 읽기만이 아니라 서로 다른 파일도 동시에 PARSE_SLOTS개까지만 읽는다."""
    paths = []
    for i in range(3):
        path = tmp_path / f"f{i}.csv"
        path.write_bytes(_csv(ROWS))
        paths.append(path)
    running, peak = 0, 0
    lock = threading.Lock()
    real = parsing.analyze_sheet

    def slow(*a, **k):
        nonlocal running, peak
        with lock:
            running += 1
            peak = max(peak, running)
        time.sleep(0.2)
        with lock:
            running -= 1
        return real(*a, **k)

    monkeypatch.setattr(parsing, "analyze_sheet", slow)
    threads = [threading.Thread(target=parsing.parse_file, args=(p, p.name)) for p in paths]
    for t in threads:
        t.start()
    for t in threads:
        t.join(timeout=30)
    assert peak == parsing.PARSE_SLOTS


def _shared_xlsx(rows: list[list[str]], merges: tuple[str, ...] = ()) -> bytes:
    """Excel처럼 공유 문자열(sharedStrings)로 저장한 xlsx — 같은 긴 글자를 여러 칸이 가리켜도 파일은 작다
    (openpyxl은 칸마다 글자를 넣어 저장해 이런 파일을 만들지 않는다)."""
    import zipfile
    from xml.sax.saxutils import escape

    ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    rel = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    pkg = "http://schemas.openxmlformats.org/package/2006/relationships"
    ct = "application/vnd.openxmlformats-officedocument.spreadsheetml"
    strings: dict[str, int] = {}
    body = "".join(
        f'<row r="{r}">'
        + "".join(
            f'<c r="{chr(65 + c)}{r}" t="s"><v>{strings.setdefault(v, len(strings))}</v></c>' for c, v in enumerate(row)
        )
        + "</row>"
        for r, row in enumerate(rows, start=1)
    )
    refs = "".join(f'<mergeCell ref="{m}"/>' for m in merges)
    merge = f'<mergeCells count="{len(merges)}">{refs}</mergeCells>' if merges else ""
    sheet = f'<worksheet xmlns="{ns}"><sheetData>{body}</sheetData>{merge}</worksheet>'
    sst = "".join(f'<si><t xml:space="preserve">{escape(s)}</t></si>' for s in strings)
    parts = {
        "[Content_Types].xml": '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        f'<Override PartName="/xl/workbook.xml" ContentType="{ct}.sheet.main+xml"/>'
        f'<Override PartName="/xl/worksheets/sheet1.xml" ContentType="{ct}.worksheet+xml"/>'
        f'<Override PartName="/xl/sharedStrings.xml" ContentType="{ct}.sharedStrings+xml"/></Types>',
        "_rels/.rels": f'<Relationships xmlns="{pkg}"><Relationship Id="rId1" Type="{rel}/officeDocument"'
        ' Target="xl/workbook.xml"/></Relationships>',
        "xl/workbook.xml": f'<workbook xmlns="{ns}" xmlns:r="{rel}"><sheets><sheet name="s" sheetId="1" r:id="rId1"/>'
        "</sheets></workbook>",
        "xl/_rels/workbook.xml.rels": f'<Relationships xmlns="{pkg}">'
        f'<Relationship Id="rId1" Type="{rel}/worksheet" Target="worksheets/sheet1.xml"/>'
        f'<Relationship Id="rId2" Type="{rel}/sharedStrings" Target="sharedStrings.xml"/></Relationships>',
        "xl/worksheets/sheet1.xml": sheet,
        "xl/sharedStrings.xml": f'<sst xmlns="{ns}" uniqueCount="{len(strings)}">{sst}</sst>',
    }
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in parts.items():
            z.writestr(name, data)
    return buf.getvalue()


def test_a_long_shared_text_repeated_in_many_cells_is_not_read(
    client: TestClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    """공유 문자열 하나를 여러 칸이 가리키면 파일·풀었을 때 크기는 작아도, 정리할 때 칸마다 글자 사본이 생긴다
    (10.7KB 파일의 32,000자 × 800칸이 읽을 때 최고 98.8MiB) — 칸마다의 글자 길이 합으로 정리 전에 멈춘다."""
    long = " " + "가" * 2_000 + " "  # 앞뒤 공백 — 정리하면 칸마다 새 글자가 된다
    rows = [["시도", "비고"], *[["서울특별시", long] for _ in range(60)]]
    data = _shared_xlsx(rows)
    assert len(data) < 10_000  # 작은 파일
    chars = len("시도") + len("비고") + 60 * (len("서울특별시") + len(long))
    monkeypatch.setattr(parsing, "MAX_FILE_TEXT_CHARS", chars - 1)
    analyzed: list[str] = []
    real = parsing.analyze_sheet
    monkeypatch.setattr(parsing, "analyze_sheet", lambda *a, **k: analyzed.append(a[0]) or real(*a, **k))
    body = client.post("/files", files=[("files", ("반복글자.xlsx", data, XLSX))]).json()[0]
    assert "글자가 너무 많아 읽지 않았습니다" in body["error"] and f"{chars:,}자" in body["error"], body
    assert analyzed == []  # 정리(칸마다 사본) 전에 멈춘다 — 기록은 '읽지 못한 파일'로 남는다 (다른 읽기 오류와 같음)
    monkeypatch.setattr(parsing, "MAX_FILE_TEXT_CHARS", chars)  # 경계: 한도와 같으면 읽는다
    body = client.post("/files", files=[("files", ("반복글자.xlsx", data, XLSX))]).json()[0]
    assert body["error"] is None and body["data_rows"] == 60, body


def test_a_huge_merged_range_in_a_tiny_xlsx_is_refused_before_loading(
    client: TestClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    """병합 범위는 openpyxl이 칸마다 객체를 만든다 — 몇 KB 파일의 큰 병합도 열기 전에 칸 수 한도로 거절한다."""
    data = _shared_xlsx([["시도", "인구"], ["서울특별시", "10"]], merges=("A3:Z40000",))  # 26 × 39,998칸
    monkeypatch.setattr(parsing, "MAX_FILE_CELLS", 1_000_000)
    opened: list[object] = []
    real = parsing.openpyxl.load_workbook
    monkeypatch.setattr(parsing.openpyxl, "load_workbook", lambda *a, **k: opened.append(a) or real(*a, **k))
    monkeypatch.setattr(parsing, "MAX_FILE_CELLS", 26 * 39_998 - 1)
    body = client.post("/files", files=[("files", ("큰병합.xlsx", data, XLSX))]).json()[0]
    assert "병합 범위만 1,039,948칸" in body["error"] and opened == [], body
