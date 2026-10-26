"""파일에서 온 이름·값의 저장 규칙 (app/name_limits.py) — SQLite와 PostgreSQL에서 같은 거절을, 쓰기 전에, 원문 없이.

PostgreSQL은 varchar(255)보다 긴 이름을 쓰는 순간 DB 오류(500)를 냈고 NUL(0x00)이 든 글자는 드라이버가 보내지 않았다
(500). SQLite는 둘 다 그대로 저장했다. 같은 시험이 두 DB에서 돈다 (MOAJWO_TEST_DATABASE_URL). 길이는 글자 수 — 한글
255자(765바이트)는 두 DB 모두 저장된다. 실제로 저장하는 값만 거절한다 — 비교에만 쓰는 전월 자료의 이름, 결과 파일에만
쓰는 '코드' 열 이름은 막지 않는다. 거절 문구에는 고정 코드와 위치(파일 번호·시트 순번·열·행)만 (이름·값 원문 없음).
"""

import io
import json
import logging
import warnings
from typing import Any

import openpyxl
import pytest
from fastapi.testclient import TestClient

from tests.conftest import POSTGRES, db_rows

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PHONE = "010-1234-5678"


def _name(n: int, head: str) -> str:
    """n자 한글 열 이름 (개인정보 형식 없음)."""
    return head + "가" * (n - len(head))


N254, N255, N256 = _name(254, "인구"), _name(255, "세대"), _name(256, "예산")
RAW = "가" * 10  # 위 긴 이름에만 있는 조각 — 응답·기록에 이 조각이 있으면 원문을 되돌려 보낸 것


@pytest.fixture(scope="module")
def client():
    from app import automation, files, integration, mapping_review, review, validation_api
    from app.main import app
    from app.mapping import MappingEngine, load_dictionary

    def no_llm(system: str, user: str, schema: dict) -> dict[str, Any]:
        return {"key": "none", "confidence": 0, "reason": "-"}

    engine = MappingEngine(load_dictionary(), None, no_llm)
    patch = pytest.MonkeyPatch()
    for module in (files, validation_api, review, mapping_review, integration, automation):
        patch.setattr(module, "get_engine", lambda: engine)
    with TestClient(app) as c:
        glossary = "코드,정식명칭,유형\n1741000,행정안전부,중앙행정기관\n"
        res = c.post("/glossary/import", files=[("file", ("g.csv", glossary.encode("utf-8"), "text/csv"))])
        assert res.json()["applied"], res.text
        yield c
    patch.undo()


def _xlsx(rows: list[list], merges: tuple[str, ...] = (), sheets: tuple[str, ...] = ("현황",)) -> bytes:
    """마지막 시트에 rows를 쓴다 (앞 시트는 안내 한 줄)."""
    wb = openpyxl.Workbook()
    with warnings.catch_warnings():  # 31자를 넘는 시트 이름은 openpyxl이 경고만 하고 저장한다
        warnings.simplefilter("ignore")
        wb.active.title = sheets[0]
        for extra in sheets[1:]:
            wb.create_sheet(extra)
    ws = wb[sheets[-1]]
    for r in rows:
        ws.append(r)
    for m in merges:
        ws.merge_cells(m)
    if len(sheets) > 1:
        wb[sheets[0]].append(["이 시트는 작성 안내입니다"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _upload(c: TestClient, name: str, data: bytes, mime: str = XLSX) -> int:
    res = c.post("/files", files=[("files", (name, data, mime))])
    assert res.status_code == 200 and res.json()[0]["id"], res.text
    return res.json()[0]["id"]


def _csv(c: TestClient, name: str, text: str) -> int:
    return _upload(c, name, text.encode("utf-8-sig"), "text/csv")


def _apply(c: TestClient, fid: int, *cols: tuple[str, str | None]) -> Any:
    items = [
        {"file_id": fid, "sheet_id": "s0", "column_id": cid, "field_key": key, "expected_confirmed": False,
         "expected_field_key": None}
        for cid, key in cols
    ]  # fmt: skip
    return c.post("/mapping/batch/apply", json={"items": items})


def _put(c: TestClient, fid: int, column_id: str, key: str | None) -> Any:
    body = {"decisions": [{"sheet_id": "s0", "column_id": column_id, "field_key": key}]}
    return c.put(f"/files/{fid}/mapping/decisions", json=body)


def _decided(fid: int) -> list[tuple]:
    return db_rows(
        "select column_name, field_key from mapping_decisions where source_file_id = ? "
        "order by length(column_name), column_name",
        fid,
    )


def _count(table: str) -> int:
    return db_rows(f"select count(*) from {table}")[0][0]


def _counts() -> dict[str, int]:
    tables = ("validation_runs", "review_sessions", "review_items", "runs", "run_changes", "mapping_decisions",
              "learned_aliases", "mapping_approvals", "standard_terms", "glossary_aliases")  # fmt: skip
    return {t: _count(t) for t in tables}


def _over_limit() -> list[tuple[str, str, int]]:
    """varchar(255) 이름 열에서 255자를 넘는 값 (표·열·건수)."""
    out = []
    for table, col in [("mapping_decisions", "column_name"), ("mapping_decisions", "sheet_name"),
                       ("review_items", '"column"'), ("review_items", "sheet"), ("review_items", "file_name"),
                       ("run_changes", '"column"'), ("run_changes", "source_column"), ("run_changes", "sheet"),
                       ("run_changes", "file_name"), ("learned_aliases", "normalized"),
                       ("standard_terms", "normalized"), ("glossary_aliases", "normalized")]:  # fmt: skip
        n = db_rows(f"select count(*) from {table} where length({col}) > 255")[0][0]
        if n:
            out.append((table, col, n))
    return out


def _approve_pending(c: TestClient, review: dict) -> None:
    for i in review["items"]:
        if i["category"] == "pending" and i["decision"] is None:
            res = c.post(f"/reviews/{review['id']}/items/{i['id']}/decision", json={"decision": "approved"})
            assert res.status_code == 200, res.text


def _ready(c: TestClient, fids: list[int]) -> dict:
    """검토를 만들고 (확인 필요 항목은 승인해) 완료한다."""
    res = c.post("/reviews", json={"file_ids": fids})
    assert res.status_code == 201, res.text
    _approve_pending(c, res.json())
    res = c.post(f"/reviews/{res.json()['id']}/complete")
    assert res.status_code == 200, res.text
    return res.json()


def _no_echo(text: str, *raws: str) -> None:
    for raw in raws:
        assert raw not in text, (raw[:12], text[:300])
    assert "\\u0000" not in text and "\\u0001" not in text  # 제어 문자가 든 원문을 되돌려 보내지 않음


# ---------------------------------------------------------------- 확정 (한 열·열 번호·일괄·공통 검토 승인)


def test_names_up_to_the_limit_are_confirmed_and_longer_ones_refused_before_writing(client) -> None:
    fid = _upload(client, "길이.xlsx", _xlsx([["기관명", N254, N255, N256], ["행정안전부", 1, 2, 3]]))
    assert _apply(client, fid, ("c1", "note")).status_code == 200
    assert _apply(client, fid, ("c2", "population")).status_code == 200
    res = _apply(client, fid, ("c3", "amount"))
    assert res.status_code == 409, res.text
    problems = res.json()["detail"]["problems"]
    assert len(problems) == 1 and problems[0].startswith(f"파일 #{fid} · 1번째 시트 · D열(머리글 1행): "), problems
    assert "열 이름이 256자로 저장 한도 255자를 넘습니다" in problems[0] and "[NAME_TOO_LONG]" in problems[0]
    _no_echo(res.text, RAW, "길이.xlsx")  # 위치만 (이름 원문·가린 파일 이름도 넣지 않음)
    # 한글 255자(765바이트)도 두 DB에서 그대로 저장된다 — 한도는 글자 수
    assert len(N255.encode("utf-8")) == 765
    assert _decided(fid) == [(N254, "note"), (N255, "population")]


def test_one_long_column_keeps_the_whole_request_unsaved_including_not_applicable(client) -> None:
    fid = _upload(client, "섞임.xlsx", _xlsx([["기관명", "섞임 메모", N256], ["행정안전부", "a", 1]]))
    before = _counts()
    # 정상 열 + 긴 열('해당 없음' — 해당 없음도 확정 행을 만든다): 아무것도 저장하지 않는다
    res = _apply(client, fid, ("c1", "note"), ("c2", None))
    assert res.status_code == 409 and len(res.json()["detail"]["problems"]) == 1, res.text
    assert _apply(client, fid, ("c2", None)).status_code == 409
    assert _decided(fid) == [] and _counts() == before
    # 정상 열만 보내면 저장된다
    assert _apply(client, fid, ("c1", "note")).status_code == 200
    assert _decided(fid) == [("섞임 메모", "note")]


def test_preview_shows_the_long_column_excluded_with_its_place(client) -> None:
    fid = _upload(client, "미리보기.xlsx", _xlsx([["기관명", N256], ["행정안전부", 1]]))
    item = {"file_id": fid, "sheet_id": "s0", "column_id": "c1", "field_key": "note"}
    plan = client.post("/mapping/batch/preview", json={"items": [item]}).json()
    entry = plan["items"][0]
    assert plan["excluded"] == 1 and entry["action"] == "excluded", plan
    assert entry["column"] == "B열(머리글 1행)" and "저장 한도 255자" in entry["reason"]
    _no_echo(json.dumps(plan, ensure_ascii=False), RAW)
    assert _decided(fid) == []
    # 원본 미리보기는 그대로 열린다 (저장하지 않는 화면)
    assert client.get(f"/files/{fid}/preview").status_code == 200


def test_api_decisions_check_the_name_found_by_column_id_too(client) -> None:
    fid = _upload(client, "직접.xlsx", _xlsx([["기관명", "직접 메모", N256], ["행정안전부", "a", 1]]))
    url = f"/files/{fid}/mapping/decisions"
    res = _put(client, fid, "c2", "note")
    assert res.status_code == 422, res.text
    assert res.json()["detail"].startswith("확정하지 않았습니다 — 1번째 시트 · C열(머리글 1행): 열 이름이 256자로")
    assert "[NAME_TOO_LONG]" in res.text
    _no_echo(res.text, RAW, "현황")
    # 예전 방식(열 이름 글자)도 같은 거절
    res = client.put(url, json={"decisions": [{"sheet": "현황", "column": N256, "field_key": None}]})
    assert res.status_code == 422 and RAW not in res.text
    # 정상 열과 함께 보내도 아무것도 저장하지 않는다
    both = [{"sheet_id": "s0", "column_id": "c1", "field_key": "note"}, {"sheet_id": "s0", "column_id": "c2",
            "field_key": None}]  # fmt: skip
    assert client.put(url, json={"decisions": both}).status_code == 422
    assert _decided(fid) == []


def test_two_level_header_joined_past_the_limit(client) -> None:
    def merged(group: int) -> bytes:  # 그룹 머리글(가로 병합) + 하위 머리글 → '그룹_하위'
        rows = [["기관명", "가" * group, None], [None, "인구", "세대"], ["행정안전부", 1, 2]]
        return _xlsx(rows, ("A1:A2", "B1:C1"))

    ok = _upload(client, "병합252.xlsx", merged(252))  # '그룹_인구' 255자
    assert _apply(client, ok, ("c1", "population"), ("c2", None)).status_code == 200
    assert [len(n) for n, _ in _decided(ok)] == [255, 255]
    long = _upload(client, "병합253.xlsx", merged(253))  # 256자
    res = _apply(client, long, ("c1", "population"))
    assert res.status_code == 409, res.text
    assert "B열(머리글 1~2행)" in res.text and "256자로 저장 한도 255자" in res.text
    assert _decided(long) == []


def test_merged_duplicate_column_is_shown_at_its_original_place(client) -> None:
    # B1:D1 병합 머리글 — C열은 B열과 값이 같아 복제로 빠지고, D열은 '<254자>_3'(256자)이 된다 (원본 D열)
    rows = [["기관명", "가" * 254, None, None], ["행정안전부", 1, 1, 2]]
    fid = _upload(client, "복제열.xlsx", _xlsx(rows, ("B1:D1",)))
    res = _apply(client, fid, ("c2", "note"))
    assert res.status_code == 409 and "D열(머리글 1행)" in res.text, res.text
    res = _put(client, fid, "c2", "note")
    assert res.status_code == 422 and "D열(머리글 1행)" in res.text, res.text
    res = client.post("/reviews", json={"file_ids": [fid]})
    assert res.status_code == 422 and f"파일 #{fid} · 1번째 시트 · D열(머리글 1행)" in res.text, res.text
    assert _decided(fid) == []


def test_sheet_names_up_to_the_limit_and_over(client) -> None:
    ok = _upload(client, "시트255.xlsx", _xlsx([["기관명", "인구수"], ["행정안전부", 1]], sheets=("시" * 255,)))
    assert _apply(client, ok, ("c0", "org_name")).status_code == 200
    assert db_rows("select length(sheet_name) from mapping_decisions where source_file_id = ?", ok) == [(255,)]
    fid = _upload(client, "시트256.xlsx", _xlsx([["기관명", "인구수"], ["행정안전부", 1]], sheets=("시" * 256,)))
    item = {"file_id": fid, "sheet_id": "s0", "column_id": "c0", "field_key": "org_name"}
    plan = client.post("/mapping/batch/preview", json={"items": [item]}).json()
    assert plan["items"][0]["sheet"] == "1번째 시트" and plan["items"][0]["action"] == "excluded", plan
    res = _apply(client, fid, ("c0", "org_name"))
    assert res.status_code == 409, res.text
    problem = res.json()["detail"]["problems"][0]
    assert problem.startswith(f"파일 #{fid} · 1번째 시트 · A열(머리글 1행): 시트 이름이 256자로"), problem
    res = _put(client, fid, "c0", "org_name")
    assert res.status_code == 422 and "1번째 시트 · A열(머리글 1행): 시트 이름이 256자로" in res.text, res.text
    res = client.post("/reviews", json={"file_ids": [fid]})
    assert res.status_code == 422 and f"파일 #{fid} · 1번째 시트: 시트 이름이 256자로" in res.text, res.text
    for text in (json.dumps(plan, ensure_ascii=False), res.text):
        _no_echo(text, "시" * 20)
    assert _decided(fid) == []


def test_sheet_place_counts_the_skipped_guide_sheet(client) -> None:
    # 파서는 '작성 안내' 시트를 건너뛰지만, 위치는 원본 통합 문서의 탭 순번으로 알린다 (두 번째 탭)
    fid = _upload(client, "안내시트.xlsx", _xlsx([["기관명", N256], ["행정안전부", 1]], sheets=("작성 안내", "현황")))
    res = client.post("/reviews", json={"file_ids": [fid]})
    assert res.status_code == 422 and f"파일 #{fid} · 2번째 시트 · B열(머리글 1행)" in res.text, res.text
    assert "2번째 시트 · B열" in _apply(client, fid, ("c1", "note")).text


def test_sheet_place_counts_a_chart_sheet_in_front(client) -> None:
    from openpyxl.chart import BarChart, Reference

    wb = openpyxl.Workbook()
    data = wb.active
    data.title = "현황"
    data.append(["기관명", N256])
    data.append(["행정안전부", 1])
    chart = BarChart()
    chart.add_data(Reference(data, min_col=2, min_row=1, max_row=2), titles_from_data=True)
    wb.create_chartsheet("차트", 0).add_chart(chart)  # 첫 탭이 차트 시트 — 읽는 워크시트는 둘째 탭
    buf = io.BytesIO()
    wb.save(buf)
    fid = _upload(client, "차트먼저.xlsx", buf.getvalue())
    res = client.post("/reviews", json={"file_ids": [fid]})
    assert res.status_code == 422 and f"파일 #{fid} · 2번째 시트 · B열(머리글 1행)" in res.text, res.text


def test_learning_skips_a_name_whose_key_grows_past_the_limit(client) -> None:
    name = "㍿" * 100  # 100자 — 정규화(NFKC) 키는 '株式会社' × 100 = 400자
    fid = _upload(client, "호환문자.xlsx", _xlsx([["기관명", name], ["행정안전부", "a"]]))
    approvals = _count("mapping_approvals")
    res = _put(client, fid, "c1", "note")
    assert res.status_code == 200 and res.json()["saved"] == 1, res.text
    assert res.json()["not_learned"][0]["reason"] == "열 이름이 길어 학습하지 않음 (저장 한도 255자)"
    assert _decided(fid) == [(name, "note")]  # 확정은 저장된다
    assert db_rows("select count(*) from learned_aliases where source_file_id = ?", fid) == [(0,)]
    assert _count("mapping_approvals") == approvals  # 승격 근거로도 남기지 않음
    # 일괄 적용 경로도 같다 (확정 저장, 학습 0)
    other = _upload(client, "호환문자2.xlsx", _xlsx([["기관명", name], ["행정안전부", "b"]]))
    res = _apply(client, other, ("c1", "note"))
    assert res.status_code == 200 and res.json()["learned"] == 0 and _decided(other) == [(name, "note")], res.text


def test_learning_key_boundary_255_is_learned_and_256_is_not(client) -> None:
    at, over = "㍿" + "가" * 251, "㍿" + "가" * 252  # 원래 252·253자 — 키 '株式会社…' 255·256자
    fid = _upload(client, "학습경계.xlsx", _xlsx([["기관명", at, over], ["행정안전부", 1, 2]]))
    res = client.put(
        f"/files/{fid}/mapping/decisions",
        json={
            "decisions": [
                {"sheet_id": "s0", "column_id": "c1", "field_key": "note"},
                {"sheet_id": "s0", "column_id": "c2", "field_key": "amount"},
            ]
        },  # fmt: skip
    )
    assert res.status_code == 200 and res.json()["saved"] == 2, res.text
    assert len(res.json()["learned"]) == 1 and len(res.json()["not_learned"]) == 1, res.json()
    rows = db_rows("select length(normalized), length(alias) from learned_aliases where source_file_id = ?", fid)
    assert rows == [(255, 252)]  # 키 255자는 두 DB에 저장된다


def test_nul_in_a_column_name_is_refused_the_same_on_both_databases(client) -> None:
    text = "기관명,인구\x00수,정상 메모\n행정안전부,1,a\n"  # NUL은 지우거나 바꾸지 않는다 — 이름을 고쳐 올리게 한다
    fid = _csv(client, "널머리글.csv", text)
    before = _counts()
    res = _put(client, fid, "c1", "population")
    assert res.status_code == 422 and "1번째 시트 · B열(머리글 1행): 열 이름에 NUL(0x00) 문자가 있습니다" in res.text
    assert "[NAME_HAS_NUL]" in res.text
    res = _apply(client, fid, ("c1", "population"), ("c2", "note"))  # 정상 열과 섞여도 아무것도 저장하지 않는다
    assert res.status_code == 409 and "[NAME_HAS_NUL]" in res.text, res.text
    item = {"file_id": fid, "sheet_id": "s0", "column_id": "c1", "field_key": "population"}
    plan = client.post("/mapping/batch/preview", json={"items": [item]}).json()
    assert plan["items"][0]["action"] == "excluded" and plan["items"][0]["column"] == "B열(머리글 1행)", plan
    res = client.post("/reviews", json={"file_ids": [fid]})
    assert res.status_code == 422 and f"파일 #{fid} · 1번째 시트 · B열(머리글 1행)" in res.text, res.text
    for text in (res.text, json.dumps(plan, ensure_ascii=False)):
        _no_echo(text)
    assert _decided(fid) == [] and _counts() == before  # 검증 기록도 남기지 않음 (PostgreSQL에서 500·고아 기록)


def test_tabs_newlines_and_hangul_are_not_refused(client) -> None:
    fid = _csv(client, "탭줄바꿈.csv", '기관명,"비고\t메모"\n행정안전부,"줄\n바꿈\t탭"\n')
    assert _apply(client, fid, ("c0", "org_name"), ("c1", "note")).status_code == 200
    review = _ready(client, [fid])
    res = client.post(f"/reviews/{review['id']}/integrate", json={})
    assert res.status_code == 201 and res.json()["status"] == "done", res.text


# ---------------------------------------------------------------- 검토·다시 검사·자동화


def test_review_is_not_created_for_long_names_and_leaves_nothing(client) -> None:
    long = _upload(client, "검토긴이름.xlsx", _xlsx([["기관명", N256], ["행정안전부", 1]]))
    before = _counts()
    res = client.post("/reviews", json={"file_ids": [long]})
    assert res.status_code == 422, res.text
    detail = res.json()["detail"]
    assert detail.startswith("이름을 저장할 수 없어 검토를 만들지 않았습니다 — ")
    assert f"파일 #{long} · 1번째 시트 · B열(머리글 1행): 열 이름이 256자로 저장 한도 255자를 넘습니다" in detail
    assert detail.endswith("[NAME_TOO_LONG]")
    _no_echo(res.text, RAW, "검토긴이름")
    assert _counts() == before  # 검증 기록·검토·항목을 남기지 않는다


def test_review_is_created_at_255_characters(client) -> None:
    fid = _upload(client, "검토255.xlsx", _xlsx([["기관명", N255], ["행정안전부", 1]]))
    res = client.post("/reviews", json={"file_ids": [fid]})
    assert res.status_code == 201, res.text
    assert db_rows('select max(length("column")) from review_items where session_id = ?', res.json()["id"]) == [(255,)]


def test_long_names_only_in_the_baseline_do_not_block_the_review(client) -> None:
    # 전월 비교 자료의 이름은 값 비교에만 쓰고 저장하지 않는다 — 두 DB 모두 저장할 수 있는 검토를 막지 않는다
    now = _upload(client, "이번달.xlsx", _xlsx([["기관명", "인구수"], ["행정안전부", 10]]))
    last = _upload(client, "지난달.xlsx", _xlsx([["기관명", "인구수", N256], ["행정안전부", 9, 1]]))
    assert _apply(client, now, ("c0", "org_name"), ("c1", "population")).status_code == 200
    res = client.post("/reviews", json={"file_ids": [now], "baseline_file_ids": [last]})
    assert res.status_code == 201, res.text
    # 한도를 넘는 값은 이 검토·두 파일 어디에도 저장되지 않았다 (공용 시험 DB의 다른 시험 자료는 보지 않음)
    longest = db_rows('select max(length("column")), max(length(sheet)), max(length(file_name)) from review_items '
                      "where session_id = ?", res.json()["id"])[0]  # fmt: skip
    assert all(n is None or n <= 255 for n in longest), longest
    assert all(len(c) <= 255 for fid in (now, last) for c, _ in _decided(fid))
    names = {i["column"] for i in res.json()["items"] if i["column"]}
    assert N256 not in names and all(len(n) <= 255 for n in names)


def test_recheck_keeps_the_old_review_when_the_new_file_has_long_names(client) -> None:
    ok = _upload(client, "다시정상.xlsx", _xlsx([["기관명", "인구수"], ["행정안전부", 1]]))
    long = _upload(client, "다시긴이름.xlsx", _xlsx([["기관명", "인구수", N256], ["행정안전부", 2, 3]]))
    assert _apply(client, ok, ("c0", "org_name"), ("c1", "population")).status_code == 200
    review = client.post("/reviews", json={"file_ids": [ok]}).json()
    before = _counts()
    res = client.post(f"/reviews/{review['id']}/recheck", json={"replace": {str(ok): long}})
    assert res.status_code == 422 and "검토를 만들지 않았습니다" in res.text and RAW not in res.text, res.text
    # 이전 파일의 확정을 새 파일로 옮긴 것도 남지 않는다 (확정 복사는 검토 생성과 한 트랜잭션)
    assert _decided(long) == [] and _counts() == before
    assert client.get(f"/reviews/{review['id']}").json()["status"] == "in_review"
    # 정상 파일로 바꾸면 확정을 이어받아 다시 검사한다
    fixed = _upload(client, "다시고침.xlsx", _xlsx([["기관명", "인구수"], ["행정안전부", 4]]))
    res = client.post(f"/reviews/{review['id']}/recheck", json={"replace": {str(ok): fixed}})
    assert res.status_code == 201, res.text
    assert _decided(fixed) == _decided(ok) == [("기관명", "org_name"), ("인구수", "population")]
    assert client.get(f"/reviews/{review['id']}").json()["status"] == "superseded"


def test_automation_run_on_long_names_records_nothing(client) -> None:
    fid = _upload(client, "자동화원본.xlsx", _xlsx([["기관명", "인구수"], ["행정안전부", 10]]))
    assert _apply(client, fid, ("c0", "org_name"), ("c1", "population")).status_code == 200
    review = _ready(client, [fid])
    run = client.post(f"/reviews/{review['id']}/integrate", json={})
    assert run.status_code == 201, run.text
    saved = client.post(f"/runs/{run.json()['id']}/automation", json={"name": f"길이 한도 자동화 {fid}"})
    assert saved.status_code == 201, saved.text
    long = _upload(client, "자동화긴이름.xlsx", _xlsx([["기관명", "인구수", N256], ["행정안전부", 11, 1]]))
    before = _counts()
    res = client.post(f"/automations/{saved.json()['id']}/run", json={"file_ids": [long]})
    assert res.status_code == 422 and "검토를 만들지 않았습니다" in res.text and RAW not in res.text, res.text
    assert _counts() == before  # 규칙 기록(확정)·검토·실행 기록을 만들지 않는다
    # 저장할 값에 NUL이 있어도 같다 — 자동화 규칙 확정을 먼저 쓰지 않는다 (검사를 마친 뒤 한 번에 기록)
    nul = _csv(client, "자동화널값.csv", "기관명,인구수\n행정안전부,  1\x00  \n")
    res = client.post(f"/automations/{saved.json()['id']}/run", json={"file_ids": [nul]})
    assert res.status_code == 422 and "[VALUE_HAS_NUL]" in res.text, res.text
    assert _decided(nul) == [] and _counts() == before


def test_automation_rules_do_not_overwrite_a_decision_made_while_the_review_is_checked(client, monkeypatch) -> None:
    """자동화 규칙 확정은 모든 검사를 마친 뒤 검토와 함께 쓴다 — 그사이(검사 중) 같은 열을 사용자가 확정했으면 그 확정을
    덮지 않는다 (규칙은 확정이 없던 열에만). 검사 기준은 규칙 값이라 '다시 검사'로 드러난다 (이름 규칙 전과 같은
    결과)."""
    from app import review as review_module
    from app.db import SessionLocal
    from app.decisions import DecisionInput, save_decisions
    from app.mapping import MappingEngine, load_dictionary
    from app.models import ReviewSession

    header = ["기관명", "담당 업무 구분", "인구수"]
    sept = _upload(client, "규칙경합9월.xlsx", _xlsx([header, ["행정안전부", "재정", 5]]))
    review = client.post("/reviews", json={"file_ids": [sept]}).json()
    for item in [i for i in review["items"] if i["category"] == "pending" and i["decision"] is None]:
        choice = item["payload"].get("field_key") if item["kind"] == "mapping" else None
        body = {"decision": "approved", "choice": choice} if choice else {"decision": "rejected"}
        assert client.post(f"/reviews/{review['id']}/items/{item['id']}/decision", json=body).status_code == 200
    assert client.post(f"/reviews/{review['id']}/complete").status_code == 200
    run = client.post(f"/reviews/{review['id']}/integrate", json={}).json()
    auto = client.post(f"/runs/{run['id']}/automation", json={"name": f"규칙 경합 {sept}"})
    assert auto.status_code == 201, auto.text
    rule = {r["column"]: r["field_key"] for r in auto.json()["column_rules"]}["담당 업무 구분"]
    oct_ = _upload(client, "규칙경합10월.xlsx", _xlsx([header, ["행정안전부", "총무", 7]]))
    user_key = "note" if rule != "note" else "remark"
    real = review_module.validate

    def validate_while_the_user_decides(*args: Any, **kwargs: Any) -> Any:
        with SessionLocal() as other:  # 다른 요청 (같은 사용자의 다른 창)이 검사 도중 확정을 커밋한다
            save_decisions(other, MappingEngine(load_dictionary(), None, None), oct_,
                           [DecisionInput("현황", "담당 업무 구분", user_key)], learn=False)  # fmt: skip
        return real(*args, **kwargs)

    monkeypatch.setattr(review_module, "validate", validate_while_the_user_decides)
    res = client.post(f"/automations/{auto.json()['id']}/run", json={"file_ids": [oct_]})
    monkeypatch.undo()
    assert res.status_code in (200, 201), res.text
    assert ("담당 업무 구분", user_key) in _decided(oct_)  # 사용자 확정이 남는다
    with SessionLocal() as db:
        reasons = review_module.recheck_reasons(db, db.get(ReviewSession, res.json()["review_id"]))
    assert reasons, "검사 기준(규칙 값)과 지금 확정이 달라 다시 검사로 드러나야 함"


def test_entity_files_listed_across_many_files_fit_the_limit(client) -> None:
    names = [f"시도별_통계자료_현황표_{i:02d}_2026년9월분.xlsx" for i in range(12)]  # 이름 약 30자 × 12개
    fids = [_upload(client, n, _xlsx([["기관명", "인구수"], ["가나다라연구기관", i]])) for i, n in enumerate(names)]
    res = client.post("/reviews", json={"file_ids": fids})
    assert res.status_code == 201, res.text  # PostgreSQL에서 이어 붙인 파일 이름(255자 초과)으로 500이 나던 경로
    shown = [i["file_name"] for i in res.json()["items"] if i["kind"] == "entity" and i["before"] == "가나다라연구기관"]
    assert shown and all(len(s) <= 255 for s in shown), shown
    assert shown[0].startswith(names[0]) and shown[0].endswith("개") and " 외 " in shown[0]


def test_automation_entity_rule_item_lists_many_files_within_the_limit(client) -> None:
    # 자동화의 이름 규칙으로 처리한 항목(category auto)도 파일 이름 목록을 한도 안으로 요약한다
    first = _upload(client, "규칙원본.xlsx", _xlsx([["기관명", "인구수"], ["마바사아연구기관", 1]]))
    assert _apply(client, first, ("c0", "org_name"), ("c1", "population")).status_code == 200
    review = client.post("/reviews", json={"file_ids": [first]}).json()
    entity = next(i for i in review["items"] if i["kind"] == "entity" and i["category"] == "pending")
    res = client.post(f"/reviews/{review['id']}/items/{entity['id']}/decision", json={"decision": "rejected"})
    assert res.status_code == 200, res.text
    _approve_pending(client, client.get(f"/reviews/{review['id']}").json())
    assert client.post(f"/reviews/{review['id']}/complete").status_code == 200
    run = client.post(f"/reviews/{review['id']}/integrate", json={})
    assert run.status_code == 201, run.text
    saved = client.post(f"/runs/{run.json()['id']}/automation", json={"name": f"규칙 파일 목록 {first}"})
    assert saved.status_code == 201, saved.text
    names = [f"규칙적용_통계자료_현황표_{i:02d}_2026년9월분.xlsx" for i in range(12)]
    fids = [_upload(client, n, _xlsx([["기관명", "인구수"], ["마바사아연구기관", i]])) for i, n in enumerate(names)]
    res = client.post(f"/automations/{saved.json()['id']}/run", json={"file_ids": fids})
    assert res.status_code in (200, 201), res.text
    session_id = res.json()["review_id"]
    lengths = db_rows("select length(file_name) from review_items where session_id = ? and kind = 'entity'", session_id)
    assert lengths and all(n <= 255 for (n,) in lengths), lengths


def test_nul_in_a_stored_cell_value_refuses_the_review_with_its_place(client) -> None:
    # 앞뒤 공백 정리 항목은 원래 칸 값을 검토 항목에 저장한다 — NUL이 있으면 아무것도 쓰기 전에 위치만 알려 거절
    fid = _csv(client, "널칸값.csv", "기관명,비고\n행정안전부,  값\x00끝  \n")
    assert _apply(client, fid, ("c0", "org_name"), ("c1", "note")).status_code == 200
    before = _counts()
    res = client.post("/reviews", json={"file_ids": [fid]})
    assert res.status_code == 422, res.text
    assert f"파일 #{fid} · 1번째 시트 · 2행 · B열" in res.text and "[VALUE_HAS_NUL]" in res.text
    _no_echo(res.text)
    assert _counts() == before  # 검증 기록도 남기지 않음 (PostgreSQL에서 500·고아 기록이던 경로)


def test_nul_in_an_entity_name_across_files_is_refused_at_its_cell_and_a_condition_is_named(client) -> None:
    # 여러 파일을 묶은 기관 이름 확인 항목은 파일 번호가 없다 — 그 값이 든 칸을 찾아 알린다. 필수 기관(사용자가 넣은
    # 검토 조건)은 원본이 아니라 조건을 고치라고 알린다
    a = _csv(client, "널기관정상.csv", "기관명,인구수\n행정안전부,1\n")
    b = _csv(client, "널기관값.csv", "기관명,인구수\n가나다\x00연구원,1\n")
    for f in (a, b):
        assert _apply(client, f, ("c0", "org_name"), ("c1", "population")).status_code == 200
    before = _counts()
    res = client.post("/reviews", json={"file_ids": [a, b]})
    assert res.status_code == 422, res.text
    assert f"파일 #{b} · 1번째 시트 · 2행 · A열" in res.text and "원본 파일에서 그 칸의 NUL" in res.text, res.text
    assert "검토 항목." not in res.text and "[VALUE_HAS_NUL]" in res.text
    _no_echo(res.text, "가나다")
    res = client.post("/reviews", json={"file_ids": [a], "required_entities": ["없는\x00기관"]})
    assert res.status_code == 422 and "검토 조건 '필수 기관'" in res.text, res.text
    assert "필수 기관 조건에서 NUL 문자를 지운 뒤" in res.text and "원본 파일에서" not in res.text, res.text
    _no_echo(res.text, "없는")
    assert _counts() == before


def test_nul_entity_names_are_looked_up_in_one_pass_over_the_cells(client, monkeypatch) -> None:
    # 파일 번호 없는 항목이 많아도 칸은 한 번만 훑는다 (항목마다 모든 칸을 훑으면 항목 수 × 칸 수 — 검토 지적 R1)
    from app import review as review_module

    calls: list[int] = []
    real = review_module._cells_holding

    def counted(targets: Any, values: set[str]) -> list[str]:
        calls.append(len(values))
        return real(targets, values)

    monkeypatch.setattr(review_module, "_cells_holding", counted)
    rows = "".join(f"가나다\x00연구원{k},{k}\n" for k in range(40))
    fid = _csv(client, "널기관여럿.csv", "기관명,인구수\n" + rows)
    assert _apply(client, fid, ("c0", "org_name"), ("c1", "population")).status_code == 200
    res = client.post("/reviews", json={"file_ids": [fid]})
    assert res.status_code == 422, res.text
    assert calls == [40], calls  # 한 번 — 값 40개를 한꺼번에
    assert f"파일 #{fid} · 1번째 시트 · 2행 · A열" in res.text and " 외 37곳" in res.text, res.text


@pytest.mark.sqlite_only("PostgreSQL은 NUL이 든 자동화 이름을 저장하지 않음 (422 — 직접 입력 글자의 안전망)")
def test_nul_from_an_automation_name_is_not_blamed_on_a_cell(client) -> None:
    # 자동화 이름의 NUL은 규칙 항목의 사유에 들어간다 — 같은 값의 칸(NUL 없음)을 지목하지 않고 설정을 고치라고 알린다
    header = "기관명,인구수\n"
    sept = _csv(client, "규칙이름9월.csv", header + "가나다라연구소,5\n")
    assert _apply(client, sept, ("c0", "org_name"), ("c1", "population")).status_code == 200
    review = client.post("/reviews", json={"file_ids": [sept]}).json()
    for item in [i for i in review["items"] if i["category"] == "pending" and i["decision"] is None]:
        body = {"decision": "rejected"}
        assert client.post(f"/reviews/{review['id']}/items/{item['id']}/decision", json=body).status_code == 200
    assert client.post(f"/reviews/{review['id']}/complete").status_code == 200
    run = client.post(f"/reviews/{review['id']}/integrate", json={}).json()
    auto = client.post(f"/runs/{run['id']}/automation", json={"name": "규칙\x00이름"})
    assert auto.status_code == 201, auto.text  # SQLite는 직접 입력 글자의 NUL을 저장한다 (남은 두 DB 차이)
    oct_ = _csv(client, "규칙이름10월.csv", header + "가나다라연구소,7\n")
    res = client.post(f"/automations/{auto.json()['id']}/run", json={"file_ids": [oct_]})
    assert res.status_code == 422, res.text
    detail = res.json()["detail"]
    assert "자동화 이름 (원본 칸이 아니라 규칙 항목의 사유에 들어가는 설정)" in detail, detail
    assert "자동화 이름에서 NUL 문자를 지운 뒤 다시 실행하세요" in detail and "원본 파일에서" not in detail, detail
    assert f"파일 #{oct_}" not in detail, detail  # NUL이 없는 칸을 지목하지 않음


def test_nul_quoted_into_a_warning_reason_is_pointed_at_its_own_cell(client) -> None:
    # 숫자로 읽지 못한 값 경고는 첫 값의 행을 가리키고 사유에 값을 3개까지 옮겨 적는다 — NUL이 둘째 값에 있으면 경고의
    # 행(첫 값의 칸)이 아니라 그 값의 칸을 알리고, 설정이 아니라 원본 파일을 고치라고 안내한다
    body = "기관명,인구수\n가나다연구원,abc\n라마바연구원,x{c}y\n사아자연구원,5\n"
    probe = _csv(client, "사유제어.csv", body.format(c="\x01"))  # NUL 대신 \x01 — 검토가 만들어져 항목 모양을 본다
    assert _apply(client, probe, ("c0", "org_name"), ("c1", "population")).status_code == 200
    items = client.post("/reviews", json={"file_ids": [probe]}).json()["items"]
    assert not any("\x01" in (i["before"] or "") + (i["after"] or "") for i in items), items  # 이전·이후 값에는 없음
    assert any(i["category"] == "warning" and i["row"] == 2 and "\x01" in (i["reason"] or "") for i in items), items
    fid = _csv(client, "사유널.csv", body.format(c="\x00"))
    assert _apply(client, fid, ("c0", "org_name"), ("c1", "population")).status_code == 200
    res = client.post("/reviews", json={"file_ids": [fid]})
    assert res.status_code == 422, res.text
    detail = res.json()["detail"]
    assert f"파일 #{fid} · 1번째 시트 · 3행 · B열" in detail, detail
    assert "2행" not in detail, detail
    assert "원본 파일에서 그 칸의 NUL 문자를 지운 뒤" in detail and "자동화" not in detail, detail


def test_nul_in_a_cell_that_is_only_read_is_refused_only_where_it_is_written(client) -> None:
    # 저장하지 않는 칸 값(검토 항목에 들어가지 않음)은 검토를 막지 않고, 결과 파일에 쓰는 통합에서 위치만 알려 거절한다
    fid = _csv(client, "널읽기.csv", "기관명,비고\n행정안전부,값\x00끝\n")
    assert _apply(client, fid, ("c0", "org_name"), ("c1", "note")).status_code == 200
    review = _ready(client, [fid])
    before = _counts()
    res = client.post(f"/reviews/{review['id']}/integrate", json={})
    assert res.status_code == 409 and f"파일 #{fid} · 1번째 시트 · 2행 · B열" in res.text, res.text
    assert "[CELL_CONTROL_CHAR]" in res.text
    _no_echo(res.text)
    assert _counts() == before


# ---------------------------------------------------------------- 검토 완료 학습


def test_review_completion_skips_learning_an_unstorable_name_with_its_own_reason(client) -> None:
    csv = "코드,정식명칭,유형\n9820001," + "다" * 255 + ",기관\n"
    assert client.post("/glossary/import", files=[("file", ("g2.csv", csv.encode("utf-8"), "text/csv"))]).json()[
        "applied"
    ]
    value = "다" * 256  # 셀 값 — 용어사전 후보(255자)와 비슷하지만 저장 한도를 넘어 별칭으로 학습할 수 없다
    fid = _upload(client, "별칭학습.xlsx", _xlsx([["기관명", "인구수"], [value, 1]]))
    assert _apply(client, fid, ("c0", "org_name"), ("c1", "population")).status_code == 200
    review = client.post("/reviews", json={"file_ids": [fid]}).json()
    target = next(i for i in review["items"] if i["kind"] == "entity" and i["before"] == value)
    assert target["category"] == "pending", target
    res = client.post(f"/reviews/{review['id']}/items/{target['id']}/decision",
                      json={"decision": "approved", "choice": "9820001"})  # fmt: skip
    assert res.status_code == 200, res.text
    _approve_pending(client, res.json())
    assert client.patch(f"/reviews/{review['id']}", json={"remember_format": True}).status_code == 200
    done = client.post(f"/reviews/{review['id']}/complete")
    assert done.status_code == 200, done.text
    learned = done.json()["learned"]
    assert learned["names"] == 0 and learned["names_not_storable"] == 1, learned
    assert learned["names_skipped"] == 0 and learned["names_pii"] == 0  # 권한·개인정보 때문이 아님
    assert db_rows("select count(*) from glossary_aliases where target_code = '9820001'") == [(0,)]
    db_rows("delete from glossary_entries where code = '9820001'")


# ---------------------------------------------------------------- 통합


def test_integration_allows_a_long_code_column_that_is_not_recorded(client) -> None:
    # 기관명 열 두 개(두 번째 248자): 결과 열 '기관명 [<248자>]'(254자)은 변경 기록에 저장되고, '… 코드'(257자)는
    # 결과 파일에만 쓴다 — 저장하지 않는 이름으로 막지 않는다 (두 DB 모두 저장)
    second = _name(248, "기관명원문")
    fid = _upload(client, "코드열.xlsx", _xlsx([["기관명", second, "인구수"], ["행정안전부", "행정안전부", 1]]))
    assert _apply(client, fid, ("c0", "org_name"), ("c1", "org_name"), ("c2", "population")).status_code == 200
    review = _ready(client, [fid])
    res = client.post(f"/reviews/{review['id']}/integrate", json={})
    assert res.status_code == 201 and res.json()["status"] == "done", res.text
    run_id = res.json()["id"]
    assert db_rows('select max(length("column")) from run_changes where run_id = ?', run_id) == [(254,)]
    wb = openpyxl.load_workbook(io.BytesIO(client.get(f"/runs/{run_id}/download").content))
    labels = [c.value for c in next(wb["통합 데이터"].iter_rows(max_row=1))]
    assert f"기관명 [{second}] 코드" in labels  # 결과 파일의 열 이름은 그대로 (257자)


def test_integration_refuses_a_recorded_result_column_name_over_the_limit_before_a_run(client) -> None:
    second = _name(251, "비고원문")  # 같은 표준 항목의 두 번째 열 → 결과 열 이름 '비고 [<251자>]' = 256자
    fid = _upload(client, "통합긴열.xlsx", _xlsx([["기관명", "메모", second], ["행정안전부", "a", "b"]]))
    assert _apply(client, fid, ("c0", "org_name"), ("c1", "note"), ("c2", "note")).status_code == 200
    review = _ready(client, [fid])
    before = _counts()
    res = client.post(f"/reviews/{review['id']}/integrate", json={})
    assert res.status_code == 409, res.text
    detail = res.json()["detail"]
    assert detail.startswith("이름을 저장할 수 없어 통합하지 않았습니다 — ")
    # 그 결과 열 이름을 만든 원본 위치(파일 번호·시트 순번·열)와 고칠 방법
    assert f"파일 #{fid} · 1번째 시트 · C열(머리글 1행): 결과 열 이름이 256자로 저장 한도 255자를 넘습니다" in detail
    assert "[NAME_TOO_LONG]" in detail and "다른 항목으로 다시 확정" in detail
    _no_echo(res.text, second[4:24], "INSERT")
    assert _counts() == before  # 실행 기록·변경 기록을 만들기 전에 거절


def test_integration_records_a_result_column_name_of_255_characters(client) -> None:
    second = _name(250, "비고원문")  # '비고 [<250자>]' = 255자
    fid = _upload(client, "통합255.xlsx", _xlsx([["기관명", "메모", second], ["행정안전부", "a", "b"]]))
    assert _apply(client, fid, ("c0", "org_name"), ("c1", "note"), ("c2", "note")).status_code == 200
    review = _ready(client, [fid])
    res = client.post(f"/reviews/{review['id']}/integrate", json={})
    assert res.status_code == 201 and res.json()["status"] == "done", res.text
    assert db_rows('select max(length("column")) from run_changes where run_id = ?', res.json()["id"]) == [(255,)]


def test_integration_refuses_control_characters_before_a_run_without_echoing_the_cell(client, caplog) -> None:
    fid = _csv(client, "제어문자.csv", f"기관명,비고\n행정안전부,담당 {PHONE}\x01메모\n")
    assert _apply(client, fid, ("c0", "org_name"), ("c1", "note")).status_code == 200
    review = _ready(client, [fid])  # 저장(DB)에는 문제가 없는 글자 — 결과 파일(엑셀)에만 쓸 수 없다
    before = _counts()
    with caplog.at_level(logging.INFO):
        res = client.post(f"/reviews/{review['id']}/integrate", json={})
    assert res.status_code == 409, res.text
    assert f"파일 #{fid} · 1번째 시트 · 2행 · B열" in res.text and "[CELL_CONTROL_CHAR]" in res.text
    _no_echo(res.text, PHONE, "담당", "제어문자")
    assert PHONE not in caplog.text and "\x01" not in caplog.text
    assert _counts() == before  # 실행 기록(runs)을 만들기 전에 거절 — 실행 기록·목록에도 원문이 남지 않는다
    assert PHONE not in client.get(f"/runs?review_id={review['id']}").text
    # 머리글(열 이름)에 든 제어 문자도 위치만
    other = _csv(client, "제어머리글.csv", "기관명,담당\x01메모\n행정안전부,a\n")
    assert _apply(client, other, ("c0", "org_name"), ("c1", "note")).status_code == 200
    res = client.post(f"/reviews/{_ready(client, [other])['id']}/integrate", json={})
    assert res.status_code == 409 and f"파일 #{other} · 1번째 시트 · B열(머리글 1행): 열 이름" in res.text, res.text
    _no_echo(res.text, "담당")


def test_control_character_from_a_glossary_name_is_reported_as_the_glossary_entry(client) -> None:
    # 칸 값이 용어사전 표시 이름으로 바뀌어 결과에 들어간다 — 그 이름에 제어 문자가 있으면 원본 칸이 아니라 용어사전
    # 항목을 고치라고 알린다 (원본을 고쳐도 남으므로). 저장은 두 DB 모두 되는 글자라 가져오기는 막지 않는다
    glossary = "코드,정식명칭,유형,별칭\n9830001,제어\x01기관,기관,제어기관약칭\n"
    res = client.post("/glossary/import", files=[("file", ("제어사전.csv", glossary.encode("utf-8"), "text/csv"))])
    assert res.json()["applied"] is True, res.text
    try:
        fid = _csv(client, "제어사전값.csv", "기관명,인구수\n제어기관약칭,3\n")
        assert _apply(client, fid, ("c0", "org_name"), ("c1", "population")).status_code == 200
        review = _ready(client, [fid])
        before = _counts()
        res = client.post(f"/reviews/{review['id']}/integrate", json={})
        assert res.status_code == 409 and "[CELL_CONTROL_CHAR]" in res.text, res.text
        detail = res.json()["detail"]
        assert f"용어사전 9830001 항목의 이름 (파일 #{fid} · 1번째 시트 · 2행 · A열에 쓰려던 값)" in detail, detail
        # 용어사전을 고치면 검사 기준(용어사전)이 바뀌므로 다시 검사를 거쳐야 통합된다 — 그 단계까지 안내
        assert "용어사전에서 고친 뒤 검토에서 다시 검사하고 통합하세요" in detail, detail
        assert "원본 파일에서" not in detail, detail
        _no_echo(res.text, "제어기관")
        assert _counts() == before
    finally:
        db_rows("delete from glossary_aliases where target_code = '9830001'")
        db_rows("delete from glossary_entries where code = '9830001'")


def test_control_character_in_a_warning_reason_names_the_row_not_another_column(client) -> None:
    # 합계 검사 불완전 경고의 사유에는 합계 행의 이름 칸 글자가 들어간다 — 경고가 가리킨 열(인구수)이 아니라 그 행을
    # 알린다
    fid = _csv(client, "경고제어.csv", "구분,인구수\n합\x1f계,9\n가,-\n나,5\n")
    assert _apply(client, fid, ("c0", None), ("c1", "population")).status_code == 200
    review = client.post("/reviews", json={"file_ids": [fid]}).json()
    assert any(i["category"] == "warning" and "\x1f" in (i["reason"] or "") for i in review["items"]), review["items"]
    _approve_pending(client, review)
    assert client.post(f"/reviews/{review['id']}/complete").status_code == 200
    res = client.post(f"/reviews/{review['id']}/integrate", json={})
    assert res.status_code == 409, res.text
    detail = res.json()["detail"]
    assert f"파일 #{fid} · 1번째 시트 · 2행 (검증 경고 사유 — 그 행의 칸에서 온 글자)" in detail, detail
    assert "B열" not in detail, detail


def test_control_character_in_a_column_name_is_not_blamed_on_a_warning_row(client) -> None:
    # 열 이름의 제어 문자는 그 열 이름의 자리로 알린다 — 그 열을 가리키는 경고가 있어도 경고의 행을 지목하지 않는다
    fid = _csv(client, "열이름경고.csv", "구분,인\x1f구수\n가,5\n나,abc\n")
    assert _apply(client, fid, ("c0", None), ("c1", "population")).status_code == 200
    review = client.post("/reviews", json={"file_ids": [fid]}).json()
    # 경고 사유에 그 열 이름(제어 문자 포함)이 들어간다
    assert any(i["category"] == "warning" and "\x1f" in (i["reason"] or "") for i in review["items"]), review["items"]
    _approve_pending(client, review)
    assert client.post(f"/reviews/{review['id']}/complete").status_code == 200
    res = client.post(f"/reviews/{review['id']}/integrate", json={})
    assert res.status_code == 409, res.text
    detail = res.json()["detail"]
    assert f"파일 #{fid} · 1번째 시트 · B열(머리글 1행): 열 이름" in detail, detail
    assert "검증 경고 사유" not in detail, detail


def test_control_character_in_a_warning_row_is_still_named_beside_a_bad_column_name(client) -> None:
    # 열 이름을 뺀 뒤에도 사유에 남은 글자(합계 행의 이름 칸)는 그 행으로 알린다 — 두 곳 모두. 합계 행은 공백류 제어
    # 문자(\x1c~\x1f)를 공백으로 읽어야 합계로 알아보므로 행 이름에는 \x1e, 열 이름에는 \x1f를 쓴다
    fid = _csv(client, "둘다제어.csv", "구분,인\x1f구수\n합\x1e계,9\n가,-\n나,5\n")
    assert _apply(client, fid, ("c0", None), ("c1", "population")).status_code == 200
    review = client.post("/reviews", json={"file_ids": [fid]}).json()
    reasons = [i["reason"] or "" for i in review["items"] if i["category"] == "warning"]
    assert any("\x1e" in r and "\x1f" in r for r in reasons), reasons
    _approve_pending(client, review)
    assert client.post(f"/reviews/{review['id']}/complete").status_code == 200
    res = client.post(f"/reviews/{review['id']}/integrate", json={})
    assert res.status_code == 409, res.text
    detail = res.json()["detail"]
    assert f"파일 #{fid} · 1번째 시트 · B열(머리글 1행): 열 이름" in detail, detail
    assert f"파일 #{fid} · 1번째 시트 · 2행 (검증 경고 사유 — 그 행의 칸에서 온 글자)" in detail, detail


def test_unexpected_integration_failure_keeps_raw_text_out_of_the_answer_record_and_log(
    client, monkeypatch, caplog
) -> None:
    from app import integration
    from app.integration import OUTPUT_DIR

    done = _upload(client, "완성결과.xlsx", _xlsx([["기관명", "인구수"], ["행정안전부", 7]]))
    assert _apply(client, done, ("c0", "org_name"), ("c1", "population")).status_code == 200
    first = client.post(f"/reviews/{_ready(client, [done])['id']}/integrate", json={})
    assert first.status_code == 201, first.text
    kept = client.get(f"/runs/{first.json()['id']}/download").content
    fid = _upload(client, "통합실패원문.xlsx", _xlsx([["기관명", "인구수"], ["행정안전부", 3]]))
    assert _apply(client, fid, ("c0", "org_name"), ("c1", "population")).status_code == 200
    review = _ready(client, [fid])

    def broken(*args: Any, **kwargs: Any) -> int:
        raise ValueError(f"담당 {PHONE} 셀 원문 cannot be used in worksheets.")

    monkeypatch.setattr(integration, "write_workbook", broken)
    with caplog.at_level(logging.INFO):
        res = client.post(f"/reviews/{review['id']}/integrate", json={})
    monkeypatch.undo()
    assert res.status_code == 409, res.text
    detail = res.json()["detail"]
    run_id, error = db_rows("select id, error from runs where review_session_id = ?", review["id"])[0]
    reason = "예기치 않은 오류 (ValueError)"
    assert detail == f"통합 결과를 만들지 못했습니다 — {reason} (실행 #{run_id}). [INTEGRATION_FAILED]"
    assert error == "예기치 않은 오류 (ValueError) [INTEGRATION_FAILED]"
    _no_echo(res.text, PHONE, "셀 원문")
    assert PHONE not in client.get(f"/runs?review_id={review['id']}").text
    assert PHONE not in caplog.text and "셀 원문" not in caplog.text
    assert f"integration run {run_id} failed: ValueError" in caplog.text  # 서버 기록: 종류·실행 번호·코드 위치만
    leftovers = [p.name for p in OUTPUT_DIR.iterdir() if p.name.endswith(".tmp")]
    assert leftovers == []  # 임시 파일은 실패를 기록한 뒤 지운다
    # 이미 완성된 다른 결과는 그대로 (지우지 않음)
    assert client.get(f"/runs/{first.json()['id']}/download").content == kept


def test_integration_failure_reason_has_no_sql(client, monkeypatch) -> None:
    from sqlalchemy.exc import DataError

    from app import runlease

    fid = _upload(client, "통합실패.xlsx", _xlsx([["기관명", "인구수"], ["행정안전부", 3]]))
    assert _apply(client, fid, ("c0", "org_name"), ("c1", "population")).status_code == 200
    review = _ready(client, [fid])

    def broken(*args: Any, **kwargs: Any) -> None:
        orig = Exception("value too long for type character varying(255)")
        raise DataError('INSERT INTO run_changes ("column") VALUES (%(column)s)', {"column": "x"}, orig)

    monkeypatch.setattr(runlease, "prepare", broken)
    res = client.post(f"/reviews/{review['id']}/integrate", json={})
    monkeypatch.undo()
    assert res.status_code == 409 and "DB 오류 (Exception)" in res.text, res.text
    assert "INSERT" not in res.text and "varying" not in res.text
    error = db_rows("select error from runs where review_session_id = ?", review["id"])[0][0]
    assert "INSERT" not in error and error.startswith("DB 오류"), error


def test_unexpected_bundle_failure_keeps_raw_text_out(client, monkeypatch, caplog) -> None:
    from app import runlease

    fid = _upload(client, "묶기실패.xlsx", _xlsx([["기관명", "인구수"], ["행정안전부", 5]]))
    assert _apply(client, fid, ("c0", "org_name"), ("c1", "population")).status_code == 200
    body = {"file_ids": [fid]}
    plan = client.post("/bundle/plan", json=body).json()
    assert plan["ready"], plan

    def broken(*args: Any, **kwargs: Any) -> None:
        raise ValueError(f"담당 {PHONE} 셀 원문")

    monkeypatch.setattr(runlease, "prepare", broken)
    with caplog.at_level(logging.INFO):
        res = client.post("/bundle/run", json={**body, "expected_basis": plan["basis"]})
    monkeypatch.undo()
    assert res.status_code == 409, res.text
    message = res.json()["detail"]["message"]
    assert message.startswith("결과 파일을 만들지 못했습니다 — 예기치 않은 오류 (ValueError) (실행 #")
    assert message.endswith("[BUNDLE_FAILED]")
    _no_echo(res.text, PHONE, "셀 원문")
    error = db_rows("select error from runs order by id desc limit 1")[0][0]
    assert PHONE not in error and error == "예기치 않은 오류 (ValueError) [BUNDLE_FAILED]", error
    assert PHONE not in caplog.text and "셀 원문" not in caplog.text


def test_bundle_keeps_the_apps_own_reading_message_when_it_fails_inside(client, monkeypatch) -> None:
    from app import bundle_api, files
    from app.parsing import ParseBusy

    fid = _upload(client, "묶기차례.xlsx", _xlsx([["기관명", "인구수"], ["행정안전부", 5]]))
    assert _apply(client, fid, ("c0", "org_name"), ("c1", "population")).status_code == 200
    body = {"file_ids": [fid]}
    plan = client.post("/bundle/plan", json=body).json()
    assert plan["ready"], plan
    inside, real_make, real_parse = {"on": False}, bundle_api.make, files.parse_file
    busy = "다른 파일을 읽는 중이라 60초 안에 읽기를 시작하지 못했습니다 — 잠시 뒤 다시 시도하세요."

    def make(*args: Any, **kwargs: Any) -> Any:
        inside["on"] = True
        try:
            return real_make(*args, **kwargs)
        finally:
            inside["on"] = False

    def parse(*args: Any, **kwargs: Any) -> Any:
        if inside["on"]:  # 결과를 만드는 중의 원본 읽기만 차례를 얻지 못함
            raise ParseBusy(busy)
        return real_parse(*args, **kwargs)

    monkeypatch.setattr(bundle_api, "make", make)
    monkeypatch.setattr(files, "parse_file", parse)
    res = client.post("/bundle/run", json={**body, "expected_basis": plan["basis"]})
    monkeypatch.undo()
    assert res.status_code == 409, res.text
    message = res.json()["detail"]["message"]
    assert message.startswith(f"결과 파일을 만들지 못했습니다 — {busy} (실행 #") and "예기치 않은 오류" not in message
    error = db_rows("select error from runs order by id desc limit 1")[0][0]
    assert error == f"{busy} [BUNDLE_FAILED]", error


# ---------------------------------------------------------------- 용어사전


def test_glossary_import_reports_unstorable_values_as_row_errors(client) -> None:
    previous = ";".join(str(1000000 + i) for i in range(63))  # 저장 형식(쉼표로 이음) 503자
    rows = [
        "코드,정식명칭,유형,별칭,이전코드,메모",
        f"9810001,{'가' * 256},기관,,,",
        f"9810002,정상기관,기관,짧은별칭;{'㍿' * 70},,",
        f"9810004,{'㈜' * 100},기관,,,",
        f"9810005,정상기관5,{'유' * 65},,,",
        f"9810006,정상기관6,기관,,,{'메' * 501}",
        f"9810007,정상기관7,기관,,{previous},",
        "9810008,널\x00기관,기관,,,",
    ]
    csv = "\n".join(rows) + "\n"
    res = client.post("/glossary/import", files=[("file", ("길이.csv", csv.encode("utf-8"), "text/csv"))])
    report = res.json()
    assert res.status_code == 200 and report["applied"] is False, res.text
    messages = {e["row"]: e["message"] for e in report["errors"]}
    assert sorted(messages) == [2, 3, 4, 5, 6, 7, 8], messages
    assert messages[2].startswith("정식명칭 256자 — 저장 한도 255자를 넘습니다")
    # 정규화로 길어진 이름은 사용자가 적은 글자 수와 늘어난 글자 수를 나눠 적고, 별칭이 여럿이면 순번을 적는다
    assert messages[3].startswith("별칭 2번째 70자 — 호환 문자(㈜·㍿ 등)를 풀어 정규화하면 280자로 저장 한도 255자를")
    assert messages[4].startswith("정식명칭 100자 — 호환 문자(㈜·㍿ 등)를 풀어 정규화하면 300자로")
    assert messages[5].startswith("유형 65자 — 저장 한도 64자를 넘습니다")
    assert messages[6].startswith("메모 501자 — 저장 한도 500자를 넘습니다")
    assert messages[7].startswith("이전코드 503자 — 저장 한도 500자를 넘습니다")
    assert messages[8].startswith("정식명칭에 NUL(0x00) 문자가 있습니다") and "[NAME_HAS_NUL]" in messages[8]
    _no_echo(res.text, "가" * 20, "유유유유유", "메메메메메")
    assert db_rows("select count(*) from glossary_entries where code like '981000%'") == [(0,)]
    # 고유 한도까지는 반영된다
    ok = f"코드,정식명칭,유형,메모\n9810003,{'나' * 255},{'형' * 64},{'모' * 500}\n"
    res = client.post("/glossary/import", files=[("file", ("한도.csv", ok.encode("utf-8"), "text/csv"))])
    assert res.json()["applied"] is True, res.text
    db_rows("delete from glossary_entries where code = '9810003'")


def test_glossary_import_row_errors_do_not_echo_the_cell(client) -> None:
    rows = [
        "코드,정식명칭,유형,시도코드,시작일,이전코드,메모",
        "98\x0010001,기관가,기관,,,,",
        f"9810013,기관나,기관,{PHONE},,,",
        "9810014,기관다,기관,,홍길동날,,",
        "9810015,기관라,기관,,,98XY,",
        "9810016,기관마,기관,,,,메\x00모",
    ]
    csv = "\n".join(rows) + "\n"
    res = client.post("/glossary/import", files=[("file", ("되돌림.csv", csv.encode("utf-8"), "text/csv"))])
    assert res.status_code == 200 and res.json()["applied"] is False, res.text
    messages = {e["row"]: e["message"] for e in res.json()["errors"]}
    assert messages[2] == "코드는 7자리(숫자·영문 대문자)여야 합니다.", messages
    assert messages[3] == "시도코드는 2자리 숫자여야 합니다.", messages
    assert messages[4] == "시작일이 날짜 형식이 아닙니다 (예: 2026-01-01).", messages
    assert messages[5] == "이전코드 중 1개가 7자리(숫자·영문 대문자)가 아닙니다.", messages
    assert messages[6].startswith("메모에 NUL(0x00) 문자가 있습니다") and "[VALUE_HAS_NUL]" in messages[6], messages
    _no_echo(res.text, PHONE, "홍길동", "98XY", "10001")


def test_glossary_alias_api_refuses_unstorable_aliases(client) -> None:
    res = client.post("/glossary/aliases", json={"alias": "㍿" * 70, "target_code": "1741000"})
    assert res.status_code == 422 and "별칭을 정규화한 이름이 280자로 저장 한도 255자" in res.text, res.text
    res = client.post("/glossary/aliases", json={"alias": "널\x00별칭", "target_code": "1741000"})
    assert res.status_code == 422 and "[NAME_HAS_NUL]" in res.text, res.text
    res = client.post("/glossary/aliases", json={"alias": "정상별칭메모", "target_code": "1741000", "note": "a\x00b"})
    assert res.status_code == 422 and "[VALUE_HAS_NUL]" in res.text, res.text
    assert db_rows("select count(*) from glossary_aliases where alias like '%별칭%'") == [(0,)]


# ---------------------------------------------------------------- 업로드·남은 입력의 안전망


def test_upload_refuses_a_file_name_with_nul(client) -> None:
    # 시험 도구(httpx)는 파일 이름의 NUL을 보내지 않으므로 multipart 본문을 직접 만든다 (다른 도구는 보낼 수 있다)
    boundary = "MOAJWOBOUNDARY"
    body = (
        f'--{boundary}\r\nContent-Disposition: form-data; name="files"; filename="널\x00이름.xlsx"\r\n'
        f"Content-Type: {XLSX}\r\n\r\n".encode()
        + _xlsx([["기관명"], ["a"]])
        + f"\r\n--{boundary}--\r\n".encode()
    )
    before = _count("source_files")
    res = client.post("/files", content=body, headers={"content-type": f"multipart/form-data; boundary={boundary}"})
    result = res.json()[0]
    assert res.status_code == 200 and result["id"] is None and "[NAME_HAS_NUL]" in result["error"], res.text
    assert _count("source_files") == before  # 기록·원본 파일을 만들기 전에 거절


def test_upload_refuses_a_file_name_over_the_limit_instead_of_cutting_it(client) -> None:
    before = _count("source_files")
    name = "가" * 251 + ".xlsx"  # 256자 — 예전에는 255자로 잘려 '.xls'가 되고 읽기 오류로 기록이 남았다
    res = client.post("/files", files=[("files", (name, _xlsx([["기관명", "인구수"], ["행정안전부", 1]]), XLSX))])
    result = res.json()[0]
    assert res.status_code == 200 and result["id"] is None, res.text
    expected = "파일 이름이 256자로 저장 한도 255자를 넘어 올리지 않았습니다 — 이름을 줄여 다시 올리세요."
    assert result["error"] == f"{expected} [NAME_TOO_LONG]", result
    assert _count("source_files") == before
    fid = _upload(client, "나" * 250 + ".xlsx", _xlsx([["기관명", "인구수"], ["행정안전부", 1]]))  # 255자
    ((stored,),) = db_rows("select original_name from source_files where id = ?", fid)
    assert len(stored) == 255 and stored.endswith(".xlsx"), len(stored)  # 자르지 않고 확장자 그대로


def test_typed_text_with_nul_is_not_a_server_error_on_postgresql(client) -> None:
    """앱이 쓰기 전에 거절하지 않는 직접 입력(자동화 이름 등)의 NUL — PostgreSQL은 안전망이 입력 오류(422)로 알리고,
    SQLite는 그대로 저장한다 (남은 두 DB 차이 — app/main.py)."""
    fid = _upload(client, "안전망.xlsx", _xlsx([["기관명", "인구수"], ["행정안전부", 1]]))
    assert _apply(client, fid, ("c0", "org_name"), ("c1", "population")).status_code == 200
    run = client.post(f"/reviews/{_ready(client, [fid])['id']}/integrate", json={})
    assert run.status_code == 201, run.text
    res = client.post(f"/runs/{run.json()['id']}/automation", json={"name": f"널\x00자동화 {fid}"})
    if POSTGRES:
        assert res.status_code == 422 and "[VALUE_HAS_NUL]" in res.text, res.text
    else:
        assert res.status_code == 201, res.text


# ---------------------------------------------------------------- 이름 규칙 전에 SQLite에 쌓인 자료 (그대로 둠)


@pytest.mark.sqlite_only("SQLite에만 한도를 넘는 기존 자료가 있을 수 있음 (PostgreSQL은 쓸 수 없음)")
def test_existing_long_values_in_sqlite_are_left_as_they_are(client) -> None:
    fid = _upload(client, "기존긴이름.xlsx", _xlsx([["기관명", N256], ["행정안전부", 1]]))
    db_rows(
        "insert into mapping_decisions (source_file_id, sheet_name, column_name, field_key) values (?, '현황', ?, ?)",
        fid,
        N256,
        "note",
    )
    res = client.get(f"/files/{fid}/mapping/latest")  # 기존 확정은 그대로 읽히고 확정으로 보인다
    assert res.status_code == 200, res.text
    column = res.json()["sheets"][0]["columns"][1]
    assert column["status"] == "confirmed" and column["field_key"] == "note", column
    assert _decided(fid) == [(N256, "note")]
    # 새 확정(지금 상태를 맞게 적고 덮어쓰기를 골라도)은 저장 한도로 거절하고, 이미 있던 행은 지우거나 바꾸지 않는다
    item = {"file_id": fid, "sheet_id": "s0", "column_id": "c1", "field_key": None, "expected_confirmed": True,
            "expected_field_key": "note", "overwrite": True}  # fmt: skip
    res = client.post("/mapping/batch/apply", json={"items": [item]})
    assert res.status_code == 409 and "256자로 저장 한도 255자" in res.text, res.text
    assert _decided(fid) == [(N256, "note")]


@pytest.mark.sqlite_only("이름 규칙 전에 SQLite에서 만든 검토 항목 (PostgreSQL에는 생길 수 없음)")
def test_legacy_review_item_with_a_long_name_is_not_confirmed(client) -> None:
    fid = _upload(client, "기존검토.xlsx", _xlsx([["기관명", "옛 메모"], ["행정안전부", "a"]]))
    review = client.post("/reviews", json={"file_ids": [fid]}).json()
    item = next(i for i in review["items"] if i["kind"] == "mapping" and i["column"] == "옛 메모")
    db_rows(
        "update review_items set category = 'pending', decision = null, \"column\" = ? where id = ?", N256, item["id"]
    )
    before = _decided(fid)
    res = client.post(f"/reviews/{review['id']}/items/{item['id']}/decision", json={"decision": "rejected"})
    assert res.status_code == 409 and "256자로 저장 한도 255자" in res.text and RAW not in res.text, res.text
    assert _decided(fid) == before


@pytest.mark.sqlite_only("이름 규칙 전에 SQLite에 쌓인 확정 (PostgreSQL에는 생길 수 없음)")
def test_recheck_does_not_copy_a_legacy_long_decision(client) -> None:
    a = _upload(client, "옛확정A.xlsx", _xlsx([["기관명", "인구수"], ["행정안전부", 1]]))
    b = _upload(client, "옛확정B.xlsx", _xlsx([["기관명", "인구수"], ["행정안전부", 2]]))
    assert _apply(client, a, ("c0", "org_name"), ("c1", "population")).status_code == 200
    review = client.post("/reviews", json={"file_ids": [a]}).json()
    db_rows("insert into mapping_decisions (source_file_id, sheet_name, column_name, field_key) values (?, '현황', ?, "
            "'note')", a, N256)  # fmt: skip
    before = _counts()
    res = client.post(f"/reviews/{review['id']}/recheck", json={"replace": {str(a): b}})
    assert res.status_code == 422, res.text
    detail = res.json()["detail"]
    assert detail.startswith(f"다시 검사하지 않았습니다 — 바꾸기 전 파일 #{a}의 열 매핑 확정 중 이름을 저장할 수 없는")
    assert "열 이름이 256자로 저장 한도 255자를 넘습니다" in detail and detail.endswith("[NAME_TOO_LONG]")
    _no_echo(res.text, RAW)
    # 하나도 옮기지 않았고(정상 확정 2개도) 이전 검토는 대체되지 않았다
    assert _decided(b) == [] and _counts() == before
    assert client.get(f"/reviews/{review['id']}").json()["status"] == "in_review"


@pytest.mark.sqlite_only("이름 규칙 전에 SQLite에 쌓인 승격 후보 (PostgreSQL에는 생길 수 없음)")
def test_promotion_approval_refuses_a_legacy_long_candidate(client) -> None:
    db_rows("insert into promotion_candidates (normalized, column_name, field_key, approver_count, status) values "
            "(?, ?, 'note', 3, 'pending')", N256, N256)  # fmt: skip
    cid = db_rows("select max(id) from promotion_candidates")[0][0]
    before = _counts()
    res = client.post(f"/admin/promotions/{cid}/approve", json={})
    assert res.status_code == 409, res.text
    assert "후보의 정규화한 이름이 256자로 저장 한도 255자를 넘습니다" in res.text and "[NAME_TOO_LONG]" in res.text
    assert db_rows("select status from promotion_candidates where id = ?", cid) == [("pending",)]
    assert _counts() == before
    assert db_rows("select count(*) from standard_terms where length(normalized) > 255") == [(0,)]
    # 반려는 그대로 할 수 있다 (사전은 바꾸지 않음)
    assert client.post(f"/admin/promotions/{cid}/reject", json={}).status_code == 200


@pytest.mark.sqlite_only("이름 규칙 전에 SQLite에서 만든 검토 (PostgreSQL에는 생길 수 없음)")
def test_legacy_review_with_a_long_source_name_is_refused_at_integration_with_its_place(client, monkeypatch) -> None:
    from app import review as review_module

    fid = _upload(client, "옛검토통합.xlsx", _xlsx([["기관명", N256], ["행정안전부", "a"]]))
    assert _apply(client, fid, ("c0", "org_name")).status_code == 200
    db_rows("insert into mapping_decisions (source_file_id, sheet_name, column_name, field_key) values (?, '현황', ?, "
            "'note')", fid, N256)  # fmt: skip
    monkeypatch.setattr(review_module, "check_files", lambda *a, **k: None)  # 이름 규칙 전에 만든 검토를 흉내
    review = _ready(client, [fid])
    monkeypatch.undo()
    before = _counts()
    res = client.post(f"/reviews/{review['id']}/integrate", json={})
    assert res.status_code == 409, res.text
    detail = res.json()["detail"]
    assert f"파일 #{fid} · 1번째 시트 · B열(머리글 1행): 원본 열 이름이 256자로 저장 한도 255자를 넘습니다" in detail
    assert "검토에서 다시 검사해야 합니다" in detail
    _no_echo(res.text, RAW)
    assert _counts() == before
