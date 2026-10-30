"""검사 기준 — 검사에 적용한 매핑(확정·자동 적용)과 결과에 쓸 매핑이 같을 때만 그 검사 결과로 완료·통합한다.

재현 (54da29b — 격리 환경의 실제 업로드·확정·검토·완료·통합 API에서도 확인):
- 기준일: '자료시점'이 검사에 쓰이지 않은 채 검토를 만들고 검토에서 base_date로 승인하면, 승인이 검사 기록까지 갱신해
  다시 검사 없이 완료·통합되고 8월 자료가 나갔다 (새로 검사하면 기간 불일치로 차단).
  날짜 글자만 있는 열은 날짜 형식으로 기준일에 자동 적용되어 검사 때 이미 기간을 검사하므로, 이 경로는 날짜가 아닌
  값('미상' 등)이 섞여 자동 적용되지 않은 경우에 생긴다.
- 자동 적용: '관측값'을 학습 별칭으로 인구수에 자동 적용해 검사·완료한 뒤 다른 파일에서 건수로 확정해 별칭이 바뀌면,
  기존 검토로 건수 열이 나갔다 (인구수의 시점 값 예외로 통과한 합계 100을 건수로는 60+100=160과 비교해야 함).

실제 쿠키 로그인·DB·학습 API를 쓰고, 매핑 엔진은 시험용(임베딩 없음, LLM은 '해당 없음' 고정 응답)이다.
시험마다 새 전용 계정을 쓴다 — 학습 별칭은 계정마다 따로라 시험끼리 섞이지 않는다.
"""

import csv
import io
import itertools
from typing import Any

import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import func, select

from tests.test_accounts import owner  # noqa: F401
from tests.test_auth import _client, _provision_user, app_real_login  # noqa: F401

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
GIVEN = [["자료시점", "집행액"], ["2026-08-31", "100"], ["2026-08-31", "200"]]
MIXED = [["자료시점", "집행액"], ["2026-08-31", "100"], ["2026-08-31", "200"], ["미상", "300"]]
NORMAL = [["자료시점", "집행액"], ["2026-09-30", "100"], ["2026-09-30", "200"], ["미상", "300"]]
STOCK = [["월", "관측값"], ["합계", "100"], ["1월", "60"], ["2월", "100"]]
_users = itertools.count(1)


@pytest.fixture
def user(app_real_login, owner) -> TestClient:  # noqa: F811
    name = f"rb-user{next(_users)}"
    _provision_user(owner, app_real_login, name, f"pass-{name}", name.upper())
    return _client(app_real_login, name, f"pass-{name}")


def _upload(c: TestClient, name: str, rows: list[list], kind: str = "csv") -> int:
    """제목 행 없이 (파일·시트 이름에도 기간 없음 — 기간 근거는 열 값뿐)."""
    if kind == "csv":
        buf = io.StringIO()
        csv.writer(buf, lineterminator="\n").writerows(rows)
        content, ctype = buf.getvalue().encode("utf-8"), "text/csv"
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "자료"
        for r in rows:
            ws.append(r)
        for row in ws.iter_rows():
            for cell in row:
                cell.data_type = "s"  # 모두 글자 셀
        out = io.BytesIO()
        wb.save(out)
        content, ctype = out.getvalue(), XLSX
    res = c.post("/files", files=[("files", (f"{name}.{kind}", content, ctype))])
    assert res.status_code == 200 and res.json()[0]["id"], res.text
    return res.json()[0]["id"]


def _confirm(c: TestClient, fid: int, mapping: dict[str, str | None]) -> None:
    """열 매핑 화면의 확정 (학습 별칭도 남는다)."""
    pv = c.get(f"/files/{fid}/preview").json()
    ids = {col["name"]: col["column_id"] for col in pv["columns"]}
    items = [{"sheet_id": pv["sheet_id"], "column_id": ids[n], "field_key": k} for n, k in mapping.items()]
    assert c.put(f"/files/{fid}/mapping/decisions", json={"decisions": items}).status_code == 200


def _mapping(c: TestClient, fid: int) -> dict[str, tuple[str, str, str | None]]:
    """열 매핑 화면과 같은 판정 (모델 호출 없음) — 열 → (상태, 판정 방법, 항목)."""
    res = c.post(f"/files/{fid}/mapping?use_llm=false").json()
    cols = res["sheets"][0]["columns"] if "sheets" in res else res["columns"]
    return {m["column"]: (m["status"], m["method"], m["field_key"]) for m in cols}


def _review(c: TestClient, fids: list[int], **params: Any) -> dict:
    res = c.post("/reviews", json={"file_ids": fids, **params})
    assert res.status_code == 201, res.text
    return res.json()


def _items(review: dict, column: str) -> list[dict]:
    return [i for i in review["items"] if i["column"] == column and i["kind"] == "mapping"]


def _decide(c: TestClient, review: dict, column: str, body: dict) -> dict:
    (item,) = [i for i in _items(review, column) if i["category"] == "pending"]
    res = c.post(f"/reviews/{review['id']}/items/{item['id']}/decision", json=body)
    assert res.status_code == 200, res.text
    return res.json()


def _blocks(review: dict) -> list[tuple[str, str]]:
    return [(i["kind"], i["reason"]) for i in review["items"] if i["category"] == "block"]


def _state() -> dict[str, Any]:
    """통합 실행 기록·변경 내역·결과 파일 (거절된 요청이 아무것도 만들지 않았는지)."""
    from app.db import SessionLocal
    from app.integration import OUTPUT_DIR
    from app.models import IntegrationRun, RunChange

    with SessionLocal() as db:
        runs = db.scalar(select(func.count(IntegrationRun.id)))
        changes = db.scalar(select(func.count(RunChange.id)))
    return {"runs": runs, "changes": changes, "files": sorted(p.name for p in OUTPUT_DIR.glob("*"))}


def _excel(c: TestClient, run_id: int) -> dict[str, list[tuple[Any, str]]]:
    grid = list(openpyxl.load_workbook(io.BytesIO(c.get(f"/runs/{run_id}/download").content)).worksheets[0].iter_rows())
    return {cell.value: [(row[k].value, row[k].data_type) for row in grid[1:]] for k, cell in enumerate(grid[0])}


def _blocked(c: TestClient, review: dict) -> None:
    """완료·통합이 거절되고 아무것도 새로 생기지 않는다 (검토 상태도 그대로)."""
    before, status = _state(), c.get(f"/reviews/{review['id']}").json()["status"]
    if status == "in_review":
        done = c.post(f"/reviews/{review['id']}/complete")
        assert done.status_code == 409, done.text
    assert c.post(f"/reviews/{review['id']}/integrate", json={"include_unmapped": True}).status_code == 409
    after = c.get(f"/reviews/{review['id']}").json()
    assert (after["status"], _state()) == (status, before)
    if status == "in_review":
        assert after["completed_at"] is None


# ---------------------------------------------------------------- 기준일 승인


@pytest.mark.parametrize("kind", ["csv", "xlsx"])
def test_date_only_column_is_applied_as_base_date_when_checked(user, kind: str) -> None:
    """지정 자료(날짜 글자만 있는 '자료시점'): 날짜 형식으로 기준일에 자동 적용되어 검사 때 기간을 검사한다 —
    검토에서 승인할 항목이 아니며, 생성 단계에서 기간 불일치로 차단된다."""
    fid = _upload(user, f"기준일_지정_{kind}", GIVEN, kind)
    _confirm(user, fid, {"집행액": "expenditure"})
    assert _mapping(user, fid)["자료시점"] == ("auto", "pattern", "base_date")
    review = _review(user, [fid], expected_period="2026-09")
    assert [i["category"] for i in _items(review, "자료시점")] == ["auto"]
    assert [k for k, _ in _blocks(review)] == ["period_mismatch"]
    _blocked(user, review)


@pytest.mark.parametrize("kind", ["csv", "xlsx"])
def test_base_date_approved_in_review_needs_recheck_and_recheck_reveals_the_period(user, kind: str) -> None:
    fid = _upload(user, f"기준일_미상_{kind}", MIXED, kind)
    _confirm(user, fid, {"집행액": "expenditure"})
    assert _mapping(user, fid)["자료시점"][0] == "unresolved"  # '미상'이 섞여 날짜 형식으로 자동 적용되지 않음
    review = _review(user, [fid], expected_period="2026-09")
    (item,) = _items(review, "자료시점")
    assert item["category"] == "pending" and _blocks(review) == []  # 기간 항목이 없어 기간 검사를 하지 못함

    review = _decide(user, review, "자료시점", {"decision": "approved", "choice": "base_date"})
    reason = f"기준일_미상_{kind}.{kind} '자료시점' 열: 검사에 쓰지 않음 → 확정 기준일자"
    assert (review["needs_recheck"], review["can_complete"], review["recheck_reasons"]) == (True, False, [reason])
    assert _mapping(user, fid)["자료시점"] == ("confirmed", "user", "base_date")  # 승인 자체는 남는다
    _blocked(user, review)

    new = user.post(f"/reviews/{review['id']}/recheck", json={"replace": {}}).json()
    assert new["needs_recheck"] is False
    (block,) = _blocks(new)
    assert block[0] == "period_mismatch" and "기간 2026-08이(가) 기준 기간 2026-09과(와) 다릅니다" in block[1]
    _blocked(user, new)


def test_base_date_normal_data_completes_after_recheck(user) -> None:
    fid = _upload(user, "기준일_정상", NORMAL)
    _confirm(user, fid, {"집행액": "expenditure"})
    review = _review(user, [fid], expected_period="2026-09")
    review = _decide(user, review, "자료시점", {"decision": "approved", "choice": "base_date"})
    assert review["needs_recheck"] is True
    new = user.post(f"/reviews/{review['id']}/recheck", json={"replace": {}}).json()
    assert (_blocks(new), new["needs_recheck"], new["can_complete"]) == ([], False, True)
    assert new["params"]["expected_period"] == "2026-09"  # 저장한 검사 조건
    assert user.post(f"/reviews/{new['id']}/complete").status_code == 200
    run = user.post(f"/reviews/{new['id']}/integrate", json={"include_unmapped": True})
    assert run.status_code == 201, run.text
    assert run.json()["summary"]["expected_period"] == "2026-09"
    sheet = _excel(user, run.json()["id"])
    assert sheet["기준일자"] == [("2026-09-30", "s"), ("2026-09-30", "s"), ("미상", "s")]
    assert sheet["집행액"] == [(100, "n"), (200, "n"), (300, "n")]


# ---------------------------------------------------------------- 자동 적용 매핑 변경


def _stock_review(c: TestClient) -> tuple[int, dict]:
    """'관측값 → 인구수' 학습 별칭을 별도 파일 확정으로 준비하고, 대상 파일은 월만 확정해 검토·완료한다."""
    _confirm(c, _upload(c, "별칭준비_인구", STOCK), {"관측값": "population", "월": "month"})
    fid = _upload(c, "대상_관측값", STOCK)
    _confirm(c, fid, {"월": "month"})
    assert _mapping(c, fid)["관측값"] == ("auto", "learned", "population")  # 파일별 확정 없이 자동 적용
    review = _review(c, [fid])
    assert _blocks(review) == []  # 인구수(시점 값): 합계 100 = 마지막 달 100 → 누적 예외
    assert c.post(f"/reviews/{review['id']}/complete").status_code == 200
    return fid, review


def test_auto_mapping_change_after_check_blocks_integration_until_recheck(user) -> None:
    fid, review = _stock_review(user)
    assert user.get(f"/reviews/{review['id']}").json()["needs_recheck"] is False
    _confirm(user, _upload(user, "별칭변경_건수", STOCK), {"관측값": "count"})  # 같은 계정의 다른 파일에서 확정
    assert _mapping(user, fid)["관측값"] == ("auto", "learned", "count")  # 자동 적용이 바뀜
    assert _mapping(user, fid)["월"] == ("confirmed", "user", "month")  # 대상 파일의 확정은 그대로

    got = user.get(f"/reviews/{review['id']}").json()
    reason = "대상_관측값.csv '관측값' 열: 자동 적용 인구수 — 자동 적용 근거가 바뀜 (학습 별칭 인구수 → 건수)"
    assert (got["needs_recheck"], got["recheck_reasons"]) == (True, [reason])
    before = _state()
    res = user.post(f"/reviews/{review['id']}/integrate", json={"include_unmapped": True})
    assert res.status_code == 409 and reason in res.json()["detail"], res.text
    assert _state() == before

    new = user.post(f"/reviews/{review['id']}/recheck", json={"replace": {}}).json()
    (block,) = _blocks(new)  # 건수(발생 실적)로 검사하면 누적 예외가 없다
    assert block[0] == "sum_mismatch" and "하위 2행 합계 160" in block[1] and "(차이 -60)" in block[1]
    _blocked(user, new)


def test_auto_mapping_drift_seen_only_when_integrating_is_rejected_before_output(user, monkeypatch) -> None:
    """조회 때 알 수 없는 자동 적용 변화(임베딩 유사도 등)는 통합이 다시 계산한 매핑으로, 결과를 만들기 전에 막는다."""
    from app import integration

    fid, review = _stock_review(user)
    real = integration.get_engine()

    class Drifted:
        fields = real.fields

        def map_columns(self, *args: Any, **kwargs: Any) -> Any:
            result = real.map_columns(*args, **kwargs)
            for m in result.columns:
                if m.column == "관측값" and m.status == "auto":
                    m.field_key, m.field_label = "count", "건수"
            return result

    monkeypatch.setattr(integration, "get_engine", lambda: Drifted())
    assert user.get(f"/reviews/{review['id']}").json()["needs_recheck"] is False
    before = _state()
    res = user.post(f"/reviews/{review['id']}/integrate", json={"include_unmapped": True})
    assert res.status_code == 409, res.text
    assert "결과에 쓸 열 매핑이 검사 때와 다릅니다" in res.json()["detail"]
    assert "'관측값' 열: 검사 때 자동 적용 인구수 → 지금 자동 적용 건수" in res.json()["detail"]
    assert _state() == before


def test_alias_change_during_integration_uses_the_checked_mapping(user, monkeypatch) -> None:
    """통합이 학습 별칭을 읽은 뒤 다른 요청이 별칭을 바꿔도 이 실행은 검사한 매핑(인구수)으로 만든다 — 다음은 거절."""
    from app import integration
    from app.db import SessionLocal
    from app.models import LearnedAlias

    fid, review = _stock_review(user)
    original, changed = integration.load_parsed, []

    def change_then_load(db, file_id):  # 검사 기준 비교 다음, 파일을 읽는 순간에 끼어든 다른 요청
        if not changed:
            with SessionLocal() as other:
                mine = LearnedAlias.alias == "관측값", LearnedAlias.user_id == review["created_by"]
                alias = other.scalar(select(LearnedAlias).where(*mine))
                alias.field_key = "count"
                other.commit()
            changed.append(file_id)
        return original(db, file_id)

    monkeypatch.setattr(integration, "load_parsed", change_then_load)
    run = user.post(f"/reviews/{review['id']}/integrate", json={"include_unmapped": True})
    assert run.status_code == 201, run.text
    assert changed == [fid]
    sheet = _excel(user, run.json()["id"])
    assert sheet["인구수"] == [(60, "n"), (100, "n")] and "건수" not in sheet
    assert user.get(f"/reviews/{review['id']}").json()["needs_recheck"] is True
    before = _state()
    assert user.post(f"/reviews/{review['id']}/integrate", json={"include_unmapped": True}).status_code == 409
    assert _state() == before


# ---------------------------------------------------------------- 검토 안 승인: 검사 입력이 그대로인지


def test_review_approvals_reuse_the_check_only_when_the_check_input_is_unchanged(user) -> None:
    rows = [
        ["분류명", "메모가", "메모나", "빈칸메모", "숫자메모", "관할구역명"],
        # 이 시험만의 지명 — 공용 시험 DB에 다른 시험이 넣은 용어사전 항목('강서구' 등)에 따라 결과가 바뀌지 않게
        ["가", "참고1", "참고4", "있음", "1", "승인가구"],
        ["나", "참고2", "참고5", "", "2", "승인나구"],
        ["다", "참고3", "참고6", "있음", "3", "승인다구"],
    ]
    fid = _upload(user, "승인의존", rows)
    review = _review(user, [fid])
    steps = [
        ("메모가", "note", False),  # 검사가 읽지 않는 항목 + 빈 칸·숫자 없는 글자 열
        ("분류명", None, False),  # '해당 없음'
        ("빈칸메모", "manager", True),  # 빈 칸 → 공백 경고
        ("숫자메모", "vehicle_type", True),  # 숫자 → 합계·전월·단위 비교
        ("관할구역명", "sigungu", True),  # 기관·지역 항목 → 필수 기관·이름 확인
        ("메모나", "note", True),  # 같은 항목 열이 둘 → 전월·단위 비교 대상이 바뀜
    ]
    flagged = []
    for column, key, needs in steps:
        body = {"decision": "approved", "choice": key} if key else {"decision": "rejected"}
        review = _decide(user, review, column, body)
        if needs:
            flagged.append(column)
        assert sorted(r.split("'")[1] for r in review["recheck_reasons"]) == sorted(flagged), (column, review)
    assert review["needs_recheck"] is True
    new = user.post(f"/reviews/{review['id']}/recheck", json={"replace": {}}).json()
    assert new["needs_recheck"] is False

    def found(r: dict) -> set[tuple[str, str]]:
        return {(i["kind"], i["before"] or i["column"]) for i in r["items"] if i["category"] in ("pending", "warning")}

    # 다시 검사하면 승인한 매핑이 검사에 쓰인다 — 시군구 이름 확인과 빈 칸 경고는 이전 검사 결과에 없던 항목
    expected = {("entity", "승인가구"), ("entity", "승인나구"), ("entity", "승인다구"), ("blank_value", "빈칸메모")}
    assert expected <= found(new) - found(review)


def test_reasons_mask_personal_data_in_names(user) -> None:
    rows = [["담당 kim@example.com", "집행액"], ["2026-08-31", "100"], ["미상", "200"]]
    fid = _upload(user, "가림확인", rows)
    _confirm(user, fid, {"집행액": "expenditure"})
    review = _review(user, [fid], expected_period="2026-09")
    pending = [i for i in review["items"] if i["category"] == "pending" and i["kind"] == "mapping"]
    res = user.post(f"/reviews/{review['id']}/items/{pending[0]['id']}/decision",
                    json={"decision": "approved", "choice": "base_date"})  # fmt: skip
    reasons = res.json()["recheck_reasons"]
    assert len(reasons) == 1 and "kim@example.com" not in reasons[0] and "'담당" in reasons[0]


# ---------------------------------------------------------------- 전월 비교 자료


def test_baseline_file_changes_are_compared(user) -> None:
    """전월 비교에 쓴 파일의 확정·자동 적용 근거가 바뀌어도 검사 결과를 쓸 수 없다 (통합하지 않아도 검사에 썼음)."""
    _confirm(user, _upload(user, "전월별칭준비", [["기관", "실적값"], ["가", "1"]]), {"실적값": "expenditure"})
    rows = [["기관명", "집행액", "실적값"], ["행정안전부", "100", "5"], ["기상청", "200", "6"]]
    target = _upload(user, "이번달", rows)
    base = _upload(user, "전월", rows)
    for fid in (target, base):
        _confirm(user, fid, {"기관명": "org_name", "집행액": "expenditure"})
    review = _review(user, [target], baseline_file_ids=[base])
    assert review["needs_recheck"] is False

    _confirm(user, base, {"집행액": "budget"})
    got = user.get(f"/reviews/{review['id']}").json()
    assert got["recheck_reasons"] == ["전월 자료 전월.csv '집행액' 열: 확정 집행액 → 확정 예산액"]
    _confirm(user, base, {"집행액": "expenditure"})  # 되돌리면 다시 같은 상태
    assert user.get(f"/reviews/{review['id']}").json()["needs_recheck"] is False

    _confirm(user, _upload(user, "전월별칭변경", [["기관", "실적값"], ["가", "1"]]), {"실적값": "count"})
    reasons = user.get(f"/reviews/{review['id']}").json()["recheck_reasons"]
    assert (
        "전월 자료 전월.csv '실적값' 열: 자동 적용 집행액 — 자동 적용 근거가 바뀜 (학습 별칭 집행액 → 건수)" in reasons
    )
    assert "이번달.csv '실적값' 열: 자동 적용 집행액 — 자동 적용 근거가 바뀜 (학습 별칭 집행액 → 건수)" in reasons
