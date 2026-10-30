"""과거 분석 결과를 조회·내려받을 때 지금 개인정보 가림 규칙을 적용한다 (저장값·수치·구조는 그대로).

저장된 결과는 저장 당시 규칙으로 가린 글자라, 지금 규칙으로는 가려야 할 글자(예: 한글에 붙은 이메일)가
남아 있을 수 있다.
지금 코드로는 원문이 저장되지 않으므로, 실제 흐름(업로드 → 확정 → 확인 → 실행)으로 결과를 만든 뒤
그 결과의 저장 JSON 글자 칸에 원문 개인정보 형식(합성)을 **직접 넣어** 과거 결과를 꾸민다.
"""

import copy
import io
import json
import re
import zipfile
from collections.abc import Iterator
from typing import Any

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app.pii import mask_text
from tests.conftest import db_connect
from tests.test_accounts import owner  # noqa: F401
from tests.test_analysis import POP, THREE, _body, _file, _run, _upload
from tests.test_auth import _client, _provision_user, app_real_login  # noqa: F401

GLUE = "인구수kim@example.com"  # 수정 전 규칙이 놓친 형식 (한글에 붙은 이메일)
GLUE_U = "인구수_kim@example.com"
MAIL = "kim@example.com"
PHONE = "010-1234-5678"
RRN = "900101-1234567"
# 지금 규칙으로 가리면 서로 같아지는 이름들 (그룹·행·구간이 합쳐지거나 사라지면 안 됨)
ORGS = {
    "기관A": "기관kim@example.com",
    "기관B": "기관kix@example.com",
    "기관C": PHONE,
    "기관D": RRN,
    "기관E": "기관k***@example.com",  # 이미 가린 값
    "기관F": "=1+1 kim@example.com",  # 수식처럼 시작하는 글자
}
SHEETS = {"s0": "현황kim@example.com", "s1": "현황kix@example.com", "s2": "설명_010-9876-5432"}
FILE_NAME = "인구수_kim@example.com.xlsx"
BINS = ["구간kim@example.com", "구간kix@example.com"]
# 화면·API·Excel 어디에도 남으면 안 되는 원문 조각
RAW = ("kim@example.com", "kix@example.com", PHONE, "010-9876-5432", "1234567", "9876-5432")
HEX = "ab12cd01098765432" + "e" * 47  # 16진수 해시 안에 전화번호 형태의 숫자열 → 구조값이라 그대로 둬야 함
# 조회 키·식별자·해시·코드 (가리지 않음)
KEEP = {"sha256", "sheet_id", "column_id", "field_key", "policy", "status", "reason", "region_level"}
KEEP |= {"group_by", "group_fields", "aggregates", "on"}
PATHS = ("", "/groups", "/excluded", "/detail", "/download")


def _db(sql: str, *args: Any) -> list[tuple]:
    with db_connect() as con:
        return con.execute(sql, args).fetchall()


def _stored(rid: int) -> dict[str, Any]:
    (row,) = _db(
        "select summary, conditions, inputs, result, created_at, created_by, partial, field_key "
        "from analysis_runs where id = ?",
        rid,
    )
    return dict(zip(("summary", "conditions", "inputs", "result", "created_at", "created_by", "partial", "field_key"),
                    row, strict=True))  # fmt: skip


def _put_stored(rid: int, **cols: Any) -> None:
    with db_connect() as con:
        for k, v in cols.items():
            text = v if isinstance(v, str) else json.dumps(v, ensure_ascii=False)
            con.execute(f"update analysis_runs set {k} = ? where id = ?", (text, rid))


def _restore(rid: int, st: dict[str, Any]) -> None:
    """꾸민 원문을 저장 당시 글자 그대로 되돌린다.
    시험 DB는 세션 공용이라, 'DB 어디에도 원문 없음'을 보는 다른 시험이 이 자료를 보지 않게 한다."""
    _put_stored(rid, **{k: st[k] for k in ("summary", "conditions", "inputs", "result")})


def expect(value: Any, key: str | None = None) -> Any:
    """기대값: 저장값의 글자만 지금 규칙(mask_text)으로 가리고, 구조값·숫자·참거짓·None은 그대로."""
    if isinstance(value, dict):
        return {k: v if k in KEEP else expect(v, k) for k, v in value.items()}
    if isinstance(value, list):
        return [expect(v, key) for v in value]
    return mask_text(value) if isinstance(value, str) else value


def same(a: Any, b: Any, where: str = "") -> None:
    """값과 형(int·float·bool·None·str)까지 같은지 (1 == 1.0 == True 를 같다고 보지 않음)."""
    assert type(a) is type(b), f"{where}: {type(a).__name__} != {type(b).__name__} ({a!r} / {b!r})"
    if isinstance(a, dict):
        assert list(a) == list(b), f"{where}: 키 {list(a)} != {list(b)}"
        for k in a:
            same(a[k], b[k], f"{where}.{k}")
    elif isinstance(a, list):
        assert len(a) == len(b), f"{where}: 길이 {len(a)} != {len(b)}"
        for i, (x, y) in enumerate(zip(a, b, strict=True)):
            same(x, y, f"{where}[{i}]")
    else:
        assert a == b, f"{where}: {a!r} != {b!r}"


def no_raw(text: str, where: str) -> None:
    found = [r for r in RAW if r in text]
    assert not found, f"{where}: 원문 조각 {found}"


def _xlsx_sheets(sheets: dict[str, list[list]]) -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _confirm_sheet(c: TestClient, fid: int, sheet_id: str, mapping: dict[str, str | None]) -> None:
    preview = c.get(f"/files/{fid}/preview", params={"sheet_id": sheet_id}).json()
    cols = {x["name"]: x["column_id"] for x in preview["columns"]}
    items = [{"sheet_id": sheet_id, "column_id": cols[n], "field_key": k} for n, k in mapping.items()]
    assert c.put(f"/files/{fid}/mapping/decisions", json={"decisions": items}).status_code == 200


def _all(c: TestClient, rid: int, kind: str, limit: int = 200) -> list[dict]:
    items: list[dict] = []
    offset = 0
    while True:
        page = c.get(f"/analysis/runs/{rid}/{kind}", params={"offset": offset, "limit": limit}).json()
        items += page["items"]
        offset += limit
        if offset >= page["total"]:
            return items


def _taint(fid: int, rid: int) -> None:
    """실제로 만든 결과의 저장 JSON 글자 칸에 원문 형식을 넣어 '수정 전 규칙으로 저장된 과거 결과'를 꾸민다."""
    st = _stored(rid)
    summary, cond, inp, res = (json.loads(st[k]) for k in ("summary", "conditions", "inputs", "result"))
    unit = f"명 {MAIL}"
    summary.update(variant=GLUE_U, unit=unit)
    cond.update(variant=GLUE_U, unit=unit, region_level="lower", records=[f"#{fid} {SHEETS['s0']}"])
    cond["filters"]["dims"] = {"org_name": [ORGS["기관A"], ORGS["기관B"], PHONE]}
    for r, label in zip(cond["bins"]["ranges"], BINS, strict=True):
        r["label"] = label
    cond["warnings"] = [f"열 이름 '{GLUE}' 확인", "열 이름 '인구수kix@example.com' 확인"]
    cond["level_conflicts"] = [
        {"period": "2026-09", "upper": {"target": f"서울특별시 · {PHONE}", "file_id": fid, "row": 3, "value": 100,
                                        "rows": 1},
         "lower_count": 2, "lower_sum": 90, "lower_missing": 0, "difference": 10}
    ]  # fmt: skip
    for t in inp:
        t.update(file_name=FILE_NAME, sheet=SHEETS[t["sheet_id"]], sha256=HEX)
        if t["status"] == "ok":
            t.update(column=GLUE if t["sheet_id"] == "s0" else GLUE_U, variant=GLUE_U, unit=unit)
            t["candidates"] = [{**x, "column": t["column"], "variant": GLUE_U} for x in t["candidates"]]
            t["decisions"] = [[GLUE if n == "인구수" else n, k] for n, k in t["decisions"]] + [[RRN, None]]
        t["reasons"] = [*t["reasons"], f"열 '{GLUE}' 확인 필요", "열 '인구수kix@example.com' 확인 필요"]
    inp[0]["repeated"] = {"keys": 1, "rows": 2, "same_value": 0, "different_value": 1,
                          "samples": [{"period": "2026-09", "target": f"서울특별시 · {GLUE}", "rows": [3, 4],
                                       "values": [100, f"담당 {MAIL}"]}]}  # fmt: skip
    for g in res["groups"]:
        g["keys"] = [ORGS.get(k, k) for k in g["keys"]]
        if g["bin_index"] is not None:
            g["bin"] = BINS[g["bin_index"]]
    for b in res["bins"]:
        if b["index"] is not None:
            b["label"] = BINS[b["index"]]
    by_sheet = {t["sheet"]: t["sheet_id"] for t in json.loads(st["inputs"])}
    for x in res["detail"] + res["excluded"]:
        x.update(file_name=FILE_NAME, sheet=SHEETS[by_sheet[x["sheet"]]])
        if "keys" in x:
            x["keys"] = [ORGS.get(k, k) for k in x["keys"]]
    res["excluded"][0]["value"] = f"담당 {PHONE}"
    res["excluded_targets"] = [
        {**t, "file_name": FILE_NAME, "sheet": SHEETS[t["sheet_id"]], "reasons": [*t["reasons"], f"담당 {MAIL}"]}
        for t in res["excluded_targets"]
    ]
    _put_stored(rid, summary=summary, conditions=cond, inputs=inp, result=res)


@pytest.fixture(scope="module")
def past(app_real_login, owner) -> Iterator[dict[str, Any]]:  # noqa: F811
    users = {}
    for name, display in (("am-user1", "분석자kim@example.com"), ("am-user2", "AM-USER2")):
        _provision_user(owner, app_real_login, name, "pass-" + name, display)
        users[name] = _client(app_real_login, name, "pass-" + name)
    u = users["am-user1"]
    head = [[POP], ["시도", "기관명", "인구수"]]
    data = _xlsx_sheets(
        {
            "현황A": [*head, ["서울특별시", "기관A", 100], ["서울특별시", "기관B", "-"], ["부산광역시", "기관C", 300]],
            "현황B": [*head, ["대구광역시", "기관D", 400], ["대구광역시", "기관E", "x"], ["광주광역시", "기관F", 600]],
            "참고": [*head, ["울산광역시", "기관G", 700]],  # 확정하지 않은 시트 → 자료 제외
        }
    )
    fid = _upload(u, "가림.xlsx", data)
    for sid in ("s0", "s1"):
        _confirm_sheet(u, fid, sid, {"시도": "sido", "기관명": "org_name", "인구수": "population"})
    bins = {"on": "sum", "ranges": [{"label": "작음", "max": 350}, {"label": "큼", "min": 350}]}
    targets = [{"file_id": fid, "sheet_id": s} for s in ("s0", "s1", "s2")]
    r = _run(u, {**_body([fid], group_by=["org_name"], bins=bins), "targets": targets})
    original = _stored(r["id"])
    _taint(fid, r["id"])
    plain = _run(u, _body([_file(u, "보통.xlsx", THREE, {"시도": "sido", "인구수": "population"})]))
    yield {"users": users, "owner": owner, "fid": fid, "id": r["id"], "plain": plain["id"]}
    _restore(r["id"], original)
    with db_connect() as con:
        con.execute("update users set display_name = 'AM-USER1' where username = 'am-user1'")


def _detail_expect(st: dict[str, Any]) -> dict[str, Any]:
    res = json.loads(st["result"])
    meta = {k: v for k, v in res.items() if k not in ("groups", "detail", "excluded")}
    return {
        "conditions": expect(json.loads(st["conditions"])),
        "inputs": expect(json.loads(st["inputs"])),
        "result": expect({**meta, "groups_total": len(res["groups"])}),
    }


# ---------------------------------------------------------------- 재현 자료가 실제로 원문을 담고 있는지


def test_stored_past_result_really_contains_raw_text(past) -> None:
    st = _stored(past["id"])
    text = " ".join(st[k] for k in ("summary", "conditions", "inputs", "result"))
    for raw in (GLUE, GLUE_U, MAIL, PHONE, RRN, "010-9876-5432", *ORGS.values(), *SHEETS.values(), *BINS, HEX):
        assert raw in text, raw
    assert mask_text(HEX) != HEX  # 구조값을 가렸다면 해시가 바뀌었을 것


# ---------------------------------------------------------------- 6개 경로


def test_list_masks_summary_and_keeps_numbers(past) -> None:
    u = past["users"]["am-user1"]
    st = _stored(past["id"])
    (item,) = [x for x in u.get("/analysis/runs").json() if x["id"] == past["id"]]
    no_raw(json.dumps(item, ensure_ascii=False), "GET /analysis/runs")
    same({k: v for k, v in item.items() if k not in ("id", "created_at")}, expect(json.loads(st["summary"])), "summary")
    assert item["variant"] == "인구수_***@example.com" and item["unit"] == "명 k***@example.com"


def test_detail_masks_conditions_inputs_meta_and_creator(past) -> None:
    u = past["users"]["am-user1"]
    st = _stored(past["id"])
    body = u.get(f"/analysis/runs/{past['id']}").json()
    no_raw(json.dumps(body, ensure_ascii=False), "GET /analysis/runs/{id}")
    for k, v in _detail_expect(st).items():
        same(body[k], v, k)
    assert body["created_by"] == "분석자k***@example.com" and body["read_only"] is False
    c, inp = body["conditions"], body["inputs"]
    assert c["filters"]["dims"]["org_name"] == ["기관k***@example.com", "기관k***@example.com", "010-****-5678"]
    assert [r["label"] for r in c["bins"]["ranges"]] == ["구간k***@example.com", "구간k***@example.com"]
    assert [(r["min"], r["max"]) for r in c["bins"]["ranges"]] == [(None, 350), (350, None)]
    assert c["level_conflicts"][0]["upper"]["target"] == "서울특별시 · 010-****-5678"
    assert (c["level_conflicts"][0]["lower_sum"], c["level_conflicts"][0]["difference"]) == (90, 10)
    assert [t["sheet"] for t in inp] == ["현황k***@example.com", "현황k***@example.com", "설명_010-****-5432"]
    assert {t["sha256"] for t in inp} == {HEX} and [t["sheet_id"] for t in inp] == ["s0", "s1", "s2"]
    assert inp[0]["decisions"][-1] == ["900101-1******", None]
    assert inp[0]["repeated"]["samples"][0]["values"] == [100, "담당 k***@example.com"]


def test_group_pages_mask_names_and_keep_colliding_groups(past) -> None:
    u = past["users"]["am-user1"]
    res = json.loads(_stored(past["id"])["result"])
    first = u.get(f"/analysis/runs/{past['id']}/groups", params={"limit": 2}).json()
    pages = _all(u, past["id"], "groups", limit=2)
    assert first["total"] == len(res["groups"]) == 6 and first["offset"] == 0 and len(first["items"]) == 2
    same(pages, expect(res["groups"]), "groups")
    same(_all(u, past["id"], "groups"), pages, "groups(한 번에)")
    no_raw(json.dumps(pages, ensure_ascii=False), "groups")
    keys = [g["keys"][0] for g in pages]
    assert keys.count("기관k***@example.com") == 3 and "=1+1 k***@example.com" in keys
    assert [g["sum"] for g in pages] == [g["sum"] for g in res["groups"]]  # 순서·수치 그대로
    assert sorted(g["bin"] for g in pages if g["bin_index"] is not None) == ["구간k***@example.com"] * 4
    bins = u.get(f"/analysis/runs/{past['id']}").json()["result"]["bins"]
    assert [(b["index"], b["groups"]) for b in bins] == [(b["index"], b["groups"]) for b in res["bins"]]


@pytest.mark.parametrize("kind", ["excluded", "detail"])
def test_row_pages_mask_names_and_keep_rows(past, kind) -> None:
    u = past["users"]["am-user1"]
    res = json.loads(_stored(past["id"])["result"])
    first = u.get(f"/analysis/runs/{past['id']}/{kind}", params={"limit": 1}).json()
    rows = _all(u, past["id"], kind, limit=1)
    assert first["all"] == res[f"{kind}_total"] and first["total"] == len(res[kind]) and len(first["items"]) == 1
    same(rows, expect(res[kind]), kind)
    no_raw(json.dumps(rows, ensure_ascii=False), kind)
    # 서로 다른 시트의 같은 행 번호가 가린 뒤 같은 이름이 되어도 두 행 모두 남는다
    seen = [(x["file_id"], x["sheet"], x["row"]) for x in rows]
    assert len(seen) == len(res[kind]) and len(set(seen)) < len(seen)


def test_download_uses_the_same_masked_values_in_all_four_sheets(past) -> None:
    u = past["users"]["am-user1"]
    rid = past["id"]
    raw = u.get(f"/analysis/runs/{rid}/download")
    assert raw.status_code == 200
    with zipfile.ZipFile(io.BytesIO(raw.content)) as z:
        for name in z.namelist():
            xml = z.read(name).decode("utf-8")
            no_raw(xml, name)
            assert not re.search(r"<f[ >]", xml), f"{name}: 수식"
    wb = openpyxl.load_workbook(io.BytesIO(raw.content))
    assert wb.sheetnames == ["결과", "조건", "제외", "근거 행"]
    detail = u.get(f"/analysis/runs/{rid}").json()
    groups, excluded, rows = (_all(u, rid, k) for k in ("groups", "excluded", "detail"))
    aggs = detail["result"]["aggregates"]
    table = list(wb["결과"].iter_rows(min_row=2))
    assert len(table) == len(groups)
    for cells, g in zip(table, groups, strict=True):
        values = [c.value for c in cells]
        assert values[0] == g["keys"][0] and values[1 + len(aggs)] == g["bin"]
        for j, a in enumerate(aggs):
            cell = cells[1 + j]
            if g[a] is None:
                assert cell.value == "계산할 값 없음" and cell.data_type == "s"
            else:
                assert cell.value == g[a] and cell.data_type == "n"
    formula_like = [c for c in wb["결과"]["A"] if isinstance(c.value, str) and c.value.startswith("=")]
    assert [c.value for c in formula_like] == ["=1+1 k***@example.com"] and formula_like[0].data_type == "s"
    near = [list(r) for r in wb["근거 행"].iter_rows(min_row=2, values_only=True)]
    assert near == [[x["file_id"], x["file_name"], x["sheet"], x["row"], x["period"], *x["keys"], x["value"]]
                    for x in rows]  # fmt: skip
    listed = [list(r) for r in wb["제외"].iter_rows(values_only=True)]
    start = next(i for i, r in enumerate(listed) if r[0] == "파일 번호")
    assert [(r[0], r[1], r[2], r[3], r[4], r[6]) for r in listed[start + 1 :]] == [
        (x["file_id"], x["file_name"], x["sheet"], x["row"], x["period"], x["value"]) for x in excluded
    ]
    whole = next(r for r in listed if r[0] == detail["result"]["excluded_targets"][0]["file_id"])
    assert whole[1:3] == [FILE_NAME.replace("_kim", "_***"), "설명_010-****-5432"]
    cond = {r[0]: r[1] for r in wb["조건"].iter_rows(values_only=True) if r[0] and r[1] is not None}
    c = detail["conditions"]
    assert cond["열 이름의 구분"] == c["variant"] and cond["단위"] == c["unit"] and cond["분석 번호"] == rid
    assert cond["생성 시각 (KST)"] == detail["created_at"][:19].replace("T", " ")
    assert "기관k***@example.com, 기관k***@example.com, 010-****-5678" in cond["필터 — 지역·기관·구분"]
    assert cond["지역 수준 차이"].startswith("2026-09 서울특별시 · 010-****-5678: 상위 값 100")
    assert cond["반복 행"].endswith(f"#{past['fid']} 현황k***@example.com")
    labels = [r[0] for r in wb["조건"].iter_rows(values_only=True) if str(r[0]).startswith("구간 '")]
    assert labels == ["구간 '구간k***@example.com'", "구간 '구간k***@example.com'"]
    files = [r for r in wb["조건"].iter_rows(values_only=True) if r[0] == past["fid"]]
    assert [(r[2], r[3], r[9]) for r in files] == [
        ("현황k***@example.com", "인구수k***@example.com", HEX),
        ("현황k***@example.com", "인구수_***@example.com", HEX),
        ("설명_010-****-5432", "-", HEX),
    ]
    assert "900101-1****** → 해당 없음" in files[0][10]


# ---------------------------------------------------------------- 권한·되풀이·부작용·호환


def test_permissions_unchanged_and_admin_sees_the_same_masking(past, app_real_login) -> None:  # noqa: F811
    rid = past["id"]
    mine, other, admin = past["users"]["am-user1"], past["users"]["am-user2"], past["owner"]
    anon = _client(app_real_login)
    for path in ("", "/groups", "/excluded", "/detail"):
        a, b = mine.get(f"/analysis/runs/{rid}{path}").json(), admin.get(f"/analysis/runs/{rid}{path}").json()
        if path == "":
            assert (a.pop("read_only"), b.pop("read_only")) == (False, True)
        same(b, a, f"관리자{path}")
    wa = openpyxl.load_workbook(io.BytesIO(mine.get(f"/analysis/runs/{rid}/download").content))
    wb = openpyxl.load_workbook(io.BytesIO(admin.get(f"/analysis/runs/{rid}/download").content))
    for name in wa.sheetnames:
        assert list(wa[name].values) == list(wb[name].values), name
    for path in PATHS:
        assert other.get(f"/analysis/runs/{rid}{path}").status_code == 404, path
        assert anon.get(f"/analysis/runs/{rid}{path}").status_code == 401, path
    assert all(x["id"] != rid for x in other.get("/analysis/runs").json())
    assert all(x["id"] != rid for x in admin.get("/analysis/runs").json())  # 목록은 본인 것만 (기존 그대로)
    assert anon.get("/analysis/runs").status_code == 401


def test_repeated_views_are_idempotent_without_recompute_parse_or_model(past, monkeypatch) -> None:
    from app import analysis, analysis_api, files, mapping, mapping_review, parsing

    def boom(*a: Any, **k: Any) -> Any:
        raise AssertionError("저장 결과 조회가 다시 계산·파싱·모델 호출을 했음")

    for mod, name in ((analysis_api, "evaluate"), (analysis_api, "compute"), (analysis_api, "conditions"),
                      (analysis_api, "inputs"), (analysis, "file_view"), (mapping_review, "file_view"),
                      (files, "load_parsed"), (mapping_review, "load_parsed"), (parsing, "parse_file"),
                      (mapping, "get_engine")):  # fmt: skip
        monkeypatch.setattr(mod, name, boom, raising=False)
    u = past["users"]["am-user1"]
    rid = past["id"]
    before = _stored(rid)
    count = _db("select count(*) from analysis_runs")[0][0]

    def views() -> list[Any]:
        out: list[Any] = [u.get("/analysis/runs").json()]
        for path in ("", "/groups", "/excluded", "/detail"):
            out.append(u.get(f"/analysis/runs/{rid}{path}", params={"limit": 3} if path else None).json())
        wb = openpyxl.load_workbook(io.BytesIO(u.get(f"/analysis/runs/{rid}/download").content))
        out.append({n: [[(c.value, c.data_type) for c in r] for r in wb[n].iter_rows()] for n in wb.sheetnames})
        return out

    one, two = views(), views()
    same(two, one, "다시 조회")
    assert _stored(rid) == before  # 요약·조건·입력·결과·생성 시각 그대로
    assert _db("select count(*) from analysis_runs")[0][0] == count


def test_normal_result_is_unchanged(past) -> None:
    """개인정보 형식이 없는 결과는 저장값 그대로 (가림이 아무것도 바꾸지 않음)."""
    u = past["users"]["am-user1"]
    rid = past["plain"]
    st = _stored(rid)
    body = u.get(f"/analysis/runs/{rid}").json()
    res = json.loads(st["result"])
    same(body["conditions"], json.loads(st["conditions"]), "conditions")
    same(body["inputs"], json.loads(st["inputs"]), "inputs")
    same(_all(u, rid, "groups"), res["groups"], "groups")
    same(_all(u, rid, "detail"), res["detail"], "detail")
    (item,) = [x for x in u.get("/analysis/runs").json() if x["id"] == rid]
    same({k: v for k, v in item.items() if k not in ("id", "created_at")}, json.loads(st["summary"]), "summary")


def test_older_results_without_optional_fields_are_not_filled(past, request) -> None:
    """선택 필드가 없던 예전 결과: 없는 칸을 지금 매핑·규칙으로 채우지 않고 그대로 보여 준다."""
    u = past["users"]["am-user1"]
    rid = _run(u, _body([_file(u, "예전.xlsx", THREE, {"시도": "sido", "인구수": "population"})]))["id"]
    st = _stored(rid)
    request.addfinalizer(lambda: _restore(rid, st))
    cond, inp = json.loads(st["conditions"]), json.loads(st["inputs"])
    for k in ("region_level", "level_conflicts", "base_dates", "records"):
        cond.pop(k, None)
    for t in inp:
        t.pop("repeated", None)
        t.pop("rows_are_records", None)
        t["column"] = GLUE
    _put_stored(rid, conditions=cond, inputs=inp)
    body = u.get(f"/analysis/runs/{rid}").json()
    assert not {"region_level", "level_conflicts", "base_dates", "records"} & set(body["conditions"])
    assert all("repeated" not in t and "rows_are_records" not in t for t in body["inputs"])
    assert [t["column"] for t in body["inputs"]] == ["인구수k***@example.com"]
    res = u.get(f"/analysis/runs/{rid}/download")
    assert res.status_code == 200 and GLUE not in json.dumps(
        [list(ws.values) for ws in openpyxl.load_workbook(io.BytesIO(res.content))], ensure_ascii=False, default=str
    )


def test_masking_does_not_touch_stored_copy(past) -> None:
    """가림은 내보낼 사본에만 한다 — 저장 JSON(원문 포함)은 바이트까지 그대로."""
    rid = past["id"]
    before = copy.deepcopy(_stored(rid))
    u = past["users"]["am-user1"]
    for path in PATHS:
        assert u.get(f"/analysis/runs/{rid}{path}").status_code == 200
    assert _stored(rid) == before and GLUE in before["inputs"]
