"""개인정보 형식이 든 파일·시트·열 이름과 개인정보 탐지 위치(pii.items)의 출력 경계.

업로드·파일 목록·개인정보 확인·미리보기·묶기 결과 구성 확인(plan) 응답에 이름 원문이 나오지 않는지 본다.
다른 사용자 권한, 저장값(파일 이름·탐지 결과·확인 상태), 포함을 확인한 결과 데이터 시트의 원본 값은 그대로인지도 본다.
과거에 저장된 탐지 결과(원래 시트 이름이 든 항목)도 내보낼 때 지금 규칙으로 가린다. 합성 자료만 쓴다.
"""

import io
import json
from dataclasses import asdict
from typing import Any

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app.pii import mask_text
from tests.conftest import db_connect
from tests.test_accounts import owner  # noqa: F401
from tests.test_auth import _client, _provision_user, app_real_login  # noqa: F401

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
FILE = "kjp1@example.test.xlsx"  # 파일 이름 — 이메일 형식
SHEET = "kjp2@example.test"  # 시트 이름 — 이메일 형식
COLUMN = "010-7701-4455"  # 열 이름 — 휴대전화 형식
EMAIL_CELL, PHONE_CELL = "kjp3@example.test", "010-7702-5566"  # 칸 값
INFO_SHEET = "설명 kjp4@example.test"  # 앱이 읽지 않는 설명 시트 (건너뛴 시트 목록에만 나옴)
RAW = ("kjp1@", "kjp2@", "kjp3@", "kjp4@", "7701-4455", "7702-5566")  # 원문 일부라도 응답에 나오면 안 됨
HEAD = ["시도", "시군구", COLUMN, "담당 메일"]
ROWS = [["서울특별시", "종로구", 10, EMAIL_CELL], ["서울특별시", "용산구", 20, PHONE_CELL]]
SHOWN_FILE, SHOWN_SHEET = "k***@example.test.xlsx", "k***@example.test"
RESULT_SHEET = "k___@example.test"  # 결과 워크북의 시트 이름: 가린 이름의 금지 글자(*)를 _로 (bundle.FORBIDDEN)
# 탐지 위치: 머리글의 열 이름(1행 C), 칸 값 두 개 — 종류·행·열 문자·가린 값은 가리기 전과 같아야 한다
FOUND = [
    ("mobile", 1, "C", "010-****-4455"),
    ("email", 2, "D", "k***@example.test"),
    ("mobile", 3, "D", "010-****-5566"),
]
# 가림을 두 번 적용해도 같아야 하는 값 (합성 이름·한글에 붙은 이메일·주민등록번호 형식·위치 설명 글자)
TWICE = [FILE, SHEET, COLUMN, EMAIL_CELL, PHONE_CELL, INFO_SHEET, "인구kjp2@example.test", "900101-1234567 명부",
         f"{SHEET} 시트 D2 ({EMAIL_CELL})"]  # fmt: skip


def _xlsx() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET
    for row in [HEAD, *ROWS]:
        ws.append(row)
    wb.create_sheet(INFO_SHEET).append(["합성 시험 자료입니다."])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _raw_in(text: str) -> list[str]:
    return [r for r in RAW if r in text]


def _upload(c: TestClient, name: str = FILE) -> dict[str, Any]:
    res = c.post("/files", files=[("files", (name, _xlsx(), XLSX))])
    assert res.status_code == 200, res.text
    body = res.json()[0]
    assert body["id"] and body["error"] is None, body
    return body


def _db(sql: str, *args: Any) -> list[tuple]:
    with db_connect() as con:
        return con.execute(sql, args).fetchall()


def _db_write(sql: str, *args: Any) -> None:
    with db_connect() as con:
        con.execute(sql, args)


def _found(items: list[dict[str, Any]]) -> list[tuple]:
    return [(i["type"], i["row"], i["column"], i["masked"]) for i in items]


@pytest.fixture(scope="module")
def px(app_real_login, owner) -> dict[str, TestClient]:  # noqa: F811
    made = {"owner": owner}
    for name in ("pio-user1", "pio-user2"):
        _provision_user(owner, app_real_login, name, "pass-" + name, name.upper())
        made[name] = _client(app_real_login, name, "pass-" + name)
    return made


# ---------------------------------------------------------------- API 응답 (TestClient)


def test_upload_response_masks_names_and_pii_locations(px) -> None:
    """업로드 응답: 파일 이름·시트 이름·건너뛴 설명 시트·열 이름·탐지 위치의 시트 이름을 가린다 (식별은 id·순번)."""
    res = px["pio-user1"].post("/files", files=[("files", (FILE, _xlsx(), XLSX))])
    assert res.status_code == 200
    (body,) = res.json()
    assert _raw_in(res.text) == []
    assert body["original_name"] == SHOWN_FILE
    assert [s["name"] for s in body["sheets"]] == [SHOWN_SHEET]
    assert body["skipped_sheets"] == ["설명 k***@example.test"]
    assert body["sheets"][0]["columns"] == ["시도", "시군구", "010-****-4455", "담당 메일"]
    assert body["pii"]["counts"] == {"rrn": 0, "mobile": 2, "email": 1} and body["pii"]["acknowledged"] is False
    assert [i["sheet"] for i in body["pii"]["items"]] == [SHOWN_SHEET] * 3
    assert _found(body["pii"]["items"]) == FOUND  # 종류·행·열 문자·가린 값은 그대로


def test_upload_refusals_do_not_echo_the_raw_file_name(px) -> None:
    """저장 전에 거절한 업로드(형식·이름 길이)도 같은 출력 경계를 지난다 — 거절 문구에도 이름 원문이 없다."""
    long_name = FILE.replace(".xlsx", "가" * 260 + ".xlsx")
    res = px["pio-user1"].post(
        "/files",
        files=[("files", ("kjp1@example.test.txt", b"x", "text/plain")), ("files", (long_name, _xlsx(), XLSX))],
    )
    assert res.status_code == 200
    bad_ext, too_long = res.json()
    assert bad_ext["id"] is None and bad_ext["error"] and bad_ext["original_name"] == "k***@example.test.txt"
    assert too_long["id"] is None and "NAME_TOO_LONG" in too_long["error"]
    assert too_long["original_name"] == mask_text(long_name)
    assert _raw_in(res.text) == []


def test_file_list_original_name_is_the_masked_display_name(px) -> None:
    """파일 목록: 호환 필드 original_name도 display_name과 같은 가린 이름, 탐지 위치의 시트 이름도 가린다.
    DB의 원래 이름·탐지 결과는 그대로 (화면 응답만 가림)."""
    c = px["pio-user1"]
    fid = _upload(c)["id"]
    res = c.get("/files")
    assert res.status_code == 200
    item = next(f for f in res.json() if f["id"] == fid)
    assert _raw_in(res.text) == []
    assert item["original_name"] == item["display_name"] == SHOWN_FILE
    assert [i["sheet"] for i in item["pii"]["items"]] == [SHOWN_SHEET] * 3 and _found(item["pii"]["items"]) == FOUND
    ((name, summary, acknowledged),) = _db(
        "select original_name, pii_summary, pii_acknowledged_at from source_files where id = ?", fid
    )
    assert name == FILE and acknowledged is None
    assert [i["sheet"] for i in json.loads(summary)["items"]] == [SHEET] * 3


def test_acknowledge_response_masks_pii_locations(px) -> None:
    """개인정보 포함 확인 응답의 탐지 위치도 가린다. 확인 시각만 기록하고 탐지 결과(저장값)는 바꾸지 않는다."""
    c = px["pio-user1"]
    fid = _upload(c)["id"]
    res = c.post(f"/files/{fid}/pii/acknowledge")
    assert res.status_code == 200, res.text
    body = res.json()
    assert _raw_in(res.text) == []
    assert body["acknowledged"] is True and [i["sheet"] for i in body["items"]] == [SHOWN_SHEET] * 3
    assert _found(body["items"]) == FOUND
    ((summary, acknowledged),) = _db("select pii_summary, pii_acknowledged_at from source_files where id = ?", fid)
    assert acknowledged is not None and [i["sheet"] for i in json.loads(summary)["items"]] == [SHEET] * 3


def test_preview_and_mapping_responses_have_no_raw_names(px) -> None:
    """미리보기·열 매핑 응답 (예전부터 가림 — 같은 합성 이름으로 비교해 둔다)."""
    c = px["pio-user1"]
    fid = _upload(c)["id"]
    for res in (c.get(f"/files/{fid}/preview"), c.post(f"/files/{fid}/mapping"), c.get(f"/files/{fid}/mapping/latest")):
        assert res.status_code == 200, res.text
        assert _raw_in(res.text) == [], str(res.url)
    preview = c.get(f"/files/{fid}/preview").json()
    assert (preview["original_name"], preview["sheet"]) == (SHOWN_FILE, SHOWN_SHEET)
    assert [col["name"] for col in preview["columns"]] == ["시도", "시군구", "010-****-4455", "담당 메일"]


def test_bundle_plan_pii_summary_matches_the_masked_file_list(px) -> None:
    """묶기 결과 구성 확인(plan)의 pii_summary는 파일 목록과 같은 형식·같은 가린 값이다."""
    c = px["pio-user1"]
    fid = _upload(c)["id"]
    res = c.post("/bundle/plan", json={"file_ids": [fid]})
    assert res.status_code == 200, res.text
    assert _raw_in(res.text) == []
    (f,) = res.json()["files"]
    assert f["pii"] == "needs_ack" and [i["sheet"] for i in f["pii_summary"]["items"]] == [SHOWN_SHEET] * 3
    listed = next(x for x in c.get("/files").json() if x["id"] == fid)
    assert f["pii_summary"] == listed["pii"]


def test_other_users_permissions_are_unchanged(px) -> None:
    """다른 일반 사용자에게는 없는 파일(404), 최고 관리자는 조회만(미리보기·plan 200, 개인정보 확인 403) — 그대로.
    거절·조회 응답 어디에도 이름 원문이 없고, 남의 확인 요청으로 확인 상태가 바뀌지 않는다."""
    u1, u2, admin = px["pio-user1"], px["pio-user2"], px["owner"]
    fid = _upload(u1)["id"]
    assert fid not in {f["id"] for f in u2.get("/files").json()}
    assert fid not in {f["id"] for f in admin.get("/files").json()}  # 목록은 본인 파일(과 계정 이전 파일)만
    for res in (
        u2.post(f"/files/{fid}/pii/acknowledge"),
        u2.get(f"/files/{fid}/preview"),
        u2.post("/bundle/plan", json={"file_ids": [fid]}),
    ):
        assert res.status_code == 404, res.text
        assert _raw_in(res.text) == []
    preview = admin.get(f"/files/{fid}/preview")
    assert preview.status_code == 200 and _raw_in(preview.text) == []
    ack = admin.post(f"/files/{fid}/pii/acknowledge")
    assert ack.status_code == 403 and _raw_in(ack.text) == []
    plan = admin.post("/bundle/plan", json={"file_ids": [fid]})
    assert plan.status_code == 200, plan.text
    (f,) = plan.json()["files"]
    assert f["own"] is False and any("다른 사용자의 파일" in p for p in f["problems"])
    assert _raw_in(plan.text) == []
    ((acknowledged,),) = _db("select pii_acknowledged_at from source_files where id = ?", fid)
    assert acknowledged is None


def test_legacy_stored_pii_summary_is_masked_with_current_rules(px) -> None:
    """과거에 저장된 탐지 결과(원래 시트 이름, 예전 규칙이 놓친 한글에 붙은 이메일, 위치 설명 글자)도 내보낼 때
    지금 규칙으로 가린다. DB의 저장값은 바꾸지 않는다."""
    c = px["pio-user1"]
    fid = _upload(c)["id"]
    ((saved,),) = _db("select pii_summary from source_files where id = ?", fid)
    legacy = {
        "counts": {"rrn": 0, "mobile": 1, "email": 1},
        "items": [
            {"type": "mobile", "sheet": SHEET, "row": 1, "column": "C", "masked": "010-****-4455"},
            {"type": "email", "sheet": "인구kjp2@example.test", "row": 2, "column": "D", "masked": "k***@example.test",
             "where": f"{SHEET} 시트 D2 ({EMAIL_CELL})"},
        ],
    }  # fmt: skip
    stored = json.dumps(legacy, ensure_ascii=False)
    _db_write("update source_files set pii_summary = ? where id = ?", stored, fid)
    try:
        listed = c.get("/files")
        plan = c.post("/bundle/plan", json={"file_ids": [fid]})
        ack = c.post(f"/files/{fid}/pii/acknowledge")
        for res in (listed, plan, ack):
            assert res.status_code == 200, res.text
            assert _raw_in(res.text) == [], str(res.url)
        items = ack.json()["items"]
        assert [i["sheet"] for i in items] == [SHOWN_SHEET, "인구k***@example.test"]
        assert items[1]["where"] == "k***@example.test 시트 D2 (k***@example.test)"
        assert _found(items) == [("mobile", 1, "C", "010-****-4455"), ("email", 2, "D", "k***@example.test")]
        assert next(x for x in listed.json() if x["id"] == fid)["pii"]["items"] == items
        assert plan.json()["files"][0]["pii_summary"]["items"] == items
        ((after,),) = _db("select pii_summary from source_files where id = ?", fid)
        assert after == stored  # 저장값은 그대로
    finally:
        _db_write("update source_files set pii_summary = ? where id = ?", saved, fid)


def test_result_data_sheet_keeps_original_values_after_acknowledgement(px) -> None:
    """포함을 확인한 파일의 묶기 결과 데이터 시트는 원본 값(열 이름·칸 값)을 그대로 담는다 — 화면 가림과 다른 정책.
    결과 시트 이름·원본 메모와 결과 구성 확인·실행 응답은 가린 이름이다 (기존 그대로)."""
    c = px["pio-user1"]
    fid = _upload(c)["id"]
    assert c.post(f"/files/{fid}/pii/acknowledge").status_code == 200
    plan = c.post("/bundle/plan", json={"file_ids": [fid]})
    assert plan.status_code == 200 and plan.json()["ready"], plan.text
    assert _raw_in(plan.text) == []
    run = c.post("/bundle/run", json={"file_ids": [fid], "expected_basis": plan.json()["basis"]})
    assert run.status_code == 201, run.text
    assert _raw_in(run.text) == []
    dl = c.get(f"/runs/{run.json()['id']}/download")
    assert dl.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(dl.content))
    assert wb.sheetnames == [RESULT_SHEET]  # 파일 이름(확장자 제외)을 가린 값에서 금지 글자만 바꾼 것 (기존 규칙)
    ws = wb[RESULT_SHEET]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    assert rows[0] == ["원본 행", "행 종류", *HEAD]  # 결과 데이터: 열 이름 원문
    assert [r[2:] for r in rows[1:]] == ROWS  # 결과 데이터: 칸 값 원문
    assert _raw_in(ws["A1"].comment.text) == []


# ---------------------------------------------------------------- 함수 (가림 규칙)


def _parsed_items() -> list[dict[str, Any]]:
    """실제 파서가 저장하는 형식의 탐지 위치 (원래 시트 이름)."""
    import tempfile
    from pathlib import Path

    from app.parsing import parse_file

    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "x.xlsx"
        path.write_bytes(_xlsx())
        return [asdict(i) for i in parse_file(path, FILE).pii.items]


def test_shown_pii_items_masks_every_text_and_keeps_positions() -> None:
    from app.pii import shown_pii_items

    items = _parsed_items()
    assert [i["sheet"] for i in items] == [SHEET] * 3  # 저장 형식은 원래 시트 이름
    legacy = {"type": "email", "sheet": "명부", "row": 7, "column": "B", "masked": "k***@example.test",
              "where": f"{SHEET} B7"}  # fmt: skip
    shown = shown_pii_items([*items, legacy])
    assert [i["sheet"] for i in shown] == [SHOWN_SHEET] * 3 + ["명부"]
    assert _found(shown) == [*FOUND, ("email", 7, "B", "k***@example.test")]
    assert shown[3]["where"] == "k***@example.test B7"
    assert items[0]["sheet"] == SHEET and legacy["where"] == f"{SHEET} B7"  # 넘겨받은 값은 바꾸지 않음
    assert shown_pii_items(None) == [] and shown_pii_items({"sheet": SHEET}) == []


@pytest.mark.parametrize("text", TWICE)
def test_mask_text_is_idempotent_on_masked_names(text: str) -> None:
    """가린 값을 한 번 더 가려도 같다 — 공용 출력 경계에서 가린 탐지 위치를 다른 화면이 다시 가려도 표시가 같다."""
    once = mask_text(text)
    assert once != text and mask_text(once) == once


def test_masked_pii_items_are_unchanged_by_a_second_masking() -> None:
    """공용 출력 경계(shown_pii_items)에서 가린 항목에 시트 이름만 다시 가리는 처리(공통 key 연결 plan)를 더해도
    같다."""
    from app.pii import shown_pii_items

    once = shown_pii_items(_parsed_items())
    assert shown_pii_items(once) == once
    assert [{**it, "sheet": mask_text(it["sheet"])} for it in once] == once
