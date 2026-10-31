"""공통 key 연결: 실제 업로드(CSV·XLSX·XLS) → 연결 확인(plan) → 만들기(run) → 결과 엑셀·실행 기록·내려받기로 확인한다.

- 6-1 합성 자료의 기대값은 제품 함수를 부르지 않고 손으로 센 상수다 (EXPECTED_*).
- 6-2 경계: 다른 key 열 이름·같은 파일/열 이름·여러 시트·행 순서·복합 key·'001'과 1·긴 숫자·대소문자·공백·빈/중복 key·
  공통 없음·일부에만·가림 충돌·개인정보 확인·다른 계정·제어 문자·긴 이름·수식 글자·plan 뒤 변경·중복 생성 요청·410·예산.
- 거절(422·404·409·503)은 실행 기록·결과 파일을 남기지 않는다 (_state로 단언). 합성 자료만 쓴다.
"""

import csv
import hashlib
import io
import json
import threading
import uuid
from datetime import datetime
from typing import Any

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app.config import DATA_DIR
from app.integration import OUTPUT_DIR
from tests.conftest import db_connect, db_dump
from tests.test_accounts import owner  # noqa: F401
from tests.test_auth import _client, _provision_user, app_real_login  # noqa: F401

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
XLS = "application/vnd.ms-excel"
CSV = "text/csv"

# ---------------------------------------------------------------- 6-1 합성 자료 (설계 7.1) — 기대값은 손으로 센 상수
A_ROWS = [["사번", "부서", "비고"], ["001", "총무", "메모A1"], ["002", "기획", "메모A2"], ["003", "정보", "메모A3"]]
B_ROWS = [
    ["사용자번호", "교육 건수", "메모"],
    ["001", "10", "메모-b1"],
    ["003", "30", "메모-b3"],
    ["004", "40", "메모-b4"],
]
C_ROWS = [["사용자ID", "출장 건수", "기타"], ["001", "1", "기타-c1"], ["002", "2", "기타-c2"], ["005", "5", "기타-c5"]]
D_ROWS = [
    ["직원코드", "처리 건수", "설명"],
    ["001", "100", "설명-d1"],
    ["003", "300", "설명-d3"],
    ["006", "600", "설명-d6"],
]
EXPECTED_HEADER = ["사번", "부서", "교육 건수", "출장 건수", "처리 건수"]
EXPECTED_BASE = [("001", "총무", 10, 1, 100), ("002", "기획", None, 2, None), ("003", "정보", 30, None, 300)]
EXPECTED_ALL = [
    *EXPECTED_BASE,
    ("004", None, 40, None, None),
    ("005", None, None, 5, None),
    ("006", None, None, None, 600),
]
EXPECTED_COMMON = [("001", "총무", 10, 1, 100)]
EXPECTED_SOURCES_BASE = [[1, 2, 2, 2, 2], [2, 3, None, 3, None], [3, 4, 3, None, 3]]
EXPECTED_PREVIEW_BASE = [
    {"key": ["001"], "cells": ["총무", 10, 1, 100],
     "sources": [{"input": 0, "row": 2}, {"input": 1, "row": 2}, {"input": 2, "row": 2}, {"input": 3, "row": 2}]},
    {"key": ["002"], "cells": ["기획", None, 2, None], "sources": [{"input": 0, "row": 3}, {"input": 2, "row": 3}]},
    {"key": ["003"], "cells": ["정보", 30, None, 300], "sources": [{"input": 0, "row": 4}, {"input": 1, "row": 3},
                                                                 {"input": 3, "row": 3}]},
]  # fmt: skip
# 선택하지 않은 열의 이름·값 — 워크북 어디에도(모든 시트·메모·머리글), plan JSON·실행 기록에도 없어야 한다
UNPICKED = [
    "비고",
    "메모",
    "기타",
    "설명",
    "메모A1",
    "메모A2",
    "메모A3",
    "메모-b1",
    "메모-b3",
    "메모-b4",
    "기타-c1",
    "기타-c2",
    "기타-c5",
    "설명-d1",
    "설명-d3",
    "설명-d6",
]
SHEETS = ["연결 결과", "출처", "연결 요약"]


# ---------------------------------------------------------------- 합성 파일 만들기·올리기


def _csv_bytes(rows: list[list[Any]]) -> bytes:
    buf = io.StringIO()
    csv.writer(buf, lineterminator="\n").writerows(rows)
    return buf.getvalue().encode("utf-8-sig")


def _xlsx_bytes(sheets: dict[str, list[list[Any]]], merges: dict[str, tuple[str, ...]] | None = None) -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
        for m in (merges or {}).get(name, ()):
            ws.merge_cells(m)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _xls_bytes(rows: list[list[Any]], sheet: str = "자료") -> bytes:
    """xls 합성 자료 (글자 셀은 글자 그대로, 숫자는 xls 특성상 실수 칸)."""
    xlwt = pytest.importorskip("xlwt")
    wb = xlwt.Workbook(encoding="utf-8")
    ws = wb.add_sheet(sheet)
    for r, row in enumerate(rows):
        for j, v in enumerate(row):
            if v is not None:
                ws.write(r, j, v)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _upload(c: TestClient, name: str, data: bytes, media: str = XLSX) -> int:
    res = c.post("/files", files=[("files", (name, data, media))])
    assert res.status_code == 200 and res.json()[0]["id"], res.text
    return res.json()[0]["id"]


def _csv(c: TestClient, name: str, rows: list[list[Any]]) -> int:
    return _upload(c, name, _csv_bytes(rows), CSV)


def _xlsx(c: TestClient, name: str, rows: list[list[Any]], sheet: str = "자료") -> int:
    return _upload(c, name, _xlsx_bytes({sheet: rows}))


# ---------------------------------------------------------------- 요청·응답 도우미


def _input(
    fid: int,
    key: str | list[str] = "c0",
    cols: tuple[str, ...] = ("c1",),
    sheet: str = "s0",
    names: dict[str, str | None] | None = None,
    rows: dict[str, list[int]] | None = None,
) -> dict[str, Any]:
    keys = [key] if isinstance(key, str) else list(key)
    columns = [{"column_id": col, "out_name": (names or {}).get(col)} for col in cols]
    body: dict[str, Any] = {"file_id": fid, "sheet_id": sheet, "key_columns": keys, "columns": columns}
    if rows is not None:
        body["rows"] = {"include": rows.get("include", []), "exclude": rows.get("exclude", [])}
    return body


def _body(inputs: list[dict[str, Any]], scope: str | None = "base", base: dict[str, Any] | None = None,
          key_names: list[str] | None = None) -> dict[str, Any]:  # fmt: skip
    body: dict[str, Any] = {
        "inputs": inputs,
        "base": base or {"file_id": inputs[0]["file_id"], "sheet_id": inputs[0]["sheet_id"]},
    }
    if scope is not None:
        body["scope"] = scope
    if key_names is not None:
        body["key_names"] = key_names
    return body


def _plan(c: TestClient, body: dict[str, Any], code: int = 200) -> dict[str, Any]:
    res = c.post("/keyjoin/plan", json=body)
    assert res.status_code == code, res.text
    return res.json()


def _run(c: TestClient, body: dict[str, Any], basis: str, request_id: str | None = None) -> Any:
    return c.post("/keyjoin/run", json={**body, "expected_basis": basis, "request_id": request_id or uuid.uuid4().hex})


def _make(c: TestClient, body: dict[str, Any]) -> tuple[dict[str, Any], openpyxl.Workbook, dict[str, Any]]:
    plan = _plan(c, body)
    assert plan["ready"], plan["problems"]
    res = _run(c, body, plan["basis"])
    assert res.status_code == 201, res.text
    run = res.json()
    dl = c.get(f"/runs/{run['id']}/download")
    assert dl.status_code == 200
    return run, openpyxl.load_workbook(io.BytesIO(dl.content)), plan


def _db(sql: str, *args: Any) -> list[tuple]:
    with db_connect() as con:
        cur = con.execute(sql, args)
        return cur.fetchall() if cur.description is not None else []  # INSERT는 돌려줄 행이 없다 (PostgreSQL은 오류)


def _state() -> dict[str, Any]:
    """실행 기록 수·'실행 중' 수와 결과 폴더의 파일 — 거절 전후가 같아야 한다 (담당 표시 파일 제외)."""
    runs = _db("SELECT count(*), sum(CASE WHEN status = 'running' THEN 1 ELSE 0 END) FROM runs")[0]
    files = sorted(p.name for p in OUTPUT_DIR.glob("*") if not p.name.startswith(".")) if OUTPUT_DIR.is_dir() else []
    return {"runs": runs[0], "running": runs[1] or 0, "files": files}


def _grid(ws: Any) -> list[list[Any]]:
    return [[c.value for c in row] for row in ws.iter_rows()]


def _texts(wb: openpyxl.Workbook) -> str:
    """워크북 전체(모든 시트의 값·메모·시트 이름)를 글자 하나로 — '없어야 할 것' 검사용."""
    parts = list(wb.sheetnames)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None:
                    parts.append(str(cell.value))
                if cell.comment is not None:
                    parts.append(cell.comment.text)
    return "\n".join(parts)


def _codes(plan: dict[str, Any], level: str = "problems") -> list[str]:
    return [p["code"] for p in plan[level]]


def _first(plan: dict[str, Any], code: str, level: str = "problems") -> dict[str, Any]:
    return next(p for p in plan[level] if p["code"] == code)


@pytest.fixture(scope="module")
def kj(app_real_login, owner) -> dict[str, TestClient]:  # noqa: F811
    made = {"owner": owner}
    for name in ("kj-user1", "kj-user2"):
        _provision_user(owner, app_real_login, name, "pass-" + name, name.upper())
        made[name] = _client(app_real_login, name, "pass-" + name)
    return made


@pytest.fixture(scope="module")
def abcd(kj) -> dict[str, Any]:
    """6-1 합성 자료 A.csv·B.xlsx·C.xls·D.csv (kj-user1 소유)와 기준 A의 요청 본문."""
    u = kj["kj-user1"]
    a = _csv(u, "A.csv", A_ROWS)
    b = _upload(u, "B.xlsx", _xlsx_bytes({"교육": B_ROWS}))
    c = _upload(u, "C.xls", _xls_bytes(C_ROWS, "출장"), XLS)
    d = _csv(u, "D.csv", D_ROWS)
    inputs = [_input(a), _input(b), _input(c), _input(d)]
    return {"fids": [a, b, c, d], "inputs": inputs, "body": _body(inputs)}


# ================================================================ 6-1: 기대값 상수와 셀 단위 검사


def test_6_1_base_scope_matches_the_constants_cell_by_cell(kj, abcd) -> None:
    u = kj["kj-user1"]
    plan = _plan(u, abcd["body"])
    assert plan["ready"] and plan["mode"] == "keyjoin" and plan["policy"] == "keyjoin-policy-2"  # 2: 행 포함 정책
    assert plan["scope"] == "base" and plan["base"] == {"file_id": abcd["fids"][0], "sheet_id": "s0", "index": 0}
    assert plan["header"] == EXPECTED_HEADER
    assert plan["result"] == {"rows": 3, "key_columns": 1, "columns": 5, "union": 6, "intersection": 1, "excluded": 0}
    assert plan["preview"]["rows"] == EXPECTED_PREVIEW_BASE
    assert (plan["preview"]["shown"], plan["preview"]["total"], plan["preview"]["pool"]) == (3, 3, 200)
    names = [ip["key_columns"][0]["name"] for ip in plan["inputs"]]
    assert names == ["사번", "사용자번호", "사용자ID", "직원코드"]  # key 열 이름이 달라도 된다
    labels = zip(plan["inputs"], abcd["fids"], strict=True)
    assert all(ip["label"].startswith(f"입력 {i + 1} · #{f} ") for i, (ip, f) in enumerate(labels))
    assert plan["inputs"][2]["label"].endswith("C.xls / 출장 (1번째 시트)")
    kinds = ("text", "int", "float", "bool", "date", "empty", "filled", "missing_mark", "precision", "illegal")
    assert plan["inputs"][0]["key_columns"][0]["kinds"] == dict.fromkeys(kinds, 0) | {"text": 3}
    assert plan["inputs"][0]["rows"] == {"data": 3, "subtotal": 0, "note": 0, "empty": 0}
    assert [ip["keys"]["linked"] for ip in plan["inputs"]] == [3, 2, 2, 2]
    assert [ip["keys"]["unlinked"] for ip in plan["inputs"]] == [0, 1, 1, 1]
    assert [ip["picked"][0]["out_name"] for ip in plan["inputs"]] == EXPECTED_HEADER[1:]
    assert all(ip["picked"][0]["default_name"] for ip in plan["inputs"])
    assert _codes(plan, "warnings") == ["KEYJOIN_UNLINKED"] * 3 and plan["problems"] == []
    # 글자 추정: 머리글 19자 + key 열 '001'·'002'·'003' × 4입력 36자 + 부서 열 6자 = 61자 (건수 열은 숫자라 0)
    assert plan["budget"] == {"output_cells": 134, "index_cells": 18, "text_chars": 61, "cells_weighted": 152,
                              "text_weighted": 122}  # fmt: skip
    assert UNPICKED[0] not in json.dumps(plan, ensure_ascii=False) and "메모A1" not in json.dumps(
        plan, ensure_ascii=False
    )
    before = _state()
    run, wb, _ = _make(u, abcd["body"])
    assert wb.sheetnames == SHEETS
    ws = wb["연결 결과"]
    assert [c.value for c in ws[1]] == EXPECTED_HEADER and ws.max_column == 5
    rows = list(ws.iter_rows(min_row=2))
    assert [tuple(c.value for c in r) for r in rows] == EXPECTED_BASE
    for r in rows:
        assert r[0].data_type == "s"  # key 칸은 항상 글자 (선행 0 보존)
        for c in r[1:]:
            assert (c.value is None and c.data_type == "n") or isinstance(c.value, int | str)
            if isinstance(c.value, int):
                assert c.data_type == "n"
    assert _grid(wb["출처"])[0] == [
        "결과 행",
        *[f"입력 {i + 1} · #{f} · 1번째 시트 원본 행" for i, f in enumerate(abcd["fids"])],
    ]
    assert _grid(wb["출처"])[1:] == EXPECTED_SOURCES_BASE
    text = _texts(wb)
    assert not any(x in text for x in UNPICKED)
    summary_cells = [v for row in _grid(wb["연결 요약"]) for v in row if v is not None]
    assert not any(v in ("001", "002", "003", "총무", "기획", "정보") for v in summary_cells)
    assert "key 열은 글자 형식" in text and "0 아님" in text and "선택하지 않은 열" in text
    # 실행 기록
    s = run["summary"]
    assert s["mode"] == "keyjoin" and s["mode_label"] == "공통 key 연결" and s["policy"] == "keyjoin-policy-2"
    assert (
        s["scope"] == "base" and s["input_count"] == 4 and len(s["inputs"]) == 4 and s["basis"] == plan_basis(u, abcd)
    )
    assert s["result"] == {"rows": 3, "columns": 5, "key_columns": 1} and s["counts"] == {"union": 6, "intersection": 1,
                                                                                            "excluded": 0}  # fmt: skip
    assert s["header"] == EXPECTED_HEADER and s["sheet_order"] == SHEETS and s["base"]["index"] == 0
    assert s["key_columns"] == [{"name": "사번", "letter": "A"}]
    assert s["inputs"][1]["picked"] == [{"column_id": "c1", "out_name": "교육 건수"}]
    assert [i["linked"] for i in s["inputs"]] == [3, 2, 2, 2] and s["inputs"][1]["file_name"] == "B.xlsx"
    assert (
        run["row_count"] == 3 and run["input_file_count"] == 4 and run["change_count"] == 0 and run["status"] == "done"
    )
    assert run["warning_count"] == 3 and run["file_name"].startswith("key연결_") and run["review_session_id"] is None
    assert json.loads(_db("SELECT file_ids FROM runs WHERE id = ?", run["id"])[0][0]) == abcd["fids"]
    assert not any(x in json.dumps(s, ensure_ascii=False) for x in UNPICKED)
    listed = {r["id"]: r for r in u.get("/runs").json()}
    assert listed[run["id"]]["summary"]["mode"] == "keyjoin" and listed[run["id"]]["status"] == "done"
    assert u.get(f"/runs/{run['id']}").json()["summary"]["sheet_order"] == SHEETS
    after = _state()
    assert after["runs"] == before["runs"] + 1 and len(after["files"]) == len(before["files"]) + 1


def plan_basis(u: TestClient, abcd: dict[str, Any]) -> str:
    return _plan(u, abcd["body"])["basis"]


def test_6_1_all_and_common_scopes(kj, abcd) -> None:
    u = kj["kj-user1"]
    _, wb, plan = _make(u, _body(abcd["inputs"], "all"))
    assert plan["result"]["rows"] == 6 and plan["result"]["union"] == 6
    assert [tuple(c.value for c in r) for r in wb["연결 결과"].iter_rows(min_row=2)] == EXPECTED_ALL
    assert [r["key"] for r in plan["preview"]["rows"]] == [[k[0]] for k in EXPECTED_ALL]
    assert [ip["keys"]["unlinked"] for ip in plan["inputs"]] == [3, 3, 3, 3]
    assert _grid(wb["출처"])[1:] == [*EXPECTED_SOURCES_BASE, [4, None, 4, None, None], [5, None, None, 4, None],
                                     [6, None, None, None, 4]]  # fmt: skip
    run, wb, plan = _make(u, _body(abcd["inputs"], "common"))
    assert plan["result"] == {"rows": 1, "key_columns": 1, "columns": 5, "union": 6, "intersection": 1, "excluded": 5}
    assert [tuple(c.value for c in r) for r in wb["연결 결과"].iter_rows(min_row=2)] == EXPECTED_COMMON
    assert (
        _first(plan, "KEYJOIN_EXCLUDED", "warnings")["count"] == 5
        and "5개" in _first(plan, "KEYJOIN_EXCLUDED", "warnings")["text"]
    )
    assert all(ip["keys"]["unlinked"] == 0 for ip in plan["inputs"])
    assert run["summary"]["counts"]["excluded"] == 5 and run["row_count"] == 1
    assert _plan(u, _body(abcd["inputs"], None))["scope"] == "base"  # scope 생략 = base
    other = _body(abcd["inputs"], "base", base={"file_id": abcd["fids"][1], "sheet_id": "s0"})
    plan = _plan(u, other)
    assert [r["key"] for r in plan["preview"]["rows"]] == [["001"], ["003"], ["004"]]  # 기준을 바꾸면 집합·순서가 바뀜
    assert plan["header"] == ["사용자번호", *EXPECTED_HEADER[1:]]


# ================================================================ 요청 형식 (422)


def test_request_format_errors_are_422_with_codes_and_no_echo(kj, abcd, monkeypatch) -> None:
    from app import keyjoin

    u = kj["kj-user1"]
    a, b = abcd["fids"][:2]
    before = _state()

    def bad(body: dict[str, Any], code: str) -> str:
        res = u.post("/keyjoin/plan", json=body)
        assert res.status_code == 422, (code, res.text)
        detail = res.json()["detail"]
        assert isinstance(detail, str) and detail.endswith(f"[{code}]"), (code, detail)
        return detail

    bad(_body([_input(a)]), "KEYJOIN_INPUTS_TOO_FEW")
    bad(_body([_input(a), _input(a)]), "KEYJOIN_INPUT_DUPLICATE")
    bad(_body([_input(a), _input(b, key=["c0", "c1"], cols=())]), "KEYJOIN_KEY_COUNT_MISMATCH")
    bad(_body([_input(a, key=["c0", "c0"]), _input(b, key=["c0", "c1"])]), "KEYJOIN_COLUMN_DUPLICATE")
    bad(_body([_input(a, cols=("c1", "c1")), _input(b)]), "KEYJOIN_COLUMN_DUPLICATE")
    bad(_body([_input(a, cols=("c0",)), _input(b)]), "KEYJOIN_KEY_IN_COLUMNS")
    bad(_body([_input(a), _input(b)], base={"file_id": 999_999, "sheet_id": "s0"}), "KEYJOIN_BASE_INVALID")
    bad(_body([_input(a), _input(b)], key_names=["x", "y"]), "KEYJOIN_KEY_NAMES_COUNT")
    bad(_body([_input(a, names={"c1": "  "}), _input(b)]), "KEYJOIN_NAME_BLANK")
    bad(_body([_input(a, names={"c1": "가" * 256}), _input(b)]), "KEYJOIN_NAME_TOO_LONG")
    secret = "비밀\x00이름"
    detail = bad(
        _body([_input(a, names={"c1": secret}), _input(b)]), "KEYJOIN_NAME_ILLEGAL"
    )  # NUL — 두 DB 같게, 저장 전
    assert "비밀" not in detail
    bad(_body([_input(a), _input(b)], key_names=["a\x01"]), "KEYJOIN_NAME_ILLEGAL")
    bad(_body([_input(a, sheet="s7"), _input(b)]), "KEYJOIN_SHEET_UNKNOWN")
    bad(_body([_input(a, cols=("c9",)), _input(b)]), "KEYJOIN_COLUMN_UNKNOWN")
    monkeypatch.setattr(keyjoin, "MAX_RESULT_COLUMNS", 2)
    bad(_body([_input(a, cols=("c1", "c2")), _input(b)]), "KEYJOIN_TOO_MANY_COLUMNS")
    monkeypatch.undo()
    assert u.post("/keyjoin/plan", json={**_body([_input(a), _input(b)]), "unknown": 1}).status_code == 422
    assert u.post("/keyjoin/plan", json={"inputs": [_input(a), _input(b)]}).status_code == 422  # base 필수
    missing = {k: v for k, v in _input(a).items() if k != "sheet_id"}
    base_b = {"file_id": b, "sheet_id": "s0"}
    assert u.post("/keyjoin/plan", json=_body([missing, _input(b)], base=base_b)).status_code == 422  # sheet_id 필수
    res = u.post(
        "/keyjoin/run", json={**_body([_input(a), _input(b)]), "expected_basis": "x" * 16, "request_id": "nope"}
    )
    assert res.status_code == 422
    assert _state() == before


# ================================================================ 같은 이름·여러 시트·행 순서·복합 key


def test_same_file_and_column_names_are_told_apart_and_values_are_not_overwritten(kj) -> None:
    u = kj["kj-user1"]
    rows1 = [["사번", "건수"], ["001", 1], ["002", 2]]
    rows2 = [["사번", "건수"], ["001", 10], ["002", 20]]
    f1, f2 = _xlsx(u, "건수.xlsx", rows1), _xlsx(u, "건수.xlsx", rows2)
    body = _body([_input(f1), _input(f2)])
    run, wb, plan = _make(u, body)
    assert plan["header"] == ["사번", "건수", f"건수 (입력 2 · #{f2})"]
    assert plan["inputs"][1]["picked"][0] == {"column_id": "c1", "name": "건수", "letter": "B",
                                              "out_name": f"건수 (입력 2 · #{f2})", "default_name": True}  # fmt: skip
    assert [tuple(c.value for c in r) for r in wb["연결 결과"].iter_rows(min_row=2)] == [("001", 1, 10), ("002", 2, 20)]
    # 사용자가 적은 이름끼리 같거나 key 이름과 같으면 차단, 사용자 이름과 같은 서버 기본 이름은 접미로 구분
    plan = _plan(u, _body([_input(f1, names={"c1": "같음"}), _input(f2, names={"c1": "같음"})]))
    assert (
        not plan["ready"]
        and _codes(plan) == ["KEYJOIN_OUT_NAME_DUPLICATE"]
        and _first(plan, "KEYJOIN_OUT_NAME_DUPLICATE")["input"] == 1
    )
    plan = _plan(u, _body([_input(f1, names={"c1": "사번"}), _input(f2)]))
    assert not plan["ready"] and _codes(plan) == ["KEYJOIN_OUT_NAME_DUPLICATE"]
    plan = _plan(u, _body([_input(f1, names={"c1": "건수"}), _input(f2)]))
    assert plan["ready"] and plan["header"] == ["사번", "건수", f"건수 (입력 2 · #{f2})"]
    plan = _plan(u, _body([_input(f1), _input(f2)], key_names=["직원"]))
    assert plan["ready"] and plan["header"][0] == "직원"
    # 같은 파일의 두 시트를 두 입력으로 (시트는 식별자로 구분, 라벨에 '2번째 시트')
    two = _upload(u, "상하반기.xlsx", _xlsx_bytes({"상반기": [["사번", "건수"], ["001", 5]],
                                                   "하반기": [["사번", "건수"], ["001", 7], ["002", 9]]}))  # fmt: skip
    body = _body([_input(two, sheet="s0"), _input(two, sheet="s1")])
    run, wb, plan = _make(u, body)
    assert "2번째 시트" in plan["inputs"][1]["label"] and plan["inputs"][1]["tab"] == 1 and run["input_file_count"] == 1
    assert [tuple(c.value for c in r) for r in wb["연결 결과"].iter_rows(min_row=2)] == [("001", 5, 7)]
    assert _grid(wb["출처"])[0][2] == f"입력 2 · #{two} · 2번째 시트 원본 행"
    assert plan["result"]["union"] == 2  # 빠진 파일(요청에 없는 입력)은 그냥 없는 것 — 서버가 따로 기억하지 않는다
    assert _plan(u, _body([_input(two, sheet="s1"), _input(f1)]))["result"]["rows"] == 2


def test_row_order_does_not_change_values_and_output_follows_the_base(kj) -> None:
    u = kj["kj-user1"]
    a = _csv(u, "순서a.csv", [["사번", "부서"], ["003", "정보"], ["001", "총무"], ["002", "기획"]])
    b1 = _xlsx(u, "순서b1.xlsx", [["사번", "건수"], ["001", 10], ["002", 20], ["003", 30]])
    b2 = _xlsx(u, "순서b2.xlsx", [["사번", "건수"], ["003", 30], ["002", 20], ["001", 10]])
    g1 = _grid(_make(u, _body([_input(a), _input(b1)]))[1]["연결 결과"])
    g2 = _grid(_make(u, _body([_input(a), _input(b2)]))[1]["연결 결과"])
    assert g1 == g2 == [["사번", "부서", "건수"], ["003", "정보", 30], ["001", "총무", 10], ["002", "기획", 20]]


def test_composite_key_is_a_tuple_and_counts_must_match(kj) -> None:
    u = kj["kj-user1"]
    a = _csv(u, "복합a.csv", [["기관", "사번", "값"], ["K1", "001", "a1"], ["K2", "001", "a2"]])  # 같은 사번, 다른 기관
    b = _csv(
        u, "복합b.csv", [["기관코드", "번호", "값"], ["K1", "001", "b1"], ["K2", "001", "b2"], ["K1", "002", "b3"]]
    )
    body = _body([_input(a, key=["c0", "c1"], cols=("c2",)), _input(b, key=["c0", "c1"], cols=("c2",))])
    run, wb, plan = _make(u, body)
    assert plan["header"] == ["기관", "사번", "값", f"값 (입력 2 · #{b})"] and plan["result"]["key_columns"] == 2
    assert _grid(wb["연결 결과"])[1:] == [["K1", "001", "a1", "b1"], ["K2", "001", "a2", "b2"]]
    assert plan["preview"]["rows"][0]["key"] == ["K1", "001"] and run["summary"]["result"]["key_columns"] == 2
    assert run["summary"]["key_columns"] == [{"name": "기관", "letter": "A"}, {"name": "사번", "letter": "B"}]
    swapped = _body([_input(a, key=["c0", "c1"], cols=("c2",)), _input(b, key=["c1", "c0"], cols=("c2",))])
    other = _plan(u, swapped)
    assert other["basis"] != plan["basis"] and other["inputs"][1]["keys"]["linked"] == 0  # 순서가 다르면 다른 key
    assert _grid(_make(u, swapped)[1]["연결 결과"])[1:] == [["K1", "001", "a1", None], ["K2", "001", "a2", None]]
    res = u.post("/keyjoin/plan", json=_body([_input(a, key=["c0", "c1"], cols=("c2",)), _input(b, cols=("c2",))]))
    assert res.status_code == 422 and "KEYJOIN_KEY_COUNT_MISMATCH" in res.json()["detail"]


# ================================================================ key 정확성: 정수/글자·선행 0·공백·대소문자·긴 숫자


def test_int_cell_links_to_text_123_but_not_to_00123_space_or_case_variants(kj) -> None:
    u = kj["kj-user1"]
    long1, long2 = "12345678901234567890123", "12345678901234567890124"  # 23자리, 끝자리만 다름
    a = _csv(u, "정확a.csv", [["id", "값"], ["123", "a123"], ["00123", "a00123"], [" 123", "a_sp"], ["abc", "aabc"],
                              ["ABC", "aABC"], [long1, "aL1"], [long2, "aL2"], ["001", "a001"]])  # fmt: skip
    b = _xlsx(u, "정확b.xlsx", [["id", "값"], [123, "b123"], ["abc", "babc"], [1, "b1"], [long1, "bL1"]])
    run, wb, plan = _make(u, _body([_input(a), _input(b)]))
    assert _grid(wb["연결 결과"])[1:] == [
        ["123", "a123", "b123"],  # 정수 칸 123 ↔ 글자 '123'
        ["00123", "a00123", None],  # '00123'은 다른 key
        [" 123", "a_sp", None],  # 앞 공백은 다른 key
        ["abc", "aabc", "babc"],
        ["ABC", "aABC", None],  # 대소문자 구분
        [long1, "aL1", "bL1"],  # 긴 글자형 숫자는 끝자리까지 보존
        [long2, "aL2", None],
        ["001", "a001", None],  # '001'과 정수 1은 다른 key
    ]
    ws = wb["연결 결과"]
    assert all(r[0].data_type == "s" for r in ws.iter_rows(min_row=2))
    warns = {p["code"]: p for p in plan["warnings"]}
    assert warns["KEYJOIN_INT_TEXT_LINK"]["count"] == 1 and warns["KEYJOIN_WHITESPACE_VARIANTS"]["count"] == 1
    assert "KEYJOIN_MIXED_KEY_KINDS" in warns and warns["KEYJOIN_MIXED_KEY_KINDS"]["input"] == 1
    kinds = plan["inputs"][1]["key_columns"][0]["kinds"]
    assert (kinds["int"], kinds["text"]) == (2, 2)
    # 공백 변형 묶음('123'·' 123')에 든 key가 있는 입력마다 1 (입력 2의 정수 123도 그 묶음의 '123'이다)
    assert [ip["keys"]["whitespace_variants"] for ip in plan["inputs"]] == [1, 1]
    assert not any(v in warns["KEYJOIN_INT_TEXT_LINK"]["text"] for v in ("a123", "b123", "abc", long1))
    assert run["warning_count"] == len(plan["warnings"])


def test_xls_numeric_key_links_to_xlsx_int_and_types_are_refused(kj) -> None:
    u = kj["kj-user1"]
    c = _upload(u, "숫자키.xls", _xls_bytes([["id", "값"], [123, "c123"], [1, "c1"], ["007", "c007"]]), XLS)
    b = _xlsx(u, "숫자키b.xlsx", [["id", "값"], [123, "b123"], ["001", "b001"], ["007", "b007"]])
    run, wb, plan = _make(u, _body([_input(c), _input(b)]))
    assert _grid(wb["연결 결과"])[1:] == [["123", "c123", "b123"], ["1", "c1", None], ["007", "c007", "b007"]]
    assert plan["inputs"][0]["key_columns"][0]["kinds"]["int"] == 2  # xls 숫자 칸(실수)도 정수 경로
    bad = _upload(u, "자료형.xlsx", _xlsx_bytes({"k": [["id", "값"], [datetime(2026, 1, 1), 7], [True, 8], [1.5, 9],
                                                      [1234567890123456, 10], [2**53 + 2, 11]]}))  # fmt: skip
    plan = _plan(u, _body([_input(bad), _input(b)]))
    assert not plan["ready"] and _codes(plan) == ["KEYJOIN_KEY_TYPE", "KEYJOIN_KEY_PRECISION"]
    typed = _first(plan, "KEYJOIN_KEY_TYPE")
    assert (typed["rows"], typed["count"], typed["column"], typed["input"]) == ([2, 3, 4], 3, "A", 0)
    assert "문자형 식별자로 저장" in typed["text"] and "2026" not in typed["text"]
    prec = _first(plan, "KEYJOIN_KEY_PRECISION")
    assert prec["rows"] == [5, 6] and "텍스트로 저장" in prec["text"] and "1234567890123456" not in prec["text"]
    kinds = plan["inputs"][0]["key_columns"][0]["kinds"]
    assert (kinds["date"], kinds["bool"], kinds["float"], kinds["precision"]) == (1, 1, 1, 2)
    xls_big = _upload(u, "큰실수.xls", _xls_bytes([["id", "값"], [2**53, 1], [12345678901234567, 2]]), XLS)
    plan = _plan(u, _body([_input(xls_big), _input(b)]))
    assert _codes(plan) == ["KEYJOIN_KEY_PRECISION"] and _first(plan, "KEYJOIN_KEY_PRECISION")["count"] == 2


# ================================================================ 중복·빈 key·결측 표기·병합 채움·소계/주석 — 차단


def test_duplicate_and_empty_keys_block_the_run_without_any_record(kj) -> None:
    u = kj["kj-user1"]
    a = _csv(u, "중복a.csv", [["사번", "부서"], ["KEY-DUP-7731", "총무"], ["002", "기획"], ["KEY-DUP-7731", "감사"]])
    b = _csv(u, "중복b.csv", [["사번", "건수"], ["KEY-DUP-7731", "10"], ["002", "20"]])
    outputs = {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in OUTPUT_DIR.glob("*.xlsx")}
    before = _state()
    plan = _plan(u, _body([_input(a), _input(b)]))
    dup = _first(plan, "KEYJOIN_KEY_DUPLICATE")
    assert not plan["ready"] and plan["inputs"][0]["keys"]["duplicate_keys"] == 1
    assert (dup["input"], dup["file_id"], dup["sheet_id"], dup["column"]) == (0, a, "s0", "A")
    assert (dup["rows"], dup["count"], dup["groups"]) == ([2, 4], 2, 1) and "KEY-DUP" not in dup["text"]
    assert dup["text"].startswith(f"입력 1 (#{a} · 1번째 시트 · A열): key가 같은 데이터 행이 2행(1묶음, 2행·4행)")
    assert plan["inputs"][0]["problems"] == [dup]
    # key 값은 (가린) 미리보기에만 — 문제 문구·경고·실행 응답에는 없다
    assert "KEY-DUP" not in json.dumps([*plan["problems"], *plan["warnings"], plan["inputs"]], ensure_ascii=False)
    res = _run(u, _body([_input(a), _input(b)]), plan["basis"])
    assert res.status_code == 409 and res.json()["detail"]["code"] == "KEYJOIN_NOT_READY"
    assert res.json()["detail"]["problems"] == [dup["text"]] and "KEY-DUP" not in res.text
    assert _state() == before
    assert {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in OUTPUT_DIR.glob("*.xlsx")} == outputs
    # 빈 key·공백만·결측 표기·복합 key 일부 누락 — 데이터 행이면 차단, 빈 값끼리 연결하지 않음
    e1 = _csv(u, "빈키1.csv", [["사번", "부서"], ["", "총무"], ["   ", "기획"], ["-", "정보"], ["004", "감사"]])
    e2 = _csv(u, "빈키2.csv", [["사번", "건수"], ["", "10"], ["004", "40"]])
    plan = _plan(u, _body([_input(e1), _input(e2)]))
    assert not plan["ready"]
    assert sorted(_codes(plan)) == ["KEYJOIN_KEY_EMPTY", "KEYJOIN_KEY_EMPTY", "KEYJOIN_KEY_MISSING_MARK"]
    empty = _first(plan, "KEYJOIN_KEY_EMPTY")
    assert (empty["rows"], empty["count"], empty["input"]) == ([2, 3], 2, 0) and plan["inputs"][0]["keys"][
        "empty_rows"
    ] == 2
    assert _first(plan, "KEYJOIN_KEY_MISSING_MARK")["rows"] == [4]
    p1 = _csv(u, "부분1.csv", [["기관", "사번", "값"], ["K1", "", "x"], ["K1", "002", "y"]])
    p2 = _csv(u, "부분2.csv", [["기관", "사번", "값"], ["K1", "002", "z"]])
    plan = _plan(u, _body([_input(p1, key=["c0", "c1"], cols=("c2",)), _input(p2, key=["c0", "c1"], cols=("c2",))]))
    assert _codes(plan) == ["KEYJOIN_KEY_PARTIAL"] and _first(plan, "KEYJOIN_KEY_PARTIAL")["column"] == "B"
    assert plan["inputs"][0]["keys"]["partial_rows"] == 1
    assert _state() == before


def test_merged_key_cells_block_and_subtotal_note_rows_need_a_decision(kj) -> None:
    """key 칸의 병합 채움은 차단. 소계·주석으로 추정된 행은 key 칸이 있으면 '확인 필요'(차단 — 결정 뒤에 만듦), key 칸이
    비어 있으면 구조 행이라 자동 제외하고 경고한다."""
    u = kj["kj-user1"]
    merged = _upload(u, "병합.xlsx", _xlsx_bytes({"k": [["사번", "값"], ["001", 1], [None, 2], ["003", 3]]},
                                                 {"k": ("A2:A3",)}))  # fmt: skip
    b = _xlsx(u, "병합b.xlsx", [["사번", "건수"], ["001", 10]])
    plan = _plan(u, _body([_input(merged), _input(b)]))
    filled = _first(plan, "KEYJOIN_KEY_FILLED")
    assert not plan["ready"] and filled["rows"] == [3] and "위 행의 값을 복사" in filled["text"]
    assert plan["inputs"][0]["key_columns"][0]["kinds"]["filled"] == 1
    rows = [["사번", "값"], ["001", 1], ["002", 2], ["합계", 3], ["※ 자료: 시험", None], [None, "소계"]]
    sub = _xlsx(u, "소계.xlsx", rows)
    before = _state()
    plan = _plan(u, _body([_input(sub), _input(b)]))
    pending = _first(plan, "KEYJOIN_ROWS_UNRESOLVED")
    assert not plan["ready"] and pending["rows"] == [4, 5] and pending["count"] == 2 and pending["input"] == 0
    assert "시험" not in json.dumps(plan, ensure_ascii=False)  # 주석 행의 글자(원문)는 어디에도 없다 — 위치만
    review = plan["inputs"][0]["row_review"]
    assert (
        review["total"] == 2 and review["unresolved"] == 2 and review["source_rows"] == 4 and review["link_rows"] == 2
    )
    assert [(r["row"], r["column"], r["reason"], r["in_key"], r["decision"]) for r in review["items"]] == [
        (4, "A", "subtotal", True, None),
        (5, None, "note", False, None),
    ]
    skipped = _first(plan, "KEYJOIN_NONDATA_ROWS_SKIPPED", "warnings")  # 6행: key 칸이 비어 있는 소계 행 — 구조 행
    assert skipped["rows"] == [6] and skipped["column"] == "B" and "자동 제외" in skipped["text"]
    assert plan["inputs"][0]["rows"] == {"data": 2, "subtotal": 2, "note": 1, "empty": 0}
    assert _state() == before
    decided = _body([_input(sub, rows={"exclude": [4, 5]}), _input(b)])
    run, wb, plan = _make(u, decided)
    excluded = _first(plan, "KEYJOIN_ROWS_EXCLUDED", "warnings")
    assert plan["ready"] and excluded["rows"] == [4, 5] and excluded["count"] == 2
    assert plan["inputs"][0]["row_review"]["excluded"] == 2 and plan["inputs"][0]["row_review"]["unresolved"] == 0
    assert _grid(wb["연결 결과"])[1:] == [["001", 1, 10], ["002", 2, None]]
    summary = _grid(wb["연결 요약"])
    assert any(r[0] == "제외한 행" and r[1] == "입력 1" and r[2] == "4·5" for r in summary if r and len(r) > 2)
    assert run["summary"]["inputs"][0]["skipped_rows"] == 1 and run["summary"]["inputs"][0]["excluded_rows"] == 2
    assert run["summary"]["inputs"][0]["review_rows"] == 2 and run["summary"]["policy"] == "keyjoin-policy-2"


def test_no_common_users_and_partial_presence(kj) -> None:
    u = kj["kj-user1"]
    a = _csv(u, "공통a.csv", [["사번", "부서"], ["001", "총무"], ["002", "기획"]])
    b = _csv(u, "공통b.csv", [["사번", "건수"], ["003", "30"], ["004", "40"]])
    before = _state()
    plan = _plan(u, _body([_input(a), _input(b)], "common"))
    assert not plan["ready"] and _codes(plan) == ["KEYJOIN_NO_COMMON"] and plan["result"]["rows"] == 0
    assert plan["result"]["intersection"] == 0 and plan["preview"]["rows"] == []
    assert plan["inputs"][0]["keys"]["indexed"] == 2  # 읽고 판정한 뒤의 0행 안내 (실패 실행 기록 없음)
    res = _run(u, _body([_input(a), _input(b)], "common"), plan["basis"])
    assert res.status_code == 409 and res.json()["detail"]["code"] == "KEYJOIN_NOT_READY" and _state() == before
    _, wb, plan = _make(u, _body([_input(a), _input(b)]))  # 기준 파일의 사용자만: 다른 파일에만 있는 사용자는 빈 칸
    assert _grid(wb["연결 결과"])[1:] == [["001", "총무", None], ["002", "기획", None]]
    assert _first(plan, "KEYJOIN_UNLINKED", "warnings")["count"] == 2
    _, wb, plan = _make(u, _body([_input(a), _input(b)], "all"))
    assert _grid(wb["연결 결과"])[1:] == [
        ["001", "총무", None],
        ["002", "기획", None],
        ["003", None, 30],
        ["004", None, 40],
    ]


# ================================================================ 개인정보·권한·가림·원문 비노출


def _no_raw(text: str, raws: tuple[str, ...]) -> None:
    for raw in raws:
        assert raw not in text


def test_masked_collision_keys_stay_apart_and_raw_values_never_leave_the_result_file(kj, caplog) -> None:
    u = kj["kj-user1"]
    # 원문 표식은 이 시험만의 값으로 둔다 — 공유 시험 DB에는 다른 시험(가린 구분 분석 등)이 kim1@example.com 같은 글자를
    # 열 이름으로 남겨 두므로, db_dump 전체에서 '없음'을 보려면 다른 시험과 겹치지 않는 글자여야 한다.
    raws = ("kjz1@example.test", "kjz2@example.test", "010-7701-4455", "010-9909-4455")
    a = _csv(u, "가림a.csv", [["이메일", "값"], ["가팀 kjz1@example.test", "a1"], ["가팀 kjz2@example.test", "a2"],
                              ["010-7701-4455", "a3"], ["010-9909-4455", "a4"]])  # fmt: skip
    b = _csv(u, "가림b.csv", [["mail", "값"], ["가팀 kjz2@example.test", "b2"], ["010-9909-4455", "b4"]])
    caplog.set_level("DEBUG")
    for fid in (a, b):
        plan = _plan(u, _body([_input(a), _input(b)]))
        assert not plan["ready"] and "KEYJOIN_PII_ACK_REQUIRED" in _codes(plan)
        assert u.post(f"/files/{fid}/pii/acknowledge").status_code == 200
    run, wb, plan = _make(u, _body([_input(a), _input(b)]))
    assert plan["result"]["rows"] == 4 and "KEYJOIN_PII_IN_RESULT" in _codes(plan, "warnings")
    collision = _first(plan, "KEYJOIN_MASKED_COLLISION", "warnings")
    assert collision["count"] == 2 and "가림 구분" in collision["text"]
    assert [r["key"] for r in plan["preview"]["rows"]] == [
        ["가팀 k***@example.test (가림 구분 1)"], ["가팀 k***@example.test (가림 구분 2)"],
        ["010-****-4455 (가림 구분 1)"], ["010-****-4455 (가림 구분 2)"],
    ]  # fmt: skip
    assert [r["cells"] for r in plan["preview"]["rows"]] == [["a1", None], ["a2", "b2"], ["a3", None], ["a4", "b4"]]
    assert _grid(wb["연결 결과"])[1:] == [
        ["가팀 kjz1@example.test", "a1", None], ["가팀 kjz2@example.test", "a2", "b2"],
        ["010-7701-4455", "a3", None], ["010-9909-4455", "a4", "b4"],
    ]  # fmt: skip
    _no_raw(json.dumps(plan, ensure_ascii=False), raws)
    _no_raw(json.dumps(run, ensure_ascii=False), raws)
    _no_raw(u.get(f"/runs/{run['id']}").text + u.get("/runs").text, raws)
    _no_raw(_db("SELECT summary FROM runs WHERE id = ?", run["id"])[0][0], raws)
    _no_raw(db_dump(), raws)
    _no_raw(caplog.text, raws)
    summary_text = "\n".join(str(v) for row in _grid(wb["연결 요약"]) for v in row if v is not None)
    _no_raw(summary_text, raws)
    assert "k***@example.test" not in summary_text  # 요약 시트에는 가린 key도 없다 (개수·설정만)


def test_pii_acknowledgement_is_required_and_checked_without_parsing(kj, monkeypatch) -> None:
    from app import keyjoin

    u = kj["kj-user1"]
    a = _csv(u, "확인a.csv", [["사번", "연락처"], ["001", "010-1234-5678"], ["002", "010-2222-3333"]])
    b = _csv(u, "확인b.csv", [["사번", "건수"], ["001", "10"]])
    plan = _plan(u, _body([_input(a), _input(b)]))
    blocked = _first(plan, "KEYJOIN_PII_ACK_REQUIRED")
    basis_before = plan["basis"]
    assert not plan["ready"] and blocked["input"] == 0 and plan["inputs"][0]["pii"] == "needs_ack"
    assert plan["inputs"][0]["keys"] is None and plan["inputs"][1]["keys"] is None and plan["preview"]["rows"] == []
    assert (
        plan["inputs"][0]["key_columns"][0]["kinds"] is None
        and plan["inputs"][0]["pii_summary"]["counts"]["mobile"] == 2
    )
    calls: list[int] = []
    real = keyjoin.load_parsed

    def counting(session: Any, file_id: int) -> Any:
        calls.append(file_id)
        return real(session, file_id)

    monkeypatch.setattr(keyjoin, "load_parsed", counting)
    plan = _plan(u, _body([_input(a), _input(b)]))
    assert not plan["ready"] and calls == []  # 차단된 요청은 key 분석(파싱)을 하지 않는다 (file_info 캐시 적중)
    res = _run(u, _body([_input(a), _input(b)]), plan["basis"])
    assert res.status_code == 409 and res.json()["detail"]["code"] == "KEYJOIN_NOT_READY" and calls == []
    assert u.post(f"/files/{a}/pii/acknowledge").status_code == 200
    run, wb, plan = _make(u, _body([_input(a), _input(b)]))
    assert plan["inputs"][0]["pii"] == "acknowledged" and "KEYJOIN_PII_IN_RESULT" in _codes(plan, "warnings")
    assert wb["연결 결과"]["B2"].value == "010-1234-5678"  # 확인한 파일은 결과에 원본 값 (묶기·통합과 같은 정책)
    assert "010-1234-5678" not in json.dumps(plan, ensure_ascii=False) and "010-1234-5678" not in json.dumps(run)
    assert plan["basis"] != basis_before and calls  # 확인 상태가 바뀌면 basis도 바뀐다 (아래 변경 시험과 같이)


def test_other_accounts_and_the_manager_follow_the_existing_policy(kj, abcd) -> None:
    u1, u2, own = kj["kj-user1"], kj["kj-user2"], kj["owner"]
    body = abcd["body"]
    before = _state()
    assert u2.post("/keyjoin/plan", json=body).status_code == 404  # 남의 파일은 없는 것처럼
    res = u2.post("/keyjoin/run", json={**body, "expected_basis": "x" * 16, "request_id": uuid.uuid4().hex})
    assert res.status_code == 404
    mine = _csv(u2, "u2.csv", [["사번", "값"], ["001", "u2"]])
    mixed = _body([_input(mine), _input(abcd["fids"][0])])
    assert u2.post("/keyjoin/plan", json=mixed).status_code == 404 and _state() == before
    plan = _plan(own, body)  # 최고 관리자는 볼 수 있지만 연결할 수는 없다 (조회만)
    assert not plan["ready"] and _codes(plan) == ["KEYJOIN_FILE_NOT_OWN"] * 4
    assert all(ip["keys"] is None and not ip["own"] for ip in plan["inputs"]) and plan["preview"]["rows"] == []
    assert _run(own, body, plan["basis"]).status_code == 409 and _state() == before
    run, _, _ = _make(u1, body)
    assert u2.get(f"/runs/{run['id']}").status_code == 404 and u2.get(f"/runs/{run['id']}/download").status_code == 404
    assert own.get(f"/runs/{run['id']}/download").status_code == 200  # 최고 관리자는 기존 정책대로 조회·내려받기


def test_learning_confirmation_and_analysis_engines_are_never_touched(kj, abcd, monkeypatch) -> None:
    from app import analysis, decisions, mapping, promotion

    def boom(*a: Any, **k: Any) -> Any:
        raise AssertionError("key 연결은 학습·확정·분석 엔진을 부르지 않아야 함")

    counts = _db("SELECT (SELECT count(*) FROM mapping_decisions), (SELECT count(*) FROM learned_aliases), "
                 "(SELECT count(*) FROM mapping_approvals)")  # fmt: skip
    for module, name in ((decisions, "save_decisions"), (promotion, "record_approval"), (mapping, "get_engine"),
                         (analysis, "evaluate"), (analysis, "compute")):  # fmt: skip
        monkeypatch.setattr(module, name, boom)
    run, wb, plan = _make(kj["kj-user1"], abcd["body"])
    assert plan["ready"] and run["status"] == "done" and wb.sheetnames == SHEETS
    assert _db("SELECT (SELECT count(*) FROM mapping_decisions), (SELECT count(*) FROM learned_aliases), "
               "(SELECT count(*) FROM mapping_approvals)") == counts  # fmt: skip


# ================================================================ 제어 문자·긴 이름·긴 key·수식 글자·0과 빈 칸·수식 칸


def test_illegal_characters_are_refused_by_position_only(kj) -> None:
    u = kj["kj-user1"]
    b = _csv(u, "제어b.csv", [["사번", "건수"], ["001", "10"]])
    bad_key = _upload(u, "제어키.xls", _xls_bytes([["사번", "값"], ["00\x011", "x"], ["002", "y"]]), XLS)
    plan = _plan(u, _body([_input(bad_key), _input(b)]))
    illegal = _first(plan, "KEYJOIN_KEY_ILLEGAL")
    assert not plan["ready"] and illegal["rows"] == [2] and illegal["column"] == "A" and "\x01" not in json.dumps(plan)
    bad_value = _csv(u, "제어값.csv", [["사번", "값"], ["001", "앞￾뒤"], ["002", "정상"]])
    before = _state()
    plan = _plan(u, _body([_input(bad_value), _input(b)]))
    illegal = _first(plan, "KEYJOIN_VALUE_ILLEGAL")
    assert (
        not plan["ready"] and illegal["rows"] == [2] and "2행 · B열" in illegal["text"] and "￾" not in json.dumps(plan)
    )
    assert (
        "앞" not in illegal["text"] and _run(u, _body([_input(bad_value), _input(b)]), plan["basis"]).status_code == 409
    )
    bad_name = _upload(u, "제어이름.xls", _xls_bytes([["사번", "값\x01"], ["001", "x"]]), XLS)
    plan = _plan(u, _body([_input(bad_name), _input(b)]))
    source = _first(plan, "KEYJOIN_NAME_ILLEGAL_SOURCE")
    assert not plan["ready"] and source["column"] == "B" and plan["inputs"][0]["keys"] is None
    assert _state() == before


def test_long_names_long_keys_and_formula_like_text_are_kept_as_text(kj) -> None:
    u = kj["kj-user1"]
    long_key = "K" * 5000
    a = _csv(u, "긴a.csv", [["사번", "값"], [long_key, "=SUM(A1)"], ["002", "+ 1"], ["003", "-x"], ["004", "@a"]])
    b = _csv(u, "긴b.csv", [["사번", "건수"], [long_key, "0"], ["002", "-"]])
    name255 = "열" + "이" * 254
    run, wb, plan = _make(u, _body([_input(a, names={"c1": name255}), _input(b)]))
    assert plan["header"] == ["사번", name255, "건수"] and len(name255) == 255
    ws = wb["연결 결과"]
    assert ws["A2"].value == long_key and ws["A2"].data_type == "s"  # 긴 key를 자르지 않는다
    assert [(c.value, c.data_type) for c in ws["B"][1:]] == [("=SUM(A1)", "s"), ("+ 1", "s"), ("-x", "s"), ("@a", "s")]
    assert ws["C2"].value == 0 and ws["C3"].value == "-" and ws["C4"].value is None  # 실제 0은 0, 결측 표기는 그대로
    assert ws["B1"].value == name255 and run["summary"]["header"][1] == name255
    assert u.post("/keyjoin/plan", json=_body([_input(a, names={"c1": name255 + "름"}), _input(b)])).status_code == 422


def _with_cached(data: bytes, cell: str, value: str) -> bytes:
    """openpyxl은 수식 결과값을 저장하지 않으므로 결과값이 저장된 수식 칸을 시험용으로 만든다 (test_bundle 선례)."""
    import re
    import zipfile

    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for info in src.infolist():
            body = src.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                body = re.sub(rf'(<c r="{cell}"[^>]*>)<f>([^<]*)</f><v\s*/?>(</v>)?'.encode(),
                              rf"\1<f>\2</f><v>{value}</v>".encode(), body)  # fmt: skip
                assert f"<v>{value}</v>".encode() in body
            z.writestr(info, body)
    return out.getvalue()


def test_missing_values_stay_blank_formula_gaps_get_a_note_and_long_numbers_become_text(kj) -> None:
    u = kj["kj-user1"]
    data = _with_cached(_xlsx_bytes({"k": [["사번", "값", "비율"], ["001", "=1+1", 0.125], ["002", "=2+2", 0],
                                           ["003", 7, None]]}), "B2", "2")  # fmt: skip
    a = _upload(u, "수식.xlsx", data)
    # 17자리 정수는 CSV로 (openpyxl이 쓴 xlsx 숫자 칸은 이미 16자리로 줄어 있다 — docs/NUMBERS.md, 복원하지 않음)
    b = _csv(u, "수식b.csv", [["사번", "건수"], ["001", "0"], ["003", "12345678901234567"]])
    run, wb, plan = _make(u, _body([_input(a, cols=("c1", "c2")), _input(b)]))
    warns = {p["code"]: p for p in plan["warnings"]}
    assert warns["KEYJOIN_FORMULA_MISSING"]["count"] == 1 and warns["KEYJOIN_LONG_NUMBERS"]["count"] == 1
    ws = wb["연결 결과"]
    assert ws["B2"].value == 2 and ws["B2"].comment is None  # 결과값이 저장된 수식 → 그 값
    assert ws["B3"].value is None and "0으로 채우지 않음" in ws["B3"].comment.text  # 결과값 없는 수식 → 빈 칸 + 메모
    assert ws["B4"].value == 7
    assert (
        ws["C2"].value == 0.125 and ws["C3"].value == 0 and ws["C4"].value is None
    )  # 소수 그대로, 0은 0, 빈 칸은 빈 칸
    assert ws["D2"].value == 0 and ws["D3"].value is None  # 실제 0은 0, 미연결은 0이 아니라 빈 칸
    assert ws["D4"].value == "12345678901234567" and ws["D4"].data_type == "s"  # 17자리 정수는 글자로 (자릿수 보존)
    assert run["summary"]["formula_missing"] == 1 and run["summary"]["long_numbers"] == 1
    assert plan["preview"]["rows"][1]["cells"] == [None, 0, None]
    assert plan["preview"]["rows"][2]["cells"] == [7, None, "12345678901234567"]


# ================================================================ plan 뒤 변경(basis)·중복 생성 요청·장애·410


def test_changes_after_the_plan_are_rejected_before_any_side_effect(kj, abcd, monkeypatch) -> None:
    from app import keyjoin

    u = kj["kj-user1"]
    inputs, fids = abcd["inputs"], abcd["fids"]
    plan = _plan(u, abcd["body"])
    before = _state()

    def stale(body: dict[str, Any]) -> None:
        res = _run(u, body, plan["basis"])
        assert res.status_code == 409 and res.json()["detail"]["code"] == "KEYJOIN_BASIS_CHANGED", res.text
        assert "다시 확인" in res.json()["detail"]["message"]

    stale(_body(inputs, "all"))  # 범위
    stale(_body(inputs, key_names=["직원"]))  # key 이름
    stale(_body([_input(fids[0], names={"c1": "부서명"}), *inputs[1:]]))  # 결과 이름
    stale(_body([_input(fids[0], cols=("c1", "c2")), *inputs[1:]]))  # 가져올 열
    stale(_body([*inputs[1:], inputs[0]], base={"file_id": fids[0], "sheet_id": "s0"}))  # 입력 순서
    stale(_body(inputs[:3]))  # 입력 제거
    extra = _csv(u, "추가.csv", [["사번", "값"], ["001", "e"]])
    stale(_body([*inputs, _input(extra)]))  # 입력 추가
    stale(_body(inputs, base={"file_id": fids[1], "sheet_id": "s0"}))  # 기준 입력
    monkeypatch.setattr(keyjoin, "MAX_ROWS", 299_999)  # 한도(정책) 변경
    stale(abcd["body"])
    monkeypatch.undo()
    assert _state() == before
    assert _plan(u, abcd["body"])["basis"] == plan["basis"]  # 같은 구성은 같은 basis
    assert _run(u, abcd["body"], plan["basis"]).status_code == 201


def test_same_request_id_returns_the_same_run_and_mismatch_or_failure_is_409(kj, abcd, monkeypatch) -> None:
    from app import keyjoin

    u = kj["kj-user1"]
    body = abcd["body"]
    plan = _plan(u, body)
    before = _state()
    rid = uuid.uuid4().hex
    first = _run(u, body, plan["basis"], rid)
    assert first.status_code == 201
    second = _run(u, body, plan["basis"], rid)
    assert second.status_code == 200 and second.json()["id"] == first.json()["id"]  # 재전송 — 새 실행 없음
    assert second.json()["summary"]["request_id"] == rid
    mid = _state()
    assert mid["runs"] == before["runs"] + 1 and len(mid["files"]) == len(before["files"]) + 1
    other = _plan(u, _body(abcd["inputs"], "all"))
    res = _run(u, _body(abcd["inputs"], "all"), other["basis"], rid)  # 같은 번호, 다른 설정
    assert res.status_code == 409 and res.json()["detail"]["code"] == "KEYJOIN_REQUEST_MISMATCH"
    assert res.json()["detail"]["run_id"] == first.json()["id"] and _state() == mid

    # 생성 중 오류 → 실패 기록, 같은 번호는 409 USED, 새 번호는 201
    def fail(*a: Any, **k: Any) -> None:
        raise OSError(28, "injected: no space left on device")

    monkeypatch.setattr(keyjoin, "_write_summary", fail)
    rid2 = uuid.uuid4().hex
    res = _run(u, body, plan["basis"], rid2)
    assert res.status_code == 409 and res.json()["detail"]["code"] == "KEYJOIN_FAILED"
    assert "파일 오류" in res.json()["detail"]["message"] and "KEYJOIN_FAILED" in res.json()["detail"]["message"]
    failed = _db("SELECT status, summary, error, id FROM runs ORDER BY id DESC LIMIT 1")[0]
    assert failed[0] == "failed" and json.loads(failed[1])["mode"] == "keyjoin" and "KEYJOIN_FAILED" in failed[2]
    assert res.json()["detail"]["run_id"] == failed[3]  # 실패로 기록된 실행 번호를 함께 알린다
    assert json.loads(failed[1])["request_id"] == rid2 and "inputs" in json.loads(failed[1])
    assert len(_state()["files"]) == len(mid["files"])  # 임시 파일 없음
    monkeypatch.undo()
    res = _run(u, body, plan["basis"], rid2)
    assert res.status_code == 409 and res.json()["detail"]["code"] == "KEYJOIN_REQUEST_USED"
    assert res.json()["detail"]["run_id"] is not None
    assert _run(u, body, plan["basis"]).status_code == 201
    listed = u.get("/runs").json()
    assert {r["status"] for r in listed if r["summary"].get("mode") == "keyjoin"} >= {"done", "failed"}


def test_a_second_request_while_a_result_is_being_made_is_refused_and_leaves_nothing(kj, abcd) -> None:
    from app.output_slot import output_slot

    u = kj["kj-user1"]
    plan = _plan(u, abcd["body"])
    held, release = threading.Event(), threading.Event()

    def busy() -> None:
        with output_slot():
            held.set()
            release.wait(30)

    th = threading.Thread(target=busy, daemon=True)
    th.start()
    assert held.wait(10)
    before = _state()
    try:
        for res in (_run(u, abcd["body"], plan["basis"]), u.post("/keyjoin/plan", json=abcd["body"])):
            assert res.status_code == 503, res.text
            assert res.json()["detail"]["code"] == "output_busy" and res.headers.get("retry-after") == "10"
        assert _state() == before
    finally:
        release.set()
        th.join(10)
    assert _run(u, abcd["body"], plan["basis"]).status_code == 201


def test_missing_result_file_is_410_and_other_users_data_does_not_block(kj, abcd) -> None:
    u1, u2 = kj["kj-user1"], kj["kj-user2"]
    run, _, _ = _make(u1, abcd["body"])
    path = DATA_DIR / run["output_path"].removeprefix("data/")
    path.chmod(0o600)
    path.unlink()
    res = u1.get(f"/runs/{run['id']}/download")
    assert res.status_code == 410 and "결과 파일이 서버에 없습니다" in res.json()["detail"]
    a = _csv(u2, "u2a.csv", A_ROWS)
    b = _csv(u2, "u2b.csv", [["사번", "건수"], ["001", "10"]])
    assert _make(u2, _body([_input(a), _input(b)]))[0]["status"] == "done"  # 다른 사용자의 같은 자료는 서로 막지 않음


# ================================================================ 예산 — 경계와 무부수효과, 한도 초과는 파싱하지 않음


def test_budgets_block_just_above_and_allow_at_the_limit(kj, abcd, monkeypatch) -> None:
    from app import keyjoin

    u = kj["kj-user1"]
    body = abcd["body"]
    plan = _plan(u, body)
    weighted = plan["budget"]["cells_weighted"]
    assert weighted == 152 and plan["limits"]["cells"] == keyjoin.MAX_OUTPUT_CELLS
    calls: list[int] = []
    real = keyjoin.load_parsed

    def counting(session: Any, file_id: int) -> Any:
        calls.append(file_id)
        return real(session, file_id)

    monkeypatch.setattr(keyjoin, "load_parsed", counting)
    before = _state()
    text_weighted = plan["budget"]["text_weighted"]
    limits = (("MAX_OUTPUT_CELLS", weighted), ("MAX_ROWS", 12), ("MAX_OUTPUT_TEXT_CHARS", text_weighted))
    for name, ok in limits:
        over = ok - 1
        monkeypatch.setattr(keyjoin, name, over)
        calls.clear()
        refused = _plan(u, body)
        assert not refused["ready"] and len(refused["problems"]) == 1, (name, refused["problems"])
        code = refused["problems"][0]["code"]
        assert code in ("KEYJOIN_TOO_MANY_CELLS", "KEYJOIN_TOO_MANY_ROWS", "KEYJOIN_TOO_MANY_CHARS")
        assert "나눠 만드세요" in refused["problems"][0]["text"] and calls == []  # 한도를 넘은 입력은 읽지 않는다
        assert refused["inputs"][0]["keys"] is None and refused["budget"]["cells_weighted"] == weighted
        res = _run(u, body, refused["basis"])
        assert res.status_code == 409 and res.json()["detail"]["code"] == "KEYJOIN_NOT_READY" and calls == []
        assert _state() == before
        monkeypatch.setattr(keyjoin, name, ok)
        allowed = _plan(u, body)
        assert allowed["ready"] and allowed["limits"] and calls, name  # 경계: 한도와 같으면 만든다
        assert _run(u, body, allowed["basis"]).status_code == 201
        before = _state()
        monkeypatch.undo()
        monkeypatch.setattr(keyjoin, "load_parsed", counting)
    assert _first(_plan(u, body), "KEYJOIN_UNLINKED", "warnings")  # 거절 뒤 다음 정상 요청
    monkeypatch.setattr(keyjoin, "INDEX_WEIGHT", 0.5)
    halved = _plan(u, body)
    assert halved["budget"]["cells_weighted"] == 134 + 9 and halved["basis"] != plan["basis"]  # 가중치도 basis에


def test_output_cells_count_the_source_sheet_and_the_summary(kj, abcd) -> None:
    """칸 예산: 결과 행 × (key + 가져온 열 + 1 + 입력 수) + 요약 고정분. 출처 시트가 포함된다 (행마다 1 + N칸)."""
    plan = _plan(kj["kj-user1"], abcd["body"])
    n, k, picked, rows = 4, 1, 4, 3
    assert plan["budget"]["output_cells"] == rows * (k + picked + 1 + n) + 40 + 16 * n
    assert plan["budget"]["index_cells"] == 3 * 3 * (k + 1)  # 기준이 아닌 입력 3개 × 3행 × (key 1 + 선택 열 1)


# ---------------------------------------------------------------- 독립 검토 보완 (2026-10-26 — 반박 검증을 거친 지적)


def test_file_id_out_of_range_and_zero_width_names_are_422_before_any_read(kj, abcd) -> None:
    """PRIVACY-3·PRIVACY-4: 정수 범위 밖 파일 번호는 DB에 닿기 전에 422(500 아님), 폭 0 서식 문자만의 이름은 빈 이름."""
    u = kj["kj-user1"]
    inputs = abcd["inputs"]
    before = _state()
    huge = [{**inputs[0], "file_id": 2**70}, inputs[1]]
    res = u.post("/keyjoin/plan", json=_body(huge, base={"file_id": inputs[1]["file_id"], "sheet_id": "s0"}))
    assert res.status_code == 422, res.text
    zero = u.post("/keyjoin/plan", json={**_body(inputs[:2]), "key_names": ["\u200b\ufeff"]})
    assert zero.status_code == 422 and "KEYJOIN_NAME_BLANK" in zero.text and "\u200b" not in zero.text
    named = _body([_input(inputs[0]["file_id"], names={"c1": "\u2060"}), inputs[1]])
    res = u.post("/keyjoin/plan", json=named)
    assert res.status_code == 422 and "KEYJOIN_NAME_BLANK" in res.text
    assert _state() == before


def test_error_code_text_stays_text_and_merged_value_cells_are_announced(kj) -> None:
    """KEYS-3: '#N/A' 같은 글자는 오류 칸이 아니라 글자 칸. KEYS-2: 가져올 열의 병합 채움은 값을 두되 경고로 알린다."""
    u = kj["kj-user1"]
    a = _csv(u, "오류코드.csv", [["사번", "상태"], ["001", "#N/A"], ["002", "#REF!"], ["003", "N/A"]])
    b = _csv(u, "오류코드b.csv", [["사번", "건수"], ["001", "1"], ["002", "2"], ["003", "3"]])
    _, wb, plan = _make(u, _body([_input(a), _input(b)]))
    ws = wb["연결 결과"]
    assert [(c.value, c.data_type) for c in ws["B"][1:]] == [("#N/A", "s"), ("#REF!", "s"), ("N/A", "s")]
    assert plan["preview"]["rows"][0]["cells"][0] == "#N/A"
    rows = [["사번", "부서"], ["001", "총무"], ["002", None], ["003", "기획"]]
    merged = _upload(u, "값병합.xlsx", _xlsx_bytes({"k": rows}, {"k": ("B2:B3",)}))
    _, wb, plan = _make(u, _body([_input(merged), _input(b)]))
    filled = _first(plan, "KEYJOIN_VALUE_FILLED", "warnings")
    assert filled["rows"] == [3] and filled["count"] == 1 and filled["column"] == "B" and filled["input"] == 0
    assert "병합" in filled["text"] and "총무" not in filled["text"]  # 값 원문은 문구에 없다
    assert _grid(wb["연결 결과"])[1:] == [["001", "총무", 1], ["002", "총무", 2], ["003", "기획", 3]]
    assert "병합 칸" in _texts(wb)  # 연결 요약 시트의 '주의'·'안내'에 채움 규칙이 적힌다
    assert plan["inputs"][0]["key_columns"][0]["kinds"]["filled"] == 0  # key 열의 병합 채움과는 별개


def test_a_total_in_an_unselected_column_asks_instead_of_dropping_the_row(kj) -> None:
    """선택하지 않은 메모 열의 'Total' 때문에 소계로 추정된 정상 행: 자동으로 빠지지 않고(차단) 포함하면 연결, 제외하면
    출처·요약에 남는다. 결정은 요청 내용·basis에 들어가므로 결정이 다르면 다른 basis·다른 요청이다."""
    u = kj["kj-user1"]
    a = _csv(
        u,
        "구분소계.csv",
        [["사번", "부서", "구분"], ["001", "총무", "일반"], ["002", "기획", "Total"], ["003", "정보", "일반"]],
    )
    b = _csv(u, "구분소계b.csv", [["사번", "건수"], ["001", "1"], ["002", "2"], ["003", "3"]])
    plan = _plan(u, _body([_input(a), _input(b)]))
    pending = _first(plan, "KEYJOIN_ROWS_UNRESOLVED")
    assert not plan["ready"] and pending["rows"] == [3] and "판정" not in pending["text"]
    item = plan["inputs"][0]["row_review"]["items"][0]
    assert item == {"row": 3, "column": "C", "reason": "subtotal", "in_key": False, "decision": None}
    assert "Total" not in json.dumps(plan, ensure_ascii=False) and "기획" not in json.dumps(plan, ensure_ascii=False)
    with_row = _body([_input(a, rows={"include": [3]}), _input(b)])
    run, wb, plan_in = _make(u, with_row)
    assert _grid(wb["연결 결과"])[1:] == [["001", "총무", 1], ["002", "기획", 2], ["003", "정보", 3]]
    assert plan_in["inputs"][0]["row_review"]["included"] == 1 and plan_in["inputs"][0]["row_review"]["link_rows"] == 3
    assert "KEYJOIN_ROWS_EXCLUDED" not in _codes(plan_in, "warnings")
    without = _body([_input(a, rows={"exclude": [3]}), _input(b)])
    run2, wb2, plan_out = _make(u, without)
    assert _grid(wb2["연결 결과"])[1:] == [["001", "총무", 1], ["003", "정보", 3]]
    assert _grid(wb2["출처"])[1:] == [[1, 2, 2], [2, 4, 4]]
    assert plan_in["basis"] != plan_out["basis"] and run["id"] != run2["id"]
    # 같은 요청 번호로 결정만 바꿔 보내면 다른 요청 내용 — 409 (내용 해시에 결정이 들어간다)
    rid = uuid.uuid4().hex
    assert _run(u, with_row, plan_in["basis"], rid).status_code == 201
    res = _run(u, without, plan_in["basis"], rid)
    assert res.status_code == 409 and res.json()["detail"]["code"] == "KEYJOIN_REQUEST_MISMATCH"


def test_row_decisions_are_checked_against_the_current_rows_and_included_rows_still_follow_the_key_rules(kj) -> None:
    """목록에 없는 행을 가리키는 결정은 거절(차단, 다른 행에 적용하지 않음). 포함한 행도 key 중복 검사를 그대로 거친다.
    포함·제외 양쪽에 같은 행을 적으면 422. 소계 행이 많으면(2,000행 넘게) 차단한다."""
    u = kj["kj-user1"]
    a = _csv(
        u,
        "결정.csv",
        [["사번", "부서", "구분"], ["001", "총무", "일반"], ["001", "기획", "Total"], ["003", "정보", "소계"]],
    )
    b = _csv(u, "결정b.csv", [["사번", "건수"], ["001", "1"], ["003", "3"]])
    stale = _plan(u, _body([_input(a, rows={"include": [3], "exclude": [9]}), _input(b)]))
    bad = _first(stale, "KEYJOIN_ROW_DECISION_STALE")
    assert not stale["ready"] and bad["rows"] == [9] and bad["count"] == 1
    assert "KEYJOIN_ROWS_UNRESOLVED" in _codes(stale)  # 4행(소계)은 아직 미결정
    dup = _plan(u, _body([_input(a, rows={"include": [3, 4]}), _input(b)]))
    assert "KEYJOIN_KEY_DUPLICATE" in _codes(dup) and _first(dup, "KEYJOIN_KEY_DUPLICATE")["rows"] == [2, 3]
    both = u.post("/keyjoin/plan", json=_body([_input(a, rows={"include": [3], "exclude": [3]}), _input(b)]))
    assert both.status_code == 422 and "KEYJOIN_ROW_DECISION_INVALID" in both.text
    twice = u.post("/keyjoin/plan", json=_body([_input(a, rows={"include": [3, 3]}), _input(b)]))
    assert twice.status_code == 422 and "KEYJOIN_ROW_DECISION_INVALID" in twice.text
    run, wb, plan = _make(u, _body([_input(a, rows={"include": [4], "exclude": [3]}), _input(b)]))
    assert _grid(wb["연결 결과"])[1:] == [["001", "총무", 1], ["003", "정보", 3]]


def test_too_many_suspected_rows_block_with_a_count(kj, monkeypatch) -> None:
    from app import keyjoin

    monkeypatch.setattr(keyjoin, "ROW_REVIEW_MAX", 3)
    u = kj["kj-user1"]
    rows = [["사번", "값", "구분"]] + [[f"{n:03d}", n, "Total"] for n in range(1, 6)]
    a = _csv(u, "소계많음.csv", rows)
    b = _csv(u, "소계많음b.csv", [["사번", "건수"], ["001", "1"]])
    plan = _plan(u, _body([_input(a), _input(b)]))
    many = _first(plan, "KEYJOIN_ROWS_TOO_MANY")
    assert not plan["ready"] and many["count"] == 5 and "5행" in many["text"] and "3행" in many["text"]
    review = plan["inputs"][0]["row_review"]
    assert review["total"] == 5 and review["listed"] == 3 and review["unresolved"] == 5


def test_resend_must_be_well_formed_and_a_request_being_made_is_refused_even_with_two_slots(
    kj, abcd, monkeypatch
) -> None:
    """PRIVACY-2: 같은 번호의 재전송도 형식 검사(422)가 먼저. SAFETY-1: 같은 번호가 만들어지는 중이면
    차례가 2개여도 둘째는 503 하나 — 실행·파일은 하나, 끝난 뒤 재전송은 200."""
    from app import keyjoin_api, output_slot

    u = kj["kj-user1"]
    body = abcd["body"]
    plan = _plan(u, body)
    rid = uuid.uuid4().hex
    assert _run(u, body, plan["basis"], rid).status_code == 201
    mid = _state()
    bad = u.post("/keyjoin/run", json={**body, "inputs": [], "expected_basis": plan["basis"], "request_id": rid})
    assert bad.status_code == 422 and "KEYJOIN_INPUTS_TOO_FEW" in bad.text
    assert _state() == mid

    monkeypatch.setattr(output_slot, "_SLOTS", threading.BoundedSemaphore(2))
    entered, release = threading.Event(), threading.Event()
    real = keyjoin_api.build_plan

    def slow_plan(*a: Any, **k: Any) -> Any:
        if not entered.is_set():  # 첫 요청만 계획 단계에서 붙잡아 둔다 (둘째가 그 사이에 온다)
            entered.set()
            release.wait(30)
        return real(*a, **k)

    monkeypatch.setattr(keyjoin_api, "build_plan", slow_plan)
    rid2 = uuid.uuid4().hex
    first: list[Any] = []
    th = threading.Thread(target=lambda: first.append(_run(u, body, plan["basis"], rid2)), daemon=True)
    th.start()
    assert entered.wait(10)
    try:
        second = _run(u, body, plan["basis"], rid2)
        assert second.status_code == 503, second.text
        assert (
            second.json()["detail"]["code"] == "output_busy" and "같은 요청 번호" in second.json()["detail"]["message"]
        )
    finally:
        release.set()
        th.join(60)
    assert first and first[0].status_code == 201, first and first[0].text
    after = _state()
    assert after["runs"] == mid["runs"] + 1 and len(after["files"]) == len(mid["files"]) + 1
    monkeypatch.setattr(keyjoin_api, "build_plan", real)
    assert _run(u, body, plan["basis"], rid2).status_code == 200  # 끝난 뒤의 재전송은 그 실행


def test_pii_item_sheet_names_are_masked_and_run_makes_no_preview(kj, abcd, monkeypatch) -> None:
    """PRIVACY-1: 개인정보 탐지 항목의 시트 이름도 가린다. SAFETY-5: 미리보기는 확인(plan)에만 만든다."""
    from app import keyjoin

    u = kj["kj-user1"]
    sheet = "kjr1@example.test"
    rows = [["사번", "이메일"], ["001", "kjr2@example.test"], ["002", "kjr3@example.test"]]
    p = _upload(u, "연락처.xlsx", _xlsx_bytes({sheet: rows}))
    b = _csv(u, "연락처b.csv", [["사번", "건수"], ["001", "1"]])
    plan = _plan(u, _body([_input(p), _input(b)]))
    assert plan["inputs"][0]["pii"] == "needs_ack" and "KEYJOIN_PII_ACK_REQUIRED" in _codes(plan)
    items = plan["inputs"][0]["pii_summary"]["items"]
    assert items and all(it["sheet"] != sheet for it in items)
    assert sheet not in json.dumps(plan, ensure_ascii=False) and "kjr2@example.test" not in json.dumps(plan)

    calls: list[str] = []
    real = keyjoin._preview

    def counted(*a: Any, **k: Any) -> Any:
        calls.append("preview")
        return real(*a, **k)

    monkeypatch.setattr(keyjoin, "_preview", counted)
    plan = _plan(u, abcd["body"])
    assert calls == ["preview"] and plan["preview"]["shown"] == 3
    assert _run(u, abcd["body"], plan["basis"]).status_code == 201
    assert calls == ["preview"]  # run은 미리보기를 만들지 않는다 (가림 충돌 경고는 두 경로 모두)


# ---------------------------------------------------------------- 요청 번호 보완 (2026-10-27 — 0021 색인·내용 해시)

RUN_ROW = (
    "INSERT INTO runs (user_id, status, trigger, input_file_count, file_ids, row_count, change_count, warning_count,"
    " summary, request_kind, request_id, request_hash) VALUES (?, ?, 'manual', 0, '[]', 0, 0, 0, ?, ?, ?, ?)"
)


def _user_id(name: str) -> int:
    return _db("SELECT id FROM users WHERE username = ?", name)[0][0]


def test_resend_is_found_however_many_runs_came_between(kj, abcd) -> None:
    """같은 번호는 최근 몇 건이 아니라 이 사용자의 실행 전부(고유 색인)에서 찾는다 — 사이에 다른 실행 60건이 있어도."""
    u = kj["kj-user1"]
    body = abcd["body"]
    plan = _plan(u, body)
    rid = uuid.uuid4().hex
    first = _run(u, body, plan["basis"], rid)
    assert first.status_code == 201
    uid = _user_id("kj-user1")
    for n in range(60):  # 사이의 다른 실행 (묶기·통합 모양 — 요청 번호 없음)
        _db(RUN_ROW, uid, "done", json.dumps({"mode": "bundle", "n": n}), None, None, None)
    mid = _state()
    again = _run(u, body, plan["basis"], rid)
    assert again.status_code == 200 and again.json()["id"] == first.json()["id"]
    assert _state() == mid


def test_same_request_id_compares_the_request_content_not_the_basis_the_client_sent(kj, abcd) -> None:
    """같은 번호·처음 expected_basis 그대로 열·범위·이름·기준 입력을 바꾸면 409(내용 비교). 같은 내용이면 basis가
    달라도 그 실행(200) — 완료된 같은 요청은 원본을 다시 읽지 않는다."""
    u = kj["kj-user1"]
    inputs = abcd["inputs"]
    body = _body(inputs[:2])
    plan = _plan(u, body)
    rid = uuid.uuid4().hex
    first = _run(u, body, plan["basis"], rid)
    assert first.status_code == 201
    mid = _state()
    a, b = inputs[0]["file_id"], inputs[1]["file_id"]
    variants = {
        "가져올 열": _body([_input(a, cols=("c1", "c2")), inputs[1]]),
        "범위": _body(inputs[:2], "all"),
        "결과 이름": _body([_input(a, names={"c1": "부서명"}), inputs[1]]),
        "기준 입력": _body(inputs[:2], base={"file_id": b, "sheet_id": "s0"}),
        "key 결과 이름": _body(inputs[:2], key_names=["사번2"]),
    }
    for label, changed in variants.items():
        res = _run(u, changed, plan["basis"], rid)  # 처음 basis를 그대로 보냄
        assert res.status_code == 409 and res.json()["detail"]["code"] == "KEYJOIN_REQUEST_MISMATCH", (label, res.text)
        assert res.json()["detail"]["run_id"] == first.json()["id"]
        assert _state() == mid, label
    same = _run(u, body, "f" * 64, rid)  # 같은 내용 — 화면이 보낸 basis는 비교하지 않는다
    assert same.status_code == 200 and same.json()["id"] == first.json()["id"] and _state() == mid


def test_a_request_id_from_an_older_version_without_a_hash_is_refused(kj, abcd) -> None:
    """이전 판(0021 전)에서 만든 실행은 해시가 없어 내용을 비교할 수 없다 — 409 LEGACY, 새 실행·파일 없음."""
    u = kj["kj-user1"]
    body = abcd["body"]
    plan = _plan(u, body)
    rid = uuid.uuid4().hex
    _db(RUN_ROW, _user_id("kj-user1"), "done", json.dumps({"mode": "keyjoin", "request_id": rid}), "keyjoin", rid, None)
    before = _state()
    res = _run(u, body, plan["basis"], rid)
    assert res.status_code == 409 and res.json()["detail"]["code"] == "KEYJOIN_REQUEST_LEGACY", res.text
    assert _state() == before


def test_a_resent_request_of_another_user_with_the_same_id_is_not_mixed_up(kj, abcd) -> None:
    """요청 번호는 사용자 안에서만 고유하다 — 다른 사용자가 같은 번호로 보내면 자기 실행이 새로 만들어지고(201),
    남의 실행이 돌아오거나(200) 남의 실행과 비교되지(409) 않는다. 남의 파일을 가리킨 요청은 파일 없음(404)이고
    각자의 재전송은 각자의 실행을 돌려준다."""
    u1, u2 = kj["kj-user1"], kj["kj-user2"]
    body = abcd["body"]
    plan = _plan(u1, body)
    rid = uuid.uuid4().hex
    first = _run(u1, body, plan["basis"], rid)
    assert first.status_code == 201, first.text
    own = _body([_input(_csv(u2, "A.csv", A_ROWS)), _input(_upload(u2, "B.xlsx", _xlsx_bytes({"교육": B_ROWS})))])
    plan2 = _plan(u2, own)
    second = _run(u2, own, plan2["basis"], rid)  # 같은 번호 — 다른 사용자의 자기 파일
    assert second.status_code == 201 and second.json()["id"] != first.json()["id"], second.text
    res = _run(u2, _body(abcd["inputs"][:2]), plan["basis"], uuid.uuid4().hex)  # 남의 파일을 가리킴
    assert res.status_code == 404 and "run_id" not in str(res.json()["detail"]), res.text
    assert _run(u1, body, plan["basis"], rid).json()["id"] == first.json()["id"]
    assert _run(u2, own, plan2["basis"], rid).json()["id"] == second.json()["id"]
    assert _db("SELECT count(*) FROM runs WHERE request_id = ?", rid) == [(2,)]  # 사용자마다 한 행


def test_when_another_process_inserts_the_same_request_first_the_second_returns_that_run(kj, abcd, monkeypatch) -> None:
    """실행 행을 넣는 순간 고유 색인이 막으면(다른 프로세스가 먼저) 결과 파일을 만들지 않고 그 실행을 돌려준다."""
    from app import keyjoin, keyjoin_api

    u = kj["kj-user1"]
    body = abcd["body"]
    plan = _plan(u, body)
    rid = uuid.uuid4().hex
    uid = _user_id("kj-user1")
    real = keyjoin_api.build_plan
    inserted: list[int] = []

    def race(db: Any, user: Any, req: Any, keep_values: bool = False) -> Any:  # noqa: ANN401
        made = real(db, user, req, keep_values=keep_values)
        if (
            keep_values and not inserted
        ):  # 계획 다시 판정 뒤·실행 행을 넣기 전에 다른 프로세스가 같은 번호로 먼저 넣었다
            _db(RUN_ROW, uid, "running", json.dumps({"mode": "keyjoin", "request_id": rid}), "keyjoin", rid,
                keyjoin.request_hash(req))  # fmt: skip
            inserted.append(_db("SELECT id FROM runs WHERE request_id = ?", rid)[0][0])
        return made

    monkeypatch.setattr(keyjoin_api, "build_plan", race)
    before = _state()
    res = _run(u, body, plan["basis"], rid)
    assert res.status_code == 200 and res.json()["id"] == inserted[0] and res.json()["status"] == "running", res.text
    after = _state()
    assert (
        after["runs"] == before["runs"] + 1 and after["files"] == before["files"]
    )  # 넣어 둔 행 하나뿐, 결과 파일 없음
    assert _db("SELECT count(*) FROM runs WHERE request_id = ?", rid) == [(1,)]
