"""로그인·권한 분리·표준 사전 승격·대시보드.

이 모듈만 실제 로그인(쿠키)으로 검사한다. 사용자마다 별도 TestClient(쿠키 저장소)를 쓴다.
"""

import io
import json
import sqlite3
from collections.abc import Iterator
from typing import Any

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app.config import DATA_DIR

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PROMO_COLUMN = "세부 추진 내역"  # 기본 표준 사전에 없는 열 이름


def _xlsx(rows: list[list]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "현황"
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


ROWS = [
    ["기관명", "집행액", PROMO_COLUMN],
    ["행정안전부", 1000, "정상 추진"],
    ["기상청", 2500, "검토 중"],
]


@pytest.fixture(scope="module")
def app_real_login(logged_in_owner) -> Iterator[Any]:
    """기존 테스트용 로그인 대체를 잠시 끄고 실제 쿠키 로그인으로 검사한다."""
    from app import admin_api, automation, files, integration, review, validation_api
    from app.auth import current_user
    from app.main import app
    from app.mapping import MappingEngine, load_dictionary

    def no_llm(system: str, user: str, schema: dict) -> dict[str, Any]:
        return {"key": "none", "confidence": 0, "reason": "해당 없음"}

    engine = MappingEngine(load_dictionary(), None, no_llm)
    patch = pytest.MonkeyPatch()
    for module in (files, validation_api, review, integration, automation, admin_api):
        patch.setattr(module, "get_engine", lambda: engine)
    override = app.dependency_overrides.pop(current_user)
    yield app
    app.dependency_overrides[current_user] = override
    patch.undo()


def _client(app: Any, username: str | None = None, password: str = "test-password-1") -> TestClient:
    c = TestClient(app)
    if username is not None:
        res = c.post("/auth/login", json={"username": username, "password": password})
        assert res.status_code == 200, res.text
    return c


@pytest.fixture(scope="module")
def accounts(app_real_login) -> dict[str, TestClient]:
    owner = _client(app_real_login, "test-owner")
    made = {"owner": owner}
    # 0016: 역할은 owner(최고 관리자)·user 두 가지.
    # 예전 admin1이 맡던 '전체 조회·남의 것 변경 불가' 검사는 owner로 옮겼다
    for username in ["user1", "user2", "user3", "user4"]:
        res = owner.post(
            "/users",
            json={
                "username": username,
                "display_name": username.upper(),
                "status": "active",
                "password": "pass-" + username,
            },
        )
        assert res.status_code == 201, res.text
        made[username] = _client(app_real_login, username, "pass-" + username)
    return made


def _upload(c: TestClient, rows: list[list] = ROWS, name: str = "자료.xlsx") -> int:
    res = c.post("/files", files=[("files", (name, _xlsx(rows), XLSX))])
    assert res.status_code == 200, res.text
    return res.json()[0]["id"]


def test_login_required_and_session_lifecycle(app_real_login, accounts) -> None:
    anon = _client(app_real_login)
    for method, path in [("get", "/files"), ("get", "/runs"), ("get", "/dashboard"), ("get", "/users")]:
        assert getattr(anon, method)(path).status_code == 401, path
    assert anon.get("/health").status_code == 200
    assert anon.get("/auth/status").json() == {"setup_required": False, "user": None}
    # 계정이 이미 있으면 최초 설정은 막힌다
    res = anon.post("/auth/setup", json={"username": "evil", "display_name": "x", "password": "password-123"})
    assert res.status_code == 409

    bad = anon.post("/auth/login", json={"username": "user1", "password": "wrong-password"})
    unknown = anon.post("/auth/login", json={"username": "nobody", "password": "wrong-password"})
    assert bad.status_code == unknown.status_code == 401
    assert bad.json()["detail"] == unknown.json()["detail"]  # 아이디가 있는지 알려주지 않는다

    assert accounts["user1"].get("/auth/status").json()["user"]["username"] == "user1"
    me = accounts["user1"].get("/auth/me").json()
    assert me["role"] == "user" and me["role_label"] == "일반 사용자"
    cookie = accounts["user1"].cookies.get("moajwo_session")
    assert cookie and "user1" not in cookie
    with sqlite3.connect(DATA_DIR / "moajwo.db") as con:
        stored = [r[0] for r in con.execute("select token_hash from auth_sessions")]
        hashes = [r[0] for r in con.execute("select password_hash from users")]
    assert cookie not in stored  # DB에는 토큰 해시만
    assert all(h.startswith("scrypt$") and "pass-" not in h for h in hashes)

    temp = _client(app_real_login, "user4", "pass-user4")
    assert temp.post("/auth/logout").status_code == 204
    assert temp.get("/auth/me").status_code == 401


def test_lockout_after_repeated_failures(app_real_login, accounts) -> None:
    accounts["owner"].post(
        "/users", json={"username": "locky", "display_name": "L", "status": "active", "password": "pass-locky"}
    )
    c = _client(app_real_login)
    for _ in range(5):
        assert c.post("/auth/login", json={"username": "locky", "password": "nope-nope"}).status_code == 401
    res = c.post("/auth/login", json={"username": "locky", "password": "pass-locky"})
    assert res.status_code == 401 and "잠" in res.json()["detail"]


def test_role_gates(accounts) -> None:
    user, owner = accounts["user1"], accounts["owner"]
    # 일반 사용자: 사전 관리·승격 심사·사용자 관리·설정·관리 화면 API 불가
    assert user.post("/glossary/aliases", json={"alias": "행안", "target_code": "1741000"}).status_code == 403
    assert user.get("/admin/promotions").status_code == 403
    assert user.get("/admin/standard-terms").status_code == 403
    assert user.get("/users").status_code == 403
    assert user.get("/settings").status_code == 403
    assert user.get("/admin/overview").status_code == 403
    assert user.get("/admin/audit").status_code == 403
    assert user.get("/glossary/resolve", params={"name": "행정안전부"}).status_code == 200  # 검증에 쓰는 조회는 가능
    # 0016: 중간 관리자(admin)·추가 소유자는 만들 수 없다 (예전: 관리자는 403, 소유자는 가능 → 이제 역할 값 자체를 거부)
    for role in ("admin", "owner"):
        boss = {
            "username": f"boss-{role}",
            "display_name": "B",
            "role": role,
            "status": "active",
            "password": "pass-boss2",
        }
        assert owner.post("/users", json=boss).status_code == 422
    assert not [u for u in owner.get("/users").json() if u["username"].startswith("boss-")]
    owner_id = next(u["id"] for u in owner.get("/users").json() if u["username"] == "test-owner")
    # 최고 관리자: 설정 가능, 자기 계정의 이용 중지·일반 사용자로 변경은 불가 (역할·활성은 PATCH로 바꿀 수 없음)
    assert owner.get("/settings").status_code == 200
    assert owner.post(f"/users/{owner_id}/suspend", json={}).status_code == 409
    assert owner.patch(f"/users/{owner_id}", json={"role": "user"}).status_code == 422
    assert owner.patch(f"/users/{owner_id}", json={"is_active": False}).status_code == 422
    me = owner.get("/auth/me").json()
    assert me["role"] == "owner" and me["role_label"] == "최고 관리자(소유자)"


def test_users_see_only_their_own_records(accounts) -> None:
    u1, u2, admin = accounts["user1"], accounts["user2"], accounts["owner"]  # 0016: 전체 조회 = 최고 관리자
    fid = _upload(u1, name="u1자료.xlsx")
    assert fid in [f["id"] for f in u1.get("/files").json()]
    assert fid not in [f["id"] for f in u2.get("/files").json()]
    assert u2.post(f"/files/{fid}/mapping", params={"use_llm": "false"}).status_code == 404
    assert u2.post("/reviews", json={"file_ids": [fid]}).status_code == 404

    review = u1.post("/reviews", json={"file_ids": [fid]}).json()
    for item in [i for i in review["items"] if i["category"] == "pending"]:
        res = u1.post(
            f"/reviews/{review['id']}/items/{item['id']}/decision",
            json={
                "decision": "approved" if item["payload"].get("field_key") else "rejected",
                "choice": item["payload"].get("field_key"),
            },
        )
        assert res.status_code == 200, res.text
    assert u1.post(f"/reviews/{review['id']}/complete").status_code == 200
    run = u1.post(f"/reviews/{review['id']}/integrate").json()
    assert run["status"] == "done" and run["user_name"] == "USER1"

    # 사용자는 자기 실행 기록만
    assert [r["id"] for r in u1.get("/runs").json()] == [run["id"]]
    assert run["id"] not in [r["id"] for r in u2.get("/runs").json()]
    assert u2.get(f"/runs/{run['id']}").status_code == 404
    assert u2.get(f"/runs/{run['id']}/download").status_code == 404
    assert u2.get(f"/reviews/{review['id']}").status_code == 404
    # 최고 관리자는 전체 조회(읽기 전용)
    assert run["id"] in [r["id"] for r in admin.get("/runs").json()]
    assert admin.get(f"/runs/{run['id']}/download").status_code == 200
    seen = admin.get(f"/reviews/{review['id']}").json()
    assert seen["can_edit"] is False and seen["created_by_name"] == "USER1"


def test_admin_cannot_modify_result_data(accounts) -> None:
    u1, admin = accounts["user1"], accounts["owner"]  # 0016: 예전 관리자·소유자 둘 다 → 최고 관리자
    run = u1.get("/runs").json()[0]
    for client in (admin,):
        for method in ("put", "patch", "delete"):
            assert getattr(client, method)(f"/runs/{run['id']}").status_code == 405
        assert client.post(f"/runs/{run['id']}/changes", json={}).status_code == 405
        assert client.delete(f"/runs/{run['id']}/download").status_code == 405
    # 남의 검토를 대신 승인·완료·통합하지 않는다
    fid = _upload(u1, name="u1자료2.xlsx")
    review = u1.post("/reviews", json={"file_ids": [fid]}).json()
    pending = [i for i in review["items"] if i["category"] == "pending"]
    assert pending
    res = admin.post(f"/reviews/{review['id']}/items/{pending[0]['id']}/decision", json={"decision": "rejected"})
    assert res.status_code == 403
    assert admin.post(f"/reviews/{review['id']}/complete").status_code == 403
    assert admin.post(f"/reviews/{review['id']}/integrate").status_code == 403
    # 결과 파일은 읽기 전용
    path = DATA_DIR / run["output_path"].removeprefix("data/")
    assert path.is_file()
    with pytest.raises(PermissionError):
        path.open("r+b")


def _confirm_promo_column(c: TestClient) -> None:
    fid = _upload(c, name="승격.xlsx")
    res = c.put(
        f"/files/{fid}/mapping/decisions",
        json={"decisions": [{"sheet": "현황", "column": PROMO_COLUMN, "field_key": "note"}]},
    )
    assert res.status_code == 200, res.text


def test_promotion_after_three_users_and_admin_approval(accounts) -> None:
    admin = accounts["owner"]  # 0016: 승격 심사 = 최고 관리자
    _confirm_promo_column(accounts["user1"])
    _confirm_promo_column(accounts["user1"])  # 같은 사용자가 여러 번 승인해도 1명
    _confirm_promo_column(accounts["user2"])
    assert not [c for c in admin.get("/admin/promotions").json()["items"] if c["column_name"] == PROMO_COLUMN]
    _confirm_promo_column(accounts["user3"])
    queue = admin.get("/admin/promotions").json()
    assert queue["threshold"] == 3
    candidate = next(c for c in queue["items"] if c["column_name"] == PROMO_COLUMN)
    assert candidate["status"] == "pending" and candidate["approver_count"] == 3
    assert sorted(candidate["approvers"]) == ["USER1", "USER2", "USER3"]
    assert candidate["created_at"].endswith("+09:00")  # 화면에는 한국 시간으로

    # 반영 전: 4번째 사용자에게는 아직 자동 적용되지 않는다
    u4 = accounts["user4"]
    fid = _upload(u4, name="u4.xlsx")
    before = u4.post(f"/files/{fid}/mapping", params={"use_llm": "false"}).json()["sheets"][0]["columns"]
    assert next(c for c in before if c["column"] == PROMO_COLUMN)["method"] != "standard"

    res = admin.post(f"/admin/promotions/{candidate['id']}/approve", json={"note": "공통 비고 열"})
    assert res.status_code == 200 and res.json()["status"] == "approved"
    terms = admin.get("/admin/standard-terms").json()
    assert any(t["term"] == PROMO_COLUMN and t["field_key"] == "note" for t in terms)
    assert admin.post(f"/admin/promotions/{candidate['id']}/approve", json={}).status_code == 409  # 이미 처리됨

    after = u4.post(f"/files/{fid}/mapping", params={"use_llm": "false"}).json()["sheets"][0]["columns"]
    promoted = next(c for c in after if c["column"] == PROMO_COLUMN)
    assert (promoted["method"], promoted["status"], promoted["field_key"]) == ("standard", "auto", "note")


def test_dashboard_counts_and_savings(accounts) -> None:
    u1, u2, owner = accounts["user1"], accounts["user2"], accounts["owner"]
    mine = u1.get("/dashboard").json()
    assert mine["scope"] == "mine"
    assert mine["month_runs"] == len(u1.get("/runs").json()) >= 1
    assert mine["validated_rows"] >= 2 and mine["validated_files"] >= 1
    assert mine["baseline_minutes"] is None and mine["saved_seconds"] == 0 and mine["runs_without_baseline"] >= 1
    assert u2.get("/dashboard").json()["month_runs"] == 0

    # 소유자가 수동 기준 소요 시간을 정하면 절약 시간 = 기준 - 실제 처리 시간
    assert u2.put("/settings", json={"manual_baseline_minutes": 30}).status_code == 403  # 0016: admin1 → 일반 사용자
    assert owner.put("/settings", json={"manual_baseline_minutes": 30}).json()["manual_baseline_minutes"] == 30
    runs = u1.get("/runs").json()
    done = [r for r in runs if r["status"] == "done"]
    expected = sum(max(0, 1800 - r["duration_ms"] / 1000) for r in done)
    mine = u1.get("/dashboard").json()
    assert mine["saved_runs"] == len(done) and mine["runs_without_baseline"] == 0
    assert abs(mine["saved_seconds"] - expected) < 1
    everyone = owner.get("/dashboard").json()
    assert everyone["scope"] == "all" and everyone["month_runs"] >= mine["month_runs"]

    # 통합 전 발견한 오류: 합계 불일치(차단)가 있는 자료 → 다시 검사해도 같은 오류는 한 번만
    rows = [["기관명", "집행액"], ["행정안전부", 100], ["기상청", 200], ["합계", 999]]
    fid = _upload(u2, rows, "합계틀림.xlsx")
    review = u2.post("/reviews", json={"file_ids": [fid]}).json()
    blocks = review["counts"]["block"]
    assert blocks >= 1
    u2.post(f"/reviews/{review['id']}/recheck", json={})
    stats = u2.get("/dashboard").json()
    assert stats["blocks_found"] == blocks
    assert stats["errors_found"] == stats["blocks_found"] + stats["warnings_found"]
    owner.put("/settings", json={"manual_baseline_minutes": None})


def test_glossary_name_learning_needs_manager(accounts) -> None:
    """'이 형식 기억하기'로 기관명을 공용 용어사전에 학습하는 것은 관리자·소유자만.

    일반 사용자는 열 이름만 개인 별칭으로 학습한다.
    """
    u3, owner = accounts["user3"], accounts["owner"]
    with sqlite3.connect(DATA_DIR / "moajwo.db") as con:
        had_entry = con.execute("select count(*) from glossary_entries where code = '1550000'").fetchone()[0]
    csv = "\n".join(["코드,정식명칭,유형", "1550000,국가유산청,중앙행정기관", ""]).encode()
    assert owner.post("/glossary/import", files=[("file", ("g.csv", csv, "text/csv"))]).json()["applied"]
    rows = [["기관명", "집행액"], ["국가유산", 100]]
    fid = _upload(u3, rows, "이름.xlsx")
    review = u3.post("/reviews", json={"file_ids": [fid]}).json()
    u3.patch(f"/reviews/{review['id']}", json={"remember_format": True})
    named = [i for i in review["items"] if i["kind"] == "entity" and i["category"] == "pending"]
    assert named and named[0]["payload"]["candidates"], "용어사전 후보가 있는 처음 보는 이름이 확인 필요로 나와야 함"
    for item in [i for i in review["items"] if i["category"] == "pending"]:
        candidates = item["payload"].get("candidates") or []
        choice = (
            item["payload"].get("field_key")
            if item["kind"] == "mapping"
            else (candidates[0]["code"] if candidates else None)
        )
        u3.post(
            f"/reviews/{review['id']}/items/{item['id']}/decision",
            json={"decision": "approved" if choice else "rejected", "choice": choice},
        )
    done = u3.post(f"/reviews/{review['id']}/complete").json()
    with sqlite3.connect(DATA_DIR / "moajwo.db") as con:
        learned = con.execute("select count(*) from glossary_aliases where alias = '국가유산'").fetchone()[0]
    assert learned == 0
    assert done["learned"]["names"] == 0 and done["learned"]["names_skipped"] == 1
    if not had_entry:  # 같은 테스트 DB를 쓰는 용어사전 테스트에 영향이 없도록
        with sqlite3.connect(DATA_DIR / "moajwo.db") as con:
            con.execute("delete from glossary_entries where code = '1550000'")


def test_dashboard_counts_same_content_once(accounts) -> None:
    """같은 파일을 다시 올려 검사해도(내용 sha256이 같으면) '검증한 행 수'에 한 번만 더한다."""
    u = accounts["user4"]
    rows = [["기관명", "집행액"], ["행정안전부", 1], ["기상청", 2], ["인사혁신처", 3]]
    before = u.get("/dashboard").json()["validated_rows"]
    data = _xlsx(
        rows
    )  # 한 번 만든 바이트를 두 번 올린다 (xlsx는 저장 시각이 들어가 다시 만들면 sha256이 달라질 수 있음)
    first, second = (
        u.post("/files", files=[("files", (name, data, XLSX))]).json()[0]["id"]
        for name in ("같은내용.xlsx", "같은내용_다시.xlsx")
    )
    assert u.post("/validation", json={"file_ids": [first]}).status_code == 200
    mid = u.get("/dashboard").json()["validated_rows"]
    assert u.post("/validation", json={"file_ids": [second]}).status_code == 200
    after = u.get("/dashboard").json()["validated_rows"]
    assert mid - before == 3 and after == mid


def _decision_rows(fid: int) -> list[tuple]:
    with sqlite3.connect(DATA_DIR / "moajwo.db") as con:
        return con.execute(
            "select column_name, field_key from mapping_decisions where source_file_id = ?", (fid,)
        ).fetchall()


def _pii_ack(fid: int) -> object:
    with sqlite3.connect(DATA_DIR / "moajwo.db") as con:
        return con.execute("select pii_acknowledged_at from source_files where id = ?", (fid,)).fetchone()[0]


def test_review_cannot_change_other_users_file(accounts) -> None:
    """교차 검증 ③: 관리자·소유자가 남의 파일로 자기 검토를 만들어 승인해도 원본의 매핑·개인정보 확인은 그대로."""
    u3, owner = accounts["user3"], accounts["owner"]  # 0016: 예전 관리자·소유자 둘 다 → 최고 관리자
    column = "교차검증 전용 메모"  # 사전·승격 이름에 없는 열 → 확인 필요 항목
    rows = [
        ["기관명", column, "담당 연락처"],
        ["행정안전부", "정상", "010-1234-5678"],
        ["기상청", "검토", "010-2345-6789"],
    ]
    fid = _upload(u3, rows, "u3개인정보.xlsx")
    # 직접 변경은 막혀 있다 (기존 동작)
    direct = owner.put(
        f"/files/{fid}/mapping/decisions",
        json={"decisions": [{"sheet": "현황", "column": column, "field_key": "note"}]},
    )
    assert direct.status_code in (403, 404)
    for client in (owner,):
        review = client.post("/reviews", json={"file_ids": [fid]})
        assert review.status_code == 201, review.text  # 조회용 검토는 만들 수 있다
        review = review.json()
        pending = {i["kind"]: i for i in review["items"] if i["category"] == "pending"}
        assert {"mapping", "pii"} <= set(pending), pending
        for kind in ("mapping", "pii"):
            res = client.post(
                f"/reviews/{review['id']}/items/{pending[kind]['id']}/decision",
                json={"decision": "approved", "choice": "note"} if kind == "mapping" else {"decision": "approved"},
            )
            assert res.status_code == 403, (kind, res.text)
        assert _decision_rows(fid) == [] and _pii_ack(fid) is None  # 부분 변경 없음
    # 파일 주인은 자기 검토에서 그대로 처리할 수 있다
    review = u3.post("/reviews", json={"file_ids": [fid]}).json()
    item = next(i for i in review["items"] if i["category"] == "pending" and i["kind"] == "mapping")
    res = u3.post(
        f"/reviews/{review['id']}/items/{item['id']}/decision", json={"decision": "approved", "choice": "note"}
    )
    assert res.status_code == 200, res.text
    assert (column, "note") in _decision_rows(fid)


def test_recheck_and_automation_cannot_write_other_users_file(accounts) -> None:
    """교차 검증 ③: 재검사의 파일 교체와 자동화 실행으로도 남의 파일 매핑을 바꾸지 못한다."""
    from app.db import SessionLocal
    from app.models import Automation, User
    from app.review import column_rule_key

    u4, admin = accounts["user4"], accounts["owner"]  # 0016: admin1 → 최고 관리자
    victim = _upload(u4, name="u4대상.xlsx")
    mine = _upload(admin, name="admin본인.xlsx")
    own = admin.put(  # 관리자 자기 파일의 확정 → 교체 시 남의 파일로 복사되면 안 된다
        f"/files/{mine}/mapping/decisions",
        json={"decisions": [{"sheet": "현황", "column": "집행액", "field_key": "budget"}]},
    )
    assert own.status_code == 200, own.text
    review = admin.post("/reviews", json={"file_ids": [mine]}).json()
    res = admin.post(f"/reviews/{review['id']}/recheck", json={"replace": {str(mine): victim}})
    assert res.status_code == 403, res.text
    assert _decision_rows(victim) == []

    with SessionLocal() as db:
        admin_id = db.query(User.id).filter(User.username == "test-owner").scalar()
        automation = Automation(
            name="우회 확인",
            owner_id=admin_id,
            column_rules=json.dumps([{"key": column_rule_key(PROMO_COLUMN), "field_key": "note"}]),
            entity_rules="{}",
            validation="{}",
            output_format="{}",
        )
        db.add(automation)
        db.commit()
        automation_id = automation.id
    res = admin.post(f"/automations/{automation_id}/run", json={"file_ids": [victim]})
    assert res.status_code in (201, 403, 422), res.text
    assert _decision_rows(victim) == []


PII_ROWS = [
    ["기관명", "alice@example.com", "bob@example.com", "집행액"],
    ["행정안전부", "정상", "검토", 1000],
    ["기상청", "보류", "정상", 2500],
]


def test_pii_column_names_masked_in_api_and_ids_keep_decisions(accounts) -> None:
    """후속 검증 ④: 업로드·매핑·확정·검토 응답에 개인정보 열 이름 원문이 없고, 마스킹이 같아지는 두 열도 column_id로
    따로 확정되며 재분석에서 유지된다. 무권한 사용자는 볼 수 없다."""
    u1, u2 = accounts["user1"], accounts["user2"]
    upload = u1.post("/files", files=[("files", ("개인정보머리글.xlsx", _xlsx(PII_ROWS), XLSX))])
    assert upload.status_code == 200, upload.text
    assert "alice@example.com" not in upload.text and "bob@example.com" not in upload.text
    fid = upload.json()[0]["id"]
    mapping = u1.post(f"/files/{fid}/mapping?use_llm=false")
    assert "alice@example.com" not in mapping.text and "bob@example.com" not in mapping.text
    cols = mapping.json()["sheets"][0]["columns"]
    ids = [c["column_id"] for c in cols]
    assert len(set(ids)) == len(ids) == 4
    a, b = cols[1], cols[2]
    res = u1.put(
        f"/files/{fid}/mapping/decisions",
        json={"decisions": [
            {"sheet": "현황", "column_id": a["column_id"], "field_key": "note"},
            {"sheet": "현황", "column_id": b["column_id"], "field_key": None},
            {"sheet": "현황", "column": "집행액", "field_key": "expenditure"},  # 예전 방식(원래 이름)도 호환
        ]},
    )  # fmt: skip
    assert res.status_code == 200, res.text
    assert "alice@example.com" not in res.text
    assert sorted(_decision_rows(fid)) == sorted(
        [("alice@example.com", "note"), ("bob@example.com", None), ("집행액", "expenditure")]
    )
    again = {c["column_id"]: c for c in u1.post(f"/files/{fid}/mapping?use_llm=false").json()["sheets"][0]["columns"]}
    assert (again[a["column_id"]]["status"], again[a["column_id"]]["field_key"]) == ("confirmed", "note")
    assert (again[b["column_id"]]["status"], again[b["column_id"]]["field_key"]) == ("confirmed", None)
    bad = u1.put(
        f"/files/{fid}/mapping/decisions",
        json={"decisions": [{"sheet": "현황", "column_id": "c99", "field_key": None}]},
    )
    assert bad.status_code == 422 and "example.com" not in bad.text
    assert u2.post(f"/files/{fid}/mapping?use_llm=false").status_code == 404  # 무권한
    # 검토 기록(DB에는 연결용 원문)도 화면 응답에서는 가려진다
    review = u1.post("/reviews", json={"file_ids": [fid]})
    assert review.status_code == 201 and "alice@example.com" not in review.text
    again_review = u1.get(f"/reviews/{review.json()['id']}")
    assert "alice@example.com" not in again_review.text and "bob@example.com" not in again_review.text


def test_review_items_report_read_only_per_file(accounts) -> None:
    """후속 검증 ⑤: 검토 응답이 항목별 변경 가능 여부(can_change)를 알려 준다 — 여러 사람의 파일이 섞인 검토 포함."""
    u3, admin = accounts["user3"], accounts["owner"]  # 0016: admin1 → 최고 관리자
    rows = [["기관명", "후속검증 전용 메모"], ["행정안전부", "정상"], ["기상청", "검토"]]
    other = _upload(u3, rows, "u3후속.xlsx")
    mine = _upload(admin, [["기관명", "관리자 전용 메모"], ["행정안전부", "정상"]], "admin후속.xlsx")
    review = admin.post("/reviews", json={"file_ids": [other, mine]}).json()
    by_file = {}
    for item in review["items"]:
        if item["category"] == "pending" and item["kind"] == "mapping":
            by_file[item["file_id"]] = item["can_change"]
    assert by_file == {other: False, mine: True}, by_file
    own_item = next(
        i for i in review["items"] if i["category"] == "pending" and i["kind"] == "mapping" and i["file_id"] == mine
    )
    res = admin.post(f"/reviews/{review['id']}/items/{own_item['id']}/decision", json={"decision": "rejected"})
    assert res.status_code == 200, res.text


PII_RAW = ("carol@example.com", "010-3456-7890", "3456-7890", "900101-1234567", "1234567", "dave@example.com")


def _no_raw(text: str) -> bool:
    return not any(raw in text for raw in PII_RAW)


def test_pii_values_masked_in_validation_review_and_history(accounts, caplog) -> None:
    """병합 전 보완 ③: 기관·지역 값, 열 이름, 오류 근거(필수 기관), 참고, 검토 payload, 과거 기록 조회, 로그 어디에도
    개인정보 원문이 없다. 개인정보 형식 이름은 별칭 등록 불가로 표시되고, 정상 이름은 그대로 등록할 수 있다."""
    import logging

    u1, u2 = accounts["user1"], accounts["user2"]
    rows = [
        ["기관명", "dave@example.com", "집행액"],
        ["carol@example.com", "x", 1],
        ["010-3456-7890 담당", "y", 2],
        ["900101-1234567", "z", 3],
        ["기상청", "w", 4],
    ]
    fid = _upload(u1, rows, "개인정보값.xlsx")
    caplog.set_level(logging.DEBUG)
    run = u1.post(
        "/validation",
        json={"file_ids": [fid], "required_entities": ["carol@example.com", "기상청"], "entity_field": "org_name"},
    )
    assert run.status_code == 200 and _no_raw(run.text), run.text
    body = run.json()
    regs = {e["value"]: e["registrable"] for e in body["entities"]}
    assert regs and all(v is False for k, v in regs.items() if "*" in k) and regs.get("기상청", True) is True
    past = u1.get(f"/validation/{body['id']}")
    assert past.status_code == 200 and _no_raw(past.text)
    assert u2.get(f"/validation/{body['id']}").status_code in (403, 404)  # 다른 사용자의 기록
    review = u1.post("/reviews", json={"file_ids": [fid], "required_entities": ["carol@example.com"]})
    assert review.status_code == 201 and _no_raw(review.text)
    again = u1.get(f"/reviews/{review.json()['id']}")
    assert _no_raw(again.text)
    assert _no_raw(caplog.text)
    # 다른 사용자가 이 검토의 항목 id를 자기 검토 경로로 보내도 막힌다
    item = review.json()["items"][0]
    own = u2.post("/reviews", json={"file_ids": [_upload(u2, ROWS, "u2자료.xlsx")]}).json()
    res = u2.post(f"/reviews/{own['id']}/items/{item['id']}/decision", json={"decision": "rejected"})
    assert res.status_code == 404
    assert (
        u2.post(
            f"/reviews/{review.json()['id']}/items/{item['id']}/decision", json={"decision": "rejected"}
        ).status_code
        == 404
    )


def test_result_workbook_pii_policy(accounts) -> None:
    """병합 전 보완 ④: 머리글에만 개인정보가 있어도 탐지·확인 절차에 들어가고, 포함을 거부한 파일은 결과·변경 내역에
    없으며, 승인한 파일의 원문은 결과 엑셀에 들어갈 수 있다(정책). 결과 다운로드·실행 기록은 권한이 있는 사람만."""
    u4, u2, admin = accounts["user4"], accounts["user2"], accounts["owner"]  # 0016: admin1 → 최고 관리자
    header_only = _upload(
        u4, [["기관명", "erin@example.com", "집행액"], ["행정안전부", "a", 1], ["기상청", "b", 2]], "머리글만.xlsx"
    )
    body_pii = _upload(u4, [["기관명", "연락처", "집행액"], ["인사혁신처", "010-5678-1234", 3]], "본문.xlsx")
    review = u4.post("/reviews", json={"file_ids": [header_only, body_pii]}).json()
    pii_items = {i["file_id"]: i for i in review["items"] if i["category"] == "pending" and i["kind"] == "pii"}
    assert set(pii_items) == {header_only, body_pii}  # 머리글만 있는 파일도 확인 대상
    for item in [i for i in review["items"] if i["category"] == "pending"]:
        decision = "rejected" if item["kind"] == "pii" and item["file_id"] == body_pii else "approved"
        if item["kind"] == "mapping":
            decision = "rejected"  # '해당 없음'으로 확정
        if item["kind"] == "entity":
            decision = "rejected"
        res = u4.post(f"/reviews/{review['id']}/items/{item['id']}/decision", json={"decision": decision})
        assert res.status_code == 200, res.text
    assert u4.post(f"/reviews/{review['id']}/complete").status_code == 200
    run = u4.post(f"/reviews/{review['id']}/integrate", json={"include_unmapped": True})
    assert run.status_code == 201, run.text
    run_id = run.json()["id"]
    download = u4.get(f"/runs/{run_id}/download")
    assert download.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(download.content))
    texts = [str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value is not None] + [
        c.comment.text for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.comment
    ]
    joined = "\n".join(texts)
    assert "인사혁신처" not in joined and "010-5678-1234" not in joined  # 포함 거부 파일은 결과·변경 내역에 없음
    assert "머리글만.xlsx" in joined and "erin@example.com" in joined  # 승인한 파일의 원문 머리글은 결과에 남을 수 있음
    assert u2.get(f"/runs/{run_id}/download").status_code == 404  # 다른 일반 사용자
    assert u2.get(f"/runs/{run_id}").status_code == 404
    assert admin.get(f"/runs/{run_id}").status_code == 200  # 관리자·소유자는 조회 가능 (기존 정책)
    assert "erin@example.com" not in admin.get(f"/runs/{run_id}").text  # 실행 기록 응답(화면)은 마스킹


def test_numeric_rrn_never_reaches_responses_and_history(accounts) -> None:
    """최종 확인 ③: 정수로 저장된 주민등록번호가 합계 검사·과거 검증 결과·저장된 검토 항목·자동화 규칙 키로 새지 않는다.
    다른 파일의 13자리 금액, 같은 파일의 일반 숫자·식별자는 그대로다."""
    from sqlalchemy import select

    from app.db import SessionLocal
    from app.models import Automation, ReviewItem, ReviewSession, User, ValidationRun

    u3 = accounts["user3"]
    raw_int, raw_grouped = 9001011234567, "9,001,011,234,567"
    rows = [["기관명", "번호값", "집행액"], ["합계", raw_int, 3], [8502022345678, 8502022345678, 1], ["기상청", 1, 2]]
    fid = _upload(u3, rows, "숫자주민번호.xlsx")
    run = u3.post("/validation", json={"file_ids": [fid]})
    assert run.status_code == 200 and str(raw_int) not in run.text and raw_grouped not in run.text
    assert any("개인정보가 탐지되어 합계 검사에서 제외" in n for n in run.json()["notes"])
    # 예전 방식으로 저장된 과거 기록(원문 expected/actual·메시지)도 조회 응답에서는 가린다
    with SessionLocal() as db:
        user_id = db.scalar(select(User.id).where(User.username == "user3"))
        legacy = {
            "blocks": [{"code": "sum_mismatch", "label": "합계 불일치", "file_id": fid,
                        "file_name": "숫자주민번호.xlsx",
                        "sheet": "현황", "row": 2, "column": "번호값", "expected": 8502022345679.0, "actual": raw_int,
                        "message": f"'합계' 행의 '번호값' 값 {raw_grouped}이(가) 다릅니다."},
                       {"code": "sum_mismatch", "label": "합계 불일치", "file_id": fid + 999, "file_name": "금액.xlsx",
                        "sheet": "현황", "row": 2, "column": "예산", "expected": 1, "actual": raw_int,
                        "message": f"금액 {raw_grouped}"}],
            "warnings": [], "notes": [], "entities": [],
            "pii": [{"file_id": fid, "file_name": "숫자주민번호.xlsx", "counts": {"rrn": 2, "mobile": 0, "email": 0},
                     "items": [], "acknowledged": False}],
        }  # fmt: skip
        record = ValidationRun(
            status="blocked", params="{}", result=json.dumps(legacy, ensure_ascii=False), created_by=user_id
        )
        db.add(record)
        session = ReviewSession(status="in_review", params="{}", file_ids=json.dumps([fid]), excluded_file_ids="[]",
                                remember_format=False, created_by=user_id)  # fmt: skip
        db.add(session)
        db.flush()
        db.add(ReviewItem(session_id=session.id, category="block", kind="sum_mismatch", file_id=fid,
                          reason=f"값 {raw_grouped}",
                          reason_source="rule", payload=json.dumps({"rows": [2], "value": raw_int}), before=raw_grouped,
                          after="8,502,022,345,679 (하위 행 합계)"))  # fmt: skip
        db.add(Automation(name="키확인", owner_id=user_id, column_rules="[]", validation="{}", output_format="{}",
                          entity_rules=json.dumps({"850202-2345678": "1741000", "기상청": "1360000"})))  # fmt: skip
        db.commit()
        run_id, review_id = record.id, session.id
    past = u3.get(f"/validation/{run_id}").json()
    first = json.dumps(past["blocks"][0], ensure_ascii=False)  # 주민등록번호가 탐지된 파일의 기록
    assert raw_grouped not in first and str(raw_int) not in first
    assert past["blocks"][0]["expected"] != 8502022345679.0  # 주민등록번호에서 계산된 13자리 값도 가린다
    assert (
        past["blocks"][1]["actual"] == raw_int and raw_grouped in past["blocks"][1]["message"]
    )  # 다른 파일의 금액은 그대로
    shown = u3.get(f"/reviews/{review_id}").text
    assert raw_grouped not in shown and str(raw_int) not in shown
    autos = u3.get("/automations").text
    assert "850202-2345678" not in autos and "기상청" in autos
