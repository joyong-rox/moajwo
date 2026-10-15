"""데이터 분석 1차: 실제 업로드 → 확정 → 대상 확인 → 실행 → 결과·Excel 흐름으로 확인한다 (합성 자료만)."""

import hashlib
import io
import sqlite3
from typing import Any

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app.config import DATA_DIR
from tests.test_accounts import owner  # noqa: F401
from tests.test_auth import _client, app_real_login  # noqa: F401

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
POP = "주민등록 인구 현황 2026년 9월 (단위: 명)"


def _xlsx(rows: list[list], title: str | None = POP, sheet: str = "현황") -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    if title:
        ws.append([title])
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.data_type = "s"  # 시험 자료: '=1+1'을 수식이 아닌 글자로 둔다
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _upload(c: TestClient, name: str, data: bytes) -> int:
    res = c.post("/files", files=[("files", (name, data, XLSX))])
    assert res.status_code == 200 and res.json()[0]["id"], res.text
    return res.json()[0]["id"]


def _confirm(c: TestClient, fid: int, mapping: dict[str, str | None]) -> None:
    cols = {col["name"]: col["column_id"] for col in c.get(f"/files/{fid}/preview").json()["columns"]}
    items = [{"sheet_id": "s0", "column_id": cols[n], "field_key": k} for n, k in mapping.items()]
    res = c.put(f"/files/{fid}/mapping/decisions", json={"decisions": items})
    assert res.status_code == 200, res.text


def _file(c: TestClient, name: str, rows: list[list], mapping: dict[str, str | None], **kw: Any) -> int:
    fid = _upload(c, name, _xlsx(rows, **kw))
    _confirm(c, fid, mapping)
    return fid


def _body(fids: list[int], **kw: Any) -> dict[str, Any]:
    return {"field_key": "population", "targets": [{"file_id": f} for f in fids], **kw}


def _preview(c: TestClient, body: dict) -> dict:
    res = c.post("/analysis/preview", json=body)
    assert res.status_code == 200, res.text
    return res.json()


def _run(c: TestClient, body: dict) -> dict:
    pv = _preview(c, body)
    res = c.post("/analysis/run", json={**body, "expected_fingerprint": pv["fingerprint"]})
    assert res.status_code == 201, res.text
    rid = res.json()["id"]
    detail = c.get(f"/analysis/runs/{rid}").json()
    groups = c.get(f"/analysis/runs/{rid}/groups", params={"limit": 200}).json()["items"]
    return {"id": rid, "detail": detail, "groups": groups, "preview": pv}


def _db(sql: str, *args: Any) -> list[tuple]:
    with sqlite3.connect(DATA_DIR / "moajwo.db") as con:
        return con.execute(sql, args).fetchall()


@pytest.fixture(scope="module")
def an(app_real_login, owner) -> dict[str, TestClient]:  # noqa: F811
    made = {"owner": owner}
    for name in ("an-user1", "an-user2"):
        res = owner.post("/users", json={"username": name, "display_name": name.upper(), "status": "active",
                                         "password": "pass-" + name})  # fmt: skip
        assert res.status_code == 201, res.text
        made[name] = _client(app_real_login, name, "pass-" + name)
    return made


THREE = [["시도", "인구수"], ["서울특별시", 10000], ["부산광역시", 20000], ["대구광역시", 30000]]


# ---------------------------------------------------------------- 핵심 결과


def test_three_regions_sum_mean_min_max_count_and_bins(an) -> None:
    u = an["an-user1"]
    fid = _file(u, "세지역_인구.xlsx", THREE, {"시도": "sido", "인구수": "population"})
    r = _run(u, _body([fid]))
    (g,) = r["groups"]
    assert (g["sum"], g["mean"], g["min"], g["max"], g["count"]) == (60000, 20000, 10000, 30000, 3)
    assert g["partial"] is False and r["detail"]["result"]["partial"] is False
    cond = r["detail"]["conditions"]
    assert cond["unit"] == "명" and cond["periods"] == ["2026-09"] and "유효한 숫자의 개수" in cond["mean_denominator"]
    bins = {"on": "sum", "ranges": [{"max": 20000, "max_inclusive": False}, {"min": 20000, "min_inclusive": True}]}
    r2 = _run(u, _body([fid], group_by=["sido"], bins=bins))
    by = {tuple(g["keys"]): g for g in r2["groups"]}
    assert {k[0]: g["bin"] for k, g in by.items()} == {
        "서울특별시": "20,000 미만",
        "부산광역시": "20,000 이상",
        "대구광역시": "20,000 이상",
    }
    counts = {b["label"]: b["groups"] for b in r2["detail"]["result"]["bins"]}
    assert counts["20,000 미만"] == 1 and counts["20,000 이상"] == 2 and counts["미분류"] == 0


def test_different_column_names_with_same_confirmed_meaning_are_combined(an) -> None:
    u = an["an-user1"]
    a = _file(u, "이름A.xlsx", [["시도", "인구수"], ["서울특별시", 10000]], {"시도": "sido", "인구수": "population"})
    b = _file(
        u,
        "이름B.xlsx",
        [["지역", "주민등록인구"], ["부산광역시", 20000]],
        {"지역": "sido", "주민등록인구": "population"},
    )
    r = _run(u, _body([a, b]))
    assert r["groups"][0]["sum"] == 30000 and r["groups"][0]["count"] == 2
    cols = {t["column"] for t in r["detail"]["inputs"]}
    assert cols == {"인구수", "주민등록인구"}


# ---------------------------------------------------------------- 단위·기간·구분어


def test_units_are_not_mixed(an) -> None:
    u = an["an-user1"]
    a = _file(u, "단위명.xlsx", THREE[:2], {"시도": "sido", "인구수": "population"})
    b = _file(u, "단위천명.xlsx", [["시도", "인구수"], ["부산광역시", 20]], {"시도": "sido", "인구수": "population"},
              title="주민등록 인구 현황 2026년 9월 (단위: 천명)")  # fmt: skip
    pv = _preview(u, _body([a, b]))
    assert any("단위가 다른 자료" in p for p in pv["problems"])
    res = u.post("/analysis/run", json={**_body([a, b]), "expected_fingerprint": pv["fingerprint"]})
    assert res.status_code == 409
    r = _run(u, _body([a, b], unit="명"))
    assert r["groups"][0]["sum"] == 10000
    excluded = {t["file_id"]: t for t in r["detail"]["result"]["excluded_targets"]}
    assert "단위가 다름" in excluded[b]["reasons"][0]


def test_periods_are_split_not_summed(an) -> None:
    u = an["an-user1"]
    a = _file(u, "8월.xlsx", [["시도", "인구수"], ["서울특별시", 100]], {"시도": "sido", "인구수": "population"},
              title="인구 현황 2026년 8월 (단위: 명)")  # fmt: skip
    b = _file(u, "9월.xlsx", [["시도", "인구수"], ["서울특별시", 101]], {"시도": "sido", "인구수": "population"},
              title="인구 현황 2026년 9월 (단위: 명)")  # fmt: skip
    r = _run(u, _body([a, b]))
    assert r["detail"]["result"]["period_split"] is True
    assert {tuple(g["keys"]): g["sum"] for g in r["groups"]} == {("2026-08",): 100, ("2026-09",): 101}


def test_qualifier_variants_are_not_merged(an) -> None:
    u = an["an-user1"]
    a = _file(u, "남.xlsx", [["시도", "인구_남"], ["서울특별시", 5]], {"시도": "sido", "인구_남": "population"})
    b = _file(u, "여.xlsx", [["시도", "인구_여"], ["부산광역시", 6]], {"시도": "sido", "인구_여": "population"})
    pv = _preview(u, _body([a, b]))
    assert any("구분이 다른 열" in p for p in pv["problems"]) and len(pv["variants"]) == 2
    male = next(v for v in pv["variants"] if "남" in v)
    r = _run(u, _body([a, b], variant=male))
    assert r["groups"][0]["sum"] == 5
    assert any("구분" in x for x in r["detail"]["result"]["excluded_targets"][0]["reasons"])


def test_missing_period_or_unit_needs_user_confirmation(an) -> None:
    u = an["an-user1"]
    fid = _file(u, "정보없음.xlsx", THREE[:2], {"시도": "sido", "인구수": "population"}, title=None)
    pv = _preview(u, _body([fid]))
    (t,) = pv["targets"]
    assert (
        t["status"] == "needs_input"
        and any("기간" in x for x in t["reasons"])
        and any("단위" in x for x in t["reasons"])
    )
    assert u.post("/analysis/run", json={**_body([fid]), "expected_fingerprint": pv["fingerprint"]}).status_code == 409
    body = {
        "field_key": "population",
        "targets": [{"file_id": fid, "sheet_id": "s0", "unit": "명", "period": "2026-09"}],
    }
    r = _run(u, body)
    assert r["groups"][0]["sum"] == 10000
    (inp,) = r["detail"]["inputs"]
    assert (inp["unit_source"], inp["period_source"]) == ("사용자 확인", "사용자 확인")


def test_auto_suggestion_is_not_used(an) -> None:
    u = an["an-user1"]
    fid = _upload(u, "확정없음.xlsx", _xlsx(THREE))
    _confirm(u, fid, {"시도": "sido"})  # 인구수는 확정하지 않음
    pv = _preview(u, _body([fid]))
    assert pv["targets"][0]["status"] == "excluded" and "확정된 열" in pv["targets"][0]["reasons"][0]


# ---------------------------------------------------------------- 합계·소계·상위 지역·중복


def test_totals_subtotals_and_upper_regions_are_not_double_counted(an) -> None:
    u = an["an-user1"]
    rows = [["시도", "인구수"], ["서울특별시", 10], ["부산광역시", 20], ["합계", 30], ["계", 30]]
    fid = _file(u, "합계행.xlsx", rows, {"시도": "sido", "인구수": "population"})
    r = _run(u, _body([fid]))
    assert r["groups"][0]["sum"] == 30
    reasons = r["detail"]["result"]["reason_counts"]
    assert sum(reasons.get(k, 0) for k in ("subtotal", "total_like")) == 2
    rows = [
        ["시도", "시군구", "인구수"],
        ["서울특별시", None, 100],
        ["서울특별시", "종로구", 40],
        ["서울특별시", "중구", 60],
    ]
    fid2 = _file(u, "상하위.xlsx", rows, {"시도": "sido", "시군구": "sigungu", "인구수": "population"})
    r2 = _run(u, _body([fid2]))
    assert r2["groups"][0]["sum"] == 100 and r2["detail"]["result"]["reason_counts"].get("upper_region") == 1


def test_duplicate_files_and_conflicting_values_block_the_run(an) -> None:
    u = an["an-user1"]
    data = _xlsx(THREE)
    a, b = _upload(u, "중복1.xlsx", data), _upload(u, "중복2.xlsx", data)
    for fid in (a, b):
        _confirm(u, fid, {"시도": "sido", "인구수": "population"})
    pv = _preview(u, _body([a, b]))
    assert any("같은 내용의 파일" in p for p in pv["problems"])
    assert u.post("/analysis/run", json={**_body([a, b]), "expected_fingerprint": pv["fingerprint"]}).status_code == 409
    c = _file(
        u,
        "충돌.xlsx",
        [["시도", "인구수"], ["서울특별시", 15000], ["광주광역시", 1]],
        {"시도": "sido", "인구수": "population"},
    )
    pv2 = _preview(u, _body([a, c]))
    (ov,) = pv2["overlaps"]
    assert ov["count"] == 1 and ov["different_value"] == 1
    assert ov["samples"][0]["a"]["value"] == 10000 and ov["samples"][0]["b"]["value"] == 15000
    assert (
        u.post("/analysis/run", json={**_body([a, c]), "expected_fingerprint": pv2["fingerprint"]}).status_code == 409
    )


# ---------------------------------------------------------------- 결측·비공개·0·잘못된 숫자


def test_missing_values_are_not_zero_and_marked_partial(an) -> None:
    u = an["an-user1"]
    rows = [["시도", "인구수"], ["서울특별시", 100], ["부산광역시", "-"], ["대구광역시", None], ["인천광역시", "x"],
            ["광주광역시", 0], ["대전광역시", "비공개"], ["울산광역시", "12a"]]  # fmt: skip
    fid = _file(u, "결측.xlsx", rows, {"시도": "sido", "인구수": "population"})
    r = _run(u, _body([fid]))
    (g,) = r["groups"]
    assert (g["count"], g["sum"], g["mean"], g["min"], g["max"]) == (2, 100, 50, 0, 100)
    assert g["partial"] is True and r["detail"]["result"]["partial"] is True
    reasons = r["detail"]["result"]["reason_counts"]
    assert reasons == {"missing_marker": 3, "blank": 1, "not_number": 1}
    r2 = _run(u, _body([fid], group_by=["sido"]))
    by = {g["keys"][0]: g for g in r2["groups"]}
    assert by["부산광역시"]["count"] == 0 and by["부산광역시"]["sum"] is None and by["부산광역시"]["mean"] is None
    xlsx = openpyxl.load_workbook(io.BytesIO(u.get(f"/analysis/runs/{r2['id']}/download").content))
    sheet = {row[0]: row for row in xlsx["결과"].iter_rows(min_row=2, values_only=True)}
    assert sheet["부산광역시"][1] == "계산할 값 없음" and sheet["광주광역시"][1] == 0


# ---------------------------------------------------------------- 변경·권한·보존·Excel


def test_mapping_change_after_preview_is_rejected(an) -> None:
    u = an["an-user1"]
    fid = _file(u, "변경.xlsx", THREE, {"시도": "sido", "인구수": "population"})
    pv = _preview(u, _body([fid]))
    _confirm(u, fid, {"시도": "region"})
    res = u.post("/analysis/run", json={**_body([fid]), "expected_fingerprint": pv["fingerprint"]})
    assert res.status_code == 409 and "바뀌었습니다" in res.json()["detail"]["message"]


def test_other_users_files_and_results_are_blocked(an) -> None:
    u1, u2, owner_c = an["an-user1"], an["an-user2"], an["owner"]
    fid = _file(u1, "권한.xlsx", THREE, {"시도": "sido", "인구수": "population"})
    assert u2.post("/analysis/preview", json=_body([fid])).status_code == 404
    r = _run(u1, _body([fid]))
    for path in ("", "/groups", "/excluded", "/detail", "/download"):
        assert u2.get(f"/analysis/runs/{r['id']}{path}").status_code == 404
    seen = owner_c.get(f"/analysis/runs/{r['id']}")
    assert seen.status_code == 200 and seen.json()["read_only"] is True
    assert all(x["id"] != r["id"] for x in u2.get("/analysis/runs").json())


def test_no_side_effects_on_mappings_learning_or_files(an, monkeypatch) -> None:
    from app import mapping

    u = an["an-user1"]
    fid = _file(u, "보존.xlsx", THREE, {"시도": "sido", "인구수": "population"})
    tables = ("mapping_decisions", "learned_aliases", "mapping_approvals", "mapping_analyses", "runs")
    before = {t: _db(f"select count(*) from {t}")[0][0] for t in tables}
    path = _db("select stored_path, sha256 from source_files where id = ?", fid)[0]

    def boom(*a: Any, **k: Any) -> Any:
        raise AssertionError("분석이 매핑 엔진(모델)을 불렀음")

    monkeypatch.setattr(mapping, "get_engine", boom)
    monkeypatch.setattr(mapping.MappingEngine, "map_columns", boom)
    r = _run(u, _body([fid]))
    u.get(f"/analysis/runs/{r['id']}/download")
    assert {t: _db(f"select count(*) from {t}")[0][0] for t in tables} == before
    assert hashlib.sha256((DATA_DIR / path[0]).read_bytes()).hexdigest() == path[1]


def test_excel_matches_screen_and_stores_text_not_formulas(an) -> None:
    u = an["an-user1"]
    rows = [["기관명", "인구수"], ["=1+1", 7], ["@기관A", 8], ["기상청", "-"]]
    fid = _file(u, "수식.xlsx", rows, {"기관명": "org_name", "인구수": "population"})
    r = _run(u, _body([fid], group_by=["org_name"]))
    raw = u.get(f"/analysis/runs/{r['id']}/download").content
    wb = openpyxl.load_workbook(io.BytesIO(raw))
    ws = wb["결과"]
    cells = {ws.cell(row=i, column=1).value: ws.cell(row=i, column=1) for i in range(2, ws.max_row + 1)}
    assert cells["=1+1"].data_type == "s" and cells["@기관A"].data_type == "s"
    sheet = {row[0]: row for row in ws.iter_rows(min_row=2, values_only=True)}
    for g in r["groups"]:
        assert sheet[g["keys"][0]][1] == (g["sum"] if g["sum"] is not None else "계산할 값 없음")
    reasons = {row[0]: row[1] for row in wb["제외"].iter_rows(min_row=2, max_row=4, values_only=True) if row[0]}
    assert (
        reasons.get("비공개·해당 없음 표기 ('-', 'x', '…' 등)")
        == r["detail"]["result"]["reason_counts"]["missing_marker"]
    )
    conditions = {row[0]: row[1] for row in wb["조건"].iter_rows(min_row=1, values_only=True) if row[0]}
    assert conditions["단위"] == "명" and conditions["결과 상태"].startswith("부분 집계")


def test_result_pages_do_not_recompute(an, monkeypatch) -> None:
    from app import analysis

    u = an["an-user1"]
    fid = _file(u, "페이지.xlsx", [["시도", "인구수"], *[[f"지역{i}", i] for i in range(1, 121)]],
                {"시도": "sido", "인구수": "population"})  # fmt: skip
    body = {**_body([fid], group_by=["sido"]), "targets": [{"file_id": fid, "sheet_id": "s0"}]}
    r = _run(u, body)
    calls = []
    monkeypatch.setattr(analysis, "file_view", lambda *a, **k: calls.append(1))
    pages = [u.get(f"/analysis/runs/{r['id']}/groups", params={"offset": o, "limit": 50}).json() for o in (0, 50, 100)]
    assert [len(p["items"]) for p in pages] == [50, 50, 20] and pages[0]["total"] == 120 and not calls


def test_input_limits_and_bin_overlaps_are_rejected(an, monkeypatch) -> None:
    from app import analysis

    u = an["an-user1"]
    fid = _file(u, "제한.xlsx", THREE, {"시도": "sido", "인구수": "population"})
    overlap = {"on": "sum", "ranges": [{"min": 0, "max": 10, "max_inclusive": True}, {"min": 10, "max": 20}]}
    res = u.post("/analysis/preview", json=_body([fid], bins=overlap))
    assert res.status_code == 422 and "겹칩니다" in res.json()["detail"]
    ok = {"on": "sum", "ranges": [{"min": 0, "max": 10}, {"min": 10, "max": 20}]}
    assert u.post("/analysis/preview", json=_body([fid], bins=ok)).status_code == 200
    monkeypatch.setattr(analysis, "MAX_ROWS", 2)
    res = u.post("/analysis/preview", json=_body([fid]))
    assert res.status_code == 413 and "행이" in res.json()["detail"]


# ---------------------------------------------------------------- 보완: 상위·하위 지역의 비교 범위


def _ok_run(c: TestClient, body: dict) -> int:
    return c.post("/analysis/run", json={**body, "expected_fingerprint": _preview(c, body)["fingerprint"]}).status_code


def test_upper_region_is_kept_when_sub_regions_are_in_another_period(an) -> None:
    u = an["an-user1"]
    a = _file(u, "8월_서울전체.xlsx", [["시도", "인구수"], ["서울특별시", 100]],
              {"시도": "sido", "인구수": "population"}, title="인구 현황 2026년 8월 (단위: 명)")  # fmt: skip
    b = _file(u, "9월_종로구.xlsx", [["시도", "시군구", "인구수"], ["서울특별시", "종로구", 40]],
              {"시도": "sido", "시군구": "sigungu", "인구수": "population"},
              title="인구 현황 2026년 9월 (단위: 명)")  # fmt: skip
    r = _run(u, _body([a, b]))
    got = {tuple(g["keys"]): (g["sum"], g["count"]) for g in r["groups"]}
    assert got == {("2026-08",): (100, 1), ("2026-09",): (40, 1)}
    assert "upper_region" not in r["detail"]["result"]["reason_counts"]


def test_upper_and_lower_with_different_vehicle_types_are_not_compared(an) -> None:
    u = an["an-user1"]
    rows = [
        ["시도", "시군구", "차종", "등록대수"],
        ["서울특별시", None, "승용", 100],
        ["서울특별시", "종로구", "화물", 40],
    ]
    mapping = {"시도": "sido", "시군구": "sigungu", "차종": "vehicle_type", "등록대수": "vehicles_total"}
    fid = _file(u, "차종별.xlsx", rows, mapping, title="자동차 등록 현황 2026년 9월 (단위: 대)")
    r = _run(u, {**_body([fid], group_by=["vehicle_type"]), "field_key": "vehicles_total"})
    assert {g["keys"][0]: g["sum"] for g in r["groups"]} == {"승용": 100, "화물": 40}
    assert "upper_region" not in r["detail"]["result"]["reason_counts"]


def test_partial_sub_regions_need_an_aggregation_level_choice(an) -> None:
    u = an["an-user1"]
    rows = [["시도", "시군구", "인구수"], ["서울특별시", None, 100], ["서울특별시", "종로구", 40]]
    fid = _file(u, "일부하위.xlsx", rows, {"시도": "sido", "시군구": "sigungu", "인구수": "population"})
    pv = _preview(u, _body([fid]))
    assert any("집계 수준" in p for p in pv["problems"]) and pv["region_levels"]["conflicts"]
    assert _ok_run(u, _body([fid])) == 409  # 상위 행을 자동으로 빼고 '전체 집계'처럼 만들지 않음
    low = _run(u, _body([fid], region_level="lower"))
    res = low["detail"]["result"]
    assert low["groups"][0]["sum"] == 40 and res["level_mismatch"] == 1 and res["reason_counts"] == {"upper_region": 1}
    assert any("하위 지역 합계" in w for w in low["detail"]["conditions"]["warnings"])
    up = _run(u, _body([fid], region_level="upper"))
    assert up["groups"][0]["sum"] == 100 and up["detail"]["result"]["reason_counts"] == {"lower_region": 1}
    # 하위 합계가 상위 값과 같으면 상위 행만 중복으로 뺀다 (기존 사례 유지 — 한 열에 시도·시군구를 함께 적은 표도 같게)
    whole = [["행정구역", "인구수"], ["서울특별시", 100], ["서울특별시 종로구", 40], ["서울특별시 중구", 60]]
    fid2 = _file(u, "한열구조.xlsx", whole, {"행정구역": "region", "인구수": "population"})
    r = _run(u, _body([fid2]))
    assert r["groups"][0]["sum"] == 100 and r["detail"]["result"]["reason_counts"] == {"upper_region": 1}


# ---------------------------------------------------------------- 보완: 반복 행·표현 차이에 따른 중복


def test_repeated_rows_in_one_file_need_confirmation(an) -> None:
    u = an["an-user1"]
    rows = [["시도", "인구수"], ["서울특별시", 100], ["서울특별시", 100], ["부산광역시", 50], ["부산광역시", 70]]
    fid = _file(u, "반복행.xlsx", rows, {"시도": "sido", "인구수": "population"})
    pv = _preview(u, _body([fid]))
    (t,) = pv["targets"]
    assert t["status"] == "needs_input" and any("반복" in x for x in t["reasons"])
    assert t["repeated"]["keys"] == 2 and t["repeated"]["same_value"] == 1 and t["repeated"]["different_value"] == 1
    assert _ok_run(u, _body([fid])) == 409  # 반복 행을 임의로 합치거나 지우지 않는다
    body = {"field_key": "population", "targets": [{"file_id": fid, "sheet_id": "s0", "rows_are_records": True}]}
    r = _run(u, body)  # 사용자가 '서로 다른 기록'이라고 확인한 뒤에만 모두 더한다
    assert r["groups"][0]["sum"] == 320 and r["groups"][0]["count"] == 4


def test_repeated_in_one_file_and_once_in_another_is_still_an_overlap(an) -> None:
    u = an["an-user1"]
    a = _file(u, "반복A.xlsx", [["시도", "인구수"], ["서울특별시", 100], ["서울특별시", 100]],
              {"시도": "sido", "인구수": "population"})  # fmt: skip
    b = _file(u, "한번B.xlsx", [["시도", "인구수"], ["서울특별시", 100]], {"시도": "sido", "인구수": "population"})
    plain = _preview(u, _body([a, b]))  # 반복 확인 전에도 다른 자료와의 중복을 함께 보여 준다
    assert any("같은 대상" in p for p in plain["problems"]) and plain["targets"][0]["status"] == "needs_input"
    targets = [{"file_id": a, "sheet_id": "s0", "rows_are_records": True}, {"file_id": b}]
    body = {"field_key": "population", "targets": targets}
    pv = _preview(u, body)
    (ov,) = pv["overlaps"]
    assert ov["count"] == 1 and any("같은 대상" in p for p in pv["problems"])
    assert _ok_run(u, body) == 409


def test_same_region_in_different_mapping_structures_is_an_overlap(an) -> None:
    u = an["an-user1"]
    a = _file(u, "구조_시도.xlsx", [["시도", "시군구", "인구수"], ["서울특별시", None, 100],
                                   ["서울특별시", "종로구", 40]],
              {"시도": "sido", "시군구": "sigungu", "인구수": "population"})  # fmt: skip
    b = _file(u, "구조_지역.xlsx", [["지역", "인구수"], ["서울특별시", 120], ["서울특별시 종로구", 40]],
              {"지역": "region", "인구수": "population"})  # fmt: skip
    pv = _preview(u, _body([a, b], region_level="lower"))
    (ov,) = pv["overlaps"]
    assert ov["count"] == 2 and ov["same_value"] == 1 and ov["different_value"] == 1
    assert _ok_run(u, _body([a, b], region_level="lower")) == 409


# ---------------------------------------------------------------- 보완: 기간 정보 충돌


def test_conflicting_period_columns_block_the_run(an) -> None:
    u = an["an-user1"]
    a = _file(u, "연도월충돌.xlsx", [["연도", "월", "시도", "인구수"], [2025, "2026-09", "서울특별시", 100]],
              {"연도": "year", "월": "month", "시도": "sido", "인구수": "population"})  # fmt: skip
    pv = _preview(u, _body([a]))
    assert any("기간 열" in p and "2025" in p and "2026-09" in p for p in pv["problems"])
    assert _ok_run(u, _body([a])) == 409
    b = _file(u, "기준일연도충돌.xlsx", [["기준일", "연도", "시도", "인구수"], ["2026-09-30", 2025, "서울특별시", 100]],
              {"기준일": "base_date", "연도": "year", "시도": "sido", "인구수": "population"})  # fmt: skip
    assert any("기간 열" in p for p in _preview(u, _body([b]))["problems"])


def test_file_name_date_is_not_compared_as_the_statistics_date(an) -> None:
    u = an["an-user1"]
    fid = _file(u, "인구현황_2026년10월발행.xlsx", THREE, {"시도": "sido", "인구수": "population"})
    (t,) = _preview(u, _body([fid]))["targets"]
    assert t["status"] == "ok" and t["period"] == "2026-09"
    assert "제목" in t["period_source"] and "파일명" in t["period_source"]


def test_stock_values_with_different_base_dates_in_one_month_are_not_summed(an) -> None:
    u = an["an-user1"]
    a = _file(u, "기준일15.xlsx", [["기준일", "시도", "인구수"], ["2026-09-15", "서울특별시", 100]],
              {"기준일": "base_date", "시도": "sido", "인구수": "population"})  # fmt: skip
    b = _file(u, "기준일30.xlsx", [["기준일", "시도", "인구수"], ["2026-09-30", "부산광역시", 50]],
              {"기준일": "base_date", "시도": "sido", "인구수": "population"})  # fmt: skip
    pv = _preview(u, _body([a, b]))
    assert any("기준일" in p and "2026-09-15" in p and "2026-09-30" in p for p in pv["problems"])
    assert _ok_run(u, _body([a, b])) == 409
    flow = [["일자", "기관명", "집행액"], ["2026-09-01", "기관A", 10], ["2026-09-15", "기관B", 5]]
    c = _file(u, "집행.xlsx", flow, {"일자": "base_date", "기관명": "org_name", "집행액": "expenditure"},
              title="집행 실적 (단위: 원)")  # fmt: skip
    r = _run(u, {**_body([c]), "field_key": "expenditure"})  # 발생 실적은 같은 달끼리 더해도 된다
    assert r["groups"][0]["sum"] == 15 and r["detail"]["conditions"]["periods"] == ["2026-09"]


# ---------------------------------------------------------------- 보완: 새 입력·출력 경로의 개인정보 마스킹


def test_user_inputs_are_masked_in_results_conditions_and_excel(an) -> None:
    u = an["an-user1"]
    fid = _file(u, "마스킹.xlsx", THREE, {"시도": "sido", "인구수": "population"})
    raw = ["aa@example.com", "ab@example.com", "010-1234-5678", "hong@example.com"]
    ranges = [
        {"label": raw[0], "max": 20000},
        {"label": raw[1], "min": 20000, "max": 30000},
        {"label": "같은이름", "min": 30000, "max": 40000},
        {"label": "같은이름", "min": 40000},
    ]
    filters = {"dims": {"sido": ["서울특별시", "부산광역시", "대구광역시", raw[2]]}}
    r = _run(u, _body([fid], group_by=["sido"], bins={"on": "sum", "ranges": ranges}, filters=filters))
    counts = r["detail"]["result"]["bins"]
    assert [b["groups"] for b in counts[:4]] == [1, 1, 1, 0]  # 가린 이름·같은 이름이 같아도 구간끼리 합치지 않음
    shown = [b["label"] for b in counts[:4]]
    assert len(set(shown)) == 4 and all("a***@example.com" in s for s in shown[:2])
    stored = _db("select summary || conditions || inputs || result from analysis_runs where id = ?", r["id"])[0][0]
    xlsx = openpyxl.load_workbook(io.BytesIO(u.get(f"/analysis/runs/{r['id']}/download").content))
    cells = [str(v) for ws in xlsx for row in ws.iter_rows(values_only=True) for v in row if v is not None]
    for text in raw[:3]:
        assert text not in stored and text not in str(r) and not any(text in c for c in cells)
    by = {row[0]: row for row in xlsx["결과"].iter_rows(min_row=2, values_only=True)}
    assert by["대구광역시"][1] == 30000 and isinstance(by["대구광역시"][2], int | float)  # 숫자는 숫자 그대로
    nounit = _file(u, "단위없음.xlsx", THREE[:2], {"시도": "sido", "인구수": "population"}, title=None)
    target = {"file_id": nounit, "sheet_id": "s0", "unit": raw[3], "period": "2026-09"}
    r2 = _run(u, {"field_key": "population", "unit": raw[3], "targets": [target]})
    stored2 = _db("select conditions || inputs || summary from analysis_runs where id = ?", r2["id"])[0][0]
    assert raw[3] not in stored2 and "h***@example.com" in stored2
    bad_bins = {"on": "sum", "ranges": [{"label": raw[3], "min": 0, "max": 10}, {"label": raw[0], "min": 5}]}
    for bad in ({"group_by": [raw[3]]}, {"filters": {"dims": {raw[3]: ["x"]}}}, {"bins": bad_bins}):
        res = u.post("/analysis/preview", json=_body([fid], **bad))
        assert res.status_code == 422 and raw[3] not in res.text and raw[0] not in res.text, res.text


# ---------------------------------------------------------------- 보완: 최고 관리자는 남의 자료로 새 결과를 만들지 않음


def test_manager_cannot_save_new_results_from_other_users_files(an) -> None:
    u1, owner_c = an["an-user1"], an["owner"]
    fid = _file(u1, "남의파일.xlsx", THREE, {"시도": "sido", "인구수": "population"})
    mine = _run(u1, _body([fid]))
    before = _db("select count(*) from analysis_runs")[0][0]
    assert owner_c.post("/analysis/preview", json=_body([fid])).status_code == 403
    res = owner_c.post("/analysis/run", json={**_body([fid]), "expected_fingerprint": mine["preview"]["fingerprint"]})
    assert res.status_code == 403 and _db("select count(*) from analysis_runs")[0][0] == before
    assert owner_c.get(f"/analysis/runs/{mine['id']}").json()["read_only"] is True  # 기존 결과 조회는 그대로
    assert owner_c.get(f"/analysis/runs/{mine['id']}/download").status_code == 200


def test_run_with_conditions_changed_after_preview_is_rejected(an) -> None:
    u = an["an-user1"]
    fid = _file(u, "조건변경.xlsx", THREE, {"시도": "sido", "인구수": "population"})
    pv = _preview(u, _body([fid]))
    changed = _body([fid], group_by=["sido"])
    res = u.post("/analysis/run", json={**changed, "expected_fingerprint": pv["fingerprint"]})
    assert res.status_code == 409 and "바뀌었습니다" in res.json()["detail"]["message"]
