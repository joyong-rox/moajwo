"""수량·식별자 판정 — 임의의 열 이름을 확정하는 API 흐름 (전용 시험 계정, 실제 쿠키 로그인).

학습 별칭은 사용자별이라 전용 계정을 쓰면 다른 시험 계정의 학습 기록과 섞이지 않는다.
재현: 수량 이름 목록 밖의 '규모'를 expenditure로 확정해도 긴 값 때문에 글자가 되어 합계 검사·분석에서 빠졌고,
'매출식별부호'는 '매출' 때문에 숫자로 남아 브라우저(JSON 숫자)와 Excel에서 끝자리만 다른 식별자가 같아졌다.
"""

import csv
import io
import re
from typing import Any

import openpyxl
import pytest
from fastapi.testclient import TestClient

from tests.test_accounts import owner  # noqa: F401
from tests.test_auth import _client, app_real_login  # noqa: F401

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
TITLE = "2026년 9월 사업별 집행 현황 (단위: 원)"
BAD = [["합계", "2000000000000000"], ["사업A", "1000000000000000"], ["사업B", "100"], ["사업C", "200"]]
BAD_COMMA = [["합계", "2,000,000,000,000,000"], ["사업A", "1,000,000,000,000,000"], ["사업B", "100"], ["사업C", "200"]]
OK = [["합계", "1000000000000010"], ["사업A", "1000000000000000"], ["사업B", "12.5"], ["사업C", "-2.5"],
      ["사업D", "0"], ["사업E", "-"]]  # fmt: skip
LONG = ["12345678901234567890123", "12345678901234567890124", "12345678901234567890125"]
SHORT = ["1741000", "1360000", "1310000"]


@pytest.fixture(scope="module")
def pu(app_real_login, owner) -> TestClient:  # noqa: F811
    res = owner.post("/users", json={"username": "pid-user1", "display_name": "PID-USER1", "status": "active",
                                     "password": "pass-pid-user1"})  # fmt: skip
    assert res.status_code == 201, res.text
    return _client(app_real_login, "pid-user1", "pass-pid-user1")


def _file(rows: list[list], kind: str, title: str | None = TITLE) -> tuple[bytes, str]:
    data = ([[title]] if title else []) + rows
    if kind == "csv":
        buf = io.StringIO()
        csv.writer(buf, lineterminator="\n").writerows(data)
        return buf.getvalue().encode("utf-8"), "text/csv"
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in data:
        ws.append(r)
    for row in ws.iter_rows(min_row=2 if title else 1):
        for cell in row:
            cell.data_type = "s"  # 모두 글자 셀
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), XLSX


def _upload(c: TestClient, name: str, data: bytes, ctype: str) -> int:
    res = c.post("/files", files=[("files", (name, data, ctype))])
    assert res.status_code == 200 and res.json()[0]["id"], res.text
    return res.json()[0]["id"]


def _confirm(c: TestClient, fid: int, mapping: dict[str, str | None]) -> dict:
    pv = c.get(f"/files/{fid}/preview").json()
    ids = {col["name"]: col["column_id"] for col in pv["columns"]}
    items = [{"sheet_id": pv["sheet_id"], "column_id": ids[n], "field_key": k} for n, k in mapping.items()]
    assert c.put(f"/files/{fid}/mapping/decisions", json={"decisions": items}).status_code == 200
    return pv


def _integrate(c: TestClient, review: dict) -> list:
    rid = review["id"]
    for item in review["items"]:
        if item["category"] == "pending":
            body = {"decision": "rejected"} if item["kind"] == "entity" else {"decision": "approved"}
            assert c.post(f"/reviews/{rid}/items/{item['id']}/decision", json=body).status_code == 200
    assert c.post(f"/reviews/{rid}/complete").status_code == 200
    run = c.post(f"/reviews/{rid}/integrate", json={"include_unmapped": True})
    assert run.status_code == 201, run.text
    return list(
        openpyxl.load_workbook(io.BytesIO(c.get(f"/runs/{run.json()['id']}/download").content))
        .worksheets[0]
        .iter_rows()
    )


def _excel_column(grid: list, name: str) -> list[tuple[Any, str]]:
    hi = next(k for k, row in enumerate(grid) if any(cell.value in (name, f"{name} (원본 열)") for cell in row))
    k = next(i for i, cell in enumerate(grid[hi]) if cell.value in (name, f"{name} (원본 열)"))
    return [(row[k].value, row[k].data_type) for row in grid[hi + 1 :] if row[k].value is not None]


def _amount(c: TestClient, column: str, rows: list[list], kind: str, tag: str) -> dict:
    data, ctype = _file([["구분", column], *rows], kind)
    fid = _upload(c, f"{tag}.{kind}", data, ctype)
    pv = _confirm(c, fid, {column: "expenditure", "구분": "department"})
    j = [col["name"] for col in pv["columns"]].index(column)
    validation = c.post("/validation", json={"file_ids": [fid]}).json()
    body = {"field_key": "expenditure", "targets": [{"file_id": fid}]}
    fingerprint = c.post("/analysis/preview", json=body).json()["fingerprint"]
    run = c.post("/analysis/run", json={**body, "expected_fingerprint": fingerprint})
    assert run.status_code == 201, run.text
    group = c.get(f"/analysis/runs/{run.json()['id']}/groups").json()["items"][0]
    excluded = c.get(f"/analysis/runs/{run.json()['id']}/excluded").json()["items"]
    return {
        "fid": fid,
        "parsed": [r["cells"][j] for r in pv["rows"]],
        "blocks": [b["code"] for b in validation["blocks"]],
        "block_messages": [b["message"] for b in validation["blocks"]],
        "status": validation["status"],
        "warnings": [w["code"] for w in validation["warnings"]],
        "sum": group["sum"],
        "count": group["count"],
        "excluded": sorted(e["reason"] for e in excluded),
        "review": c.post("/reviews", json={"file_ids": [fid]}).json(),
    }


# ---------------------------------------------------------------- 수량 이름 목록 밖의 확정 수량 열


@pytest.mark.parametrize("column", ["규모", "집행액"])
@pytest.mark.parametrize("rows", [BAD, BAD_COMMA], ids=["쉼표 없음", "쉼표 있음"])
@pytest.mark.parametrize("kind", ["csv", "xlsx"])
def test_confirmed_amount_total_mismatch_is_blocked(pu, column: str, rows: list[list], kind: str) -> None:
    r = _amount(pu, column, rows, kind, f"합계오류_{column}_{kind}_{len(rows[0][1])}")
    assert r["parsed"] == [2000000000000000, 1000000000000000, 100, 200]  # 이름·쉼표와 관계없이 같은 값·자료형
    assert r["blocks"] == ["sum_mismatch"]  # 2,000,000,000,000,000 ≠ 1,000,000,000,000,300
    assert (r["sum"], r["count"], r["excluded"]) == (1000000000000300, 3, ["subtotal"])
    assert r["review"]["can_complete"] is False
    assert pu.post(f"/reviews/{r['review']['id']}/integrate", json={}).status_code == 409


@pytest.mark.parametrize("column", ["규모", "집행액"])
@pytest.mark.parametrize("kind", ["csv", "xlsx"])
def test_confirmed_amount_normal_data_analyzes_and_integrates_as_numbers(pu, column: str, kind: str) -> None:
    r = _amount(pu, column, OK, kind, f"정상_{column}_{kind}")
    assert r["parsed"] == [1000000000000010, 1000000000000000, 12.5, -2.5, 0, "-"]
    assert r["blocks"] == []
    assert (r["sum"], r["count"], r["excluded"]) == (1000000000000010, 4, ["missing_marker", "subtotal"])
    cells = _excel_column(_integrate(pu, r["review"]), "집행액")  # 확정 항목 이름으로 나감
    assert cells == [(1000000000000000, "n"), (12.5, "n"), (-2.5, "n"), (0, "n"), ("-", "s")]


def test_rereading_confirmed_file_keeps_decisions_and_numbers(pu) -> None:
    """업로드 뒤 다시 열 때마다 원본을 다시 읽는다 — 확정값·'다시 분석 필요' 표시와 수량 해석이 그대로."""
    r = _amount(pu, "규모", BAD, "csv", "재읽기")
    for _ in range(2):
        latest = pu.get(f"/files/{r['fid']}/mapping/latest").json()
        confirmed = {
            col["column"]: col["field_key"] for col in latest["sheets"][0]["columns"] if col["status"] == "confirmed"
        }
        assert confirmed == {"구분": "department", "규모": "expenditure"} and latest["stale"] is False
        again = pu.post("/validation", json={"file_ids": [r["fid"]]}).json()
        assert [b["code"] for b in again["blocks"]] == ["sum_mismatch"]


# ---------------------------------------------------------------- 수량 단어가 들어간 긴 식별자


@pytest.mark.parametrize("column", ["매출식별부호", "고유식별부호"])
@pytest.mark.parametrize("kind", ["csv", "xlsx"])
def test_identifier_with_quantity_word_keeps_text_through_api_and_excel(pu, column: str, kind: str) -> None:
    data, ctype = _file(
        [[column, "식별부호", "수량"], *[[LONG[i], SHORT[i], str(i + 1)] for i in range(3)]], kind, None
    )
    fid = _upload(pu, f"식별자_{column}.{kind}", data, ctype)
    res = pu.get(f"/files/{fid}/preview")
    raw = res.text
    for value in LONG:
        assert f'"{value}"' in raw, value  # JSON 글자 — 브라우저 Number로 바뀌지 않음
    for value in SHORT:
        assert re.search(rf"[\[,]{value}[,\]]", raw), value  # 짧은 식별부호는 기존처럼 숫자
    _confirm(pu, fid, {col["name"]: None for col in res.json()["columns"]})  # 모두 '해당 없음'
    grid = _integrate(pu, pu.post("/reviews", json={"file_ids": [fid]}).json())
    assert _excel_column(grid, column) == [(v, "s") for v in LONG]  # 글자 셀, 끝자리만 다른 값 각각 유지
    assert _excel_column(grid, "식별부호") == [(int(v), "n") for v in SHORT]
    assert _excel_column(grid, "수량") == [(1, "n"), (2, "n"), (3, "n")]


# ---------------------------------------------------------------- 이름에 '식별'이 든 열을 수량으로 확정

CORRECT = [["합계", "1000000000000300"], ["사업A", "1000000000000000"], ["사업B", "100"], ["사업C", "200"]]


@pytest.mark.parametrize("column", ["비식별처리비용", "식별부호"])
@pytest.mark.parametrize("kind", ["csv", "xlsx"])
def test_confirmed_quantity_with_identifier_name_is_blocked_on_wrong_total(pu, column: str, kind: str) -> None:
    r = _amount(pu, column, BAD, kind, f"식별이름_합계오류_{column}_{kind}")
    assert r["parsed"] == ["2000000000000000", "1000000000000000", "100", "200"]  # 미리보기는 원문 그대로 (글자)
    assert r["blocks"] == ["sum_mismatch"]  # 확정한 수량 열은 계산용 값으로 합계 검사
    assert "1,000,000,000,000,300" in r["block_messages"][0] and "차이 999,999,999,999,700" in r["block_messages"][0]
    assert (r["sum"], r["count"], r["excluded"]) == (1000000000000300, 3, ["subtotal"])  # 상세 행이 빠지지 않음
    assert r["review"]["can_complete"] is False
    assert pu.post(f"/reviews/{r['review']['id']}/integrate", json={}).status_code == 409  # 통합 차단


@pytest.mark.parametrize("column", ["비식별처리비용", "식별부호"])
@pytest.mark.parametrize("kind", ["csv", "xlsx"])
def test_confirmed_quantity_with_identifier_name_integrates_numbers(pu, column: str, kind: str) -> None:
    r = _amount(pu, column, CORRECT, kind, f"식별이름_정상_{column}_{kind}")
    assert (r["status"], r["blocks"], r["warnings"]) == ("passed", [], [])
    assert (r["sum"], r["count"], r["excluded"]) == (1000000000000300, 3, ["subtotal"])
    cells = _excel_column(_integrate(pu, r["review"]), "집행액")
    assert cells == [(1000000000000000, "n"), (100, "n"), (200, "n")]  # 확정 항목 이름·숫자 셀


def test_confirmed_quantity_with_identifier_name_mixed_values(pu) -> None:
    r = _amount(pu, "비식별처리비용", OK, "csv", "식별이름_혼합")
    assert (r["status"], r["blocks"]) == ("passed", [])
    assert (r["sum"], r["count"], r["excluded"]) == (1000000000000010, 4, ["missing_marker", "subtotal"])
    cells = _excel_column(_integrate(pu, r["review"]), "집행액")
    assert cells == [(1000000000000000, "n"), (12.5, "n"), (-2.5, "n"), (0, "n"), ("-", "s")]  # 결측 표기는 그대로


def _validate_status(c: TestClient, fid: int) -> tuple[str, list[str]]:
    body = c.post("/validation", json={"file_ids": [fid]}).json()
    return body["status"], [b["code"] for b in body["blocks"]]


def _preview_cells(c: TestClient, fid: int, column: str) -> list[Any]:
    pv = c.get(f"/files/{fid}/preview").json()
    j = [col["name"] for col in pv["columns"]].index(column)
    return [r["cells"][j] for r in pv["rows"]]


def test_reread_follows_latest_confirmation_without_changing_cached_original(pu) -> None:
    """같은 파일을 다시 올리지 않고 확정만 바꾼다: 확정 없음(원문 보존) → 수량 확정 → 다른 항목 → '해당 없음'."""
    c = pu
    raw = ["1000000000000300", "1000000000000000", "100", "200"]
    bad_data, ctype = _file([["구분", "계약식별부호"], *BAD], "csv")
    ok_data, _ = _file([["구분", "계약식별부호"], *CORRECT], "csv")
    bad, ok = _upload(c, "재읽기_합계오류.csv", bad_data, ctype), _upload(c, "재읽기_정상.csv", ok_data, ctype)
    for fid in (bad, ok):
        _confirm(c, fid, {"구분": "department"})

    def integrated(column: str) -> list[tuple[Any, str]]:
        return _excel_column(_integrate(c, c.post("/reviews", json={"file_ids": [ok]}).json()), column)

    # 1) 확정 없음 — 식별자 원문 보존: 계산하지 않음 (통합하려면 검토에서 이 열을 직접 결정해야 하므로 3)·4)에서 확인)
    assert _validate_status(c, bad) == ("passed", [])
    assert _preview_cells(c, ok, "계약식별부호") == raw
    # 2) 수량 항목으로 확정 — 같은 파일을 다시 읽어 계산용 값으로 검사·분석·통합
    for fid in (bad, ok):
        _confirm(c, fid, {"계약식별부호": "expenditure"})
    assert _validate_status(c, bad) == ("blocked", ["sum_mismatch"])
    body = {"field_key": "expenditure", "targets": [{"file_id": ok}]}
    run = c.post(
        "/analysis/run",
        json={**body, "expected_fingerprint": c.post("/analysis/preview", json=body).json()["fingerprint"]},
    )
    group = c.get(f"/analysis/runs/{run.json()['id']}/groups").json()["items"][0]
    assert (group["sum"], group["count"]) == (1000000000000300, 3)
    assert integrated("집행액") == [(1000000000000000, "n"), (100, "n"), (200, "n")]
    assert _preview_cells(c, ok, "계약식별부호") == raw  # 분석이 미리보기 캐시를 읽은 뒤에도 원문 그대로
    # 3) 다른 항목(비고)으로 변경 — 수량 해석을 하지 않고, 원문 자릿수가 그대로 남음
    for fid in (bad, ok):
        _confirm(c, fid, {"계약식별부호": "note"})
    assert _validate_status(c, bad) == ("passed", [])
    assert integrated("비고") == [(v, "s") for v in raw[1:]]
    # 4) '해당 없음'으로 확정 — 마찬가지로 원문 그대로
    for fid in (bad, ok):
        _confirm(c, fid, {"계약식별부호": None})
    assert _validate_status(c, bad) == ("passed", [])
    assert integrated("계약식별부호") == [(v, "s") for v in raw[1:]]
    assert (
        _preview_cells(c, ok, "계약식별부호") == raw and _preview_cells(c, bad, "계약식별부호")[0] == "2000000000000000"
    )


def test_learned_alias_auto_quantity_keeps_original_and_warns(pu) -> None:
    """학습된 이름(자동 적용)만으로는 원문을 숫자로 바꾸지 않는다 — 대신 계산하지 못했음을 알린다."""
    c = pu
    data, ctype = _file([["구분", "협약식별부호"], *BAD], "csv")
    first = _upload(c, "학습_첫파일.csv", data, ctype)
    _confirm(c, first, {"구분": "department", "협약식별부호": "expenditure"})  # 학습 별칭이 생김
    second = _upload(c, "학습_둘째파일.csv", data, ctype)  # 같은 이름, 이 파일에서는 확정하지 않음
    cols = {col["column"]: col for col in c.post(f"/files/{second}/mapping").json()["sheets"][0]["columns"]}
    assert (cols["협약식별부호"]["status"], cols["협약식별부호"]["method"]) == ("auto", "learned")
    body = c.post("/validation", json={"file_ids": [second]}).json()
    assert body["status"] == "passed_with_warnings" and body["blocks"] == []  # 검사를 건너뛴 채 '통과'로 두지 않음
    (warning,) = [w for w in body["warnings"] if w["code"] == "quantity_not_number"]
    assert "자동 적용" in warning["message"] and warning["label"] == "수량으로 계산하지 못한 값"
    review = c.post("/reviews", json={"file_ids": [second]}).json()
    assert "quantity_not_number" in [i["kind"] for i in review["items"] if i["category"] == "warning"]
    assert _preview_cells(c, second, "협약식별부호")[1] == "1000000000000000"  # 원문 그대로
    first_body = c.post("/validation", json={"file_ids": [first]}).json()
    assert [b["code"] for b in first_body["blocks"]] == ["sum_mismatch"]  # 확정한 파일은 계산용 값으로 검사
