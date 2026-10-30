"""D-01: 합계 항목의 산식 확인 — 같은 이름의 합계라도 자료마다 산식이 다를 수 있어 값으로 본다.

시연 자료의 통일부 '발신건수합계' 셀은 =SUM(F2:I2)로 Mail·외부에 시스템·대용량 발신까지 더하고, 다른 부처는 =E2+F2
(Mail+외부)다. 차이를 입력 오류로 단정하지 않고 어느 산식과 맞는지 근거를 보여 준다.
정책 (2026-10-23): 원본 값을 옮기는 통합표·파일별 시트는 경고만 하고, 같은 뜻의 통계로 더하는 데이터 통계 분석·묶기
통계 요약은 다시 확정하거나 명시적으로 뺄 때까지 막는다. 구성 항목이 없어 확인하지 못한 경우는 막지 않고
'검증 불가'로 적는다.
"""

import io
from pathlib import Path

import openpyxl
import pytest
from fastapi import HTTPException
from fastapi.testclient import TestClient

from app import analysis
from app.bundle import ItemIn, SummaryIn, double_count_notes
from app.formulas import check_totals
from app.validation import ValidationParams, validate
from tests.test_accounts import owner  # noqa: F401
from tests.test_analysis import _upload, _xlsx
from tests.test_auth import _client, _provision_user, app_real_login  # noqa: F401
from tests.test_mail_statistics import evaluate
from tests.test_validation import make_input, make_xlsx

# 통일부처럼: 발신건수합계 = 대용량 + 시스템 + Mail + 외부
UNIF = ["통계일자", "기관코드", "조직코드", "조직 명", "대용량파일전송발신수", "시스템발신수"]
UNIF += ["발신Mail건수", "발신외부건수", "발신건수합계"]
UNIF_ROWS = [
    ["2026-10-20", 1250000, 1250370, "운영지원과", 3, 10, 600, 80, 693],
    ["2026-10-20", 1250000, 1250663, "운영지원과", 1, 4, 350, 50, 405],
]
# 다른 부처처럼: 발신건수합계 = Mail + 외부
PLAIN = ["날짜", "기관코드", "조직코드", "조직명", "발신Mail건수", "발신외부건수", "발신건수합계"]
PLAIN_ROWS = [
    ["2026-10-20", 1741000, 1741001, "기획조정실", 100, 20, 120],
    ["2026-10-20", 1741000, 1741002, "운영지원과", 50, 5, 55],
]
UNIF_COLS = {"mail_large_file_sent_count": 4, "mail_system_sent_count": 5, "mail_sent_count": 6,
             "mail_external_sent_count": 7, "mail_sent_total_count": 8}  # fmt: skip


def test_check_totals_names_the_formula_the_values_follow() -> None:
    [check] = [c for c in check_totals(list(enumerate(UNIF_ROWS, 2)), UNIF_COLS) if c.total == "mail_sent_total_count"]
    assert check.checked == 2 and [row for row, _ in check.differ] == [2, 3] and check.difference == 18
    assert check.alternative == ("mail_all_sent_total_count", 2)


def test_missing_parts_and_unreadable_cells_are_not_guessed() -> None:
    # 구성 항목 열이 하나라도 없으면 그 합계는 보지 않는다
    assert check_totals([(2, [1, 2, 3])], {"mail_sent_total_count": 2, "mail_sent_count": 0}) == []
    # 숫자로 읽을 수 없는 칸이 있는 행은 비교하지 않는다 (0으로 보지 않음)
    cols = {"mail_sent_count": 0, "mail_external_sent_count": 1, "mail_sent_total_count": 2}
    [check] = check_totals([(2, [100, "-", 999]), (3, [50, 5, 55])], cols)
    assert check.checked == 1 and not check.differ


def test_validation_warns_only_where_the_formula_differs(tmp_path: Path) -> None:
    unif = make_input(make_xlsx(tmp_path / "unif.xlsx", [UNIF, *UNIF_ROWS]), 1)
    plain = make_input(make_xlsx(tmp_path / "plain.xlsx", [PLAIN, *PLAIN_ROWS]), 2)
    result = validate([unif, plain], ValidationParams())
    found = [w for w in result.warnings if w.code == "formula_mismatch"]
    assert [w.file_id for w in found] == [1]  # 산식이 맞는 자료는 경고하지 않는다
    assert "2행이 다릅니다" in found[0].message and "'발신전체건수합계'(으)로 확정할지" in found[0].message
    assert (found[0].expected, found[0].actual) == (1080, 1098)  # 구성 항목 합 → 합계 열 값 (차이 18)
    assert "통합표에는 원본 값 그대로" in found[0].message  # 통계에서는 막힌다는 것도 함께 알린다
    assert result.status != "blocked"  # 원본 값을 옮기는 통합은 경고만 한다


def test_analysis_blocks_a_total_whose_formula_differs() -> None:
    ev = evaluate([(UNIF, UNIF_ROWS, {}), (PLAIN, PLAIN_ROWS, {})], "mail_sent_total_count")
    unif, plain = ev.targets
    assert unif.status == "needs_input" and unif.formula["state"] == "mismatch"
    assert (unif.formula["checked"], unif.formula["differ"], unif.formula["alternative"]) == (2, 2, "발신전체건수합계")
    assert any("다시 확정하거나 이 시트를 계산에서 빼세요" in r for r in unif.reasons), unif.reasons
    assert plain.status == "ok" and plain.formula["state"] == "verified" and plain.formula["checked"] == 2
    with pytest.raises(HTTPException) as exc:  # 다른 산식의 합계를 같은 항목으로 더하지 않는다 (값을 고치지도 않음)
        analysis.compute(ev)
    assert exc.value.status_code == 409 and any("#1" in p for p in exc.value.detail["problems"])


def test_explicitly_excluding_that_sheet_runs_and_keeps_the_exclusion() -> None:
    ev = evaluate([(UNIF, UNIF_ROWS, {}), (PLAIN, PLAIN_ROWS, {})], "mail_sent_total_count", {1: {"exclude": True}})
    unif, plain = ev.targets
    assert unif.status == "excluded" and unif.reasons == [analysis.USER_EXCLUDED] and unif.formula is None
    res = analysis.compute(ev)
    assert res["groups"][0]["sum"] == 120 + 55  # 뺀 시트의 값은 더하지 않는다
    assert [t["reasons"] for t in res["excluded_targets"]] == [[analysis.USER_EXCLUDED]]  # 결과에 뺀 사실이 남는다


def test_confirming_the_wider_total_for_that_file_matches_its_formula() -> None:
    # 파일별 사용자 확정으로 통일부 열만 '발신전체건수합계'로 — 그 산식과 맞으므로 막지도 알리지도 않는다
    ev = evaluate([(UNIF, UNIF_ROWS, {"발신건수합계": "mail_all_sent_total_count"})], "mail_all_sent_total_count")
    assert not [w for w in ev.warnings if "산식" in w]
    assert ev.targets[0].status == "ok" and ev.targets[0].formula["state"] == "verified"
    assert analysis.compute(ev)["groups"][0]["sum"] == 693 + 405
    assert not [w for w in evaluate([(PLAIN, PLAIN_ROWS, {})], "mail_sent_total_count").warnings if "산식" in w]


def test_unverifiable_total_is_not_blocked_and_not_called_verified() -> None:
    # 구성 항목 열을 '해당 없음'으로 두면 산식을 확인할 수 없다 — 오류로 단정하지 않아 막지 않지만 '검증 불가'로 적는다
    ev = evaluate([(PLAIN, PLAIN_ROWS, {"발신Mail건수": None, "발신외부건수": None})], "mail_sent_total_count")
    [target] = ev.targets
    assert target.status == "ok" and target.formula["state"] == "unverifiable"
    assert "발신Mail건수, 발신외부건수" in target.formula["text"] and "검증 완료 아님" in target.formula["text"]
    assert any("산식 검증 불가" in w for w in ev.warnings)
    assert analysis.compute(ev)["groups"][0]["sum"] == 120 + 55
    # 숫자로 읽지 못한 칸이 있는 행은 비교하지 않고, 확인 범위에 그 행 수를 적는다
    rows = [PLAIN_ROWS[0], ["2026-10-20", 1741000, 1741002, "운영지원과", 50, "-", 55]]
    [target] = evaluate([(PLAIN, rows, {})], "mail_sent_total_count").targets
    assert target.formula["state"] == "verified" and (target.formula["checked"], target.formula["rows"]) == (1, 2)
    assert "1행은 비교하지 않음" in target.formula["text"]


def test_summary_notes_double_counting_of_a_total_and_its_parts() -> None:
    items = [ItemIn(field_key="mail_sent_total_count", aggregates=["sum"]), ItemIn(field_key="mail_sent_count",
                                                                                   aggregates=["sum"])]  # fmt: skip
    [note] = double_count_notes(SummaryIn(items=items))
    assert "'발신Mail건수'" in note and "두 번" in note
    nested = [ItemIn(field_key="mail_all_sent_total_count", aggregates=["sum"]),
              ItemIn(field_key="mail_sent_total_count", aggregates=["sum"])]  # fmt: skip
    assert any("'발신건수합계'" in n for n in double_count_notes(SummaryIn(items=nested)))  # 합계 안의 합계
    alone = [ItemIn(field_key="mail_sent_count", aggregates=["sum"]), ItemIn(field_key="mail_received_count",
                                                                             aggregates=["sum"])]  # fmt: skip
    assert double_count_notes(SummaryIn(items=alone)) == []


@pytest.fixture(scope="module")
def fm(app_real_login, owner) -> TestClient:  # noqa: F811
    _provision_user(owner, app_real_login, "fm-user1", "pass-fm-user1", "FM-USER1")
    return _client(app_real_login, "fm-user1", "pass-fm-user1")


def _confirm_all(c: TestClient, fid: int) -> None:
    """자동 제안을 그대로 확정한다 (시연처럼 모든 열을 확정 — 분석은 확정한 매핑만 쓴다)."""
    cols = c.post(f"/files/{fid}/mapping?use_llm=false").json()["sheets"][0]["columns"]
    items = [{"sheet_id": "s0", "column_id": x["column_id"], "field_key": x["field_key"]} for x in cols]
    assert c.put(f"/files/{fid}/mapping/decisions", json={"decisions": items}).status_code == 200


def _two_files(c: TestClient) -> list[int]:
    sources = (("unif.xlsx", [UNIF, *UNIF_ROWS]), ("plain.xlsx", [PLAIN, *PLAIN_ROWS]))
    fids = [_upload(c, name, _xlsx(rows, title=None)) for name, rows in sources]
    for fid in fids:
        _confirm_all(c, fid)
    return fids


def test_bundle_summary_blocks_the_mismatched_total_until_that_sheet_is_left_out(fm: TestClient) -> None:
    c = fm
    fids = _two_files(c)
    links = [{"file_id": f, "sheet_id": "s0", "unit": "건"} for f in fids]
    keys = ("mail_sent_total_count", "mail_sent_count")  # 합계와 그 구성 항목을 함께 고름
    items = [{"field_key": k, "aggregates": ["sum"], "links": links} for k in keys]
    body = {"file_ids": fids, "summary": {"items": items}}
    plan = c.post("/bundle/plan", json=body).json()
    assert any("두 번 셉니다" in w for w in plan["warnings"]), plan["warnings"]
    total = plan["summary"]["items"][0]
    assert not total["ready"] and any("산식 불일치" in p for p in total["problems"]), total["problems"]
    assert [lk["formula"]["state"] for lk in total["links"]] == ["mismatch", "verified"]
    before = c.get("/runs").json()
    res = c.post("/bundle/run", json={**body, "expected_basis": plan["basis"]})
    assert res.status_code == 409 and any("산식 불일치" in p for p in res.json()["detail"]["problems"]), res.text
    assert c.get("/runs").json() == before  # 거절한 요청은 실행 기록·결과 파일을 남기지 않는다
    # 그 시트를 합계 항목에서만 명시적으로 뺀다 (파일별 시트에는 그대로 들어감) — 뺀 사실과 이유가 요약 시트에 남는다
    left_out = [{**links[0], "exclude": True, "exclude_reason": "산식이 다른 합계"}, links[1]]
    body["summary"]["items"][0]["links"] = left_out
    plan = c.post("/bundle/plan", json=body).json()
    assert plan["summary"]["items"][0]["ready"], plan["summary"]["items"][0]["problems"]
    res = c.post("/bundle/run", json={**body, "expected_basis": plan["basis"]})
    assert res.status_code == 201, res.text
    wb = openpyxl.load_workbook(io.BytesIO(c.get(f"/runs/{res.json()['id']}/download").content))
    rows = list(wb["통계 요약"].iter_rows(values_only=True))
    assert any(r[0] == "주의" and "두 번 셉니다" in r[1] for r in rows)
    assert any(r[7] == "뺌 (사용자)" and r[8] == "산식이 다른 합계" for r in rows if len(r) > 8), rows
    assert any(r[9] and str(r[9]).startswith("산식 확인") for r in rows if len(r) > 9)  # 포함한 시트는 확인 범위


def test_analysis_api_keeps_the_exclusion_and_the_formula_check_in_the_download(fm: TestClient) -> None:
    c = fm
    fids = _two_files(c)
    body = {"field_key": "mail_sent_total_count", "targets": [{"file_id": f, "unit": "건"} for f in fids]}
    pv = c.post("/analysis/preview", json=body).json()
    assert [t["formula"]["state"] for t in pv["targets"]] == ["mismatch", "verified"]
    before = c.get("/analysis/runs").json()
    blocked = c.post("/analysis/run", json={**body, "expected_fingerprint": pv["fingerprint"]})
    assert blocked.status_code == 409 and c.get("/analysis/runs").json() == before  # 막힌 실행은 기록을 남기지 않는다
    targets = [
        {"file_id": fids[0], "sheet_id": "s0", "exclude": True},
        {"file_id": fids[1], "sheet_id": "s0", "unit": "건"},
    ]
    body = {"field_key": "mail_sent_total_count", "targets": targets}
    pv = c.post("/analysis/preview", json=body).json()
    res = c.post("/analysis/run", json={**body, "expected_fingerprint": pv["fingerprint"]})
    assert res.status_code == 201, res.text
    wb = openpyxl.load_workbook(io.BytesIO(c.get(f"/analysis/runs/{res.json()['id']}/download").content))
    rows = list(wb["조건"].iter_rows(values_only=True))
    start = next(i for i, r in enumerate(rows) if r[0] == "파일 번호")
    col = rows[start].index("산식 확인 (합계 항목)")
    shown = {r[0]: (r[7], r[8], r[col]) for r in rows[start + 1 :] if r[0] in fids}
    assert shown[fids[0]][:2] == ("제외", analysis.USER_EXCLUDED)
    assert shown[fids[1]][0] == "사용" and shown[fids[1]][2].startswith("산식 확인 —")
    assert any(r[0] == fids[0] and r[3] == analysis.USER_EXCLUDED for r in wb["제외"].iter_rows(values_only=True))


def test_per_file_sheets_keep_the_original_values_with_a_warning(fm: TestClient) -> None:
    c = fm
    fids = _two_files(c)
    plan = c.post("/bundle/plan", json={"file_ids": fids}).json()  # 통계 요약 없음 — 원본 값을 옮기기만 함
    notes = [w for w in plan["warnings"] if "파일별 시트에는 원본 값 그대로" in w]
    assert len(notes) == 1 and "unif" in notes[0] and "2행에서 다릅니다" in notes[0], plan["warnings"]
    res = c.post("/bundle/run", json={"file_ids": fids, "expected_basis": plan["basis"]})
    assert res.status_code == 201, res.text  # 경고일 뿐 막지 않는다
    wb = openpyxl.load_workbook(io.BytesIO(c.get(f"/runs/{res.json()['id']}/download").content))
    values = [r for r in wb.worksheets[0].iter_rows(values_only=True)]
    assert any(693 in r for r in values)  # 통일부 합계 값을 고치지 않고 그대로 넣었다
