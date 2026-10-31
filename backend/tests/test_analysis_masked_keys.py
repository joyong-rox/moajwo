"""F-1: 가린 뒤 표시가 같아진 서로 다른 구분 값 — 분석·묶기 요약이 원문이 다른 대상을 합치거나 잘못 막지 않는지.

계산·판정은 원문 구분 표지(서버 안에서만)로 하고, 화면·저장·Excel에는 가린 값에 '(가림 구분 n)' 번호를 붙인다.
원문과 표지는 응답·저장값·Excel에 나오지 않는다.
"""

import io
import json
from typing import Any

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app.analysis import POLICY
from app.mapping_review import raw_mark
from tests.test_accounts import owner  # noqa: F401
from tests.test_analysis import _body, _db, _file, _preview, _run, _upload, _xlsx
from tests.test_auth import _client, _provision_user, app_real_login  # noqa: F401

RAW = ["kim1@example.com", "kim2@example.com", "010-1234-5678", "010-9999-5678"]
MAP = {"시도": "sido", "시군구": "sigungu", "인구수": "population"}
HDR = ["시도", "시군구", "인구수"]
A1, A2 = "가팀 k***@example.com (가림 구분 1)", "가팀 k***@example.com (가림 구분 2)"
# 원문만 다른 두 행 (시도가 달라 반복이 아님)
TWO = [HDR, ["서울특별시", "가팀 kim1@example.com", 10], ["부산광역시", "가팀 kim2@example.com", 20]]


@pytest.fixture(scope="module")
def mk(app_real_login, owner) -> TestClient:  # noqa: F811
    _provision_user(owner, app_real_login, "mk-user1", "pass-mk-user1", "MK-USER1")
    return _client(app_real_login, "mk-user1", "pass-mk-user1")


def _no_raw(*texts: str) -> None:
    joined = " ".join(texts)
    for raw in RAW:
        assert raw not in joined
        assert (
            raw_mark(f"가팀 {raw}") not in joined and raw_mark(f"나팀 {raw}") not in joined
        )  # 원문 구분 표지도 내보내지 않음


def _stored(rid: int) -> str:
    return " ".join(_db("SELECT summary, conditions, inputs, result FROM analysis_runs WHERE id = ?", rid)[0])


def _download(c: TestClient, rid: int) -> dict[str, list[list[Any]]]:
    res = c.get(f"/analysis/runs/{rid}/download")
    assert res.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    return {ws.title: [[cell.value for cell in row] for row in ws.iter_rows()] for ws in wb.worksheets}


def _groups(r: dict) -> dict[tuple[str, ...], tuple[int, float]]:
    return {tuple(g["keys"]): (g["count"], g["sum"]) for g in r["groups"]}


def test_email_collision_makes_two_groups_with_consistent_numbers(mk: TestClient) -> None:
    c = mk
    fid = _file(
        c,
        "f1_mail.xlsx",
        [HDR, ["서울특별시", "가팀 kim1@example.com", 10], ["부산광역시", "가팀 kim2@example.com", 20]],
        MAP,
    )
    r = _run(c, _body([fid], group_by=["sigungu"]))
    assert _groups(r) == {(A1,): (1, 10.0), (A2,): (1, 20.0)}
    assert r["detail"]["conditions"]["policy"] == POLICY == "analysis-policy-10"  # 9: 산식 (D-01) · 10: 정확한 수
    assert any("가림 구분" in w for w in r["detail"]["conditions"]["warnings"])
    # 쪽 이동·근거 행·Excel에서 같은 번호
    page2 = c.get(f"/analysis/runs/{r['id']}/groups", params={"offset": 1, "limit": 1}).json()["items"]
    assert [g["keys"] for g in page2] == [[A2]]
    rows = c.get(f"/analysis/runs/{r['id']}/detail", params={"limit": 50}).json()["items"]
    assert {x["row"]: x["keys"] for x in rows} == {3: [A1], 4: [A2]}
    xl = _download(c, r["id"])
    assert [row[:2] for row in xl["결과"][1:]] == [[A1, 10], [A2, 20]]
    assert {row[3]: row[5] for row in xl["근거 행"][1:]} == {3: A1, 4: A2}
    _no_raw(json.dumps(r, ensure_ascii=False), _stored(r["id"]), json.dumps(xl, ensure_ascii=False, default=str))


def test_phone_collision_is_split_too(mk: TestClient) -> None:
    c = mk
    fid = _file(
        c, "f1_tel.xlsx", [HDR, ["서울특별시", "나팀 010-1234-5678", 10], ["부산광역시", "나팀 010-9999-5678", 20]], MAP
    )
    r = _run(c, _body([fid], group_by=["sigungu"]))
    assert _groups(r) == {
        ("나팀 010-****-5678 (가림 구분 1)",): (1, 10.0),
        ("나팀 010-****-5678 (가림 구분 2)",): (1, 20.0),
    }
    _no_raw(json.dumps(r, ensure_ascii=False), _stored(r["id"]))


def test_same_raw_value_stays_one_group_without_number(mk: TestClient) -> None:
    c = mk
    fid = _file(
        c,
        "f1_same.xlsx",
        [HDR, ["서울특별시", "가팀 kim1@example.com", 10], ["서울특별시", "가팀 kim1@example.com", 20]],
        MAP,
    )
    pv = _preview(c, _body([fid], group_by=["sigungu"]))
    assert pv["targets"][0]["status"] == "needs_input"  # 같은 원문이 반복되면 기존처럼 확인을 받는다
    r = _run(c, {**_body([fid], group_by=["sigungu"]), "targets": [{"file_id": fid, "rows_are_records": True}]})
    assert _groups(r) == {("가팀 k***@example.com",): (2, 30.0)}
    assert not any("가림 구분" in w for w in r["detail"]["conditions"]["warnings"])


def test_several_group_fields_keep_the_same_number_for_the_same_raw_value(mk: TestClient) -> None:
    c = mk
    rows = [
        HDR,
        ["서울특별시", "가팀 kim1@example.com", 10],
        ["서울특별시", "가팀 kim2@example.com", 20],
        ["부산광역시", "가팀 kim1@example.com", 30],
    ]
    fid = _file(c, "f1_multi.xlsx", rows, MAP)
    r = _run(c, _body([fid], group_by=["sido", "sigungu"]))
    assert _groups(r) == {("부산광역시", A1): (1, 30.0), ("서울특별시", A1): (1, 10.0), ("서울특별시", A2): (1, 20.0)}


def test_normal_data_is_unchanged(mk: TestClient) -> None:
    c = mk
    fid = _file(
        c,
        "f1_normal.xlsx",
        [HDR, ["서울특별시", "종로구", 10], ["서울특별시", "중구", 20], ["부산광역시", "중구", 5]],
        MAP,
    )
    r = _run(c, _body([fid], group_by=["sido", "sigungu"]))
    assert _groups(r) == {
        ("부산광역시", "중구"): (1, 5.0),
        ("서울특별시", "종로구"): (1, 10.0),
        ("서울특별시", "중구"): (1, 20.0),
    }
    assert not any("가림 구분" in w for w in r["detail"]["conditions"]["warnings"])
    assert r["preview"]["filter_values"]["sigungu"] == ["종로구", "중구"]


def test_filter_picks_one_numbered_value_and_blocks_the_plain_masked_value(mk: TestClient) -> None:
    c = mk
    fid = _file(
        c,
        "f1_filter.xlsx",
        [HDR, ["서울특별시", "가팀 kim1@example.com", 10], ["부산광역시", "가팀 kim2@example.com", 20]],
        MAP,
    )
    pv = _preview(c, _body([fid], group_by=["sigungu"]))
    assert pv["filter_values"]["sigungu"] == [A1, A2]
    r = _run(c, _body([fid], group_by=["sigungu"], filters={"dims": {"sigungu": [A2]}}))
    assert _groups(r) == {(A2,): (1, 20.0)}
    before = _db("SELECT count(*) FROM analysis_runs")[0][0]
    body = _body([fid], group_by=["sigungu"], filters={"dims": {"sigungu": ["가팀 k***@example.com"]}})
    pv = _preview(c, body)
    assert any("구분할 수 없어" in p for p in pv["problems"])
    res = c.post("/analysis/run", json={**body, "expected_fingerprint": pv["fingerprint"]})
    assert res.status_code == 409 and _db("SELECT count(*) FROM analysis_runs")[0][0] == before
    _no_raw(json.dumps(pv, ensure_ascii=False), res.text)


def test_repeated_rows_and_overlaps_compare_raw_identity(mk: TestClient) -> None:
    c = mk
    rep = _file(
        c,
        "f1_rep.xlsx",
        [HDR, ["서울특별시", "가팀 kim1@example.com", 10], ["서울특별시", "가팀 kim2@example.com", 20]],
        MAP,
    )
    pv = _preview(c, _body([rep], group_by=["sigungu"]))
    assert pv["targets"][0]["status"] == "ok" and not pv["problems"]  # 원문이 다르면 반복이 아님
    a = _file(c, "f1_ov_a.xlsx", [HDR, ["부산광역시", "가팀 kim1@example.com", 30]], MAP)
    b = _file(c, "f1_ov_b.xlsx", [HDR, ["부산광역시", "가팀 kim2@example.com", 40]], MAP)
    assert not _preview(c, _body([a, b], group_by=["sigungu"]))["problems"]  # 원문이 다르면 자료 사이 중복이 아님
    same = _file(c, "f1_ov_c.xlsx", [HDR, ["부산광역시", "가팀 kim1@example.com", 31]], MAP)
    pv = _preview(c, _body([a, same], group_by=["sigungu"]))
    assert any("같은 대상" in p for p in pv["problems"])  # 원문이 같으면 기존처럼 중복
    code = _file(
        c,
        "f1_rep_code.xlsx",
        [["지역코드", "인구수"], ["kim1@example.com", 10], ["kim2@example.com", 20]],
        {"지역코드": "region_code", "인구수": "population"},
    )
    pv = _preview(c, _body([code]))
    assert pv["targets"][0]["status"] == "ok" and not pv["problems"]  # 지역 코드 칸만 있어도 원문으로 비교


def test_region_levels_compare_within_the_raw_other_identity(mk: TestClient) -> None:
    c = mk
    k1, k2 = "가팀 kim1@example.com", "가팀 kim2@example.com"
    rows = [
        ["시도", "시군구", "기관명", "인구수"],
        ["서울특별시", "", k1, 30],
        ["서울특별시", "종로구", k1, 10],
        ["서울특별시", "중구", k1, 20],
        ["서울특별시", "", k2, 100],
    ]
    fid = _file(
        c, "f1_level.xlsx", rows, {"시도": "sido", "시군구": "sigungu", "기관명": "org_name", "인구수": "population"}
    )
    r = _run(c, _body([fid], group_by=["org_name"]))
    # kim1: 하위 합계(30) = 상위(30) → 상위를 중복으로 뺌.
    # kim2: 하위 지역이 없어 상위 그대로 (두 기관을 한 비교로 섞지 않음)
    assert _groups(r) == {(A1,): (2, 30.0), (A2,): (1, 100.0)}


def _confirm_ids(c: TestClient, fid: int, items: list[tuple[str, str | None]]) -> None:
    decisions = [{"sheet_id": "s0", "column_id": cid, "field_key": k} for cid, k in items]
    res = c.put(f"/files/{fid}/mapping/decisions", json={"decisions": decisions})
    assert res.status_code == 200, res.text


def test_column_variants_that_become_the_same_after_masking_are_blocked(mk: TestClient) -> None:
    c = mk
    fids = _variant_files(c, "variant")
    before = _db("SELECT count(*) FROM analysis_runs")[0][0]
    pv = _preview(c, _body(fids))
    assert any("열 이름의 구분" in p for p in pv["problems"])
    res = c.post("/analysis/run", json={**_body(fids), "expected_fingerprint": pv["fingerprint"]})
    assert res.status_code == 409 and _db("SELECT count(*) FROM analysis_runs")[0][0] == before
    _no_raw(json.dumps(pv, ensure_ascii=False), res.text)


def test_bundle_summary_uses_raw_identity_and_keeps_raw_out(mk: TestClient) -> None:
    c = mk
    rows = [HDR, ["서울특별시", "가팀 kim1@example.com", 10], ["서울특별시", "가팀 kim2@example.com", 20]]
    fid = _file(c, "f1_bundle.xlsx", rows, MAP)
    assert (
        c.post(f"/files/{fid}/pii/acknowledge").status_code == 200
    )  # 원본 시트에 개인정보 포함 확인 (기존 정책, 이번 수정과 별개)
    off = {"file_ids": [fid]}
    plan = c.post("/bundle/plan", json=off).json()
    assert not plan["problems"]
    on = {**off, "summary": {"items": [{"field_key": "population", "aggregates": ["sum", "count"]}]}}
    plan = c.post("/bundle/plan", json=on)
    assert plan.status_code == 200 and not plan.json()["problems"], (
        plan.text
    )  # 원문이 다른 두 행을 '반복'으로 막지 않음
    res = c.post("/bundle/run", json={**on, "expected_basis": plan.json()["basis"]})
    assert res.status_code == 201, res.text
    content = c.get(f"/runs/{res.json()['id']}/download").content
    wb = openpyxl.load_workbook(io.BytesIO(content))
    summary = [[cell.value for cell in row] for row in wb["통계 요약"].iter_rows()]
    assert any(row[:3] == ["전체", 30, 2] for row in summary), summary
    run = c.get(f"/runs/{res.json()['id']}").json()
    _no_raw(json.dumps(summary, ensure_ascii=False, default=str), json.dumps(run, ensure_ascii=False))


def _variant_files(c: TestClient, tag: str) -> list[int]:
    """열 이름이 가린 뒤 같아지는 두 파일 ('인구_kim1@…'·'인구_kim2@…')."""
    fids = []
    for raw, region in (("kim1@example.com", "대전광역시"), ("kim2@example.com", "울산광역시")):
        fid = _upload(c, f"f1_{tag}_{region}.xlsx", _xlsx([["시도", f"인구_{raw}"], [region, 5]]))
        _confirm_ids(c, fid, [("c0", "sido"), ("c1", "population")])
        assert c.post(f"/files/{fid}/pii/acknowledge").status_code == 200
        fids.append(fid)
    return fids


def test_blank_none_and_zero_stay_apart_from_numbered_values(mk: TestClient) -> None:
    c = mk
    rows = [
        HDR,
        ["서울특별시", "가팀 kim1@example.com", 10],
        ["부산광역시", "가팀 kim2@example.com", 20],
        ["대구광역시", "", 0],  # 빈 구분 값 · 측정값 0 (0은 유효한 숫자)
        ["인천광역시", None, 7],  # 비어 있는 칸
    ]
    fid = _file(c, "f1_blank.xlsx", rows, MAP)
    r = _run(c, _body([fid], group_by=["sigungu"]))
    assert _groups(r) == {(A1,): (1, 10.0), (A2,): (1, 20.0), ("(값 없음)",): (2, 7.0)}
    assert sum(g["sum"] for g in r["groups"]) == 37 and sum(g["count"] for g in r["groups"]) == 4


def test_bins_classify_each_numbered_group_on_its_own(mk: TestClient) -> None:
    c = mk
    fid = _file(c, "f1_bins.xlsx", TWO, MAP)
    bins = {"on": "sum", "ranges": [{"label": "적음", "max": 15}, {"label": "많음", "min": 15}]}
    r = _run(c, _body([fid], group_by=["sigungu"], bins=bins))
    # 합친 뒤(30) 한 구간에 넣지 않고, 원문이 다른 두 그룹을 각자의 합계로 나눈다
    assert {tuple(g["keys"]): g["bin"] for g in r["groups"]} == {(A1,): "적음", (A2,): "많음"}
    assert [row[:2] + row[6:7] for row in _download(c, r["id"])["결과"][1:]] == [[A1, 10, "적음"], [A2, 20, "많음"]]


def test_owner_reads_the_same_numbers_and_other_users_get_nothing(app_real_login, owner, mk: TestClient) -> None:  # noqa: F811
    _provision_user(owner, app_real_login, "mk-user2", "pass-mk-user2", "MK-USER2")
    other = _client(app_real_login, "mk-user2", "pass-mk-user2")
    fid = _file(mk, "f1_perm.xlsx", TWO, MAP)
    rid = _run(mk, _body([fid], group_by=["sigungu"]))["id"]
    seen = owner.get(f"/analysis/runs/{rid}").json()
    assert seen["read_only"] is True
    items = owner.get(f"/analysis/runs/{rid}/groups", params={"limit": 50}).json()["items"]
    assert [(g["keys"], g["count"], g["sum"]) for g in items] == [([A1], 1, 10.0), ([A2], 1, 20.0)]
    assert [row[:2] for row in _download(owner, rid)["결과"][1:]] == [[A1, 10], [A2, 20]]
    _no_raw(json.dumps([seen, items], ensure_ascii=False))
    for path in (f"/analysis/runs/{rid}", f"/analysis/runs/{rid}/groups", f"/analysis/runs/{rid}/download"):
        assert other.get(path).status_code == 404  # 남의 결과는 있는지도 알리지 않음 (기존 권한)
    # 최고 관리자도 남의 파일로 새 결과를 만들지 않는다 (기존 권한 그대로)
    assert owner.post("/analysis/preview", json=_body([fid], group_by=["sigungu"])).status_code == 403


@pytest.mark.parametrize("bad", ["same_raw_repeated", "variant"])
def test_bundle_summary_blocked_makes_no_file_or_run(mk: TestClient, bad: str) -> None:
    c = mk
    if bad == "variant":
        fids = _variant_files(c, "bv")
    else:  # 같은 원문이 반복되면 기존처럼 확인을 받아야 해서 막힘
        rows = [HDR, ["서울특별시", "가팀 kim1@example.com", 10], ["서울특별시", "가팀 kim1@example.com", 20]]
        fids = [_file(c, "f1_bs.xlsx", rows, MAP)]
        assert c.post(f"/files/{fids[0]}/pii/acknowledge").status_code == 200
    body = {"file_ids": fids, "summary": {"items": [{"field_key": "population", "aggregates": ["sum"]}]}}
    runs = _db("SELECT count(*) FROM runs")[0][0]
    plan = c.post("/bundle/plan", json=body).json()
    expect = "열 이름의 구분" if bad == "variant" else "반복"
    assert any(expect in p for p in plan["problems"]), plan["problems"]
    res = c.post("/bundle/run", json={**body, "expected_basis": plan["basis"]})
    assert res.status_code == 409 and _db("SELECT count(*) FROM runs")[0][0] == runs
