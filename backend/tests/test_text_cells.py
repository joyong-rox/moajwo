"""Excel에 사용자 유래 글자를 쓰는 도우미들(묶기·분석)의 공통 규칙: 글자는 글자 칸 — 오류 코드 모양('#N/A')도, '='·'@'
시작도 수식·오류 칸으로 바뀌지 않는다. 숫자·날짜·참거짓·빈 칸은 그대로다."""

from datetime import date

import openpyxl

from app import analysis_api, bundle

ERROR_LIKE = ["#N/A", "#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#NUM!", "#NULL!", "=1+1", "+1", "-x", "@x", "글자"]


def _check(put) -> None:  # noqa: ANN001
    wb = openpyxl.Workbook()
    ws = wb.active
    for i, v in enumerate(ERROR_LIKE, start=1):
        put(ws, i, 1, v)
    put(ws, 1, 2, 3)
    put(ws, 2, 2, 0.5)
    put(ws, 3, 2, True)
    put(ws, 4, 2, date(2026, 10, 27))
    put(ws, 5, 2, None)
    assert [(ws.cell(i, 1).value, ws.cell(i, 1).data_type) for i in range(1, len(ERROR_LIKE) + 1)] == [
        (v, "s") for v in ERROR_LIKE
    ]
    assert [ws.cell(i, 2).data_type for i in range(1, 5)] == ["n", "n", "b", "d"] and ws.cell(5, 2).value is None


def test_bundle_put_keeps_text_as_text() -> None:
    _check(bundle._put)


def test_analysis_put_keeps_text_as_text() -> None:
    _check(analysis_api._put)
