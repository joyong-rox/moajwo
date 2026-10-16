"""문자열로 저장된 식별자의 앞자리 0과 전체 자릿수를 파싱·API·통합 Excel에서 유지한다.

재현: 이름에 코드·번호 표시가 없는 열('식별부호')의 숫자 글자는 숫자로 바뀌는데, 정수도 float를 거쳐
2^53을 넘는 자릿수가 바뀌었다 ('12345678901234567890123' → 12345678901234567741440, '9007199254740993' → …992).
끝자리만 다른 식별자가 같은 값이 되고, API(JSON 숫자 → 브라우저 Number)와 Excel(숫자 셀 15자리)에서도 끝자리가 바뀐다.
"""

import csv
import io
from pathlib import Path

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app.parsing import MAX_NUMBER_DIGITS, parse_file, to_number

FIXTURES = Path(__file__).parent / "fixtures"
LONG = ["12345678901234567890123", "12345678901234567890124", "12345678901234567890125"]  # 끝자리만 다름
NEAR_2_53 = ["9007199254740993", "9007199254740992", "9007199254740991"]
LEADING_ZERO = ["0012345", "0098765", "0000001"]


def _csv(tmp_path: Path, rows: list[list[str]], name: str = "ids.csv") -> Path:
    buf = io.StringIO()
    csv.writer(buf, lineterminator="\n").writerows(rows)
    path = tmp_path / name
    path.write_text(buf.getvalue(), encoding="utf-8")
    return path


def _columns(path: Path) -> dict[str, list]:
    sheet = parse_file(path, path.name).sheets[0]
    return {name: [r["cells"][j] for r in sheet.body] for j, name in enumerate(sheet.columns)}


# ================================================================ 새 결함: 긴 정수 글자의 자릿수 손실


@pytest.mark.parametrize(
    ("text", "expected"),
    [
        ("9007199254740993", 9007199254740993),  # 2^53 + 1 — float를 거치면 …992
        ("12345678901234567890123", 12345678901234567890123),
        ("-9007199254740993", -9007199254740993),
    ],
)
def test_to_number_keeps_integer_digits(text: str, expected: int) -> None:
    assert to_number(text) == expected and isinstance(to_number(text), int)


def test_long_digit_strings_stay_text_in_csv(tmp_path: Path) -> None:
    rows = [["고유식별부호", "연계식별부호", "수량"]] + [[LONG[i], NEAR_2_53[i], str(i + 1)] for i in range(3)]
    cols = _columns(_csv(tmp_path, rows))
    assert cols["고유식별부호"] == LONG  # 글자 그대로, 서로 다른 값
    assert cols["연계식별부호"] == NEAR_2_53
    assert cols["수량"] == [1, 2, 3]  # 수량은 숫자


def test_long_digit_strings_stay_text_in_xlsx_text_cells(tmp_path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["고유식별부호", "연계식별부호", "수량"])
    for i in range(3):
        ws.append([LONG[i], NEAR_2_53[i], i + 1])
    path = tmp_path / "ids.xlsx"
    wb.save(path)
    cols = _columns(path)
    assert (cols["고유식별부호"], cols["연계식별부호"], cols["수량"]) == (LONG, NEAR_2_53, [1, 2, 3])


def test_long_digit_strings_stay_text_in_xls_text_cells() -> None:
    cols = _columns(FIXTURES / "ids.xls")  # 식별자는 글자 셀, 수량은 숫자 셀 (합성 자료)
    assert cols["고유식별부호"] == LONG
    assert cols["연계식별부호"] == NEAR_2_53
    assert cols["식별부호"] == LEADING_ZERO and cols["관리번호"] == [*LONG[:2], "00012"]
    assert cols["수량"] == [1234, 56.7, None] and cols["금액"] == [0, -3, 7]


def test_one_long_digit_string_keeps_the_whole_column_text(tmp_path: Path) -> None:
    """열 전체를 같은 기준으로: 긴 값이 하나라도 있으면 짧은 값도 글자 (같은 열 안의 숫자·글자 뒤섞임 방지)."""
    rows = [["식별부호"], ["1234567890123456"], ["17"]]  # 16자리 (> MAX_NUMBER_DIGITS)
    assert _columns(_csv(tmp_path, rows))["식별부호"] == ["1234567890123456", "17"]


# ================================================================ 보존: 이미 정상인 사례와 수량의 기존 의미


def test_leading_zero_and_management_number_were_already_kept(tmp_path: Path) -> None:
    rows = [["식별부호", "관리번호"]] + [[LEADING_ZERO[i], LONG[i]] for i in range(3)]
    cols = _columns(_csv(tmp_path, rows))
    assert cols["식별부호"] == LEADING_ZERO  # 앞자리 0 규칙
    assert cols["관리번호"] == LONG  # 이름(번호) 규칙


def test_quantities_keep_their_meaning(tmp_path: Path) -> None:
    rows = [["수량", "금액", "비율"], ["1,234", "0", "56.7"], ["-3", "1 234", "-0.5"], ["-", "1.0", "0"]]
    cols = _columns(_csv(tmp_path, rows))
    assert cols["수량"] == [1234, -3, "-"]  # '-'는 결측 표기 그대로
    assert cols["금액"] == [0, 1234, 1.0] and isinstance(cols["금액"][2], float)
    assert cols["비율"] == [56.7, -0.5, 0]


def test_digits_up_to_the_limit_are_still_numbers(tmp_path: Path) -> None:
    fifteen = "1" * MAX_NUMBER_DIGITS
    assert _columns(_csv(tmp_path, [["건수"], [fifteen], ["2"]]))["건수"] == [int(fifteen), 2]


def test_excel_display_format_is_not_turned_into_text(tmp_path: Path) -> None:
    """숫자 12345에 표시 형식 '0000000'만 지정한 셀: 파일에 저장된 값은 12345 — 앞자리 0을 만들어 내지 않는다."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["표시형식부호", "글자부호"])
    ws.append([12345, "0012345"])
    ws.append([98765, "0098765"])
    for r in (2, 3):
        ws.cell(r, 1).number_format = "0000000"
    path = tmp_path / "fmt.xlsx"
    wb.save(path)
    cols = _columns(path)
    assert cols["표시형식부호"] == [12345, 98765]  # 저장된 숫자 그대로 (표시 형식은 재현하지 않음)
    assert cols["글자부호"] == ["0012345", "0098765"]  # 글자로 저장된 값은 그대로


# ================================================================ API: 업로드 → 미리보기·매핑 → 통합 Excel


@pytest.fixture
def api(monkeypatch: pytest.MonkeyPatch):
    from app import files, integration, review, validation_api
    from app.db import Base, engine
    from app.main import app
    from app.mapping import MappingEngine, load_dictionary

    Base.metadata.create_all(engine)
    for module in (files, review, integration, validation_api):  # 모델 호출 없음
        monkeypatch.setattr(module, "get_engine", lambda: MappingEngine(load_dictionary(), None, None))
    with TestClient(app) as c:
        yield c


def test_api_preview_mapping_and_integrated_excel_keep_identifiers(api, tmp_path: Path) -> None:
    c = api
    rows = [["식별부호", "고유식별부호", "연계식별부호", "관리번호", "수량"]] + [
        [LEADING_ZERO[i], LONG[i], NEAR_2_53[i], LONG[i], ["1,234", "56.7", "-"][i]] for i in range(3)
    ]
    data = _csv(tmp_path, rows, "식별자흐름.csv").read_bytes()
    res = c.post("/files", files=[("files", ("식별자흐름.csv", data, "text/csv"))])
    fid = res.json()[0]["id"]

    preview = c.get(f"/files/{fid}/preview")
    assert preview.status_code == 200, preview.text
    names = [col["name"] for col in preview.json()["columns"]]
    raw = preview.text  # JSON 원문: 식별자는 글자("…")로 나가야 브라우저 Number로 바뀌지 않는다
    for value in [*LONG, *NEAR_2_53, *LEADING_ZERO]:
        assert f'"{value}"' in raw, value
    cells = [r["cells"] for r in preview.json()["rows"]]
    assert [r[names.index("수량")] for r in cells] == [1234, 56.7, "-"]

    mapping = c.post(f"/files/{fid}/mapping", params={"use_llm": "false"}).json()
    samples = {col["column"]: col["samples"] for col in mapping["sheets"][0]["columns"]}
    assert samples["고유식별부호"] == LONG and samples["연계식별부호"] == NEAR_2_53

    sheet = mapping["sheets"][0]["name"]
    decisions = [{"sheet": sheet, "column": n, "field_key": None} for n in names]  # 모두 '해당 없음' (학습 없음)
    assert c.put(f"/files/{fid}/mapping/decisions", json={"decisions": decisions}).status_code == 200
    review_res = c.post("/reviews", json={"file_ids": [fid]})
    assert review_res.status_code == 201, review_res.text
    rid = review_res.json()["id"]
    for item in review_res.json()["items"]:
        if item["category"] == "pending":
            body = {"decision": "rejected"} if item["kind"] == "entity" else {"decision": "approved"}
            assert c.post(f"/reviews/{rid}/items/{item['id']}/decision", json=body).status_code == 200
    assert c.post(f"/reviews/{rid}/complete").status_code == 200
    run = c.post(f"/reviews/{rid}/integrate", json={"include_unmapped": True})
    assert run.status_code == 201, run.text
    content = c.get(f"/runs/{run.json()['id']}/download").content
    ws = openpyxl.load_workbook(io.BytesIO(content)).worksheets[0]
    grid = list(ws.iter_rows())
    header_index = next(k for k, row in enumerate(grid) if any(cell.value == "식별부호" for cell in row))
    header = [cell.value for cell in grid[header_index]]
    body = grid[header_index + 1 : header_index + 4]
    for column, expected in (("식별부호", LEADING_ZERO), ("고유식별부호", LONG), ("연계식별부호", NEAR_2_53)):
        out = [row[header.index(column)] for row in body]
        assert [cell.value for cell in out] == expected, column
        assert all(cell.data_type == "s" for cell in out), column  # 글자 셀 (Excel 숫자 15자리 제한을 받지 않음)
    qty = [row[header.index("수량")] for row in body]
    assert [cell.value for cell in qty] == [1234, 56.7, "-"] and qty[0].data_type == "n"
