import os
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from app.config import DATA_DIR
from app.db import Base, engine
from app.main import app
from tests.conftest import db_dump

FIXTURES = Path(__file__).parent / "fixtures"


@pytest.fixture(scope="module")
def client():
    Base.metadata.create_all(engine)
    with TestClient(app) as c:
        yield c


def test_upload_multiple(client: TestClient) -> None:
    csv_bytes = '이름,금액\n김,"1,000"\n합계,"1,000"\n'.encode("euc-kr")
    res = client.post(
        "/files",
        files=[
            ("files", ("members.xls", (FIXTURES / "members.xls").read_bytes(), "application/vnd.ms-excel")),
            ("files", ("거래.csv", csv_bytes, "text/csv")),
            ("files", ("memo.txt", b"hello", "text/plain")),
        ],
    )
    assert res.status_code == 200
    xls, csv, txt = res.json()

    assert xls["error"] is None and xls["sheet_count"] == 1 and xls["data_rows"] == 4
    assert xls["sheets"][0]["header_row"] == 3
    assert csv["data_rows"] == 1 and csv["encoding"] == "euc-kr"
    assert csv["sheets"][0]["name"] == "거래"  # 저장용 UUID가 아닌 원래 파일명
    assert txt["error"] and txt["id"] is None


def test_original_saved_read_only_with_new_name(client: TestClient) -> None:
    data = (FIXTURES / "members.xls").read_bytes()
    before = set((DATA_DIR / "uploads").glob("*.xls"))  # 같은 세션의 다른 시험이 올린 xls는 보지 않는다
    res = client.post("/files", files=[("files", ("../../evil.xls", data, "application/vnd.ms-excel"))])
    body = res.json()[0]
    assert body["original_name"] == "evil.xls"

    stored = [p for p in (DATA_DIR / "uploads").glob("*.xls") if p not in before]
    assert stored, "원본이 uploads 아래에 저장되어야 한다"
    for path in stored:
        assert path.read_bytes() == data
        assert not os.access(path, os.W_OK)
        assert path.name != "evil.xls"


def test_mapping_endpoint(client: TestClient, monkeypatch: pytest.MonkeyPatch) -> None:
    from app import files
    from app.mapping import MappingEngine, load_dictionary

    calls: list[str] = []

    def fake_llm(system: str, user: str, schema: dict) -> dict:
        calls.append(user)
        return {"key": "none", "confidence": 0, "reason": "해당 없음"}

    monkeypatch.setattr(files, "get_engine", lambda: MappingEngine(load_dictionary(), None, fake_llm))
    up = client.post(
        "/files",
        files=[("files", ("members.xls", (FIXTURES / "members.xls").read_bytes(), "application/vnd.ms-excel"))],
    )
    file_id = up.json()[0]["id"]

    res = client.post(f"/files/{file_id}/mapping")
    assert res.status_code == 200
    body = res.json()
    cols = {c["column"]: c for c in body["sheets"][0]["columns"]}
    assert set(cols) == {"지역", "이름", "가입일", "회비"}
    assert cols["지역"]["field_key"] == "region" and cols["지역"]["status"] == "auto"
    assert cols["지역"]["samples"] == ["서울", "부산"]
    assert body["embedding"] == "skipped" and body["llm"] == "ok"
    assert all(len(c["candidates"]) == 3 for c in cols.values())
    # LLM에는 미해결 열만 간다
    assert len(calls) == sum(1 for c in cols.values() if c["method"] == "llm")

    assert client.post("/files/999999/mapping").status_code == 404


def _xlsx_bytes(header: list[str], rows: list[list]) -> bytes:
    import io

    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "현황"
    ws.append(header)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_decisions_saved_and_learned(client: TestClient, monkeypatch: pytest.MonkeyPatch) -> None:
    from app import files
    from app.mapping import MappingEngine, load_dictionary

    engine = MappingEngine(load_dictionary(), None, None)
    monkeypatch.setattr(files, "get_engine", lambda: engine)
    xlsx = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    data = _xlsx_bytes(
        ["자치구 이름", "시도", "강수량", None],
        [["강서구", "서울특별시", 1.5, "11500"], ["강서구", "부산광역시", 2.5, "26440"]],
    )
    file_id = client.post("/files", files=[("files", ("a.xlsx", data, xlsx))]).json()[0]["id"]

    before = {c["column"]: c for c in client.post(f"/files/{file_id}/mapping").json()["sheets"][0]["columns"]}
    assert before["자치구 이름"]["method"] != "learned"
    assert (before["열4"]["method"], before["열4"]["field_key"]) == ("pattern", "region_code")

    res = client.put(
        f"/files/{file_id}/mapping/decisions",
        json={
            "decisions": [
                {"sheet": "현황", "column": "자치구 이름", "field_key": "sigungu"},
                {"sheet": "현황", "column": "자치구 이름", "field_key": "sigungu"},  # 같은 요청 중복도 안전
                {"sheet": "현황", "column": "시도", "field_key": "sido"},  # 이미 사전에 있음
                {"sheet": "현황", "column": "강수량", "field_key": None},
                {"sheet": "현황", "column": "열4", "field_key": "region_code"},  # 헤더 없는 열은 학습 안 함
            ]
        },
    )
    assert res.status_code == 200
    body = res.json()
    assert body["saved"] == 5 and body["learned"] == ["자치구 이름", "자치구 이름"]
    assert {d["column"] for d in body["not_learned"]} == {"시도", "열4"}

    after = {c["column"]: c for c in client.post(f"/files/{file_id}/mapping").json()["sheets"][0]["columns"]}
    assert all(c["status"] == "confirmed" for c in after.values())
    assert after["강수량"]["field_key"] is None

    # 다른 파일에서도 학습된 이름은 1단계에서 바로 자동 적용
    other = _xlsx_bytes(["자치구이름", "인구"], [["강남구", 100]])
    other_id = client.post("/files", files=[("files", ("b.xlsx", other, xlsx))]).json()[0]["id"]
    cols = {c["column"]: c for c in client.post(f"/files/{other_id}/mapping").json()["sheets"][0]["columns"]}
    assert (cols["자치구이름"]["status"], cols["자치구이름"]["method"], cols["자치구이름"]["field_key"]) == (
        "auto",
        "learned",
        "sigungu",
    )

    bad = client.put(
        f"/files/{file_id}/mapping/decisions",
        json={"decisions": [{"sheet": "현황", "column": "없음", "field_key": None}]},
    )
    assert bad.status_code == 422
    bad = client.put(
        f"/files/{file_id}/mapping/decisions",
        json={"decisions": [{"sheet": "현황", "column": "시도", "field_key": "nope"}]},
    )
    assert bad.status_code == 422

    fields = client.get("/mapping/fields").json()
    assert {"key": "sido", "label": "시도"}.items() <= fields[2].items()


def test_corrupt_file_reports_error(client: TestClient) -> None:
    res = client.post("/files", files=[("files", ("broken.xlsx", b"not a zip", "application/octet-stream"))])
    body = res.json()[0]
    assert body["error"] and body["id"] is not None


def test_xlsx_that_expands_beyond_the_limit_is_not_read(client: TestClient, monkeypatch: pytest.MonkeyPatch) -> None:
    """S-03: 업로드 한도는 압축된 크기라, 압축을 풀면 수백 배가 되는 xlsx를 막지 못했다 (읽을 때 메모리 약 9.5배).
    압축을 푼 크기로 막고, 읽기 전에(메모리를 쓰기 전에) 거른다. 시험은 한도만 낮춰 작은 파일로 본다."""
    import io

    import openpyxl

    from app import parsing

    wb = openpyxl.Workbook()
    for i in range(3000):
        wb.active.append(["서울특별시", i, "가나다라마바사"])
    buf = io.BytesIO()
    wb.save(buf)
    xlsx = buf.getvalue()
    xml = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    monkeypatch.setattr(parsing, "MAX_XLSX_UNCOMPRESSED_BYTES", 100_000)
    opened: list[int] = []
    real_load = parsing.openpyxl.load_workbook
    monkeypatch.setattr(parsing.openpyxl, "load_workbook", lambda *a, **k: opened.append(1) or real_load(*a, **k))
    body = client.post("/files", files=[("files", ("큰표.xlsx", xlsx, xml))]).json()[0]
    assert body["id"] is not None and "압축을 풀면" in body["error"] and "한도 0MB" not in body["error"], body
    assert opened == []  # 읽기 전에 거른다

    monkeypatch.setattr(parsing, "MAX_XLSX_UNCOMPRESSED_BYTES", 100 * 1024 * 1024)
    body = client.post("/files", files=[("files", ("보통표.xlsx", xlsx, xml))]).json()[0]
    assert body["error"] is None and body["data_rows"] == 2999 and opened == [1]


def test_validation_flow_blocks_and_pii(client: TestClient, monkeypatch: pytest.MonkeyPatch) -> None:

    from app import files, validation_api
    from app.mapping import MappingEngine, load_dictionary

    engine = MappingEngine(load_dictionary(), None, None)
    monkeypatch.setattr(files, "get_engine", lambda: engine)
    monkeypatch.setattr(validation_api, "get_engine", lambda: engine)
    xlsx = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    header = ["기관명", "예산액(백만원)", "담당자", "연락처"]
    bad = _xlsx_bytes(
        header,
        [
            ["행정안전부", 100, "홍길동", "010-1234-5678"],
            ["기상청", 200, "김철수", "kim@example.go.kr"],
            ["합계", 999, None, None],  # 합계 불일치
        ],
    )
    up = client.post("/files", files=[("files", ("2026년 9월_예산.xlsx", bad, xlsx))]).json()[0]
    assert up["pii"]["counts"] == {"rrn": 0, "mobile": 1, "email": 1}
    assert {i["masked"] for i in up["pii"]["items"]} == {"010-****-5678", "k***@example.go.kr"}
    assert up["pii"]["acknowledged"] is False

    res = client.post(
        "/validation",
        json={"file_ids": [up["id"]], "expected_period": "2026-09", "required_entities": ["행정안전부", "기상청"]},
    )
    assert res.status_code == 200
    run = res.json()
    assert run["status"] == "blocked"
    assert [b["code"] for b in run["blocks"]] == ["sum_mismatch"]
    assert run["blocks"][0]["label"] == "합계 불일치" and run["blocks"][0]["row"] == 4
    assert all(w["code"] in {"mom_outlier", "blank_value", "unit_estimate"} for w in run["warnings"])
    assert run["pii"][0]["acknowledged"] is False
    assert client.get(f"/validation/{run['id']}").json()["status"] == "blocked"

    fixed = _xlsx_bytes(
        header,
        [
            ["행정안전부", 100, "홍길동", "010-1234-5678"],
            ["기상청", 200, "김철수", "kim@example.go.kr"],
            ["합계", 300, None, None],
        ],
    )
    up2 = client.post("/files", files=[("files", ("2026년 9월_예산_수정.xlsx", fixed, xlsx))]).json()[0]
    run2 = client.post("/validation", json={"file_ids": [up2["id"]], "expected_period": "2026-09"}).json()
    assert run2["blocks"] == [] and run2["status"] == "pii_confirmation_required"

    ack = client.post(f"/files/{up2['id']}/pii/acknowledge").json()
    assert ack["acknowledged"] is True
    run3 = client.post("/validation", json={"file_ids": [up2["id"]], "expected_period": "2026-09"}).json()
    assert run3["status"] in ("passed", "passed_with_warnings")

    listed = {f["id"]: f for f in client.get("/files").json()}
    assert listed[up2["id"]]["pii"]["acknowledged"] is True

    # 원문 개인정보는 DB 어디에도 저장되지 않는다
    dump = db_dump()  # 모든 표·모든 행 (SQLite·PostgreSQL)
    assert "010-1234-5678" not in dump and "kim@example.go.kr" not in dump

    # 입력 검증
    assert client.post("/validation", json={"file_ids": []}).status_code == 422
    assert client.post("/validation", json={"file_ids": [up["id"]], "expected_period": "2026-13"}).status_code == 422
    assert client.post("/validation", json={"file_ids": [up["id"]], "baseline_file_ids": [up["id"]]}).status_code == 422
    clean = client.post("/files", files=[("files", ("c.xlsx", _xlsx_bytes(["기관명"], [["기상청"]]), xlsx))]).json()[0]
    assert client.post(f"/files/{clean['id']}/pii/acknowledge").status_code == 409
