from pathlib import Path

import openpyxl
import pytest

from app.parsing import ParseError, decode_csv, detect_header_row, normalize_cell, parse_file, to_number

FIXTURES = Path(__file__).parent / "fixtures"


@pytest.mark.parametrize(
    ("text", "expected"),
    [
        ("1,234", 1234),
        ("1,234,567.5", 1234567.5),
        ("-12,000", -12000),
        ("2 100 000", 2100000),
        ("  750000 ", 750000),
        ("0", 0),
        ("3.14", 3.14),
        ("010 1234 5678", None),  # 전화번호
        ("00123", None),  # 앞자리 0 코드값
        ("12,34", None),  # 잘못된 천 단위 구분
        ("1,234 원", None),
        ("abc", None),
    ],
)
def test_to_number(text: str, expected: object) -> None:
    assert to_number(text) == expected


def test_normalize_cell_blank_to_none() -> None:
    assert normalize_cell("   ") is None
    assert normalize_cell(3.0) == 3


@pytest.mark.parametrize(
    ("raw", "encoding"),
    [
        ("거래처,금액\n한빛,1000\n".encode("euc-kr"), "euc-kr"),
        ("거래처,금액\n똠양꿍,1000\n".encode("cp949"), "cp949"),  # EUC-KR에 없는 글자
        ("﻿거래처,금액\n".encode(), "utf-8-sig"),
    ],
)
def test_decode_csv(raw: bytes, encoding: str) -> None:
    text, used = decode_csv(raw)
    assert text.lstrip("﻿").startswith("거래처")
    assert used == encoding


def test_decode_csv_utf8_without_bom() -> None:
    text, _ = decode_csv(("거래처,금액,비고\n" * 20).encode("utf-8"))
    assert text.startswith("거래처")


def test_detect_header_skips_merged_title() -> None:
    rows = [["제목", "제목", "제목"], [None, None, None], ["이름", "부서", "금액"], ["홍길동", "총무", 100]]
    assert detect_header_row(rows) == 2


def test_detect_header_none_when_no_text() -> None:
    assert detect_header_row([[1, 2], [3, 4]]) is None


def _xlsx(tmp_path: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "매출"
    ws.append(["3분기 매출", None, None])
    ws.merge_cells("A1:C1")
    ws.append(["부서", "이름", "금액"])
    ws.append(["영업팀", "김", "1,000"])
    ws.append([None, "이", "2,000"])
    ws.append(["영업팀 소계", None, 3000])
    ws.append([None, None, None])
    ws.append(["총무팀", "박", 500])
    ws.append(["합계", None, 3500])
    ws.append([None, None, None])  # 뒤쪽 빈 행은 세지 않는다
    ws.merge_cells("A3:A4")
    path = tmp_path / "s.xlsx"
    wb.save(path)
    return path


def test_parse_xlsx(tmp_path: Path) -> None:
    path = _xlsx(tmp_path)
    before = path.read_bytes()
    result = parse_file(path)
    sheet = result.sheets[0]
    assert (sheet.header_row, sheet.data_rows, sheet.subtotal_rows, sheet.empty_rows) == (2, 3, 2, 1)
    assert sheet.columns == ["부서", "이름", "금액"]
    assert sheet.preview[1]["cells"] == ["영업팀", "이", 2000]  # 병합 ffill + 숫자 변환
    assert path.read_bytes() == before  # 원본 불변


def test_parse_xls_fixture() -> None:
    result = parse_file(FIXTURES / "members.xls")
    sheet = result.sheets[0]
    assert (sheet.header_row, sheet.data_rows, sheet.subtotal_rows, sheet.empty_rows) == (3, 4, 1, 0)
    assert [r["cells"][0] for r in sheet.preview[:3]] == ["서울", "서울", "서울"]


def test_parse_csv_euc_kr(tmp_path: Path) -> None:
    path = tmp_path / "t.csv"
    lines = ["일자,거래처,금액", '09-01,한빛,"45,000"', ",,", '09-02,가람,"1,200"', '합계,,"46,200"']
    path.write_bytes("\r\n".join(lines).encode("euc-kr"))
    result = parse_file(path)
    sheet = result.sheets[0]
    assert result.encoding == "euc-kr"
    assert (sheet.header_row, sheet.data_rows, sheet.empty_rows, sheet.subtotal_rows) == (1, 2, 1, 1)
    assert sheet.preview[0]["cells"] == ["09-01", "한빛", 45000]


def test_parse_corrupt_file(tmp_path: Path) -> None:
    path = tmp_path / "broken.xlsx"
    path.write_bytes(b"not a zip")
    with pytest.raises(ParseError):
        parse_file(path)


def _gov_stats(tmp_path: Path) -> Path:
    """정부 통계 양식: 제목 / 단위 / 2단 헤더 / 전국·시도 계 / 동명 시군구 / 결측 표기 / 하단 주석."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "시군구별"
    rows = [
        ["시군구별 인구 현황", None, None, None, None],
        ["(단위: 명, %)", None, None, None, None],
        ["행정구역", None, "인구", None, "증감률"],
        ["시도", "시군구", "2025", "2026", None],
        ["전국", None, "51,217,221", "51,159,889", "-0.1"],
        ["서울특별시", "계", "9,386,034", "9,331,828", "-0.6"],
        [None, "강서구", "556,325", "551,442", "-0.9"],
        [None, "강남구", "556,330", "559,016", "0.5"],
        ["부산광역시", "계", "3,293,362", "3,266,598", "-0.8"],
        [None, "강서구", "143,215", "145,902", "1.9"],
        [None, "중구", "38,862", "37,511", "x"],
        ["인천광역시", "중구", "160,111", "-", "…"],
        ["자료: 행정안전부 주민등록인구통계", None, None, None, None],
        ["주: 1) 2026년은 9월 말 기준", None, None, None, None],
        ["※ 외국인 제외", None, None, None, None],
    ]
    for r in rows:
        ws.append(r)
    for rng in ("A1:E1", "A2:E2", "A3:B3", "C3:D3", "E3:E4", "A6:A8", "A9:A11", "A13:E13"):
        ws.merge_cells(rng)
    path = tmp_path / "gov.xlsx"
    wb.save(path)
    return path


def test_parse_gov_stats(tmp_path: Path) -> None:
    sheet = parse_file(_gov_stats(tmp_path)).sheets[0]
    assert (sheet.header_row, sheet.header_row_end) == (3, 4)
    assert sheet.columns == ["행정구역_시도", "행정구역_시군구", "인구_2025", "인구_2026", "증감률"]
    assert (sheet.data_rows, sheet.subtotal_rows, sheet.note_rows, sheet.empty_rows) == (5, 3, 3, 0)
    assert sheet.missing_cells == 3

    kinds = [(r["row"], r["kind"]) for r in sheet.preview]
    assert kinds[0] == (5, "subtotal")  # 전국
    assert [k for _, k in kinds[-3:]] == ["note", "note", "note"]
    # 동명 시군구는 병합된 시도 값으로 구분된다
    gangseo = [r["cells"][:2] for r in sheet.preview if r["cells"][1] == "강서구"]
    assert gangseo == [["서울특별시", "강서구"], ["부산광역시", "강서구"]]
    incheon = next(r for r in sheet.preview if r["cells"][0] == "인천광역시")
    assert incheon["cells"][3:] == ["-", "…"] and incheon["missing"] == [3, 4]


def test_single_header_with_duplicate_names_not_merged_into_data(tmp_path: Path) -> None:
    """헤더에 같은 이름이 이어져도 아래 행이 금액 데이터면 2단 헤더로 보지 않는다."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["지역", "금액", "금액"])
    ws.append(["서울", 1000, 2000])
    ws.append(["부산", 3000, 4000])
    path = tmp_path / "dup.xlsx"
    wb.save(path)
    sheet = parse_file(path).sheets[0]
    assert (sheet.header_row, sheet.header_row_end, sheet.data_rows) == (1, 1, 2)
    assert sheet.columns == ["지역", "금액", "금액_2"]


@pytest.mark.parametrize(
    "text",
    ["자료: 통계청", "출처 : 기상청", "주: 1) 잠정치", "주1) 잠정치", "※ 외국인 제외", "* 추정", "(단위: 명)"],
)
def test_note_rows(text: str) -> None:
    from app.parsing import classify_row

    assert classify_row([text, text, text]) == "note"
    assert classify_row([text, None, None]) == "note"


def test_note_prefix_needs_single_value() -> None:
    from app.parsing import classify_row

    # 첫 칸이 '주'로 시작해도 다른 값이 있으면 데이터 행 (예: 주소 열)
    assert classify_row(["주: 서울", "홍길동", 100]) == "data"
    assert classify_row(["주식회사 한빛", 100]) == "data"


@pytest.mark.parametrize("group_first", [True, False])
def test_two_level_header_with_total_column(tmp_path: Path, group_first: bool) -> None:
    """'인구' 아래 '남/여/계' 하위 헤더. 하위 헤더가 고유 텍스트가 더 많아도 위 그룹 행과 합친다."""
    wb = openpyxl.Workbook()
    ws = wb.active
    if not group_first:
        ws.append(["2026년 성별 인구", None, None, None])
        ws.merge_cells("A1:D1")
    top = ws.max_row + 1 if not group_first else 1
    ws.append(["시군구", "인구", None, None])
    ws.append([None, "남", "여", "계"])
    ws.merge_cells(f"B{top}:D{top}")
    ws.merge_cells(f"A{top}:A{top + 1}")
    ws.append(["강서구", 270000, 281442, 551442])
    ws.append(["강남구", 268000, 291016, 559016])
    path = tmp_path / "two.xlsx"
    wb.save(path)
    sheet = parse_file(path).sheets[0]
    assert (sheet.header_row, sheet.header_row_end) == (top, top + 1)
    assert sheet.columns == ["시군구", "인구_남", "인구_여", "인구_계"]
    assert (sheet.data_rows, sheet.subtotal_rows) == (2, 0)


@pytest.mark.parametrize(
    ("label", "expected"),
    [
        ("합계", True),
        ("소계", True),
        ("계", True),
        ("전국", True),
        ("영업1팀 소계", True),
        ("서울특별시 합계", True),
        ("합계(전체)", True),
        ("누계 (9월)", True),
        ("소계민원센터", False),  # 실제 자료: 무인민원발급기 이름 (지명 '소계'로 시작)
        ("합계동 행정복지센터", False),
        ("총계리 마을회관", False),
        ("2025년 연간 집행 실적 총괄 합계표 작성 기준 안내", False),  # 긴 문장
    ],
)
def test_subtotal_label_only_at_end(label: str, expected: bool) -> None:
    from app.parsing import is_subtotal_label

    assert is_subtotal_label(label) is expected


def test_place_name_with_subtotal_word_is_data_row(tmp_path: Path) -> None:
    """공공데이터(무인민원발급기 설치정보)에서 발급기명이 '소계…'인 행을 소계로 오판해 통합이 막히던 문제."""
    path = tmp_path / "kiosk.csv"
    lines = ["관리기관명,발급기명,도로코드", "경상남도 창원시,소계민원센터,4100312", "경상남도 창원시,의창구청,4100313"]
    path.write_bytes("\r\n".join(lines).encode("cp949"))
    sheet = parse_file(path, path.name).sheets[0]
    assert (sheet.data_rows, sheet.subtotal_rows) == (2, 0)


@pytest.mark.parametrize(
    "text",
    [
        # 값 안의 큰따옴표("" 이스케이프) — Sniffer가 doublequote=False로 추측하면 열이 밀렸다
        'id,설명,값\n1,"안내 ""무인"" 발급기, 1층",10\n2,일반,20\n3,"쉼표, 포함",30\n',
        # 다른 열의 작은따옴표 — Sniffer가 따옴표를 '로 추측하면 "x, y"가 둘로 나뉘었다
        "id,설명,값\n1,'A',10\n2,'B',20\n3,\"x, y\",30\n",
    ],
)
def test_csv_quotes_follow_standard(tmp_path: Path, text: str) -> None:
    path = tmp_path / "q.csv"
    path.write_text(text, encoding="utf-8")
    sheet = parse_file(path, path.name).sheets[0]
    assert sheet.columns == ["id", "설명", "값"] and sheet.data_rows == 3
    assert all(len(row["cells"]) == 3 for row in sheet.body)
    values = [row["cells"][1] for row in sheet.body]
    assert '안내 "무인" 발급기, 1층' in values or "x, y" in values


def test_csv_semicolon_and_tab_delimiters_still_detected(tmp_path: Path) -> None:
    for delimiter in (";", "\t"):
        path = tmp_path / f"d{ord(delimiter)}.csv"
        lines = [delimiter.join(["시도", "인구"]), delimiter.join(["서울", "100"])]
        path.write_text("\n".join(lines) + "\n", encoding="utf-8")
        sheet = parse_file(path, path.name).sheets[0]
        assert sheet.columns == ["시도", "인구"] and sheet.data_rows == 1


def test_code_and_time_columns_stay_text_quantities_become_numbers(tmp_path: Path) -> None:
    """코드·번호·시각 열은 글자 그대로, 수량 열만 숫자로 (실제 자료: 무인민원발급기 설치정보)."""
    lines = [
        "시도코드,관리번호,평일운영시작시각,평일운영종료시각,분류,인구,금액",
        '6110000,523,0900,1800,011,"1,000",1500',
        '6110000,524,0700,2200,120,"12,345",30',
    ]
    path = tmp_path / "codes.csv"
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    sheet = parse_file(path, path.name).sheets[0]
    rows = [r["cells"] for r in sheet.body]
    assert rows[0][:5] == ["6110000", "523", "0900", "1800", "011"]  # 이름 표시(코드·번호·시각) / 앞자리 0
    assert rows[1][:5] == ["6110000", "524", "0700", "2200", "120"]  # 같은 열은 모두 글자
    assert rows[0][5:] == [1000, 1500] and rows[1][5:] == [12345, 30]  # 수량은 숫자
    kinds = {(sheet.columns[c["column"]], c["kind"]) for c in sheet.corrections}
    assert ("인구", "number") in kinds and not any(k == "number" and col.endswith(("코드", "시각")) for col, k in kinds)


def test_excel_numbers_in_code_column_are_untouched(tmp_path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["행정구역코드", "인구"])
    ws.append([11110, 100])
    path = tmp_path / "n.xlsx"
    wb.save(path)
    sheet = parse_file(path, path.name).sheets[0]
    assert sheet.body[0]["cells"] == [11110, 100]  # 엑셀에 숫자로 저장된 값은 그대로


def _xlsx_rows(tmp_path: Path, rows: list[list], merges: tuple[str, ...] = (), sheets: dict | None = None) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "데이터"
    for r in rows:
        ws.append(r)
    for m in merges:
        ws.merge_cells(m)
    for name, extra in (sheets or {}).items():
        other = wb.create_sheet(name)
        for r in extra:
            other.append(r)
    path = tmp_path / "form.xlsx"
    wb.save(path)
    return path


def test_header_found_below_long_title_block(tmp_path: Path) -> None:
    """정부 양식: 제목·작성일·빈 행·'항목 : 값' 설명 행 뒤 7~8행에 머리글 (예전 탐지 범위 1~3행 밖)."""
    rows = [
        ["자동차 통계표 조회(시군구전체)", None, None, None],
        ["다운로드 시간 : 2026-10-10", None, None, None],
        [None, "검색기간 : 202608 ~ 202608", "검색분야 : 종합일반", "단위 : 대"],
        [None, None, None, None],
        [None, None, None, None],
        [None, None, None, None],
        ["시도", "시군구", "승용", "승용"],
        ["시도", "시군구", "관용", "자가용"],
        ["전국", "합 계", 2259, 1004632],
        ["충남", "천안시", 300, 250000],
    ]
    sheet = parse_file(_xlsx_rows(tmp_path, rows)).sheets[0]
    assert (sheet.header_row, sheet.header_row_end) == (7, 8)
    assert sheet.columns == ["시도", "시군구", "승용_관용", "승용_자가용"]
    # 머리글의 '계'·'총계' 같은 글자 때문에 머리글이 합계 행으로 분류되지 않는다 (실제 합계 행만 소계)
    assert (sheet.data_rows, sheet.subtotal_rows) == (1, 1)


def test_vertical_split_header_and_spaced_labels(tmp_path: Path) -> None:
    """'구별/차종별'처럼 위·아래 줄에 흩어진 머리글을 한 줄로 합치고, '총   계' 정렬 공백을 지운다."""
    rows = [
        [None, None, None, None, None],
        [None, "2026년 8월말 자동차 등록현황", "2026년 8월말 자동차 등록현황", None, None],
        [None, None, None, None, "8.31. 현재"],
        [None, None, "차종별", "총   계", "승용차"],
        [None, "구  별", None, "총   계", "승용차"],
        [None, "총   계", "총  계", 750374, 646873],
        [None, None, "관  용", 2336, 749],
    ]
    sheet = parse_file(_xlsx_rows(tmp_path, rows, ("B2:C2",))).sheets[0]
    assert (sheet.header_row, sheet.header_row_end) == (4, 5)
    assert sheet.columns[1:] == ["구별", "차종별", "총계", "승용차"]


def test_three_level_kosis_header_and_info_sheet_skipped(tmp_path: Path) -> None:
    rows = [
        ["월별(1)", 2024, 2024, 2024],
        ["월별(1)", "총계", "총계", "승용차"],
        ["월별(1)", "계", "관용", "계"],
        ["합계", 1002889, 5860, 773919],
        ["1월", 994549, 5778, 764786],
    ]
    info = [["○ 통계표명", "자동차등록"], ["○ 출처", "기본통계"]]
    result = parse_file(_xlsx_rows(tmp_path, rows, sheets={"메타정보": info}))
    assert [s.name for s in result.sheets] == ["데이터"] and result.skipped_sheets == ["메타정보"]
    sheet = result.sheets[0]
    assert (sheet.header_row, sheet.header_row_end) == (1, 3)
    assert sheet.columns == ["월별(1)", "2024_총계_계", "2024_총계_관용", "2024_승용차_계"]
    assert (sheet.data_rows, sheet.subtotal_rows) == (1, 1)


def test_only_info_sheet_is_still_read(tmp_path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "설명"
    ws.append(["항목", "값"])
    ws.append(["가", 1])
    path = tmp_path / "info.xlsx"
    wb.save(path)
    result = parse_file(path)
    assert [s.name for s in result.sheets] == ["설명"] and result.skipped_sheets == []


def test_horizontal_merge_duplicate_columns_dropped(tmp_path: Path) -> None:
    """머리글·본문이 함께 가로 병합된 복제 열(자가용_2, _3)은 버리고, 값이 다른 열(구 이름)은 남긴다 (충남 양식)."""
    rows = [
        [None, None, None, None, "승   용", None, None, None, "승   합"],
        ["시도", "시군구", None, None, "관용", "자가용", None, "계", "계"],
        ["충남", "계", None, None, 280, 310825, None, 311105, 7423],
        ["충남", "천안시", None, "동남구", 112, 122934, None, 123046, 3407],
        ["충남", "천안시", None, "서북구", 168, 187891, None, 188059, 4016],
        ["충남", "공주시", None, None, 0, 0, None, 0, 0],
    ]
    merges = ("E1:H1", "B2:D2", "F2:G2", "B3:D3", "F3:G3", "B4:C4", "F4:G4", "B5:C5", "F5:G5", "B6:D6", "F6:G6")
    sheet = parse_file(_xlsx_rows(tmp_path, rows, merges)).sheets[0]
    assert sheet.columns == ["시도", "시군구", "시군구_3", "승용_관용", "승용_자가용", "승용_계", "승합_계"]
    data = [r["cells"] for r in sheet.body if r["kind"] == "data"]
    assert data[0] == ["충남", "천안시", "동남구", 112, 122934, 123046, 3407]
    assert data[2] == ["충남", "공주시", "공주시", 0, 0, 0, 0]
    assert len(sheet.preview[0]["cells"]) == 7 and all(c["column"] < 7 for c in sheet.corrections)


def test_same_header_names_without_merge_are_kept(tmp_path: Path) -> None:
    """병합이 아닌 같은 이름 열(CSV 등)은 값이 같아도 버리지 않는다."""
    path = tmp_path / "same.csv"
    path.write_text("구분,인구,인구\n가,1,1\n나,2,2\n", encoding="utf-8")
    assert parse_file(path).sheets[0].columns == ["구분", "인구", "인구_2"]
