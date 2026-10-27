"""통합 결과 생성: 검토가 끝난 세션으로 새 결과 엑셀(3개 시트)을 만들고 바뀐 값을 run_changes에 모두 기록한다.

- 원본 파일은 읽기만 한다. 결과는 항상 data/outputs/에 새 이름으로 만든다 (같은 이름이 있으면 실패).
- 통합 데이터의 모든 셀을 '파일에 적힌 값'과 비교해, 달라진 셀은 빠짐없이 run_changes에 기록하고 배경색을 칠한다.
  (색칠한 셀과 셀 단위 기록이 1:1인지 저장 전에 검사한다. 어긋나면 결과를 남기지 않고 실패로 기록한다.)
- 원본 열 이름 → 표준 항목 이름 매핑도 열 이름 변경으로 기록한다 (통합 행 1 = 헤더).
"""

from __future__ import annotations

import json
import logging
from dataclasses import asdict, dataclass, field
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from app import runlease
from app.check_basis import (
    BASIS_VERSION,
    CHECK_BASIS_KEY,
    LEGACY_REASON,
    applied_drift,
    approved_links,
    basis_changes,
    entity_changes,
)
from app.config import DATA_DIR, MAX_OUTPUT_CELLS, MAX_OUTPUT_TEXT_CHARS
from app.decisions import load_decisions, load_learned, load_org_codes
from app.failures import failure_reason, log_failure
from app.files import load_parsed
from app.glossary import EntityRef, Resolution, entity_ref, normalize_name
from app.glossary_api import glossary_state
from app.mapping import ColumnMapping, get_engine
from app.models import Automation, IntegrationRun, ReviewSession
from app.name_limits import (
    COLUMN_MAX,
    FILE_NAME_MAX,
    LABEL_MAX,
    SHEET_MAX,
    Problem,
    cell_place,
    column_place,
    hint,
    name_problem,
    sheet_place,
)
from app.output_slot import output_slot
from app.parsing import Cell, is_missing
from app.pii import mask_text
from app.promotion import load_standard
from app.quantity import confirmed_quantity_indexes, quantity_sheet
from app.review import mapping_reason, session_items, summarize
from app.validation import WARNING_LABEL, EntityResolver
from app.xmltext import has_illegal, printable

logger = logging.getLogger(__name__)

OUTPUT_DIR = DATA_DIR / "outputs"
# 처리 시각은 한국 시간으로 적는다. 컨테이너는 UTC이고 slim 이미지에 tzdata가 없어
# 고정 오프셋을 쓴다 (KST는 일광절약시간 없음).
KST = timezone(timedelta(hours=9), "KST")
ENTITY_FIELDS = ("org_name", "sido", "sigungu", "region")
PROVENANCE = ["원본 파일", "원본 시트", "원본 행"]
SHEET_DATA, SHEET_CHANGES, SHEET_SUMMARY = "통합 데이터", "변경 내역", "처리 요약"
CHANGE_HEADER = [
    "번호",
    "원본 파일",
    "원본 시트",
    "원본 행",
    "통합 행",
    "열",
    "원본 열",
    "이전 값",
    "이후 값",
    "종류",
    "근거",
    "확신도",
]

KIND_LABEL = {
    "ffill": "병합 셀 채움",
    "number": "숫자 변환",
    "trim": "공백 정리",
    "entity": "이름 정규화",
    "column": "열 이름 매핑",
    "other": "기타 변환",
}
# 종류별 배경색 (결과 파일 '처리 요약' 시트에 범례를 함께 넣는다)
FILLS = {
    "ffill": "DDEBF7",
    "number": "FFF2CC",
    "trim": "EDEDED",
    "entity": "E2EFDA",
    "column": "FCE4D6",
    "other": "F8CBAD",
}


class IntegrationError(Exception):
    pass


CONTROL_CHAR = "CELL_CONTROL_CHAR"  # 결과 파일(엑셀)에 쓸 수 없는 제어 문자 (NUL 포함)
INTEGRATION_FAILED = "INTEGRATION_FAILED"  # 결과를 만드는 중 예기치 않은 실패 (예외 종류만 알림)


@dataclass
class Change:
    source_file_id: int | None
    file_name: str
    sheet: str
    source_row: int | None
    output_row: int
    column: str
    source_column: str
    kind: str
    before: str | None
    after: str | None
    reason: str
    confidence: float
    output_col: int | None = None  # 통합 데이터 시트의 열 번호 (1-based, 셀 변경만)


@dataclass
class _Column:
    index: int  # 원본 시트의 열 번호 (0-based)
    source: str  # 원본 열 이름
    label: str  # 결과 열 이름
    field_key: str
    mapping: ColumnMapping


@dataclass
class _Built:
    labels: list[str] = field(default_factory=list)
    rows: list[list[Any]] = field(default_factory=list)
    changes: list[Change] = field(default_factory=list)
    files: list[dict[str, Any]] = field(default_factory=list)
    dropped_columns: list[dict[str, Any]] = field(default_factory=list)
    unmapped_columns: list[str] = field(default_factory=list)  # 원래 이름으로 넣은 '해당 없음' 열
    text_chars: int = 0  # 결과 칸·변경 기록(앞뒤 값·이유)의 글자 수 합 — 결과 글자 예산 (쌓는 중에 센다)
    # (파일 번호, 시트 이름, 원본 열 이름) → 위치 '파일 #n · N번째 시트 · C열(머리글 1행)' — 이름 없이 알릴 때
    places: dict[tuple[int, str, str], str] = field(default_factory=dict)
    sheets: dict[tuple[int, str], str] = field(default_factory=dict)  # (파일 번호, 시트 이름) → '파일 #n · N번째 시트'
    # 결과 파일에 쓸 수 없는 제어 문자가 든 곳: 처음 3곳의 위치와 전체 수 (실행 기록을 만들기 전에 거절). 그중 용어사전
    # 표시 이름에서 온 곳의 수 — 원본 파일이 아니라 용어사전을 고쳐야 한다
    control: list[str] = field(default_factory=list)
    control_count: int = 0
    control_glossary: int = 0


# ---------------------------------------------------------------- 값 비교


def _is_number(value: Any) -> bool:
    return isinstance(value, int | float) and not isinstance(value, bool)


def same(a: Any, b: Any) -> bool:
    """파일에 적힌 값과 결과 값이 같은지 (3.0과 3은 같음, '1,000'과 1000은 다름)."""
    if a is None or b is None:  # CSV의 빈 문자열('')과 빈 칸은 같은 빈 값
        return a in (None, "") and b in (None, "")
    if _is_number(a) and _is_number(b):
        return float(a) == float(b)
    if isinstance(a, datetime | date) or isinstance(b, datetime | date):
        return a == b
    return type(a) is type(b) and a == b


def text(value: Any) -> str | None:
    """run_changes·변경 내역 시트에 적을 값 표현."""
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime):
        return value.isoformat(sep=" ", timespec="seconds")
    return str(value)


# ---------------------------------------------------------------- 결과 데이터 만들기


def _columns(
    sheet_columns: list[str],
    mappings: list[ColumnMapping],
    include_unmapped: bool = False,
    standard_labels: frozenset[str] = frozenset(),
) -> tuple[list[_Column], list[dict[str, Any]]]:
    """include_unmapped: '해당 없음'으로 확정한 열(표준 항목이 없는 열)도 원래 열 이름으로 결과에 넣는다."""
    used: set[str] = set()
    columns, dropped = [], []
    for j, m in enumerate(mappings):
        if include_unmapped and not m.field_key and m.status == "confirmed":
            label = sheet_columns[j]
            if label in standard_labels or label in used:  # 표준 항목 열과 섞이지 않게 구분
                label = f"{label} (원본 열)"
            used.add(label)
            columns.append(_Column(j, sheet_columns[j], label, None, m))
            continue
        if not m.field_key or m.status not in ("confirmed", "auto"):
            reason = "'해당 없음'으로 확정" if m.status == "confirmed" else "확정되지 않은 매핑"
            dropped.append({"column": sheet_columns[j], "reason": reason})
            continue
        label = (m.field_label or m.field_key) + (f" ({m.qualifier})" if m.qualifier else "")
        if label in used:  # 같은 시트의 두 열이 같은 표준 항목이면 원본 열 이름으로 구분
            label = f"{label} [{sheet_columns[j]}]"
        used.add(label)
        columns.append(_Column(j, sheet_columns[j], label, m.field_key, m))
    return columns, dropped


class _EntityNormalizer:
    """기관·지역 열 값 → 용어사전 정식 이름·7자리 코드. 검토에서 사용자가 승인한 연결을 먼저 쓴다."""

    def __init__(self, resolver: EntityResolver, choices: dict[str, str], review_id: int) -> None:
        self.resolver = resolver
        self.choices = choices  # 이름 → 7자리 코드 (검토 승인)
        self.review_id = review_id

    def resolve(self, name: str, field_key: str, sido: str | None) -> tuple[EntityRef | None, str, float]:
        glossary = self.resolver.glossary
        if glossary is None:
            return None, "", 0.0
        lookup = f"{sido} {name}" if field_key == "sigungu" and sido else name
        # 확인 항목의 값은 row_entity_names가 만든 꼴(시도와 합친 시군구)이다 — 승인한 연결도 그 키로 찾는다
        code = self.choices.get(lookup) or self.choices.get(name)
        if code and code in glossary.entries:
            ref = entity_ref(glossary.entries[code])
            return ref, f"검토 #{self.review_id}에서 사용자가 '{name}' → '{ref.name}'({code}) 연결을 승인했습니다.", 1.0
        res: Resolution | None = self.resolver.resolve(lookup, field_key)
        if res is None or res.status != "resolved" or res.ref is None or res.note:
            return None, "", 0.0
        if res.via == "alias":
            return res.ref, f"용어사전 별칭 '{name}'에 따라 '{res.ref.name}'(으)로 바꿉니다.", 1.0
        if res.via == "short":
            return res.ref, f"'{name}'은(는) 지역 약칭이라 정식명칭 '{res.ref.name}'(으)로 바꿉니다.", 1.0
        return res.ref, "용어사전 정식명칭과 일치합니다.", 1.0


def _display_name(ref: EntityRef, field_key: str, has_sido: bool = True) -> str:
    """시군구 열에는 시도를 뺀 이름을 쓴다. 정식명칭은 '시도명 시군구명' 형식이므로 첫 공백 뒤를 쓴다.

    (자료의 시도 값이 '서울' 같은 약칭이어도 '서울특별시 강서구' → '강서구'가 되도록 접두어를 떼지 않는다)
    단, 같은 시트에 시도 열이 없으면 정식명칭 그대로 둔다
    ('중구'·'강서구'처럼 시도마다 같은 이름이 있어 줄이면 구분이 사라진다).
    """
    if field_key == "sigungu" and has_sido and " " in ref.name:
        return ref.name.split(" ", 1)[1]
    return ref.name


def _bad_text(value: Any) -> bool:
    """엑셀(XML)에 쓸 수 없는 글자(NUL·\\x01 등 제어 문자, U+FFFE·U+FFFF)가 든 글자 — 줄바꿈·탭은 쓸 수 있다."""
    return has_illegal(value)


def _control_at(built: _Built, where: str, glossary: bool = False) -> None:
    built.control_count += 1
    built.control_glossary += glossary
    if len(built.control) < 3:
        built.control.append(where)


def _few(items: list[str], limit: int = 3) -> str:
    return "; ".join(items[:limit]) + (f" 외 {len(items) - limit}건" if len(items) > limit else "")


def build(db: Session, session: ReviewSession, include_unmapped: bool = False) -> _Built:
    if session.status != "ready":
        raise IntegrationError("검토가 끝난(검토 완료) 세션만 통합할 수 있습니다.")
    params = json.loads(session.params)
    excluded = set(json.loads(session.excluded_file_ids))
    file_ids = [f for f in json.loads(session.file_ids) if f not in excluded]
    if not file_ids:
        raise IntegrationError("통합할 파일이 없습니다 (모두 제외됨).")

    basis = params.get(CHECK_BASIS_KEY)
    if not basis or basis.get("v") != BASIS_VERSION:  # 같은지 확인할 수 없으면 같다고 보지 않는다
        raise IntegrationError(f"{LEGACY_REASON} — 검토에서 다시 검사하세요.")
    engine = get_engine()
    # 확정·학습 별칭·전사 표준 사전은 한 번만 읽어 검사 기준과 비교하고, 그 값으로 결과를 만든다
    # (비교한 뒤 다른 요청이 바꿔도 이 실행은 비교한 값으로 만든다 — 바뀐 값은 다음 요청에서 거절됨)
    learned, standard = load_learned(db, session.created_by), load_standard(db)
    snapshot = {fid: load_decisions(db, fid) for fid in [*file_ids, *basis.get("baseline", [])]}
    if changed := basis_changes(db, basis, snapshot, learned, standard):
        raise IntegrationError(
            f"검사한 뒤 열 매핑이 바뀌어 이 검토 결과로 통합할 수 없습니다 — 검토에서 다시 검사하세요 "
            f"({_few(changed)})."
        )
    # 검사가 해석한 기관·지역 이름도 지금 용어사전에서 같아야 그 검사 결과로 만든다 — 한 번 읽어 비교한 이 용어사전으로
    # 결과를 만든다 (비교한 뒤 다른 요청이 바꿔도 이 실행은 비교한 값으로 — 바뀐 값은 다음 요청에서 거절됨)
    digest, load_glossary = glossary_state(db)
    glossary = load_glossary()
    items = session_items(db, session)
    record = basis.get("entities")
    if changed := entity_changes(record, digest, load_glossary, lambda: approved_links(items)):
        if record is None:
            raise IntegrationError(f"{changed[0]} — 검토에서 다시 검사하세요.")
        raise IntegrationError(
            f"검사한 뒤 기관·지역 용어사전이 바뀌어 이 검토 결과로 통합할 수 없습니다 — 검토에서 다시 검사하세요 "
            f"({_few(changed)})."
        )
    choices = {
        i.before: i.choice for i in items if i.kind == "entity" and i.decision == "approved" and i.choice and i.before
    }
    normalizer = _EntityNormalizer(
        EntityResolver(glossary, params.get("expected_period") or params.get("name_period")), choices, session.id
    )

    built = _Built(labels=list(PROVENANCE))
    standard_labels = frozenset(f.label for f in engine.fields)
    label_index: dict[str, int] = {}

    def column_at(label: str) -> int:
        if label not in label_index:
            built.labels.append(label)
            built.text_chars += len(label)  # 결과 열 이름 (통합 데이터 시트 1행·처리 요약)
            label_index[label] = len(built.labels) - 1
        return label_index[label]

    for fid in json.loads(session.file_ids):
        if fid in excluded:
            record, _ = load_parsed(db, fid)
            built.files.append({"id": fid, "name": record.original_name, "excluded": True, "sheets": []})
            if _bad_text(record.original_name):  # 처리 요약 시트에 쓴다
                _control_at(built, f"파일 #{fid}: 파일 이름")
    org_codes = load_org_codes(db)
    for fid in file_ids:
        record, parsed = load_parsed(db, fid)
        decisions = snapshot[fid]
        info: dict[str, Any] = {"id": fid, "name": record.original_name, "excluded": False, "sheets": []}
        if _bad_text(record.original_name):
            _control_at(built, f"파일 #{fid}: 파일 이름")
        for sheet in parsed.sheets:
            # 시트·열 이름은 결과 파일(출처 열·변경 내역·처리 요약)에 그대로 쓴다 — 쓰지 않는 열 이름도 요약에 남는다
            where = f"파일 #{fid} · {sheet_place(sheet)}"
            built.sheets[(fid, sheet.name)] = where
            if _bad_text(sheet.name):
                _control_at(built, f"{where}: 시트 이름")
            for j, name in enumerate(sheet.columns):
                if _bad_text(name):
                    _control_at(built, f"{where} · {column_place(sheet, j)}: 열 이름")
            mappings = engine.map_columns(
                sheet.columns,
                sheet.samples,
                use_llm=False,
                learned=learned,
                decisions=decisions.get(sheet.name),
                standard=standard,
                org_codes=org_codes,
            ).columns
            # 출력에 쓸 매핑(자동 적용 포함)이 검사에 적용한 매핑과 같아야 그 검사 결과로 만든다 — 임베딩 유사도 등
            # 조회 때는 알 수 없는 자동 적용 변화도 결과 파일을 만들기 전에 여기서 막는다
            if drift := applied_drift(basis, fid, sheet.name, mappings):
                raise IntegrationError(
                    mask_text(
                        f"{record.original_name}: 결과에 쓸 열 매핑이 검사 때와 다릅니다 — 검토에서 다시 검사하세요 "
                    )
                    + f"({_few(drift)})."
                )
            sheet = quantity_sheet(sheet, confirmed_quantity_indexes(mappings))  # 확정한 수량 열은 계산용 값
            columns, dropped = _columns(sheet.columns, mappings, include_unmapped, standard_labels)
            for c in columns:
                if c.field_key is None and c.label not in built.unmapped_columns:
                    built.unmapped_columns.append(c.label)
            for d in dropped:
                built.dropped_columns.append({"file": record.original_name, "sheet": sheet.name, **d})
                built.text_chars += sum(len(v) for v in built.dropped_columns[-1].values() if isinstance(v, str))
            for c in columns:
                built.places.setdefault((fid, sheet.name, c.source), f"{where} · {column_place(sheet, c.index)}")
                column_at(c.label)
                if c.field_key in ENTITY_FIELDS:
                    column_at(f"{c.label} 코드")
                if normalize_name(c.source) != normalize_name(c.label):
                    reason, _ = mapping_reason(c.mapping)
                    _record(
                        built,
                        Change(
                            fid,
                            record.original_name,
                            sheet.name,
                            sheet.header_row,
                            1,
                            c.label,
                            c.source,
                            "column",
                            c.source,
                            c.label,
                            f"열 이름 매핑: {reason}",
                            round(float(c.mapping.score if c.mapping.score is not None else 1.0), 4),
                        ),
                    )
            sido_col = next((c for c in columns if c.field_key == "sido"), None)
            used_rows = 0
            for entry in sheet.body:
                if entry["kind"] != "data":
                    continue
                used_rows += 1
                out_row = len(built.rows) + 2  # 1행 = 헤더
                row: list[Any] = [None] * len(built.labels)
                row[0], row[1], row[2] = record.original_name, sheet.name, entry["row"]
                sido_value = entry["cells"][sido_col.index] if sido_col else None
                for c in columns:
                    original: Cell | Any = entry["original"][c.index]
                    cell_text: Any = entry["cells"][c.index]  # 병합 칸 채움이면 원본 칸은 비고 채운 값이 여기 있다
                    value: Any = cell_text
                    ref = None
                    if isinstance(original, datetime | date):
                        value = original  # 날짜는 원본 형식 그대로 둔다 (변환하지 않음)
                    kind, reason, confidence = None, "", 1.0
                    if c.field_key in ENTITY_FIELDS and isinstance(value, str) and not is_missing(value):
                        ref, why, conf = normalizer.resolve(
                            value, c.field_key, sido_value if isinstance(sido_value, str) else None
                        )
                        if ref is not None:
                            row[label_index[f"{c.label} 코드"]] = ref.code
                            shown = _display_name(ref, c.field_key, sido_col is not None)
                            if shown != value:
                                value, kind, reason, confidence = shown, "entity", why, conf
                    row[label_index[c.label]] = value
                    place = f"{where} · {cell_place(sheet, entry['row'], c.index)}"
                    # 결과 칸·변경 내역(이전 값·사유)에 쓸 수 없는 글자 — 원본 칸(병합으로 채운 값 포함)에서 왔는지,
                    # 칸 값을 바꾼 용어사전 이름에서 왔는지 나눠 알린다 (용어사전 글자는 원본을 고쳐도 남는다). 시군구는
                    # 결과 칸에 정식명칭의 시도 뒤만 쓰지만 변경 사유(변경 내역 '근거'·셀 메모)에는 정식명칭 전체가
                    # 들어간다
                    if _bad_text(original) or _bad_text(cell_text) or (kind != "entity" and _bad_text(value)):
                        _control_at(built, place)
                    elif ref is not None and (_bad_text(value) or _bad_text(ref.name)):
                        _control_at(built, f"용어사전 {ref.code} 항목의 이름 ({place}에 쓰려던 값)", glossary=True)
                    if same(original, value):
                        continue
                    if kind is None:
                        kind, reason = _classify(original, value, c)
                    _record(
                        built,
                        Change(
                            fid,
                            record.original_name,
                            sheet.name,
                            entry["row"],
                            out_row,
                            c.label,
                            c.source,
                            kind,
                            text(original),
                            text(value),
                            reason,
                            confidence,
                            output_col=label_index[c.label] + 1,
                        ),
                    )
                built.rows.append(row)
                built.text_chars += sum(len(v) for v in row if isinstance(v, str))
                _check_size(built)  # 쌓는 중에도 — 한도를 넘는 입력은 다 읽기 전에 멈춘다 (값은 늘기만 한다)
            info["sheets"].append(
                {
                    "name": sheet.name,
                    "data_rows": sheet.data_rows,
                    "used_rows": used_rows,
                    "subtotal_rows": sheet.subtotal_rows,
                    "note_rows": sheet.note_rows,
                }
            )
        built.files.append(info)
    # 열이 나중에 늘어난 경우 앞 행들의 길이를 맞춘다
    for row in built.rows:
        row.extend([None] * (len(built.labels) - len(row)))
    if session.automation_id is not None:
        automation = db.get(Automation, session.automation_id)
        if automation is not None:
            _apply_output_format(built, json.loads(automation.output_format).get("columns", []))
    return built


def _apply_output_format(built: _Built, order: list[str]) -> None:
    """저장된 자동화의 열 순서대로 결과 열을 정렬한다 (출처 열은 맨 앞, 새 열은 뒤에 붙임)."""
    head = list(range(len(PROVENANCE)))
    rest = [i for i in range(len(PROVENANCE), len(built.labels))]
    rank = {label: n for n, label in enumerate(order)}
    rest.sort(key=lambda i: (rank.get(built.labels[i], len(rank)), i))
    new_order = head + rest
    position = {old: new for new, old in enumerate(new_order)}
    built.labels = [built.labels[i] for i in new_order]
    built.rows = [[row[i] for i in new_order] for row in built.rows]
    for ch in built.changes:
        if ch.output_col is not None:
            ch.output_col = position[ch.output_col - 1] + 1


def _classify(original: Any, value: Any, c: _Column) -> tuple[str, str]:
    if original is None and value is not None:
        return "ffill", f"'{c.source}' 열의 병합 셀(빈 칸)을 위쪽 값 '{text(value)}'(으)로 채웠습니다."
    if isinstance(original, str) and _is_number(value):
        return "number", f"숫자 문자열 '{original.strip()}'을(를) 숫자로 바꿨습니다 (천 단위 쉼표·공백 제거)."
    if isinstance(original, str) and isinstance(value, str) and original.strip() == value:
        return "trim", "앞뒤 공백을 지웠습니다."
    if isinstance(original, str) and not original.strip() and value is None:
        return "trim", "공백만 있던 칸을 빈 칸으로 정리했습니다."
    return "other", f"원본 값 '{text(original)}'을(를) '{text(value)}'(으)로 읽었습니다."


# ---------------------------------------------------------------- 엑셀 쓰기

HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
BOLD = Font(bold=True)


def _autosize(ws: Any, widths: dict[int, int]) -> None:
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = max(8, min(60, width + 2))


def _as_text(ws: Any) -> None:
    """결과 파일에는 수식·오류 칸을 쓰지 않는다. openpyxl은 '='로 시작하는 글자를 수식(f)으로, '#N/A' 같은 오류 코드
    모양 글자를 오류 칸(e)으로 저장하므로, 원본 값·열 이름·사유가 그렇게 생겼어도 글자 칸으로 둔다 — 결과를 Excel로
    열 때 HYPERLINK 같은 수식이 실행되거나 글자가 오류 값으로 바뀌지 않게 (묶기·분석 결과와 같은 규칙). 다 쓴 시트를
    한 번만 훑는다 — 행마다 ws[row]를 부르면 openpyxl이 그때마다 모든 칸에서 열 수를 다시 구해 행 수 × 칸 수가 된다."""
    for row in ws.iter_rows():
        for cell in row:
            if cell.data_type in ("f", "e"):
                cell.data_type = "s"


def _printable(value: Any) -> Any:
    """처리 요약의 설명 글(검증 경고 사유·기준 기간)에서 엑셀(XML)에 쓸 수 없는 글자를 '�'로 — 이 글은 이미
    개인정보를 가린 설명이고 원본 값이 아니다. 결과 칸·이름·변경 내역의 글자는 build가 그 자리로 거절한다."""
    return printable(value)


def write_workbook(path: Path, built: _Built, summary: dict[str, Any]) -> int:
    """3개 시트를 쓴다. 칠한 셀 수를 돌려준다 (셀 단위 변경 기록 수와 같아야 함)."""
    wb = Workbook()
    data = wb.active
    data.title = SHEET_DATA
    data.append(built.labels)
    for cell in data[1]:
        cell.font, cell.fill = BOLD, HEADER_FILL
    for row in built.rows:
        data.append(row)
    _as_text(data)
    painted = 0
    widths: dict[int, int] = {i + 1: len(str(label)) * 2 for i, label in enumerate(built.labels)}
    for ch in built.changes:
        if ch.output_col is None:
            continue
        cell = data.cell(row=ch.output_row, column=ch.output_col)
        cell.fill = PatternFill("solid", fgColor=FILLS[ch.kind])
        comment = Comment(
            f"{KIND_LABEL[ch.kind]}\n이전 값: {ch.before if ch.before is not None else '(빈 칸)'}\n{ch.reason}",
            "Moajwo",
        )
        comment.width, comment.height = 320, 120
        cell.comment = comment
        painted += 1
    for row in built.rows:
        for i, value in enumerate(row):
            widths[i + 1] = max(widths.get(i + 1, 8), len(text(value) or ""))
    _autosize(data, widths)
    data.freeze_panes = "D2"
    data.auto_filter.ref = data.dimensions

    changes = wb.create_sheet(SHEET_CHANGES)
    changes.append(CHANGE_HEADER)
    for cell in changes[1]:
        cell.font, cell.fill = BOLD, HEADER_FILL
    for n, ch in enumerate(built.changes, start=1):
        changes.append(
            [
                n,
                ch.file_name,
                ch.sheet,
                ch.source_row if ch.source_row is not None else "",
                ch.output_row,
                ch.column,
                ch.source_column,
                ch.before if ch.before is not None else "(빈 칸)",
                ch.after if ch.after is not None else "(빈 칸)",
                KIND_LABEL[ch.kind],
                ch.reason,
                ch.confidence,
            ]
        )
        changes.cell(row=n + 1, column=10).fill = PatternFill("solid", fgColor=FILLS[ch.kind])
        changes.cell(row=n + 1, column=12).number_format = "0%"
    _autosize(changes, {1: 4, 2: 24, 3: 10, 4: 6, 5: 6, 6: 16, 7: 16, 8: 16, 9: 16, 10: 10, 11: 60, 12: 6})
    for row in changes.iter_rows(min_row=2, min_col=11, max_col=11):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    changes.freeze_panes = "A2"
    changes.auto_filter.ref = changes.dimensions
    _as_text(changes)  # 이전·이후 값과 사유도 글자로

    summary_ws = wb.create_sheet(SHEET_SUMMARY)
    _write_summary(summary_ws, built, summary)
    _as_text(summary_ws)  # 파일·시트·열 이름도 글자로
    wb.save(path)
    return painted


def _write_summary(ws: Any, built: _Built, summary: dict[str, Any]) -> None:
    def title(text_: str) -> None:
        ws.append([])
        ws.append([text_])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True, size=12)

    ws.append(["처리 요약"])
    ws.cell(row=1, column=1).font = Font(bold=True, size=14)
    for key, value in [
        ("처리 시각", summary["processed_at"]),
        ("통합 실행 번호", summary["run_id"]),
        ("검토 번호", summary["review_id"]),
        ("검증 번호", summary["validation_run_id"]),
        ("기준 기간", _printable(summary["expected_period"]) or "(지정 안 함)"),
        ("통합 행 수", summary["row_count"]),
        ("변경 건수 (셀)", summary["cell_change_count"]),
        ("열 이름 매핑", summary["column_change_count"]),
        ("경고 건수", summary["warning_count"]),
        (
            "표준 항목이 없는 열",
            f"원래 이름으로 포함 ({len(built.unmapped_columns)}개)" if summary.get("include_unmapped") else "넣지 않음",
        ),
    ]:
        ws.append([key, value])

    title("입력 파일")
    ws.append(["파일", "시트", "데이터 행", "통합한 행", "소계·합계 행(제외)", "주석 행(제외)", "상태"])
    for cell in ws[ws.max_row]:
        cell.font, cell.fill = BOLD, HEADER_FILL
    for f in built.files:
        if f["excluded"]:
            ws.append([f["name"], "", "", 0, "", "", "제외 (개인정보 확인 거부)"])
            continue
        for s in f["sheets"]:
            ws.append(
                [f["name"], s["name"], s["data_rows"], s["used_rows"], s["subtotal_rows"], s["note_rows"], "통합"]
            )

    if summary.get("warnings"):
        title("검증 경고 (통합은 진행 — 확인하지 못했거나 확인이 필요한 내용)")
        ws.append(["종류", "파일", "시트", "행", "열", "내용"])
        for cell in ws[ws.max_row]:
            cell.font, cell.fill = BOLD, HEADER_FILL
        for w in summary["warnings"]:
            ws.append([w["kind"], w["file"], w["sheet"], w["row"], w["column"], _printable(w["reason"])])

    title("변경 종류별 건수")
    for kind, count in summary["change_counts"].items():
        ws.append([KIND_LABEL[kind], count])
        ws.cell(row=ws.max_row, column=1).fill = PatternFill("solid", fgColor=FILLS[kind])

    if built.dropped_columns:
        title("통합하지 않은 열")
        ws.append(["파일", "시트", "열", "이유"])
        for d in built.dropped_columns:
            ws.append([d["file"], d["sheet"], d["column"], d["reason"]])

    title("색 범례 (통합 데이터 시트)")
    for kind, label in KIND_LABEL.items():
        if kind == "column":
            continue
        ws.append([label, "셀 메모에 이전 값과 근거가 있습니다."])
        ws.cell(row=ws.max_row, column=1).fill = PatternFill("solid", fgColor=FILLS[kind])
    ws.append(["코드 열", "용어사전에서 찾은 7자리 코드를 덧붙인 열입니다 (원본 값 변경 아님, 색 없음)."])
    _autosize(ws, {1: 22, 2: 30, 3: 10, 4: 10, 5: 16, 6: 12, 7: 24})


# ---------------------------------------------------------------- 실행


def _output_path(run_id: int, now: datetime) -> Path:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    return OUTPUT_DIR / f"통합결과_{run_id:04d}_{now.strftime('%Y%m%d_%H%M%S')}.xlsx"


def _record(built: _Built, change: Change) -> None:
    """변경 기록을 더하고 그 글자를 결과 글자 예산에 센다 — 변경 내역 시트·run_changes에 그대로 들어가는 글자
    (파일·시트·열 이름, 앞뒤 값, 이유)."""
    built.changes.append(change)
    fields = (change.file_name, change.sheet, change.column, change.source_column, change.before, change.after,
              change.reason)  # fmt: skip
    built.text_chars += sum(len(v) for v in fields if v)


def _check_size(built: _Built) -> None:
    """결과에 넣을 칸·글자 — 메모리 예산(MAX_OUTPUT_CELLS·MAX_OUTPUT_TEXT_CHARS)을 넘으면 거절한다. 쌓는 중에도
    부르므로(값은 늘기만 함) 다 읽기 전에 멈추고, 실행 기록을 만들기 전이라 남는 기록·파일이 없다."""
    cells = len(built.rows) * len(built.labels) + len(built.changes) * len(CHANGE_HEADER)
    if cells > MAX_OUTPUT_CELLS:
        raise IntegrationError(
            f"결과에 넣을 칸이 {cells:,}개 이상으로 이 서버에서 한 번에 만들 수 있는 "
            f"{MAX_OUTPUT_CELLS:,}칸을 넘습니다 — 통합할 파일을 나눠 다시 검토하세요."
        )
    if built.text_chars > MAX_OUTPUT_TEXT_CHARS:
        raise IntegrationError(
            f"결과에 넣을 글자가 {built.text_chars:,}자 이상으로 이 서버에서 한 번에 만들 수 있는 "
            f"{MAX_OUTPUT_TEXT_CHARS:,}자를 넘습니다 — 긴 글자 열을 줄이거나 통합할 파일을 나눠 다시 검토하세요."
        )


def _check_names(built: _Built) -> None:
    """변경 기록(run_changes)에 실제로 저장할 이름 — 결과 열 이름·원본 열 이름·시트 이름·파일 이름 — 이 저장 규칙
    (한도·NUL)에 맞는지. 실행 기록을 만들기 전에 거절한다 (SQLite도 같은 규칙 — app/name_limits.py).

    - 결과 열 이름은 ' (원본 열)'·' [원래 이름]'이 붙어 원래 이름보다 길어질 수 있다.
    - 결과 파일에만 쓰고 기록하지 않는 이름(기관·지역 항목의 '<열> 코드' 열 등)은 보지 않는다 — 이름이 아니라
      저장하는 곳으로 가린다.
    - 원본 이름은 이름 규칙 전에 SQLite에서 만든 검토에서만 넘을 수 있다 (새 검토는 만들 때 거절됨).
    - 이름을 잘라 적지 않는다 (결과 파일의 열 이름과 기록이 달라짐). 문구에는 그 이름을 만든 원본 위치만."""
    problems: list[tuple[str, Problem]] = []
    label_problem = False
    stored = {(ch.source_file_id, ch.file_name, ch.sheet, ch.column, ch.source_column) for ch in built.changes}
    for fid, file_name, sheet, label, source in sorted(stored, key=lambda k: (k[0] or 0, k[2], k[4], k[3])):
        where = built.places.get((fid or 0, sheet, source)) or (f"파일 #{fid}" if fid is not None else "변경 기록")
        if (p := name_problem("결과 열 이름", label, LABEL_MAX)) is not None:
            problems.append((where, p))
            label_problem = True
        for p in (
            name_problem("원본 열 이름", source, COLUMN_MAX),
            name_problem("시트 이름", sheet, SHEET_MAX),
            name_problem("파일 이름", file_name, FILE_NAME_MAX),
        ):
            if p is not None:
                problems.append((where, p))
    if not problems:
        return
    unique = list(dict.fromkeys(problems))
    shown = " / ".join(f"{w}: {p.text}" for w, p in unique[:3]) + (
        f" 외 {len(unique) - 3}곳" if len(unique) > 3 else ""
    )
    why = (
        " 같은 표준 항목으로 확정한 두 번째 열과 결과에 넣은 '해당 없음' 열은 원래 열 이름이 결과 열 이름에 "
        "들어갑니다 — 그 열을 다른 항목으로 다시 확정하거나 원본에서 이름을 줄여 올린 뒤 검토에서 다시 검사하세요."
        if label_problem
        else " 검토에서 다시 검사해야 합니다."
    )
    raise IntegrationError(f"이름을 저장할 수 없어 통합하지 않았습니다 — {shown}. {hint(p for _, p in unique)}{why}")


def _check_text(built: _Built) -> None:
    """결과 파일(엑셀)에 쓸 수 없는 제어 문자가 든 칸·이름이 있으면 실행 기록을 만들기 전에 거절한다 — 위치만 알린다
    (값 원문은 응답·실행 기록·서버 기록 어디에도 넣지 않음). 값을 지우거나 바꿔 만들지 않는다."""
    if built.control_count:
        more = built.control_count - len(built.control)
        fixes = []
        if built.control_count > built.control_glossary:
            fixes.append("원본 파일에서 그 문자를 지운 뒤 다시 올리고 검토에서 다시 검사하세요")
        if built.control_glossary:
            fixes.append("용어사전 항목의 이름은 용어사전에서 고친 뒤 검토에서 다시 검사하고 통합하세요")
        raise IntegrationError(
            "통합하지 않았습니다 — 결과 파일(엑셀)에 쓸 수 없는 제어 문자(NUL 등)가 있습니다: "
            f"{' / '.join(built.control)}{f' 외 {more}곳' if more else ''}. {' · '.join(fixes)}. [{CONTROL_CHAR}]"
        )


def pending_run(db: Session, session: ReviewSession) -> IntegrationRun | None:
    """자동화 실행으로 만든 '검토 대기' 실행 기록 (다시 검사한 세션이면 이전 세션의 기록)."""
    current: ReviewSession | None = session
    while current is not None:
        run = db.scalar(
            select(IntegrationRun).where(
                IntegrationRun.review_session_id == current.id, IntegrationRun.status == "review"
            )
        )
        if run is not None:
            return run
        current = db.get(ReviewSession, current.parent_id) if current.parent_id else None
    return None


def include_unmapped_default(db: Session, session: ReviewSession) -> bool:
    """자동화로 실행한 검토면 그 자동화에 저장한 선택, 아니면 넣지 않음(기존 동작)."""
    if session.automation_id is None:
        return False
    automation = db.get(Automation, session.automation_id)
    if automation is None:
        return False
    return bool(json.loads(automation.output_format).get("include_unmapped", False))


def integrate(db: Session, session: ReviewSession, include_unmapped: bool | None = None) -> IntegrationRun:
    """include_unmapped: None이면 자동화 설정(없으면 False)을 따른다.
    결과 만들기 차례(app/output_slot.py)를 쥐고 만든다 — 차례가 없으면 OutputBusy (시작하지 않음, 기록 없음).
    결과 파일을 확정하는 순서·오류 처리는 app/runlease.py (이름을 바꾼 결과 파일은 지우지 않는다)."""
    with output_slot():
        return _integrate(db, session, include_unmapped)


def _integrate(db: Session, session: ReviewSession, include_unmapped: bool | None) -> IntegrationRun:
    started = datetime.now(KST)
    run = pending_run(db, session)
    if include_unmapped is None:
        include_unmapped = include_unmapped_default(db, session)
    built = build(db, session, include_unmapped)
    _check_size(built)  # 결과는 메모리에 다 만든 뒤 저장한다 — '실행 중' 기록을 남기기 전에 거절 (docs/MEMORY.md)
    _check_names(built)  # 변경 기록에 저장할 이름의 저장 규칙 — 같은 때 거절
    now = datetime.now(KST)
    items = session_items(db, session)
    warning_count = summarize(session, items).warnings
    excluded = set(json.loads(session.excluded_file_ids))
    # 검증 경고의 내용 (합계 검사 불완전·산식·이상치 등) — 건수만이 아니라 무엇을 확인하지 못했는지 결과에 남긴다
    shown = [i for i in items if i.category == "warning" and (i.file_id is None or i.file_id not in excluded)]
    warnings = [
        {"kind": WARNING_LABEL.get(i.kind, i.kind), "file": i.file_name, "sheet": i.sheet, "row": i.row,
         "column": i.column, "reason": mask_text(i.reason)}
        for i in shown
    ]  # fmt: skip
    # 경고 사유는 처리 요약에 쓸 때 쓸 수 없는 문자를 '�'로 바꾼다 (_printable) — 거절하지 않는다. 사유의 글자는
    # 이름·결과 칸(build가 그 자리로 거절)이거나 결과에 쓰지 않는 칸(합계 행의 이름 등 — 막을 까닭이 없음)에서 온다
    _check_text(built)  # 결과 파일에 쓸 수 없는 제어 문자 — 실행 기록을 만들기 전에 거절
    file_ids = [f["id"] for f in built.files if not f["excluded"]]
    # 자동화가 만든 '검토 대기' 행은 DB에서도 그 상태일 때만 이어 쓴다 (겹친 두 요청이 같은 행을 쓰지 않게)
    reuse = run is not None
    if run is None:
        run = IntegrationRun(
            trigger="automation" if session.automation_id else "manual",
            automation_id=session.automation_id,
            user_id=session.created_by,
            started_at=started,
        )
    run.review_session_id = session.id
    run.file_ids = json.dumps(file_ids)
    run.input_file_count = len(json.loads(session.file_ids))
    run.row_count = len(built.rows)
    run.warning_count = warning_count
    run.change_count = 0
    run.summary = "{}"
    try:
        target = runlease.begin(db, run, lambda run_id: _output_path(run_id, now), "review" if reuse else None)
    except runlease.RunError as exc:
        raise IntegrationError(str(exc)) from exc

    cell_changes = [c for c in built.changes if c.output_col is not None]
    counts: dict[str, int] = {}
    for ch in built.changes:
        counts[ch.kind] = counts.get(ch.kind, 0) + 1
    summary = {
        "processed_at": now.strftime("%Y-%m-%d %H:%M:%S (KST)"),
        "run_id": run.id,
        "review_id": session.id,
        "validation_run_id": session.validation_run_id,
        "expected_period": json.loads(session.params).get("expected_period"),
        "row_count": len(built.rows),
        "cell_change_count": len(cell_changes),
        "column_change_count": counts.get("column", 0),
        "warning_count": warning_count,
        "warnings": warnings,
        "change_counts": counts,
        "files": built.files,
        "dropped_columns": built.dropped_columns,
        "columns": built.labels,
        "include_unmapped": include_unmapped,
        "unmapped_columns": built.unmapped_columns,
    }
    tmp_path: Path | None = None
    try:
        tmp_path = runlease.new_temp(target)
        painted = write_workbook(tmp_path, built, summary)
        if painted != len(cell_changes):  # 색칠한 셀 = 셀 단위 변경 기록 (빠진 기록이 있으면 실패)
            raise IntegrationError(f"변경 기록 불일치: 칠한 셀 {painted}개, 기록 {len(cell_changes)}개")
        changes = [
            {"run_id": run.id, **{k: v for k, v in asdict(ch).items() if k != "output_col"}} for ch in built.changes
        ]
        text_summary = json.dumps(summary, ensure_ascii=False, default=str)
        runlease.prepare(db, run, tmp_path, text_summary, len(built.changes), changes)
        runlease.finish(db, run, tmp_path, target)
    except runlease.RunDone:  # 정리가 이 실행을 이미 완료로 확정했다 (결과 파일 해시 일치)
        db.refresh(run)
    except runlease.RunSettled as exc:  # 다른 곳이 실패로 정했다 — 그 이유로 알린다
        raise IntegrationError(str(exc)) from exc
    except runlease.RunUncertain:  # 결과 파일은 있고 완료 기록만 확인하지 못함 — 지우지 않고 503으로 알린다
        raise
    except Exception as exc:
        # 응답·실행 기록·서버 기록 모두 원문 없이 — 예외 종류와 실행 번호만 (app/failures.py)
        log_failure(logger, "integration", run.id, exc)
        reason = failure_reason(exc, (IntegrationError, runlease.RunError))
        try:
            runlease.fail(db, run, tmp_path, f"{reason} [{INTEGRATION_FAILED}]", "{}")
        except runlease.RunDone:
            db.refresh(run)
            return run
        except runlease.RunSettled as settled:
            raise IntegrationError(str(settled)) from exc
        raise IntegrationError(
            f"통합 결과를 만들지 못했습니다 — {reason} (실행 #{run.id}). [{INTEGRATION_FAILED}]"
        ) from exc
    finally:
        runlease.untrack(run.id)
    logger.info("integration run=%s rows=%d changes=%d", run.id, run.row_count, run.change_count)
    return run
