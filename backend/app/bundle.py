"""파일별 시트로 묶기 (취합 방식 '파일별 시트'): 고른 원본 시트를 결과 엑셀의 시트 하나씩으로 옮기고,
선택하면 '통계 요약' 시트를 더한다. 하나의 통합표(integration.py)와 따로 동작하며 검토·열 매핑을 요구하지 않는다.

- 시트 내용은 앱이 읽은 값·열만 쓴다 (서식·병합·차트·그림·수식은 옮기지 않음). 원본 파일·매핑·학습은 바꾸지 않는다.
  머리글 위 제목 행과 빈 행은 넣지 않고 몇 행인지 남긴다.
  소계·합계·주석 행은 '행 종류'와 함께 그대로 넣는다 (지우지 않음).
- 수식 칸은 원본에 저장된 결과값을 넣는다. 결과값이 저장돼 있지 않으면 비워 두고 메모로 표시한다 (0으로 채우지 않음).
- 글자는 수식으로 저장되지 않게 글자 형식으로 쓴다.
  코드·식별자 열은 파서가 글자로 둔 그대로다 (앞자리 0·긴 자릿수 보존).
- 통계 요약을 켜지 않으면 분석 엔진·모델을 부르지 않는다. 켜면 데이터 분석과 같은 엔진(analysis.evaluate·compute)으로
  항목마다 계산한다 (합계·평균을 따로 만들지 않음).
  요약을 만들 수 없으면 결과 파일을 만들지 않는다 (요약 없이 만들지 않음).
- 결과 구성 확인(plan)의 basis와 실행 때 다시 계산한 basis가 다르면 만들지 않는다
  (파일 내용·시트·순서·시트 이름·개인정보 확인 상태·요약 조건·그 시트들의 확정 매핑).
"""

from __future__ import annotations

import hashlib
import json
import logging
import os
import re
import stat
import tempfile
import threading
from collections import Counter, OrderedDict
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any, Literal

import openpyxl
from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, ConfigDict, Field
from sqlalchemy.orm import Session

from app import analysis
from app.analysis import AGG_LABEL, MEASURE_FIELDS, NO_VALUE, AnalysisIn, Evaluation, TargetIn, label_of
from app.analysis_api import LEVEL_TEXT, level_conflict_text
from app.auth import CurrentUser
from app.config import DATA_DIR
from app.decisions import load_decisions
from app.files import get_file_for, load_parsed
from app.integration import KST, OUTPUT_DIR
from app.mapping_review import sheet_id
from app.models import IntegrationRun, SourceFile
from app.parsing import FileResult
from app.pii import mask_text
from app.validation_api import pii_out

logger = logging.getLogger(__name__)

POLICY = "bundle-policy-1"  # 시트 구성·이름·쓰기 규칙이 바뀌면 올린다 (basis·실행 기록에 남음)
MODE = "sheets"
MAX_FILES = 50  # 한 번에 고를 수 있는 파일 (홈 검사 대상 한도와 같음)
MAX_SHEETS = 50  # 결과에 넣는 원본 시트
MAX_ROWS = 300_000  # 결과에 넣는 행 합계 (데이터 분석 한도와 같음)
EXCEL_MAX_ROWS = 1_048_576  # 엑셀 한 시트의 최대 행 (머리글 1행 포함)
MAX_ITEMS = 10  # 통계 요약 항목
SUMMARY_SHEET = "통계 요약"
NAME_MAX = 31  # 엑셀 시트 이름 길이 한도
EXCEL_DIGITS = 15  # 엑셀 숫자 칸이 정확히 담는 자릿수
FORBIDDEN = re.compile(r"[\[\]:*?/\\]")
RESERVED = (SUMMARY_SHEET, "처리 요약", "History")  # 요약 시트·통합표 시트 이름, 엑셀 예약 이름
FORMULA_START = ("=", "+", "-", "@", "\t", "\r")  # analysis_api와 같은 규칙: 이 글자로 시작하면 글자 형식으로 쓴다
KIND_TEXT = {"data": "데이터", "subtotal": "소계·합계", "note": "주석"}
LEAD = ["원본 행", "행 종류"]
HEAD_FILL = PatternFill("solid", fgColor="EEF1F6")
GAP_FILL = PatternFill("solid", fgColor="FFF2CC")
Aggregate = Literal["sum", "mean", "min", "max", "count"]


class BundleError(Exception):
    pass


# ---------------------------------------------------------------- 요청


class SheetRef(BaseModel):
    model_config = ConfigDict(extra="forbid")

    file_id: int
    sheet_id: str = Field(pattern=r"^s\d{1,3}$")


class LinkIn(SheetRef):
    """통계 요약 항목에서 시트 하나를 어떻게 쓸지 (확정 매핑으로 고른 열·단위·기간을 사용자가 보충하거나 뺌)."""

    column_id: str | None = Field(default=None, pattern=r"^c\d{1,4}$")  # 같은 항목·구분의 열이 여러 개일 때만
    unit: str | None = Field(default=None, min_length=1, max_length=20)
    period: str | None = Field(default=None, pattern=r"^(19|20)\d{2}(-(0[1-9]|1[0-2]))?$")
    rows_are_records: bool = False
    exclude: bool = False  # 이 항목 계산에서 뺌 (사유와 함께 요약 시트에 남음)
    exclude_reason: str | None = Field(default=None, max_length=100)


class ItemIn(BaseModel):
    model_config = ConfigDict(extra="forbid")

    field_key: str = Field(max_length=64)  # 사용자가 확정한 표준 항목 (열 위치·이름이 아니라 확정한 뜻으로 연결)
    aggregates: list[Aggregate] = Field(min_length=1, max_length=5)
    variant: str | None = Field(default=None, max_length=200)
    unit: str | None = Field(default=None, max_length=20)
    region_level: Literal["lower", "upper"] | None = None
    links: list[LinkIn] = Field(default_factory=list, max_length=MAX_SHEETS)


class SummaryIn(BaseModel):
    model_config = ConfigDict(extra="forbid")

    items: list[ItemIn] = Field(min_length=1, max_length=MAX_ITEMS)


class BundleIn(BaseModel):
    model_config = ConfigDict(extra="forbid")

    file_ids: list[int] = Field(min_length=1, max_length=MAX_FILES)
    # 결과 시트 순서. 시트가 하나뿐인 파일은 비우면 끝에 차례로 붙는다. 시트가 여러 개인 파일은 직접 골라야 한다.
    sheets: list[SheetRef] = Field(default_factory=list, max_length=MAX_SHEETS)
    summary: SummaryIn | None = None  # None = 통계 요약 없음 (분석 엔진을 부르지 않음)


class BundleRunIn(BundleIn):
    expected_basis: str = Field(min_length=16, max_length=64)


# ---------------------------------------------------------------- 파일 정보 (내용 해시로 기억)


@dataclass
class SheetInfo:
    index: int
    name: str  # 원래 시트 이름 (서버 안에서만)
    columns: int
    header_row: int | None
    two_level: bool
    kinds: dict[str, int]  # data · subtotal · note · empty
    formula_missing: int  # 수식인데 저장된 결과값이 없는 칸
    formula_cached: int  # 저장된 결과값을 쓰는 수식 칸
    long_numbers: int  # 엑셀 숫자 칸이 정확히 담지 못하는 15자리 넘는 정수


@dataclass
class FileInfo:
    file_id: int
    sha256: str
    name: str  # 원래 파일 이름 (서버 안에서만)
    ext: str
    sheets: list[SheetInfo]
    skipped: list[str]  # 앱이 읽지 않는 설명 시트 (메타정보·설명·안내)


_INFO: OrderedDict[tuple[int, str], FileInfo] = OrderedDict()
_INFO_LOCK = threading.Lock()
INFO_CACHE = 64  # 시트별 개수만 담은 작은 값이라 넉넉히 둔다


def formula_cells(path: Path) -> dict[str, set[tuple[int, int]]]:
    """xlsx 시트별 수식 칸 (1부터 센 행·열).
    파서는 저장된 결과값만 읽으므로 수식 여부는 여기서 따로 본다 (읽기만 함)."""
    if path.suffix.lower() != ".xlsx":
        return {}
    wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        out: dict[str, set[tuple[int, int]]] = {}
        for ws in wb.worksheets:
            cells: set[tuple[int, int]] = set()
            for row in ws.iter_rows():
                for cell in row:
                    if getattr(cell, "data_type", None) == "f":
                        cells.add((cell.row, cell.column))
            out[ws.title] = cells
        return out
    finally:
        wb.close()


def _gap(entry: dict[str, Any], j: int, source: list[int], formulas: set[tuple[int, int]]) -> str | None:
    """칸 하나의 수식 상태: 'missing'(결과값 없음) · 'cached'(저장된 결과값) · None(수식 아님)."""
    if (entry["row"], source[j] + 1) not in formulas:
        return None
    return "missing" if entry["original"][j] is None else "cached"


def _long(value: Any) -> bool:
    return isinstance(value, int) and not isinstance(value, bool) and len(str(abs(value))) > EXCEL_DIGITS


def _file_path(record: SourceFile) -> Path:
    return (DATA_DIR / record.stored_path).resolve()


def file_info(db: Session, record: SourceFile) -> FileInfo:
    key = (record.id, record.sha256)
    with _INFO_LOCK:
        if key in _INFO:
            _INFO.move_to_end(key)
            return _INFO[key]
    _, parsed = load_parsed(db, record.id)
    formulas = formula_cells(_file_path(record))
    sheets = []
    for i, s in enumerate(parsed.sheets):
        cells = formulas.get(s.name, set())
        kinds = {"data": s.data_rows, "subtotal": s.subtotal_rows, "note": s.note_rows, "empty": s.empty_rows}
        missing = cached = long_numbers = 0
        for entry in s.body:
            if entry["kind"] == "empty":
                continue
            for j, value in enumerate(entry["cells"]):
                state = _gap(entry, j, s.source_columns, cells) if cells else None
                missing += state == "missing"
                cached += state == "cached"
                long_numbers += _long(value)
        sheets.append(
            SheetInfo(
                i,
                s.name,
                len(s.columns),
                s.header_row,
                (s.header_row_end or 0) > (s.header_row or 0),
                kinds,
                missing,
                cached,
                long_numbers,
            )  # fmt: skip
        )
    info = FileInfo(record.id, record.sha256, record.original_name, Path(record.original_name).suffix.lower(), sheets,
                    list(parsed.skipped_sheets))  # fmt: skip
    with _INFO_LOCK:
        _INFO[key] = info
        while len(_INFO) > INFO_CACHE:
            _INFO.popitem(last=False)
    return info


# ---------------------------------------------------------------- 시트 이름


def sheet_names(bases: list[str]) -> list[tuple[str, list[str]]]:
    """결과 시트 이름 (엑셀 규칙에 맞게)과 원래 이름에서 바꾼 이유.
    예약 이름·이미 쓴 이름은 대소문자 구분 없이 피한다."""
    used = {n.casefold() for n in RESERVED}
    out = []
    for base in bases:
        reasons = []
        name = FORBIDDEN.sub("_", base)
        if name != base:
            reasons.append("시트 이름에 쓸 수 없는 글자([ ] : * ? / \\)를 _로 바꿈")
        trimmed = name.strip().strip("'").strip()
        if trimmed != name:
            reasons.append("앞뒤 공백·작은따옴표를 뺌")
            name = trimmed
        if not name:
            name = "시트"
            reasons.append("이름이 비어 '시트'로 둠")
        if len(name) > NAME_MAX:
            name = name[:NAME_MAX].rstrip()
            reasons.append(f"시트 이름 {NAME_MAX}자 제한으로 줄임")
        if name.casefold() in used:
            n = 2
            while True:
                suffix = f" ({n})"
                candidate = name[: NAME_MAX - len(suffix)].rstrip() + suffix
                if candidate.casefold() not in used:
                    break
                n += 1
            reasons.append(f"같은 이름이나 예약된 이름('{SUMMARY_SHEET}'·'처리 요약' 등)이 있어 번호를 붙임")
            name = candidate
        used.add(name.casefold())
        out.append((name, reasons))
    return out


# ---------------------------------------------------------------- 결과 구성 (plan)


@dataclass
class OutSheet:
    file: FileInfo
    sheet: SheetInfo
    name: str = ""
    renamed: list[str] = field(default_factory=list)

    @property
    def rows(self) -> int:
        return self.sheet.kinds["data"] + self.sheet.kinds["subtotal"] + self.sheet.kinds["note"]

    def out(self) -> dict[str, Any]:
        s = self.sheet
        return {
            "file_id": self.file.file_id,
            "file_name": mask_text(self.file.name),
            "sheet_id": sheet_id(s.index),
            "sheet": mask_text(s.name),
            "name": self.name,
            "renamed": self.renamed,
            "rows": self.rows,
            "kinds": {k: v for k, v in s.kinds.items() if k != "empty"},
            "empty_rows": s.kinds["empty"],
            "title_rows": max(0, (s.header_row or 1) - 1),
            "two_level_header": s.two_level,
            "columns": s.columns,
            "formula_missing": s.formula_missing,
            "formula_cached": s.formula_cached,
            "formula_checked": self.file.ext == ".xlsx",
            "long_numbers": s.long_numbers,
        }


@dataclass
class ItemPlan:
    item: ItemIn
    ev: Evaluation | None
    out: dict[str, Any]
    ready: bool


@dataclass
class Plan:
    req: BundleIn
    files: list[dict[str, Any]]
    sheets: list[OutSheet]
    items: list[ItemPlan] | None
    fields: list[dict[str, Any]]
    problems: list[str]
    warnings: list[str]
    basis: str = ""

    @property
    def ready(self) -> bool:
        return not self.problems and bool(self.sheets) and all(i.ready for i in self.items or [])

    def out(self) -> dict[str, Any]:
        total = sum(s.rows for s in self.sheets)
        return {
            "mode": MODE,
            "files": self.files,
            "sheets": [s.out() for s in self.sheets],
            "summary": None
            if self.items is None
            else {
                "sheet_name": SUMMARY_SHEET,
                "items": [i.out for i in self.items],
                "ready": all(i.ready for i in self.items),
            },
            "fields": self.fields,
            "problems": self.problems,
            "warnings": self.warnings,
            "total_rows": total,
            "sheet_count": len(self.sheets) + (1 if self.items is not None else 0),
            "limits": {
                "files": MAX_FILES,
                "sheets": MAX_SHEETS,
                "rows": MAX_ROWS,
                "excel_rows": EXCEL_MAX_ROWS,
                "items": MAX_ITEMS,
                "summary_sheets": analysis.MAX_TARGETS,
                "summary_rows": analysis.MAX_ROWS,
            },
            "basis": self.basis,
            "ready": self.ready,
        }


def _bad(message: str) -> HTTPException:
    return HTTPException(422, message)


def _detail(exc: HTTPException) -> str:
    d = exc.detail
    if isinstance(d, dict):
        return " · ".join([str(d.get("message", "")), *map(str, d.get("problems", []))]).strip(" ·")
    return str(d)


def build_plan(db: Session, user: CurrentUser, req: BundleIn) -> Plan:
    """요청한 구성으로 결과 시트·이름·한도·요약 판정을 정한다 (파일을 쓰지 않음). 실행도 같은 함수로 다시 판정한다."""
    if len(set(req.file_ids)) != len(req.file_ids):
        raise _bad("같은 파일을 두 번 골랐습니다.")
    problems: list[str] = []
    warnings: list[str] = []
    infos: dict[int, FileInfo] = {}
    records: dict[int, SourceFile] = {}
    files_out: list[dict[str, Any]] = []
    for fid in req.file_ids:
        record = get_file_for(db, user, fid)  # 볼 수 없는 파일은 없는 것처럼 404
        records[fid] = record
        info = file_info(db, record)
        infos[fid] = info
        own = record.uploaded_by == user.id or (record.uploaded_by is None and user.is_manager)
        summary = json.loads(record.pii_summary) if record.pii_summary else {}
        has_pii = any(summary.get("counts", {}).values())
        pii = "none" if not has_pii else ("acknowledged" if record.pii_acknowledged_at is not None else "needs_ack")
        shown = f"#{fid} {mask_text(record.original_name)}"
        file_problems = []
        if not own:
            file_problems.append(
                f"{shown}: 다른 사용자의 파일이라 묶을 수 없습니다 (본인 파일만 — 남의 파일은 조회만)."
            )
        if pii == "needs_ack":
            file_problems.append(
                f"{shown}: 개인정보 의심 값이 있어 포함 확인이 필요합니다 — 포함을 확인하거나 이 파일을 빼세요."
            )
        if not info.sheets:
            file_problems.append(f"{shown}: 앱이 읽은 시트가 없습니다.")
        problems.extend(file_problems)
        files_out.append(
            {
                "file_id": fid,
                "file_name": mask_text(record.original_name),
                "sheets": [
                    {
                        "sheet_id": sheet_id(s.index),
                        "sheet": mask_text(s.name),
                        "rows": s.kinds["data"] + s.kinds["subtotal"] + s.kinds["note"],
                        "columns": s.columns,
                    }
                    for s in info.sheets
                ],
                "skipped_sheets": [mask_text(n) for n in info.skipped],
                "pii": pii,
                "pii_summary": found.model_dump()
                if (found := pii_out(record))
                else None,  # 파일 목록과 같은 형식 (가린 값)
                "own": own,
                "problems": file_problems,
            }
        )
    if any(f["pii"] == "acknowledged" for f in files_out):
        warnings.append(
            "개인정보 포함을 확인한 파일이 있습니다 — 결과 파일에는 화면과 달리 원본 값이 가리지 않고 들어갑니다 "
            "(결과는 이 실행 기록을 볼 수 있는 사람만 내려받을 수 있음)."
        )
    # 시트 순서: 요청한 순서, 비어 있으면 시트가 하나뿐인 파일을 파일 순서대로
    picks: list[tuple[int, int]] = []
    for ref in req.sheets:
        if ref.file_id not in infos:
            raise _bad(f"고르지 않은 파일의 시트입니다: #{ref.file_id}")
        index = int(ref.sheet_id[1:])
        if index >= len(infos[ref.file_id].sheets):
            raise _bad(f"#{ref.file_id} 파일에 없는 시트입니다: {ref.sheet_id}")
        if (ref.file_id, index) in picks:
            raise _bad(f"같은 시트를 두 번 골랐습니다: #{ref.file_id} {ref.sheet_id}")
        picks.append((ref.file_id, index))
    for fid in req.file_ids:
        if any(p[0] == fid for p in picks):
            continue
        info = infos[fid]
        if len(info.sheets) == 1:
            picks.append((fid, 0))
        elif info.sheets:
            problems.append(
                f"#{fid} {mask_text(info.name)}: 시트가 {len(info.sheets)}개입니다 — 넣을 시트를 직접 고르세요 "
                "(고른 시트마다 결과 시트 하나)."
            )
    per_file = Counter(fid for fid, _ in picks)
    out_sheets = [OutSheet(infos[fid], infos[fid].sheets[i]) for fid, i in picks]
    # 시트가 하나뿐인 파일은 파일 이름, 여러 시트인 파일은 '파일-시트' (하나만 골라도 어느 시트인지 보이게)
    bases = [
        mask_text(Path(o.file.name).stem if len(o.file.sheets) == 1 else f"{Path(o.file.name).stem}-{o.sheet.name}")
        for o in out_sheets
    ]
    for o, (name, reasons) in zip(out_sheets, sheet_names(bases), strict=True):
        o.name, o.renamed = name, reasons
    # 한도: 넘으면 일부만 넣지 않고 만들지 않는다
    total = sum(o.rows for o in out_sheets)
    if len(out_sheets) > MAX_SHEETS:
        problems.append(
            f"결과 시트가 {len(out_sheets)}개로 한 번에 넣을 수 있는 {MAX_SHEETS}개를 넘습니다 — 나눠 만드세요."
        )
    if total > MAX_ROWS:
        problems.append(f"넣을 행이 {total:,}행으로 한 번에 넣을 수 있는 {MAX_ROWS:,}행을 넘습니다 — 나눠 만드세요.")
    for o in out_sheets:
        if o.rows + 1 > EXCEL_MAX_ROWS:
            problems.append(f"'{o.name}': {o.rows:,}행은 엑셀 한 시트에 넣을 수 있는 행 수를 넘습니다.")
    missing = sum(o.sheet.formula_missing for o in out_sheets)
    if missing:
        warnings.append(
            f"수식 결과값이 저장돼 있지 않은 칸 {missing:,}개는 비워 두고 메모로 표시합니다 (0으로 채우지 않음) — "
            "원본을 엑셀에서 열어 저장하면 결과값이 저장됩니다."
        )
    long_numbers = sum(o.sheet.long_numbers for o in out_sheets)
    if long_numbers:
        warnings.append(
            f"{EXCEL_DIGITS}자리를 넘는 숫자 {long_numbers:,}개는 엑셀이 끝자리를 정확히 담지 못합니다 "
            "(코드·식별자 열은 글자로 두어 해당 없음)."
        )
    sha: dict[str, list[int]] = {}
    for fid in req.file_ids:
        sha.setdefault(infos[fid].sha256, []).append(fid)
    for same in sha.values():
        if len(same) > 1:
            warnings.append(
                f"같은 내용의 파일을 함께 골랐습니다 ({', '.join(f'#{f}' for f in same)}) — "
                "같은 시트가 두 번 들어갑니다."
            )
    decisions = {fid: load_decisions(db, fid) for fid in per_file}
    fields = _fields(out_sheets, decisions)
    items = None if req.summary is None else _summary(db, user, req.summary, out_sheets, problems)
    plan = Plan(req, files_out, out_sheets, items, fields, problems, warnings)
    plan.basis = _basis(plan, records)
    return plan


def _fields(sheets: list[OutSheet], decisions: dict[int, dict[str, dict[str, str | None]]]) -> list[dict[str, Any]]:
    """고른 시트에서 사용자가 확정한 수량 항목 (통계 요약 항목 후보). DB만 읽는다."""
    count: Counter[str] = Counter()
    for o in sheets:
        keys = {v for v in decisions[o.file.file_id].get(o.sheet.name, {}).values() if v in MEASURE_FIELDS}
        count.update(keys)
    return sorted(({"key": k, "label": label_of(k), "sheets": n} for k, n in count.items()), key=lambda f: f["label"])


def _coverage(included: int, total: int, sheets: list[OutSheet]) -> str:
    unit = "파일" if len({o.file.file_id for o in sheets}) == len(sheets) else "시트"
    return f"{total}개 {unit} 중 {included}개 포함"


def _summary(
    db: Session, user: CurrentUser, summary: SummaryIn, sheets: list[OutSheet], problems: list[str]
) -> list[ItemPlan]:
    """통계 요약 항목마다 데이터 분석 엔진으로 판정한다 (계산은 실행 때 compute로)."""
    plans = []
    seen: set[tuple[str, str | None]] = set()
    placed = {(o.file.file_id, sheet_id(o.sheet.index)) for o in sheets}
    for item in summary.items:
        if item.field_key not in MEASURE_FIELDS:
            raise _bad(f"통계 요약에 쓸 수 없는 항목입니다: {item.field_key} (확정한 수량 항목만)")
        if len(set(item.aggregates)) != len(item.aggregates):
            raise _bad("같은 계산 방식을 두 번 골랐습니다.")
        label = label_of(item.field_key)
        item_problems: list[str] = []
        if (item.field_key, item.variant) in seen:
            item_problems.append(f"'{label}' 항목을 두 번 골랐습니다 — 하나로 합쳐 계산 방식을 고르세요.")
        seen.add((item.field_key, item.variant))
        links = {(lk.file_id, lk.sheet_id): lk for lk in item.links}
        if len(links) != len(item.links):
            raise _bad(f"'{label}': 같은 시트의 연결을 두 번 보냈습니다.")
        if stray := [k for k in links if k not in placed]:
            raise _bad(f"'{label}': 결과에 넣지 않은 시트의 연결입니다 ({', '.join(f'#{f} {s}' for f, s in stray)}).")
        dropped = {k for k, lk in links.items() if lk.exclude}
        used = [o for o in sheets if (o.file.file_id, sheet_id(o.sheet.index)) not in dropped]
        ev: Evaluation | None = None
        if len(used) > analysis.MAX_TARGETS:
            item_problems.append(
                f"'{label}': 계산할 시트가 {len(used)}개로 "
                f"통계 요약 한도({analysis.MAX_TARGETS}개, 데이터 분석과 같음)를 "
                "넘습니다 — 일부를 이 항목에서 빼거나 나눠 만드세요."
            )
        elif used:
            targets = []
            for o in used:
                lk = links.get((o.file.file_id, sheet_id(o.sheet.index)))
                targets.append(
                    TargetIn(
                        file_id=o.file.file_id,
                        sheet_id=sheet_id(o.sheet.index),
                        column_id=lk.column_id if lk else None,
                        unit=lk.unit if lk else None,
                        period=lk.period if lk else None,
                        rows_are_records=lk.rows_are_records if lk else False,
                    )
                )
            req = AnalysisIn(field_key=item.field_key, variant=item.variant, unit=item.unit, targets=targets,
                             aggregates=list(item.aggregates), region_level=item.region_level)  # fmt: skip
            try:
                ev = analysis.evaluate(db, user, req)
            except HTTPException as exc:
                item_problems.append(f"'{label}': {_detail(exc)}")
            db.rollback()  # 읽기만 한 트랜잭션을 닫는다
        else:
            item_problems.append(f"'{label}': 계산에 넣을 시트가 없습니다 (모두 뺌).")
        by_sheet = {(t.file_id, t.sheet_id): t for t in ev.targets} if ev else {}
        rows = []
        for o in sheets:
            key = (o.file.file_id, sheet_id(o.sheet.index))
            lk = links.get(key)
            base = {"file_id": key[0], "file_name": mask_text(o.file.name), "sheet_id": key[1],
                    "sheet": mask_text(o.sheet.name), "output_sheet": o.name}  # fmt: skip
            if lk is not None and lk.exclude:
                reason = (
                    mask_text(lk.exclude_reason.strip())
                    if lk.exclude_reason and lk.exclude_reason.strip()
                    else "사용자가 뺌"
                )
                rows.append({**base, "status": "user_excluded", "reasons": [], "exclude_reason": reason})
                continue
            t = by_sheet.get(key)
            if t is None:
                rows.append({**base, "status": "unchecked", "reasons": [], "exclude_reason": None})
                continue
            rows.append({**base, **t.out(), "output_sheet": o.name, "field_label": label, "exclude_reason": None})
            shown = f"'{label}' {o.name}"
            if t.status == "needs_input":
                item_problems.append(f"{shown}: {'; '.join(t.reasons)}")
            elif t.status == "excluded":
                item_problems.append(
                    f"{shown}: {'; '.join(t.reasons)} — 원본·열 매핑을 고치거나 이 항목에서 '빼기'로 확정하세요."
                )
        if ev is not None:
            item_problems.extend(f"'{label}': {p}" for p in ev.problems)
        included = sum(1 for r in rows if r["status"] == "ok")
        ready = ev is not None and not item_problems and included > 0
        pv = analysis.preview_out(ev) if ev is not None else None
        problems.extend(item_problems)
        plans.append(
            ItemPlan(
                item,
                ev,
                {
                    "field_key": item.field_key,
                    "field_label": label,
                    "aggregates": list(item.aggregates),
                    "aggregate_labels": [AGG_LABEL[a] for a in item.aggregates],
                    # 사용자가 고른 계산 조건 (고르지 않았으면 None — 엔진이 요구할 때만 화면에서 고른다)
                    "choices": {"variant": item.variant, "unit": item.unit, "region_level": item.region_level},
                    "variant": pv["variant"] if pv else item.variant,
                    "variants": pv["variants"] if pv else [],
                    "unit": pv["unit"] if pv else None,
                    "units": pv["units"] if pv else [],
                    "periods": pv["periods"] if pv else [],
                    "stock": pv["stock"] if pv else False,
                    "links": rows,
                    "coverage": {
                        "included": included,
                        "total": len(sheets),
                        "text": _coverage(included, len(sheets), sheets),
                    },
                    "rows": pv["rows"] if pv else None,
                    "reason_labels": pv["reason_labels"] if pv else {},
                    "overlaps": pv["overlaps"] if pv else [],
                    "region_levels": pv["region_levels"] if pv else None,
                    "base_dates": pv["base_dates"] if pv else {},
                    "warnings": pv["warnings"] if pv else [],
                    "problems": item_problems,
                    "ready": ready,
                },
                ready,
            )
        )
    return plans


def _basis(plan: Plan, records: dict[int, SourceFile]) -> str:
    """결과 구성 확인 뒤 바뀌었는지 보는 값.
    요약을 켜면 항목마다 분석 엔진의 fingerprint(조건·파일 해시·확정 매핑·규칙)를 더한다."""
    payload = [
        POLICY,
        MODE,
        [[fid, records[fid].sha256, records[fid].pii_acknowledged_at is not None] for fid in plan.req.file_ids],
        [[o.file.file_id, o.sheet.index, o.name] for o in plan.sheets],
        None if plan.req.summary is None else plan.req.summary.model_dump(),
        None if plan.items is None else [i.ev.fingerprint if i.ev else None for i in plan.items],
    ]
    return hashlib.sha256(json.dumps(payload, ensure_ascii=False, default=str).encode()).hexdigest()


# ---------------------------------------------------------------- 결과 파일


def _put(ws: Any, row: int, col: int, value: Any) -> Any:
    """칸에 값을 쓴다. 글자는 수식으로 저장되지 않게 항상 글자 형식으로 둔다 ('=1+1' → 글자 그대로)."""
    cell = ws.cell(row=row, column=col, value=value)
    if isinstance(value, str) and value.startswith(FORMULA_START):
        cell.data_type = "s"
    return cell


def _clean(value: Any, stats: Counter[str]) -> Any:
    """엑셀(XML)에 넣을 수 없는 제어 문자는 빼고 몇 개인지 센다 (그 밖의 값은 그대로)."""
    if isinstance(value, str) and ILLEGAL_CHARACTERS_RE.search(value):
        stats["control_chars"] += len(ILLEGAL_CHARACTERS_RE.findall(value))
        return ILLEGAL_CHARACTERS_RE.sub("", value)
    return value


def _write_sheet(ws: Any, o: OutSheet, parsed: FileResult, formulas: dict[str, set[tuple[int, int]]]) -> dict[str, Any]:
    sheet = parsed.sheets[o.sheet.index]
    cells = formulas.get(sheet.name, set())
    stats: Counter[str] = Counter()
    header = [*LEAD, *sheet.columns]
    for j, name in enumerate(header, start=1):
        c = _put(ws, 1, j, _clean(name, stats))
        c.font, c.fill = Font(bold=True), HEAD_FILL
    ws.freeze_panes = "A2"
    ws["A1"].comment = Comment(
        f"원본: #{o.file.file_id} {mask_text(o.file.name)} / {mask_text(sheet.name)}\n"
        "앱이 읽은 값·열만 옮겼습니다 (서식·병합·차트·그림·수식은 옮기지 않음). "
        "병합 칸은 앱이 읽은 대로 값을 채웠습니다.",
        "Moajwo",
    )
    out_row = 1
    for entry in sheet.body:
        if entry["kind"] == "empty":
            continue
        out_row += 1
        _put(ws, out_row, 1, entry["row"])
        _put(ws, out_row, 2, KIND_TEXT[entry["kind"]])
        for j, value in enumerate(entry["cells"]):
            original = entry["original"][j]
            if isinstance(original, datetime | date):
                value = original  # 날짜는 원본 형식 그대로 (파서가 글자로 바꾼 값 대신)
            state = _gap(entry, j, sheet.source_columns, cells) if cells else None
            if state == "missing":
                c = _put(ws, out_row, j + 3, None)
                c.fill = GAP_FILL
                c.comment = Comment("원본의 수식 결과값이 저장돼 있지 않아 비워 둠 (0으로 채우지 않음)", "Moajwo")
                stats["formula_missing"] += 1
                continue
            stats["formula_cached"] += state == "cached"
            stats["long_numbers"] += _long(value)
            _put(ws, out_row, j + 3, _clean(value, stats))
    for j, name in enumerate(header, start=1):
        ws.column_dimensions[get_column_letter(j)].width = min(40, max(8, len(str(name)) * 2 + 2))
    changed = Counter()
    for c in sheet.corrections:
        changed[c["kind"]] += c["count"]
    return {**o.out(), "written_rows": out_row - 1, "changed": dict(changed), **{k: stats[k] for k in
            ("formula_missing", "formula_cached", "long_numbers", "control_chars")}}  # fmt: skip


def _plain(v: Any) -> Any:
    if v is None:
        return NO_VALUE
    return int(v) if isinstance(v, float) and v.is_integer() and abs(v) < 2**53 else v


def _write_summary(ws: Any, items: list[tuple[ItemPlan, dict[str, Any]]], processed_at: str) -> list[dict[str, Any]]:
    row = 1

    def line(*values: Any, bold: bool = False, size: int | None = None) -> None:
        nonlocal row
        for j, v in enumerate(values, start=1):
            c = _put(ws, row, j, v)
            if bold or size:
                c.font = Font(bold=bold, size=size or 11)
        row += 1

    line(SUMMARY_SHEET, bold=True, size=14)
    line(
        f"만든 시각 {processed_at} — 만든 때의 값입니다. "
        "원본·매핑이 바뀌어도 이 시트는 자동으로 다시 계산되지 않습니다."
    )
    line(
        f"계산: Moajwo 데이터 분석과 같은 엔진·규칙({analysis.POLICY}). 평균 = 유효한 숫자의 합 ÷ 유효한 숫자 개수 "
        "(파일별 평균을 다시 평균하지 않음). 빈칸·비공개·숫자가 아닌 값은 0으로 바꾸지 않고 뺍니다."
    )
    line(
        "숫자는 계산해 넣은 값입니다 (엑셀 수식 아님). "
        "계산은 배정밀도 실수로 해 매우 큰 값의 합계는 끝자리가 다를 수 있습니다."
    )
    saved = []
    for plan, res in items:
        o = plan.out
        row += 1
        line(f"항목: {o['field_label']}", bold=True, size=12)
        cond = analysis.conditions(plan.ev) if plan.ev else {}
        line("단위", cond.get("unit") or "-")
        line("기간", ", ".join(cond.get("periods") or []) or "-")
        if cond.get("base_dates"):
            line("기준일", ", ".join(f"{k} {v}" for k, v in cond["base_dates"].items()))
        line("포함", o["coverage"]["text"])
        # 사용자가 고른 계산 조건 — 무엇을 골랐는지 그대로 적는다 ('선택한 수준' 같은 말만 남기지 않음)
        item = plan.item
        variant = cond.get("variant")
        if item.variant is not None or variant:
            chosen = "사용자 선택" if item.variant is not None else "하나뿐이라 그대로"
            line("구분", f"{variant or '구분 없음'} ({chosen})")
        if item.unit is not None:
            line("계산 단위", f"{item.unit} (사용자 선택 — 다른 단위 자료는 이 항목에서 뺌, 자동 환산 없음)")
        conflicts = [level_conflict_text(x) for x in cond.get("level_conflicts") or []]
        level_text = LEVEL_TEXT.get(item.region_level, LEVEL_TEXT[None])
        if item.region_level is not None or conflicts:
            line("집계 수준", level_text)
            for text in conflicts:
                line("", f"상위·하위 차이: {text}")
        records = cond.get("records") or []
        if records:
            line("반복 행", f"서로 다른 기록으로 확인해 모두 더함 (사용자 확인): {', '.join(records)}")
        state = "부분 집계 (자료 문제로 뺀 값이 있음)" if res["partial"] else "전체 집계 (자료 문제로 뺀 값 없음)"
        if res.get("level_mismatch"):
            state += f" · 집계 수준 {level_text} (하위 지역 합계가 상위 지역 값과 다른 곳 {res['level_mismatch']}곳)"
        line("상태", state)
        line("계산한 행", f"{res['rows_included']:,}행 포함 · {res['rows_excluded']:,}행 제외")
        labels = o["reason_labels"]
        for k, n in sorted(res["reason_counts"].items()):
            line("", f"제외 사유: {labels.get(k, k)} {n:,}행")
        header = [*(res["group_labels"] or ["범위"]), *[AGG_LABEL[a] for a in o["aggregates"]]]
        line(*header, bold=True)
        groups = []
        for g in res["groups"]:
            keys = g["keys"] or ["전체"]
            values = [_plain(g[a]) for a in o["aggregates"]]
            line(*keys, *values)
            groups.append({"keys": keys, **{a: g[a] for a in o["aggregates"]}, "partial": g["partial"]})
        line("자료", "파일", "시트", "열", "표준 항목", "단위 (근거)", "기간 (근거)", "포함", "사유", bold=True)
        for lk in o["links"]:
            unit = f"{lk.get('unit') or '-'} ({lk.get('unit_source') or '-'})" if lk.get("unit") else "-"
            period = (
                f"{lk.get('period')} ({lk.get('period_source') or '-'})"
                if lk.get("period")
                else (f"행의 {', '.join(lk.get('period_columns') or [])} 열" if lk.get("period_columns") else "-")
            )
            included = {"ok": "포함", "user_excluded": "뺌 (사용자)", "excluded": "뺌"}.get(lk["status"], lk["status"])
            reason = lk.get("exclude_reason") or "; ".join(lk.get("reasons") or [])
            line(lk["output_sheet"], f"#{lk['file_id']} {lk['file_name']}", lk["sheet"], lk.get("column") or "-",
                 lk.get("field_label") or o["field_label"], unit, period, included, reason)  # fmt: skip
        saved.append(
            {
                "field_key": o["field_key"],
                "field_label": o["field_label"],
                "unit": cond.get("unit"),
                "periods": cond.get("periods"),
                "aggregates": o["aggregates"],
                "coverage": o["coverage"],
                "partial": res["partial"],
                "rows_included": res["rows_included"],
                "rows_excluded": res["rows_excluded"],
                "variant": variant,
                "variant_chosen": item.variant is not None,
                "unit_choice": item.unit,
                "region_level": item.region_level,
                "region_level_text": level_text,
                "level_conflicts": conflicts,
                "records": records,
                "excluded": [
                    {
                        "output_sheet": lk["output_sheet"],
                        "by": "user" if lk["status"] == "user_excluded" else "engine",
                        "reason": lk.get("exclude_reason") or "; ".join(lk.get("reasons") or []),
                    }
                    for lk in o["links"]
                    if lk["status"] in ("user_excluded", "excluded")
                ],  # fmt: skip
                "groups": groups,
            }  # fmt: skip
        )
    ws.column_dimensions["A"].width = 18
    for col in "BCDEFGHI":
        ws.column_dimensions[col].width = 16
    return saved


def _output_path(run_id: int, now: datetime) -> Path:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    return OUTPUT_DIR / f"시트묶음_{run_id:04d}_{now.strftime('%Y%m%d_%H%M%S')}.xlsx"


def make(db: Session, user: CurrentUser, plan: Plan) -> IntegrationRun:
    """판정이 끝난(ready) 구성으로 결과 파일을 만든다.
    요약을 켰으면 먼저 계산하고, 계산할 수 없으면 아무것도 만들지 않는다."""
    started = datetime.now(KST)
    results: list[tuple[ItemPlan, dict[str, Any]]] = []
    for item in plan.items or []:
        assert item.ev is not None
        try:
            results.append((item, analysis.compute(item.ev)))
        except HTTPException as exc:
            raise BundleError(
                f"통계 요약 '{item.out['field_label']}'을(를) 계산할 수 없어 결과 파일을 만들지 않았습니다: "
                f"{_detail(exc)} — "
                "요약 설정을 고치거나 '통계 요약 시트 추가'를 끄고 다시 만드세요."
            ) from exc
    run = IntegrationRun(
        trigger="manual",
        user_id=user.id,
        status="running",
        started_at=started,
        file_ids=json.dumps(plan.req.file_ids),
        input_file_count=len(plan.req.file_ids),
        row_count=sum(o.rows for o in plan.sheets),
        warning_count=len(plan.warnings),
        summary=json.dumps({"mode": MODE, "stats_summary": plan.items is not None}),
    )
    db.add(run)
    db.commit()
    now = datetime.now(KST)
    processed_at = now.strftime("%Y-%m-%d %H:%M:%S (KST)")
    target = _output_path(run.id, now)
    tmp_path: Path | None = None
    moved = False  # 이번 실행이 만든 결과 파일만 지운다
    try:
        wb = Workbook()
        wb.remove(wb.active)
        parsed: dict[int, tuple[SourceFile, FileResult, dict[str, set[tuple[int, int]]]]] = {}
        sheets_out = []
        for o in plan.sheets:
            fid = o.file.file_id
            if fid not in parsed:
                record, result = load_parsed(db, fid)
                if record.sha256 != o.file.sha256:
                    raise BundleError("결과 구성 확인 뒤 파일 내용이 바뀌었습니다 — 다시 확인하세요.")
                parsed[fid] = (record, result, formula_cells(_file_path(record)))
            _, result, formulas = parsed[fid]
            sheets_out.append(_write_sheet(wb.create_sheet(o.name), o, result, formulas))
        summary_items = _write_summary(wb.create_sheet(SUMMARY_SHEET), results, processed_at) if plan.items else []
        with tempfile.NamedTemporaryFile(dir=OUTPUT_DIR, suffix=".xlsx.tmp", delete=False) as tmp:
            tmp_path = Path(tmp.name)
        wb.save(tmp_path)
        if target.exists():
            raise BundleError(f"같은 이름의 결과 파일이 이미 있습니다: {target.name}")
        os.replace(tmp_path, target)
        tmp_path, moved = None, True
        os.chmod(target, stat.S_IRUSR | stat.S_IRGRP | stat.S_IROTH)
        run.status = "done"
        run.output_path = target.relative_to(DATA_DIR).as_posix()
        run.output_sha256 = hashlib.sha256(target.read_bytes()).hexdigest()
        run.summary = json.dumps(
            {
                "mode": MODE,
                "mode_label": "파일별 시트",
                "stats_summary": plan.items is not None,
                "verified": False,  # 검토·통계 검증을 거치지 않은 묶음 (값만 옮김)
                "policy": POLICY,
                "processed_at": processed_at,
                "run_id": run.id,
                "basis": plan.basis,
                "sheet_order": [o.name for o in plan.sheets] + ([SUMMARY_SHEET] if plan.items else []),
                "sheets": sheets_out,
                "summary_items": summary_items,
                "files": [{"id": f["file_id"], "name": f["file_name"], "excluded": False} for f in plan.files],
                "warnings": plan.warnings,
                "row_count": run.row_count,
            },
            ensure_ascii=False,
            default=str,
        )
        run.finished_at = datetime.now(KST)
        db.commit()
    except Exception as exc:
        db.rollback()
        if tmp_path is not None:
            tmp_path.unlink(missing_ok=True)
        if moved and target.exists():
            os.chmod(target, stat.S_IWUSR | stat.S_IRUSR)
            target.unlink(missing_ok=True)
        failed = db.get(IntegrationRun, run.id)
        if failed is not None:
            failed.status = "failed"
            failed.error = str(exc)[:1000]
            failed.finished_at = datetime.now(KST)
            db.commit()
        logger.exception("bundle run %s failed", run.id)
        raise BundleError(f"결과 파일을 만들지 못했습니다: {exc}") from exc
    logger.info("bundle run=%s sheets=%d rows=%d summary=%s", run.id, len(plan.sheets), run.row_count, bool(plan.items))
    return run
