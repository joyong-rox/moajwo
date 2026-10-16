"""xlsx / xls / csv 파일을 시트별 격자로 읽고 헤더·행 종류를 판별한다.

원본 파일은 읽기만 한다. 이 모듈은 어떤 파일도 쓰지 않는다.
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime, time
from pathlib import Path
from typing import Any, Literal

import chardet
import openpyxl
import xlrd

from app.pii import PiiReport, mask_cell, scan_grid

SUPPORTED_EXTENSIONS = (".xlsx", ".xls", ".csv")

# 머리글 탐색 범위 (숫자 데이터가 없는 표는 위 FALLBACK_HEADER_ROWS행 규칙)
HEADER_SEARCH_ROWS = 100  # 제목·작성일·빈 행이 길게 이어지는 정부 양식도 머리글을 찾도록 (머리글 후보 수집 범위)
# 머리글 후보 뒤 첫 데이터 행을 찾는 범위 (100행째 머리글 + 101행 데이터 같은 경계를 위해 후보 범위보다 넓게).
# 글자만 있는 긴 목록을 끝까지 훑거나, 멀리 떨어진 숫자 한 줄 때문에 본문 행을 머리글로 고르지 않도록 제한한다.
DATA_SEARCH_ROWS = 200
FALLBACK_HEADER_ROWS = 3  # 숫자 데이터가 없는 표(글자 목록 등)는 예전처럼 위 3행에서 찾는다
PREVIEW_ROWS = 20
SAMPLES_PER_COLUMN = 5  # 매핑(특히 LLM)에 넘기는 열별 샘플 값 수
CHARDET_MIN_CONFIDENCE = 0.5
# chardet 판별 실패/저신뢰 시 시도 순서 (EUC-KR 우선, CP949는 EUC-KR 확장)
CSV_FALLBACK_ENCODINGS = ("euc-kr", "cp949", "utf-8")

RowKind = Literal["data", "empty", "subtotal", "note"]
Cell = str | int | float | bool | None


class ParseError(Exception):
    pass


@dataclass
class SheetResult:
    name: str
    header_row: int | None  # 1-based, 헤더를 못 찾으면 None
    columns: list[str]
    data_rows: int
    empty_rows: int
    subtotal_rows: int
    note_rows: int = 0  # 자료/주/※ 등 주석 행
    missing_cells: int = 0  # '-', 'x', '…' 등 결측 표기 셀 (데이터 행 기준)
    header_row_end: int | None = None  # 2단 헤더면 아래 헤더 행 (1-based), 아니면 header_row와 같음
    preview: list[dict[str, Any]] = field(default_factory=list)
    samples: list[list[Cell]] = field(default_factory=list)  # 열별 데이터 행 샘플 (중복·결측 제외)
    title_texts: list[str] = field(default_factory=list)  # 헤더 위 제목·단위 행 텍스트 (기간·단위 추정용)
    note_texts: list[str] = field(default_factory=list)  # 주석 행 텍스트 (단위 추정용)
    body: list[dict[str, Any]] = field(default_factory=list)  # 헤더 아래 모든 행 (검증용, API 응답에는 넣지 않음)
    pii: PiiReport = field(default_factory=PiiReport)
    header_cells: list[list[Cell]] = field(default_factory=list)  # 헤더 행 원래 셀 (2단 헤더 결합 표시용)
    corrections: list[dict[str, Any]] = field(default_factory=list)  # 파서가 바꾼 값 (열 단위 묶음, 검토 패널용)
    # 열별 원본 시트의 열 번호(0부터). 병합 복제 열을 버리면 원본 위치와 달라지므로
    # 미리보기의 열 위치(A, B…) 표시에 쓴다
    source_columns: list[int] = field(default_factory=list)


@dataclass
class FileResult:
    sheets: list[SheetResult]
    encoding: str | None = None  # csv만
    skipped_sheets: list[str] = field(default_factory=list)  # 건너뛴 설명 시트

    @property
    def pii(self) -> PiiReport:
        report = PiiReport()
        for s in self.sheets:
            report.merge(s.pii)
        return report

    @property
    def data_rows(self) -> int:
        return sum(s.data_rows for s in self.sheets)


# ---------------------------------------------------------------- 셀 정규화

_PLAIN_NUMBER = re.compile(r"^[+-]?(?:0|[1-9]\d*)(?:\.\d+)?$")
# 천 단위 구분이 올바른 경우만 숫자로 본다 ("010 1234 5678" 같은 전화번호 오변환 방지)
_GROUPED_NUMBER = re.compile(r"^[+-]?\d{1,3}(?:([, ])\d{3})(?:\1\d{3})*(?:\.\d+)?$")


def to_number(text: str) -> int | float | None:
    """'1,234', '1 234', ' 56.7 ' 같은 숫자 문자열을 숫자로. 숫자가 아니면 None."""
    s = text.strip().replace(" ", " ")
    if not s:
        return None
    if _PLAIN_NUMBER.match(s):
        pass
    elif _GROUPED_NUMBER.match(s):
        s = s.replace(",", "").replace(" ", "")
    else:
        return None
    if "." not in s:
        return int(s)  # 정수는 float를 거치지 않는다 (2^53을 넘는 자릿수가 바뀌지 않게)
    return float(s)


# 이름으로 보아 숫자가 아니라 '식별 기호'인 열 (코드·번호·구분·시각 등). '시각장애인…'은 시각(時刻)이 아니다.
_TEXT_COLUMN_HINT = re.compile(
    r"(코드|번호|구분|우편|전화|연락처|일련|시각(?!장애)|(?<![0-9A-Za-z])id(?![0-9A-Za-z]))", re.IGNORECASE
)
_DIGITS_ONLY = re.compile(r"^\d+$")
# 숫자로 정확히 다룰 수 있는 최대 자릿수 (Excel 숫자 셀은 15자리, 브라우저 Number는 2^53 미만).
# 상세 행에 이보다 긴 숫자 글자가 있고 이름이 금액·수량이 아닌 열은 식별자로 보고 글자로 둔다
# (끝자리만 다른 식별자가 같은 값이 되지 않게).
MAX_NUMBER_DIGITS = 15
# 금액·수량을 뜻하는 열 이름 — 긴 값이 있어도 식별자로 보지 않는다 (작은 값·소수·음수가 계산에서 빠지지 않게).
# '…액'은 '집행액·예산액·총액'처럼 이름 끝(또는 단위 괄호 앞)에 올 때만 본다.
_QUANTITY_COLUMN_HINT = re.compile(
    r"(금액|예산|집행|결산|지출|수입|비용|매출|실적|인구|세대수|건수|수량|대수|인원|면적|합계|소계|총계|[가-힣]액(?![가-힣]))"
)


def text_columns(columns: list[str], grid: list[list[Cell]], raw: list[list[Any]], body_start: int) -> set[int]:
    """숫자로 바꾸면 안 되는 열을 고른다 (열 전체를 같은 기준으로).

    - 이름에 코드·번호·시각 같은 표시가 있거나, 앞자리 0이 있는 숫자 글자('0900', '011')가 섞인 열
    - 상세 행(합계·소계 행 제외)에 숫자로만 된 글자가 MAX_NUMBER_DIGITS(15)자리보다 길게 있고, 이름이 금액·수량을
      뜻하지 않는 열 (숫자로 바꾸면 Excel·브라우저에서 끝자리가 바뀜). 합계 행의 큰 값은 집계값이라 근거로 보지 않는다.
    - 같은 열에서 '0900'은 글자, '1800'은 숫자가 되는 뒤섞임과 코드의 앞자리 0 손실을 막는다.
    - 천 단위 구분('1,000')이 쓰인 열은 수량으로 보고 숫자로 둔다. 엑셀에 숫자로 저장된 셀은 대상이 아니다.
    """
    keep: set[int] = set()
    for j, name in enumerate(columns):
        texts = [
            raw[r][j].strip() for r in range(body_start, len(grid)) if j < len(raw[r]) and isinstance(raw[r][j], str)
        ]
        converted = [s for s in texts if to_number(s) is not None]
        if not converted:
            continue  # 글자에서 숫자로 바뀌는 셀이 없는 열
        if any(_GROUPED_NUMBER.match(s) and not _PLAIN_NUMBER.match(s) for s in converted):
            continue
        leading_zero = any(len(s) > 1 and s.startswith("0") and _DIGITS_ONLY.match(s) for s in texts)
        detail = [
            raw[r][j].strip()
            for r in range(body_start, len(grid))
            if j < len(raw[r]) and isinstance(raw[r][j], str) and classify_row(grid[r]) != "subtotal"
        ]
        long_id = any(len(s) > MAX_NUMBER_DIGITS and _DIGITS_ONLY.match(s) for s in detail)
        too_long = long_id and not _QUANTITY_COLUMN_HINT.search(name or "")
        if leading_zero or too_long or _TEXT_COLUMN_HINT.search(name or ""):
            keep.add(j)
    return keep


def normalize_cell(value: Any) -> Cell:
    if value is None:
        return None
    if isinstance(value, bool | int | float):
        if isinstance(value, float) and value.is_integer():
            return int(value)
        return value
    if isinstance(value, datetime):
        return value.isoformat(sep=" ", timespec="seconds").removesuffix(" 00:00:00")
    if isinstance(value, date | time):
        return value.isoformat()
    text = str(value).strip()
    if not text:
        return None
    number = to_number(text)
    return number if number is not None else text


# ---------------------------------------------------------------- 파일 읽기


Filled = dict[str, set[tuple[int, int]]]  # 시트 → 병합 채움으로 값이 들어간 칸 (0-based 행, 열)


def _read_xlsx(path: Path) -> tuple[dict[str, list[list[Any]]], Filled]:
    # read_only 모드에서는 병합 셀 정보를 얻을 수 없어 일반 모드로 연다 (파일은 읽기만 함).
    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        sheets: dict[str, list[list[Any]]] = {}
        filled: Filled = {}
        for ws in wb.worksheets:
            grid = [list(row) for row in ws.iter_rows(values_only=True)]
            cells = filled.setdefault(ws.title, set())
            for rng in ws.merged_cells.ranges:
                top_left = grid[rng.min_row - 1][rng.min_col - 1] if rng.min_row - 1 < len(grid) else None
                _fill_range(grid, rng.min_row - 1, rng.max_row, rng.min_col - 1, rng.max_col, top_left, cells)
            sheets[ws.title] = grid
        return sheets, filled
    finally:
        wb.close()


def _read_xls(path: Path) -> tuple[dict[str, list[list[Any]]], Filled]:
    book = xlrd.open_workbook(path, formatting_info=True, on_demand=True)
    try:
        sheets: dict[str, list[list[Any]]] = {}
        filled: Filled = {}
        for sh in book.sheets():
            grid = [[_xls_value(book, sh.cell(r, c)) for c in range(sh.ncols)] for r in range(sh.nrows)]
            cells = filled.setdefault(sh.name, set())
            for rlo, rhi, clo, chi in sh.merged_cells:
                if rlo < len(grid):
                    value = grid[rlo][clo] if clo < len(grid[rlo]) else None
                    _fill_range(grid, rlo, rhi, clo, chi, value, cells)
            sheets[sh.name] = grid
        return sheets, filled
    finally:
        book.release_resources()


def _xls_value(book: xlrd.book.Book, cell: xlrd.sheet.Cell) -> Any:
    if cell.ctype == xlrd.XL_CELL_DATE:
        try:
            return xlrd.xldate.xldate_as_datetime(cell.value, book.datemode)
        except (ValueError, xlrd.xldate.XLDateError):
            return cell.value
    if cell.ctype == xlrd.XL_CELL_BOOLEAN:
        return bool(cell.value)
    if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR):
        return None
    return cell.value


def _fill_range(
    grid: list[list[Any]],
    r0: int,
    r1: int,
    c0: int,
    c1: int,
    value: Any,
    filled: set[tuple[int, int]] | None = None,
) -> None:
    """병합 범위 [r0, r1) x [c0, c1)를 상단 왼쪽 값으로 채운다 (ffill). 채운 칸은 filled에 기록한다."""
    for r in range(r0, min(r1, len(grid))):
        row = grid[r]
        if len(row) < c1:
            row.extend([None] * (c1 - len(row)))
        for c in range(c0, c1):
            row[c] = value
            if filled is not None and (r, c) != (r0, c0) and value is not None:
                filled.add((r, c))


def decode_csv(raw: bytes) -> tuple[str, str]:
    if raw.startswith(b"\xef\xbb\xbf"):
        return raw[3:].decode("utf-8"), "utf-8-sig"
    guess = chardet.detect(raw)
    encoding = guess.get("encoding")
    if encoding and (guess.get("confidence") or 0) >= CHARDET_MIN_CONFIDENCE:
        try:
            return raw.decode(encoding), encoding.lower()
        except (UnicodeDecodeError, LookupError):
            pass
    for candidate in CSV_FALLBACK_ENCODINGS:
        try:
            return raw.decode(candidate), candidate
        except UnicodeDecodeError:
            continue
    raise ParseError("CSV 인코딩을 판별할 수 없습니다.")


def _read_csv(path: Path, sheet_name: str) -> tuple[dict[str, list[list[Any]]], str]:
    text, encoding = decode_csv(path.read_bytes())
    # 구분자만 추측하고 따옴표 규칙은 표준(RFC 4180: " 로 감싸고 "" 로 이스케이프)으로 고정한다.
    # Sniffer의 따옴표 추측은 표본·Python 버전에 따라 달라져 쉼표가 든 값을 잘못 나눌 수 있다.
    try:
        delimiter = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|").delimiter
    except csv.Error:
        delimiter = ","
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter, quotechar='"', doublequote=True)
    rows = [list(r) for r in reader]
    return {sheet_name: rows}, encoding


# ---------------------------------------------------------------- 분석

# 합계 단어는 라벨 끝에 올 때만 본다 ('영업1팀 소계', '합계(전체)').
# '소계민원센터'처럼 이름 앞부분에 있는 것은 지명이다.
_SUBTOTAL_ANYWHERE = re.compile(r"(소계|합계|총계|누계|subtotal|grandtotal)(\([^)]*\))?$", re.IGNORECASE)
# '전국'은 시도·시군구 통계에서 전체 합계 행이다.
_SUBTOTAL_EXACT = re.compile(r"^(계|합|총|전국|total|sum)$", re.IGNORECASE)
_SUBTOTAL_MAX_LEN = 20  # 긴 문장 속 '합계' 단어는 소계 행으로 보지 않는다

# 정부 통계 하단 주석: "자료: …", "출처 : …", "주: 1) …", "주1) …", "※ …", "* …", "(단위: 명)"
_NOTE_PREFIX = re.compile(
    r"^(?:(?:자료|출처|주|비고|참고|note|notes|source)\s*\d*\s*[:：)]|[※*＊†‡]|\(?\s*단위\s*[:：])",
    re.IGNORECASE,
)

# 통계표 결측/비공개 표기. 원본 표기는 그대로 두고 표시만 한다 (0과 구분).
MISSING_MARKERS = frozenset({"-", "－", "–", "—", "x", "X", "…", "...", "·", "ㆍ", "..", "na", "n/a", "NA", "N/A"})

# 2단 헤더 아래 줄에 올 수 있는 숫자는 연도뿐이다 (그 외 숫자가 있으면 데이터 행으로 본다).
_YEAR_MIN, _YEAR_MAX = 1900, 2100


def is_missing(cell: Cell) -> bool:
    return isinstance(cell, str) and cell in MISSING_MARKERS


def is_subtotal_label(cell: Cell) -> bool:
    """'소계', '서울특별시 계'가 아닌 '계', '영업1팀 소계', '전국' 같은 합계 라벨 셀인지."""
    if not isinstance(cell, str):
        return False
    compact = re.sub(r"\s+", "", cell)
    if len(compact) > _SUBTOTAL_MAX_LEN:
        return False
    return bool(_SUBTOTAL_EXACT.match(compact) or _SUBTOTAL_ANYWHERE.search(compact))


def _is_subtotal(row: list[Cell]) -> bool:
    return any(is_subtotal_label(cell) for cell in row)


def _is_note(row: list[Cell]) -> bool:
    """값이 한 종류뿐(병합 ffill 포함)이고 주석 접두어로 시작하는 행."""
    values = {c for c in row if c is not None}
    if len(values) != 1:
        return False
    (value,) = values
    return isinstance(value, str) and bool(_NOTE_PREFIX.match(value.strip()))


def classify_row(row: list[Cell]) -> RowKind:
    if all(c is None for c in row):
        return "empty"
    if _is_note(row):
        return "note"
    if _is_subtotal(row):
        return "subtotal"
    return "data"


def _header_texts(row: list[Cell]) -> set[str]:
    return {c for c in row if isinstance(c, str) and not is_missing(c) and not _is_note([c])}


# '검색기간 : 202608 ~ 202608', '단위 : 대' 같은 '항목 : 값' 설명 칸
_META_CELL = re.compile(r"^[^:：]{1,20}[:：]\s*\S")


def _has_data_number(row: list[Cell]) -> bool:
    """연도가 아닌 숫자가 있는 행 = 데이터 행 (머리글의 '2025' 같은 연도는 제외)."""
    return any(
        isinstance(c, int | float)
        and not isinstance(c, bool)
        and not (isinstance(c, int) and _YEAR_MIN <= c <= _YEAR_MAX)
        for c in row
    )


def _is_meta_row(row: list[Cell]) -> bool:
    texts = [c for c in row if isinstance(c, str)]
    return bool(texts) and sum(1 for t in texts if _META_CELL.match(t)) * 2 >= len(set(texts))


def _is_header_candidate(row: list[Cell]) -> bool:
    """머리글이 될 수 있는 행: 서로 다른 글자 칸 2개 이상, 숫자 데이터 없음, '항목 : 값' 설명 행 아님.

    머리글의 '총계'·'계' 같은 글자는 합계 행 판정과 무관하게 열 이름으로 본다.
    """
    return len(_header_texts(row)) >= 2 and not _has_data_number(row) and not _is_meta_row(row)


def detect_header_row(rows: list[list[Cell]]) -> int | None:
    """머리글 행 (0-based).

    1) 처음 나오는 데이터 행(연도가 아닌 숫자가 있는 행)을 찾고, 그 위의 머리글 후보 중
       '서로 다른 글자 칸'이 가장 많은 행을 고른다 (동점이면 위쪽). 제목·작성일·빈 행·설명 행은 건너뛴다.
       병합된 제목 행은 ffill 후 같은 값이 반복되므로 고유 값이 1개뿐이라 후보가 되지 않는다.
    2) 숫자 데이터가 없는 표는 예전처럼 위 3행 중 글자 칸이 가장 많은 행.
    """
    candidates: list[int] = []
    for i, row in enumerate(rows[:DATA_SEARCH_ROWS]):
        if _has_data_number(row) and candidates:
            return max(candidates, key=lambda k: (len(_header_texts(rows[k])), -k))
        if i < HEADER_SEARCH_ROWS:
            if _is_header_candidate(row):
                candidates.append(i)
        elif not candidates:
            break  # 탐색 범위(100행) 안에 머리글 후보가 없으면 대체 규칙으로
        # 후보가 있으면 범위 밖까지 첫 데이터 행을 찾아 후보를 평가한다 (100행째 머리글 + 101행 데이터)
    best, best_score = None, 0
    for i, row in enumerate(rows[:FALLBACK_HEADER_ROWS]):
        score = len(_header_texts(row))
        if score > best_score:
            best, best_score = i, score
    return best


def _runs(row: list[Cell]) -> list[tuple[int, int]]:
    """같은 값이 가로로 2칸 이상 이어진 구간 [start, end). 가로 병합 셀의 ffill 흔적이다."""
    runs, i = [], 0
    while i < len(row):
        j = i + 1
        while j < len(row) and row[i] is not None and row[j] == row[i]:
            j += 1
        if j - i >= 2:
            runs.append((i, j))
        i = j
    return runs


def _is_header_like(cell: Cell) -> bool:
    if cell is None:
        return True
    if isinstance(cell, str):
        return not is_missing(cell)
    return isinstance(cell, int) and not isinstance(cell, bool) and _YEAR_MIN <= cell <= _YEAR_MAX


def _is_two_level(top: list[Cell], sub: list[Cell]) -> bool:
    """top이 가로 병합된 그룹 헤더이고 sub가 그 아래를 나눠 쓰는 하위 헤더인지.

    - top에 가로 반복 구간이 있고, 전체 폭을 덮는 한 덩어리(제목 행)는 아니어야 한다.
    - 각 구간 아래 sub 셀은 모두 채워져 있고 서로 달라야 한다.
    - sub의 모든 셀은 텍스트 또는 연도여야 한다 (금액 등 숫자가 있으면 데이터 행).
    """
    runs = _runs(top)
    if not runs or len({c for c in top if c is not None}) < 2:
        return False
    # 하위 헤더에는 '계'·'합계' 열이 흔하므로 소계 여부는 보지 않는다 (실제 소계 행은 숫자가 있어 걸러진다).
    if not all(_is_header_like(c) for c in sub) or _is_note(sub):
        return False
    for start, end in runs:
        cells = sub[start:end]
        # 하위 칸도 가로 병합(같은 값 반복)일 수 있다 — 구간을 둘 이상으로 나누기만 하면 된다
        if any(c is None for c in cells) or len(set(cells)) < 2:
            return False
    return True


def _is_vertical_split(top: list[Cell], sub: list[Cell], after: list[Cell] | None) -> bool:
    """세로로 나뉜 머리글: 위·아래 두 줄에 열 이름이 흩어져 있고 바로 아래부터 데이터.

    예) 5행 ['', '차종별', '총계', '승용차'] / 6행 ['구별', '', '총계', '승용차'] / 7행 데이터.
    아래 줄의 각 칸은 위 칸과 같거나 위 칸이 비어 있어야 하고, 위가 빈 칸에 이름이 하나 이상 있어야 한다.
    """
    if after is None or not _has_data_number(after) or not _is_header_candidate(sub):
        return False
    pairs = [(t, b) for t, b in zip(top, sub, strict=False) if b is not None]
    return bool(pairs) and all(t is None or t == b for t, b in pairs) and any(t is None for t, _ in pairs)


MAX_HEADER_ROWS = 3  # 머리글은 최대 3줄 (예: KOSIS '2024' / '총계·승용차' / '계·관용')


def detect_header_block(rows: list[list[Cell]]) -> tuple[int, int] | None:
    """헤더 구간 (top, bottom) 0-based. 1단이면 top == bottom.

    머리글 행(detect_header_row)에서 아래·위로, 가로 병합(그룹 머리글) 또는 세로 분할(위·아래 줄에 흩어진
    열 이름)로 이어지는 줄을 최대 3줄까지 묶는다. 그 위의 제목·작성일·빈 행은 머리글에 넣지 않는다.
    """
    h = detect_header_row(rows)
    if h is None:
        return None

    def connected(upper: int, lower: int) -> bool:
        after = rows[lower + 1] if lower + 1 < len(rows) else None
        if _is_two_level(rows[upper], rows[lower]):
            return True
        return bool(_header_texts(rows[upper])) and _is_vertical_split(rows[upper], rows[lower], after)

    top = bottom = h
    while bottom + 1 < len(rows) and bottom - top + 1 < MAX_HEADER_ROWS and connected(bottom, bottom + 1):
        bottom += 1
    while top > 0 and bottom - top + 1 < MAX_HEADER_ROWS and connected(top - 1, top):
        top -= 1
    return top, bottom


def _trim(grid: list[list[Cell]]) -> list[list[Cell]]:
    """뒤쪽 빈 행과 오른쪽 빈 열을 잘라내고 모든 행 길이를 맞춘다."""
    while grid and all(c is None for c in grid[-1]):
        grid.pop()
    width = max((max((i + 1 for i, c in enumerate(r) if c is not None), default=0) for r in grid), default=0)
    return [(r + [None] * width)[:width] for r in grid]


_SPACED_SYLLABLES = re.compile(r"^[가-힣](?:\s+[가-힣])+$")


def _tidy_header(value: Cell) -> str | None:
    """정렬용 공백을 지운다: '총   계' → '총계', '승   합' → '승합' (한 글자씩 띄어 쓴 경우만)."""
    if value is None:
        return None
    text = str(value).strip()
    return re.sub(r"\s+", "", text) if _SPACED_SYLLABLES.match(text) else text


def _header_names(header_rows: list[list[Cell]], width: int) -> list[str]:
    """헤더 행(1~3줄)을 위→아래로 합친 열의 머리글 의미 (같은 이름 구분 번호를 붙이기 전)."""
    names: list[str] = []
    for i in range(width):
        parts: list[str] = []
        for row in header_rows:
            value = _tidy_header(row[i])
            if value is not None and (not parts or value != parts[-1]):
                parts.append(value)
        names.append("_".join(parts) or f"열{i + 1}")
    return names


def _column_names(header_rows: list[list[Cell]], width: int) -> list[str]:
    """헤더 행(1~2줄)을 위→아래로 합쳐 열 이름을 만든다. 예: '인구' + '2025' → '인구_2025'."""
    names: list[str] = []
    seen: dict[str, int] = {}
    for name in _header_names(header_rows, width):
        if name in seen:
            seen[name] += 1
            name = f"{name}_{seen[name]}"
        else:
            seen[name] = 1
        names.append(name)
    return names


def merged_duplicate_columns(
    grid: list[list[Cell]], header: tuple[int, int] | None, columns: list[str], filled: set[tuple[int, int]]
) -> list[int]:
    """머리글이 가로 병합돼 생긴 복제 열 (버릴 열 번호).

    가로 병합은 범위 전체를 왼쪽 위 값으로 채우므로 '자가용', '자가용_2', '자가용_3'처럼 같은 이름의 열이 생긴다.
    머리글 전체 의미(구분 번호를 붙이기 전 이름)가 앞 열과 같고 머리글 칸이 병합으로 채워졌으며, 본문 값도 모든
    행에서 구간 안의 앞 열과 같으면 복제로 본다. '인구_2024'·'인구_2025'처럼 머리글 의미가 다른 열은 값이 같아도
    남기고, 값이 하나라도 다른 열(예: '시군구' 병합 아래 '천안시 | 동남구')도 실제 데이터라 남긴다.
    """
    if header is None:
        return []
    top, bottom = header
    body = grid[bottom + 1 :]
    base = _header_names(grid[top : bottom + 1], len(columns))
    drop: list[int] = []
    kept_in_run: list[int] = []
    for j in range(len(columns)):
        merged = j > 0 and any(
            (r, j) in filled and grid[r][j] is not None and grid[r][j] == grid[r][j - 1] for r in range(top, bottom + 1)
        )
        if not merged or base[j] != base[j - 1]:
            kept_in_run = [j]
            continue
        if any(all(row[j] == row[k] for row in body) for k in kept_in_run):
            drop.append(j)
        else:
            kept_in_run.append(j)
    return drop


def _corrections(
    raw: list[list[Any]], grid: list[list[Cell]], body: list[dict[str, Any]], filled: set[tuple[int, int]]
) -> list[dict[str, Any]]:
    """본문 행에서 파서가 바꾼 값을 열 단위로 묶는다: 숫자 문자열 → 숫자, 병합 셀 채움. (원본은 그대로)"""
    found: dict[tuple[str, int], dict[str, Any]] = {}
    for entry in body:
        if entry["kind"] not in ("data", "subtotal"):
            continue
        r = entry["row"] - 1
        for j, value in enumerate(grid[r]):
            original = raw[r][j] if j < len(raw[r]) else None
            if (r, j) in filled:
                kind, before = "ffill", None
            elif isinstance(original, str) and isinstance(value, int | float) and not isinstance(value, bool):
                kind, before = "number", mask_cell(original.strip())
            elif isinstance(original, str) and isinstance(value, str) and original != value:
                kind, before = "trim", mask_cell(original)
            elif isinstance(original, str) and original and not original.strip() and value is None:
                kind, before = "trim", original  # 공백만 있던 칸 → 빈 칸
            else:
                continue
            item = found.setdefault(
                (kind, j),
                {"kind": kind, "column": j, "count": 0, "rows": [], "before": before, "after": mask_cell(value)},
            )
            item["count"] += 1
            if len(item["rows"]) < 5:
                item["rows"].append(entry["row"])
    return sorted(found.values(), key=lambda c: (c["column"], c["kind"]))


def analyze_sheet(name: str, raw: list[list[Any]], filled: set[tuple[int, int]] | None = None) -> SheetResult:
    grid = _trim([[normalize_cell(v) for v in row] for row in raw])
    width = len(grid[0]) if grid else 0
    block = detect_header_block(grid)
    columns = _column_names(grid[block[0] : block[1] + 1] if block else [], width)
    drop = set(merged_duplicate_columns(grid, block, columns, filled or set()))
    source_columns = list(range(width))
    if drop:  # 머리글 가로 병합으로 생긴 복제 열은 버린다 (남는 열 이름은 그대로 — 저장된 매핑 확정과 맞춤)
        keep = [j for j in range(width) if j not in drop]
        grid = [[row[j] for j in keep] for row in grid]
        raw = [[row[j] if j < len(row) else None for j in keep] for row in raw]
        position = {j: i for i, j in enumerate(keep)}
        filled = {(r, position[c]) for r, c in (filled or set()) if c in position}
        columns = [columns[j] for j in keep]
        source_columns = keep
        width = len(keep)
    header_rows = grid[block[0] : block[1] + 1] if block else []
    body_start = block[1] + 1 if block else 0
    # 코드·번호·시각 열은 파일에 적힌 글자 그대로 둔다 (열 전체를 같은 기준으로)
    for j in text_columns(columns, grid, raw, body_start):
        for r in range(body_start, len(grid)):
            source = raw[r][j] if j < len(raw[r]) else None
            if isinstance(source, str) and isinstance(grid[r][j], int | float):
                grid[r][j] = source.strip()

    counts = {"data": 0, "empty": 0, "subtotal": 0, "note": 0}
    missing_total = 0
    preview: list[dict[str, Any]] = []
    body: list[dict[str, Any]] = []
    note_texts: list[str] = []
    samples: list[list[Cell]] = [[] for _ in range(width)]
    title_texts = [c for row in grid[: block[0] if block else 0] for c in dict.fromkeys(row) if isinstance(c, str)]
    for offset, row in enumerate(grid[body_start:]):
        kind = classify_row(row)
        counts[kind] += 1
        if kind == "data":
            for j, cell in enumerate(row):
                bucket = samples[j]
                # 샘플은 화면·LLM으로 나가므로 개인정보는 마스킹한 값으로 담는다.
                shown = mask_cell(cell)
                if (
                    len(bucket) < SAMPLES_PER_COLUMN
                    and cell is not None
                    and not is_missing(cell)
                    and shown not in bucket
                ):
                    bucket.append(shown)
        missing = [j for j, c in enumerate(row) if is_missing(c)] if kind == "data" else []
        missing_total += len(missing)
        r = body_start + offset
        source = raw[r] if r < len(raw) else []
        # 파일에 적힌 그대로의 값 (병합 채움 전 = 빈 칸). 통합 결과에서 바뀐 셀을 빠짐없이 찾는 데 쓴다.
        original = [
            None if (r, j) in (filled or ()) else (source[j] if j < len(source) else None) for j in range(width)
        ]
        entry = {"row": r + 1, "kind": kind, "cells": row, "missing": missing, "original": original}
        body.append(entry)  # 검증·통합용 값 (응답으로 내보내지 않음)
        if len(preview) < PREVIEW_ROWS:
            masked = [mask_cell(c) for c in row]
            preview.append({"row": entry["row"], "kind": kind, "cells": masked, "missing": missing})
        if kind == "note":
            note_texts.extend(c for c in dict.fromkeys(row) if isinstance(c, str))

    pii = PiiReport()
    scan_grid(name, grid, pii)
    corrections = _corrections(raw, grid, body, filled or set())

    return SheetResult(
        name=name,
        header_row=block[0] + 1 if block else None,
        header_row_end=block[1] + 1 if block else None,
        columns=columns,
        data_rows=counts["data"],
        empty_rows=counts["empty"],
        subtotal_rows=counts["subtotal"],
        note_rows=counts["note"],
        missing_cells=missing_total,
        preview=preview,
        samples=samples,
        title_texts=title_texts,
        note_texts=note_texts,
        body=body,
        pii=pii,
        header_cells=[[mask_cell(c) for c in row] for row in header_rows],
        corrections=corrections,
        source_columns=source_columns,
    )


# KOSIS 등 내려받기 파일에 딸린 통계표 설명 시트 ('메타정보', '설명', '안내')
_INFO_SHEET = re.compile(r"(메타|설명|안내)")


def parse_file(path: Path, original_name: str | None = None) -> FileResult:
    """original_name: 사용자가 올린 파일명. CSV는 시트가 없어 이 이름(확장자 제외)을 시트명으로 쓴다."""
    ext = path.suffix.lower()
    encoding = None
    filled: Filled = {}
    try:
        if ext == ".xlsx":
            sheets, filled = _read_xlsx(path)
        elif ext == ".xls":
            sheets, filled = _read_xls(path)
        elif ext == ".csv":
            sheets, encoding = _read_csv(path, Path(original_name or path.name).stem)
        else:
            raise ParseError(f"지원하지 않는 형식입니다: {ext}")
    except ParseError:
        raise
    except Exception as exc:  # 손상/암호화 파일 등 라이브러리별 예외를 사용자 메시지로 통일
        raise ParseError(f"파일을 읽을 수 없습니다 ({type(exc).__name__}).") from exc
    names = list(sheets)
    skipped = [n for n in names if _INFO_SHEET.search(n)]
    if len(skipped) == len(names):  # 시트가 설명 시트뿐이면 그대로 읽는다
        skipped = []
    sheet_results = [analyze_sheet(n, g, filled.get(n)) for n, g in sheets.items() if n not in skipped]
    return FileResult(sheets=sheet_results, encoding=encoding, skipped_sheets=skipped)
