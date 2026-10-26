"""공통 key 연결 (취합 방식 '공통 key 연결', MODE "keyjoin"): 여러 파일의 같은 사용자를 key(사번·ID 등)로 연결해
사용자별 한 행을 만들고, 입력마다 고른 열만 결과에 넣는다. 하나의 통합표(integration.py)·파일별 시트(bundle.py)와
따로 동작하며 검토·열 매핑·학습을 요구하지도 바꾸지도 않는다.

- key는 파서가 바꾸기 전 값(body[*]["original"])으로 만든 정규형 문자열의 **튜플**을 정확히 비교한다. 글자 key는 앞자리
  0·긴 자릿수·대소문자·공백을 그대로 둔다 (공백 제거·대소문자 통일·유사 일치 없음). 정수 칸 123은 글자 '123'과 연결되고
  '00123'과는 다르다. 참거짓·날짜·소수·16자리 이상 숫자 칸·병합으로 채운 칸·빈 칸·결측 표기는 key로 쓰지 않고 안내한다.
  다른 기능의 숫자 해석(cells·text_columns)은 건드리지 않는다.
- 한 입력(파일·시트) 안에서 key가 겹치면 만들지 않는다 (합치거나 고르지 않음). 입력 사이의 같은 key는 정상 연결이다.
- 범위: 기준 입력의 사용자만(base, 기본) · 모든 입력의 사용자(all) · 모든 입력에 있는 사용자(common).
- 결과: 시트 「연결 결과」(key + 고른 열, 미연결은 빈 칸 — 0 아님), 「출처」(결과 행마다 입력별 원본 행), 「연결 요약」.
  선택하지 않은 열은 어디에도 넣지 않는다. key 칸은 항상 글자 칸이다.
- 응답·실행 기록·서버 기록·오류 문구에 key 원문·가린 key·해시를 넣지 않는다 — 위치(입력 번호·파일 번호·시트 순번·
  열 문자·행 번호)·개수·코드만. 가린 뒤 같아지는 key는 합치지 않고 미리보기에 '(가림 구분 n)'을 붙인다.
- plan은 읽기 전용(파일·DB·학습·확정을 쓰지 않음). run은 같은 함수로 다시 판정해 basis가 같을 때만 만들고, 거절이면
  실행 기록·결과·임시 파일을 남기지 않는다. 결과 파일을 확정하는 순서·오류 처리는 app/runlease.py.
"""

from __future__ import annotations

import hashlib
import json
import logging
import math
from dataclasses import dataclass, field
from datetime import date, datetime, time
from pathlib import Path
from typing import Annotated, Any, Literal

from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, ConfigDict, Field
from sqlalchemy.orm import Session

from app import runlease
from app.analysis import MASKED_NO, excel_number
from app.auth import CurrentUser
from app.bundle import (
    EXCEL_DIGITS,
    EXCEL_MAX_ROWS,
    GAP_FILL,
    HEAD_FILL,
    FileInfo,
    SheetInfo,
    _bad,
    _file_path,
    _gap,
    _long,
    _put,
    file_info,
    formula_cells,
)
from app.config import MAX_OUTPUT_CELLS, MAX_OUTPUT_TEXT_CHARS
from app.failures import failure_reason, log_failure
from app.files import get_file_for, load_parsed
from app.integration import KST, OUTPUT_DIR
from app.models import IntegrationRun, SourceFile
from app.name_limits import cell_place, sheet_place
from app.parsing import MISSING_MARKERS, SheetResult
from app.pii import mask_cell, mask_rrn_number, mask_text
from app.validation_api import pii_out
from app.xmltext import has_illegal, printable

logger = logging.getLogger(__name__)

POLICY = "keyjoin-policy-1"  # key 정규형·범위·출력 규칙이 바뀌면 올린다 (basis·실행 기록에 남음)
MODE = "keyjoin"
MODE_LABEL = "공통 key 연결"
KEYJOIN_FAILED = "KEYJOIN_FAILED"  # 결과를 만드는 중 예기치 않은 실패 (예외 종류만 알림)
MAX_INPUTS = 20  # 한 번에 연결하는 입력(파일·시트) 수
MAX_KEY_COLUMNS = 5  # 복합 key의 열 수
MAX_PICK_COLUMNS = 50  # 입력 하나에서 가져오는 열 수
MAX_RESULT_COLUMNS = 200  # 결과 열 수 (key 열 + 가져온 열)
MAX_ROWS = 300_000  # 모든 입력의 데이터 행 합계 (묶기·데이터 분석 한도와 같음)
PREVIEW_ROWS = 20  # 미리보기 행
PREVIEW_POOL = 200  # 미리보기 값을 두는 key 수 (기준 입력 앞 200개, all은 입력마다 앞 200개 더)
PROBLEM_ROWS = 5  # 문제 문구에 적는 행 번호 수
OUT_NAME_MAX = 255  # 결과 열 이름 길이
INDEX_WEIGHT = 1.0  # key 인덱스 칸 하나를 결과 칸 몇 개로 셈하는지 (측정 뒤 조정 — docs/MEMORY.md)
TEXT_WEIGHT = 2  # 글자 추정치에 곱하는 안전 계수 (측정 뒤 완화 — docs/MEMORY.md)
REQUEST_LOOKBACK = 50  # 같은 요청 번호(request_id)를 찾는 최근 실행 수
SUMMARY_LINES = 40  # 연결 요약 시트의 고정 줄 수 추정 (칸 예산)
SUMMARY_PER_INPUT = 16  # 연결 요약 시트에서 입력 하나가 차지하는 칸 수 추정
RESULT_SHEET, SOURCE_SHEET, SUMMARY_SHEET = "연결 결과", "출처", "연결 요약"
Scope = Literal["base", "all", "common"]
SCOPE_LABEL = {"base": "기준 파일의 사용자만", "all": "모든 파일의 사용자", "common": "모든 파일의 공통 사용자"}
# key 칸의 종류 (key_part): 받는 것은 text·int, 나머지는 차단 사유
KINDS = ("text", "int", "float", "bool", "date", "empty", "filled", "missing_mark", "precision", "illegal")
KEY_EMPTY, KEY_MISSING, KEY_PARTIAL = "KEYJOIN_KEY_EMPTY", "KEYJOIN_KEY_MISSING_MARK", "KEYJOIN_KEY_PARTIAL"
KEY_FILLED, KEY_TYPE, KEY_PRECISION = "KEYJOIN_KEY_FILLED", "KEYJOIN_KEY_TYPE", "KEYJOIN_KEY_PRECISION"
KEY_ILLEGAL, KEY_DUPLICATE, VALUE_ILLEGAL = "KEYJOIN_KEY_ILLEGAL", "KEYJOIN_KEY_DUPLICATE", "KEYJOIN_VALUE_ILLEGAL"
KIND_CODE = {
    "float": KEY_TYPE,
    "bool": KEY_TYPE,
    "date": KEY_TYPE,
    "empty": KEY_EMPTY,
    "filled": KEY_FILLED,
    "missing_mark": KEY_MISSING,
    "precision": KEY_PRECISION,
    "illegal": KEY_ILLEGAL,
}
# 복합 key에서 여러 사유가 겹치면 앞의 것으로 알린다 (빈 값은 다른 부분이 있으면 '일부 누락')
CODE_PRIORITY = (KEY_ILLEGAL, KEY_TYPE, KEY_PRECISION, KEY_FILLED, KEY_MISSING, KEY_EMPTY)
REJECT_TEXT = {
    KEY_EMPTY: "key가 빈 데이터 행이 {n}행({rows}) 있습니다 — 빈 값끼리 연결하지 않습니다. "
    "원본에서 채우거나 그 행을 빼세요.",
    KEY_MISSING: "key가 결측 표기('-'·'x'·'N/A' 등)인 데이터 행이 {n}행({rows}) 있습니다 — 식별자로 쓰지 않습니다.",
    KEY_PARTIAL: "복합 key의 일부가 빈 데이터 행이 {n}행({rows}) 있습니다 — 일부만으로 연결하지 않습니다.",
    KEY_FILLED: "key 칸이 병합으로 채워진 데이터 행이 {n}행({rows}) 있습니다 — 위 행의 값을 복사해 연결하지 않습니다. "
    "원본에서 행마다 key를 적으세요.",
    KEY_TYPE: "참거짓·날짜·소수 값인 key 칸이 {n}행({rows}) 있습니다 — 자동으로 해석하지 않습니다. "
    "문자형 식별자로 저장해 다시 올리세요.",
    KEY_PRECISION: "16자리 이상 숫자 칸인 key가 {n}행({rows}) 있습니다 — 자릿수 보존을 확인할 수 없어 연결하지 "
    "않습니다. 텍스트로 저장한 자료를 올리세요 (표시 형식의 0은 만들지 않음).",
    KEY_ILLEGAL: "엑셀에 쓸 수 없는 글자(제어 문자 등)가 든 key가 {n}행({rows}) 있습니다 — 원본에서 지운 뒤 "
    "다시 올리세요.",
}
PREVIEW_NOTE = (
    f"미리보기는 앞 {PREVIEW_POOL}개 key 가운데 {PREVIEW_ROWS}행입니다 (결측·미연결은 빈 칸). "
    f"가린 표시가 같은 서로 다른 key에는 '({MASKED_NO} n)'을 붙여 보입니다."
)


class KeyJoinError(Exception):
    """만들 수 없음 (409) — run_id가 있으면 실패로 기록된 실행 번호."""

    run_id: int | None = None


def _failed(message: str, run_id: int) -> KeyJoinError:
    err = KeyJoinError(message)
    err.run_id = run_id
    return err


class _Gap:
    """결과값이 저장되지 않은 수식 칸 — 결과에는 빈 칸 + 메모 (0으로 채우지 않음)."""


GAP = _Gap()


# ---------------------------------------------------------------- 요청


class ColumnIn(BaseModel):
    model_config = ConfigDict(extra="forbid")

    column_id: str = Field(pattern=r"^c\d{1,4}$")
    out_name: str | None = None  # None = 원본 열 이름 (서버 기본)


class InputIn(BaseModel):
    model_config = ConfigDict(extra="forbid")

    file_id: int
    sheet_id: str = Field(pattern=r"^s\d{1,3}$")
    key_columns: list[Annotated[str, Field(pattern=r"^c\d{1,4}$")]] = Field(min_length=1, max_length=MAX_KEY_COLUMNS)
    columns: list[ColumnIn] = Field(default_factory=list, max_length=MAX_PICK_COLUMNS)


class BaseRef(BaseModel):
    model_config = ConfigDict(extra="forbid")

    file_id: int
    sheet_id: str = Field(pattern=r"^s\d{1,3}$")


class KeyJoinIn(BaseModel):
    model_config = ConfigDict(extra="forbid")

    inputs: list[InputIn] = Field(max_length=MAX_INPUTS)
    scope: Scope = "base"
    base: BaseRef  # 출력 순서와 base 범위의 사용자 집합을 정하는 입력 (모든 범위에서 필수)
    key_names: list[str] | None = None  # None = 기준 입력의 key 열 이름


class KeyJoinRunIn(KeyJoinIn):
    expected_basis: str = Field(min_length=16, max_length=64)
    request_id: str = Field(pattern=r"^[0-9a-f]{32}$")  # 같은 요청의 재전송을 서버가 알아보는 번호 (화면이 만든다)


NAME_TEXT = {
    "KEYJOIN_NAME_ILLEGAL": "에 쓸 수 없는 글자(NUL·제어 문자 등)가 있습니다",
    "KEYJOIN_NAME_BLANK": "이 비어 있습니다 (공백만)",
    "KEYJOIN_NAME_TOO_LONG": f"이 {OUT_NAME_MAX}자를 넘습니다",
}


def _name_code(name: str) -> str | None:
    """사용자가 적은 결과 이름의 형식 오류 코드 (없으면 None). 원문은 문구에 넣지 않는다."""
    if has_illegal(name):  # NUL 포함 — SQLite·PostgreSQL에서 같게, 저장 전에
        return "KEYJOIN_NAME_ILLEGAL"
    if not name.strip():
        return "KEYJOIN_NAME_BLANK"
    if len(name) > OUT_NAME_MAX:
        return "KEYJOIN_NAME_TOO_LONG"
    return None


def check_request(req: KeyJoinIn) -> None:
    """요청 형식 검사 (422) — 파일을 읽기 전에 모두 본다. 문구에는 위치·개수·코드만."""
    if len(req.inputs) < 2:
        raise _bad(f"연결할 입력이 2개 이상이어야 합니다 (지금 {len(req.inputs)}개). [KEYJOIN_INPUTS_TOO_FEW]")
    seen: set[tuple[int, str]] = set()
    k = len(req.inputs[0].key_columns)
    total = k
    for n, ip in enumerate(req.inputs, start=1):
        pair = (ip.file_id, ip.sheet_id)
        if pair in seen:
            text = f"입력 {n}: 같은 파일·시트를 두 번 골랐습니다 (#{ip.file_id} {ip.sheet_id})."
            raise _bad(f"{text} [KEYJOIN_INPUT_DUPLICATE]")
        seen.add(pair)
        if len(ip.key_columns) != k:
            raise _bad(
                f"입력 {n}: key 열이 {len(ip.key_columns)}개로 입력 1의 {k}개와 다릅니다 — 복합 key는 모든 입력이 "
                "같은 개수여야 합니다. [KEYJOIN_KEY_COUNT_MISMATCH]"
            )
        if len(set(ip.key_columns)) != len(ip.key_columns):
            raise _bad(f"입력 {n}: 같은 열을 key로 두 번 골랐습니다. [KEYJOIN_COLUMN_DUPLICATE]")
        ids = [c.column_id for c in ip.columns]
        if len(set(ids)) != len(ids):
            raise _bad(f"입력 {n}: 같은 열을 두 번 골랐습니다. [KEYJOIN_COLUMN_DUPLICATE]")
        if set(ids) & set(ip.key_columns):
            raise _bad(f"입력 {n}: key 열은 결과에 이미 들어갑니다 — 가져올 열에서 빼세요. [KEYJOIN_KEY_IN_COLUMNS]")
        for c in ip.columns:
            if c.out_name is not None and (code := _name_code(c.out_name)):
                raise _bad(f"입력 {n} {c.column_id}: 결과 열 이름{NAME_TEXT[code]}. [{code}]")
        total += len(ids)
    if (req.base.file_id, req.base.sheet_id) not in seen:
        raise _bad(f"기준 입력(#{req.base.file_id} {req.base.sheet_id})이 입력 목록에 없습니다. [KEYJOIN_BASE_INVALID]")
    if req.key_names is not None:
        if len(req.key_names) != k:
            raise _bad(f"key 결과 이름이 {len(req.key_names)}개로 key 열 {k}개와 다릅니다. [KEYJOIN_KEY_NAMES_COUNT]")
        for i, name in enumerate(req.key_names, start=1):
            if code := _name_code(name):
                raise _bad(f"key 결과 이름 {i}: 이름{NAME_TEXT[code]}. [{code}]")
    if total > MAX_RESULT_COLUMNS:
        raise _bad(
            f"결과 열이 {total}개(key {k} + 가져올 열 {total - k})로 한 번에 만들 수 있는 {MAX_RESULT_COLUMNS}개를 "
            "넘습니다 — 열을 줄이거나 나눠 만드세요. [KEYJOIN_TOO_MANY_COLUMNS]"
        )


# ---------------------------------------------------------------- 문제·위치


@dataclass
class Problem:
    code: str
    level: str  # block / warn
    text: str  # 원문·가린 이름·해시 없이 — 위치·개수·코드만
    input: int | None = None
    file_id: int | None = None
    sheet_id: str | None = None
    column: str | None = None  # 원본 열 문자
    rows: list[int] | None = None  # 처음 PROBLEM_ROWS개
    count: int | None = None
    groups: int | None = None

    def out(self) -> dict[str, Any]:
        return {
            "code": self.code,
            "level": self.level,
            "input": self.input,
            "file_id": self.file_id,
            "sheet_id": self.sheet_id,
            "column": self.column,
            "rows": self.rows,
            "count": self.count,
            "groups": self.groups,
            "text": self.text,
        }


def _block(code: str, text: str, **where: Any) -> Problem:
    return Problem(code, "block", f"{text} [{code}]", **where)


def _warn(code: str, text: str, **where: Any) -> Problem:
    return Problem(code, "warn", f"{text} [{code}]", **where)


@dataclass
class _Place:
    """name_limits의 위치 문구(sheet_place·cell_place)에 넘기는 시트 정보 (SheetInfo는 열 이름 대신 개수를 둔다)."""

    name: str
    tab: int
    columns: list[str]
    source_columns: list[int]
    header_row: int | None = None
    header_row_end: int | None = None


class _Rows:
    """행 번호 모음 — 처음 PROBLEM_ROWS개와 개수, 처음 걸린 열."""

    __slots__ = ("first", "count", "column")

    def __init__(self) -> None:
        self.first: list[int] = []
        self.count = 0
        self.column: int | None = None

    def add(self, row: int, column: int | None = None) -> None:
        self.count += 1
        if len(self.first) < PROBLEM_ROWS:
            self.first.append(row)
        if self.column is None:
            self.column = column


def _rows_text(rows: list[int]) -> str:
    return "·".join(f"{r}행" for r in rows)


# ---------------------------------------------------------------- key 정규형


def key_part(original: Any, cell: Any) -> tuple[str | None, str]:
    """key 칸 하나의 (정규형, 종류). 정규형이 None이면 그 종류(KINDS)가 차단 사유다.

    - 글자: 그대로 (strip·대소문자·NFKC 없음). 빈 글자·공백만 → empty, 결측 표기('-'·'N/A' 등) → missing_mark,
      엑셀에 쓸 수 없는 글자 → illegal
    - 정수 칸: 10진 문자열 (123 → '123', -5 → '-5'). 16자리 이상은 자릿수 보존을 확인할 수 없어 precision
    - 실수 칸(xls 숫자 칸은 항상 실수): 정수 값이면 정수와 같은 경로, 소수·NaN·무한은 float(차단), 2^53 이상은 precision
    - 참거짓·날짜·시각: 자동 해석하지 않음 (bool·date). 병합으로 채운 칸(original None·cell 있음)은 filled
    """
    if original is None:
        return (None, "filled") if cell is not None else (None, "empty")
    if isinstance(original, bool):
        return None, "bool"
    if isinstance(original, int):
        if len(str(abs(original))) > EXCEL_DIGITS:
            return None, "precision"
        return str(original), "int"
    if isinstance(original, float):
        if not math.isfinite(original) or not original.is_integer():
            return None, "float"
        if abs(original) >= 2**53 or len(str(abs(int(original)))) > EXCEL_DIGITS:
            return None, "precision"
        return str(int(original)), "int"
    if isinstance(original, datetime | date | time):
        return None, "date"
    if isinstance(original, str):
        if has_illegal(original):
            return None, "illegal"
        stripped = original.strip()
        if stripped == "":
            return None, "empty"
        if stripped in MISSING_MARKERS:
            return None, "missing_mark"
        return original, "text"
    return None, "bool"  # 그 밖의 자료형 (xlrd 오류 값 등) — 식별자로 보지 않는다


def _row_code(codes: list[tuple[str, int]], parts: int) -> tuple[str, int]:
    """복합 key 한 행의 차단 사유와 그 열 (codes: 거절된 부분의 (코드, 열 번호))."""
    for code in CODE_PRIORITY[:-1]:
        for c, j in codes:
            if c == code:
                return code, j
    j = codes[0][1]
    return (KEY_EMPTY if len(codes) == parts else KEY_PARTIAL), j


# ---------------------------------------------------------------- 입력 인덱스

Key = tuple[str, ...]


@dataclass
class KeyIndex:
    rows: dict[Key, int] = field(default_factory=dict)  # key → 원본 행 번호 (처음 나온 행)
    order: list[Key] = field(default_factory=list)  # 행 순서의 key
    values: dict[Key, list[Any]] = field(default_factory=dict)  # key → 가져올 열 값 (둔 key만)
    dups: dict[Key, list[int]] = field(default_factory=dict)  # 같은 key의 행들 (처음 행 포함)
    rejects: dict[str, _Rows] = field(default_factory=dict)  # 차단 코드 → 행
    value_illegal: _Rows = field(default_factory=_Rows)
    value_illegal_cells: list[tuple[int, int]] = field(default_factory=list)  # (행, 열) 처음 몇 개
    kinds: list[dict[str, int]] = field(default_factory=list)  # key 열별 칸 종류 수
    int_keys: set[Key] = field(default_factory=set)  # 정수 칸에서 온 key (글자 key와의 연결 수 셈용)
    skipped: dict[str, _Rows] = field(default_factory=dict)  # subtotal · note
    empty_rows: int = 0
    data_rows: int = 0
    long_numbers: int = 0
    formula_missing: int = 0
    rrn: bool = False  # 주민등록번호가 탐지된 시트 (미리보기 가림 규칙)

    @property
    def unsupported(self) -> int:
        return sum(r.count for c, r in self.rejects.items() if c not in (KEY_EMPTY, KEY_PARTIAL))

    def whitespace_variants(self, ws_map: dict[Key, set[Key]]) -> int:
        """이 입력의 key가 든, 앞뒤 공백만 다른 key 묶음 수."""
        return sum(
            1 for s, v in ws_map.items() if len(v) + (s in self.rows) >= 2 and (s in self.rows or v & self.rows.keys())
        )


def analyze_input(
    sheet: SheetResult,
    key_ids: list[int],
    pick_ids: list[int],
    formulas: set[tuple[int, int]],
    ws_map: dict[Key, set[Key]],
    keep_all: bool = False,
    keep_keys: set[Key] | None = None,
    keep_first: int = 0,
) -> KeyIndex:
    """데이터 행만 key 인덱스로 만든다 (소계·주석은 세기만, 위 행 값을 복사하지 않음). 같은 key가 다시 나오면 인덱스는
    바꾸지 않고 dups에 모은다. 값은 keep_all이거나 keep_keys에 있거나 앞 keep_first개 key일 때만 둔다 (미리보기 풀)."""
    ix = KeyIndex(kinds=[dict.fromkeys(KINDS, 0) for _ in key_ids], rrn=bool(sheet.pii.counts.get("rrn")))
    source = sheet.source_columns or list(range(len(sheet.columns)))
    for entry in sheet.body:
        kind = entry["kind"]
        if kind != "data":
            if kind == "empty":
                ix.empty_rows += 1
            else:
                ix.skipped.setdefault(kind, _Rows()).add(entry["row"])
            continue
        ix.data_rows += 1
        original, cells, row = entry["original"], entry["cells"], entry["row"]
        parts: list[str] = []
        codes: list[tuple[str, int]] = []
        numeric = False
        for pos, j in enumerate(key_ids):
            value, k = key_part(original[j] if j < len(original) else None, cells[j] if j < len(cells) else None)
            ix.kinds[pos][k] += 1
            if value is None:
                codes.append((KIND_CODE[k], j))
            else:
                parts.append(value)
                numeric = numeric or k == "int"
        if codes:
            code, j = _row_code(codes, len(key_ids))
            ix.rejects.setdefault(code, _Rows()).add(row, j)
            continue
        key = tuple(parts)
        if key in ix.rows:
            ix.dups.setdefault(key, [ix.rows[key]]).append(row)
            continue
        ix.rows[key] = row
        ix.order.append(key)
        if numeric:
            ix.int_keys.add(key)
        stripped = tuple(p.strip() for p in key)
        if stripped != key:
            ws_map.setdefault(stripped, set()).add(key)
        values = _row_values(entry, pick_ids, source, formulas, ix)
        if keep_all or len(ix.order) <= keep_first or (keep_keys is not None and key in keep_keys):
            ix.values[key] = values
    return ix


def _row_values(
    entry: dict[str, Any], pick_ids: list[int], source: list[int], formulas: set[tuple[int, int]], ix: KeyIndex
) -> list[Any]:
    """한 행의 가져올 열 값 (결과에 쓰는 규칙): 날짜는 원본 형식, 결과값 없는 수식 칸은 GAP, 16자리 이상 정수는 글자.
    쓸 수 없는 글자는 위치만 세어 차단한다 (값을 바꾸거나 자르지 않음)."""
    cells, originals, row = entry["cells"], entry["original"], entry["row"]
    out: list[Any] = []
    for j in pick_ids:
        value = cells[j] if j < len(cells) else None
        original = originals[j] if j < len(originals) else None
        if isinstance(original, datetime | date):
            value = original  # 날짜는 원본 형식 그대로 (파서가 글자로 바꾼 값 대신)
        if formulas and _gap(entry, j, source, formulas) == "missing":
            ix.formula_missing += 1
            out.append(GAP)
            continue
        if has_illegal(value):
            ix.value_illegal.add(row, j)
            if len(ix.value_illegal_cells) < PROBLEM_ROWS:
                ix.value_illegal_cells.append((row, j))
        if _long(value):
            ix.long_numbers += 1
            value = excel_number(value)
        out.append(value)
    return out


# ---------------------------------------------------------------- 결과 구성 (plan)


@dataclass
class InputPlan:
    index: int
    req: InputIn
    record: SourceFile
    info: FileInfo
    sheet: SheetInfo | None
    own: bool
    pii: str  # none / acknowledged / needs_ack
    pii_summary: dict[str, Any] | None = None  # 파일 목록과 같은 형식 (가린 값)
    key_ids: list[int] = field(default_factory=list)
    pick_ids: list[int] = field(default_factory=list)
    out_names: list[str] = field(default_factory=list)  # 가져올 열의 최종 결과 이름
    default_name: list[bool] = field(default_factory=list)
    problems: list[Problem] = field(default_factory=list)
    ix: KeyIndex | None = None
    linked: int = 0
    unlinked: int = 0
    ws_variants: int = 0

    @property
    def file_id(self) -> int:
        return self.req.file_id

    @property
    def sheet_index(self) -> int:
        return int(self.req.sheet_id[1:])

    @property
    def place(self) -> _Place:
        s = self.sheet
        assert s is not None
        return _Place(s.name, s.tab, s.column_names, s.source_columns)

    def where(self, j: int | None = None, row: int | None = None) -> str:
        """'입력 2 (#13 · 1번째 시트 · A열)' — 원문·가린 이름 없이 위치만."""
        text = f"입력 {self.index + 1} (#{self.file_id}"
        if self.sheet is not None:
            text += f" · {sheet_place(self.place)}"
            if j is not None or row is not None:
                text += f" · {cell_place(self.place, row, j)}"
        return text + ")"

    def letter(self, j: int) -> str:
        return cell_place(self.place, None, j).removesuffix("열")

    @property
    def label(self) -> str:
        text = f"입력 {self.index + 1} · #{self.file_id} {mask_text(self.info.name)}"
        if self.sheet is not None:
            text += f" / {mask_text(self.sheet.name)} ({sheet_place(self.place)})"
        return text

    def add(self, problem: Problem) -> Problem:
        problem.input, problem.file_id, problem.sheet_id = self.index, self.file_id, self.req.sheet_id
        self.problems.append(problem)
        return problem

    def out(self) -> dict[str, Any]:
        s, ix = self.sheet, self.ix
        key_columns = []
        for pos, j in enumerate(self.key_ids):
            key_columns.append(
                {
                    "column_id": self.req.key_columns[pos],
                    "name": mask_text(s.column_names[j]) if s is not None else None,
                    "letter": self.letter(j) if s is not None else None,
                    "kinds": dict(ix.kinds[pos]) if ix is not None else None,
                }
            )
        picked = []
        for pos, j in enumerate(self.pick_ids):
            picked.append(
                {
                    "column_id": self.req.columns[pos].column_id,
                    "name": mask_text(s.column_names[j]) if s is not None else None,
                    "letter": self.letter(j) if s is not None else None,
                    "out_name": mask_text(self.out_names[pos]) if self.out_names else None,
                    "default_name": self.default_name[pos] if self.default_name else None,
                }
            )
        keys = None
        if ix is not None:
            keys = {
                "indexed": len(ix.rows),
                "duplicate_keys": len(ix.dups),
                "duplicate_rows": sum(len(v) for v in ix.dups.values()),
                "empty_rows": ix.rejects[KEY_EMPTY].count if KEY_EMPTY in ix.rejects else 0,
                "partial_rows": ix.rejects[KEY_PARTIAL].count if KEY_PARTIAL in ix.rejects else 0,
                "unsupported_rows": ix.unsupported,
                "linked": self.linked,
                "unlinked": self.unlinked,
                "whitespace_variants": self.ws_variants,
            }
        return {
            "index": self.index,
            "file_id": self.file_id,
            "file_name": mask_text(self.info.name),
            "sheet_id": self.req.sheet_id,
            "sheet": mask_text(s.name) if s is not None else None,
            "tab": s.tab if s is not None else None,
            "label": self.label,
            "own": self.own,
            "pii": self.pii,
            "pii_summary": self.pii_summary,
            "key_columns": key_columns,
            "picked": picked,
            "rows": dict(s.kinds) if s is not None else None,
            "keys": keys,
            "problems": [p.out() for p in self.problems],
        }


@dataclass
class Plan:
    req: KeyJoinIn
    user_id: int
    inputs: list[InputPlan]
    base: InputPlan
    key_names: list[str]
    header: list[str]  # key 이름 + 가져올 열의 결과 이름 (원문 — 응답·기록에는 가려서)
    problems: list[Problem]  # 차단·경고 모두 (level로 구분)
    budget: dict[str, int] = field(default_factory=dict)
    result: dict[str, int] = field(default_factory=dict)
    preview: dict[str, Any] = field(default_factory=dict)
    analyzed: bool = False
    out_keys: list[Key] | None = None  # 결과 행 순서의 key (run의 keep_values=True에서 값과 함께)
    basis: str = ""

    @property
    def blocks(self) -> list[Problem]:
        return [p for p in self.problems if p.level == "block"]

    @property
    def warnings(self) -> list[Problem]:
        return [p for p in self.problems if p.level == "warn"]

    @property
    def ready(self) -> bool:
        return self.analyzed and not self.blocks

    @property
    def file_ids(self) -> list[int]:
        return list(dict.fromkeys(ip.file_id for ip in self.inputs))

    def out(self) -> dict[str, Any]:
        return {
            "mode": MODE,
            "policy": POLICY,
            "scope": self.req.scope,
            "base": {"file_id": self.base.file_id, "sheet_id": self.base.req.sheet_id, "index": self.base.index},
            "inputs": [ip.out() for ip in self.inputs],
            "header": [mask_text(h) for h in self.header],
            "result": self.result,
            "preview": self.preview,
            "problems": [p.out() for p in self.blocks],
            "warnings": [p.out() for p in self.warnings],
            "limits": {
                "inputs": MAX_INPUTS,
                "key_columns": MAX_KEY_COLUMNS,
                "columns": MAX_PICK_COLUMNS,
                "result_columns": MAX_RESULT_COLUMNS,
                "rows": MAX_ROWS,
                "cells": MAX_OUTPUT_CELLS,
                "text_chars": MAX_OUTPUT_TEXT_CHARS,
                "preview_rows": PREVIEW_ROWS,
                "excel_rows": EXCEL_MAX_ROWS,
            },
            "budget": self.budget,
            "basis": self.basis,
            "ready": self.ready,
        }


def _pii_state(record: SourceFile) -> str:
    summary = json.loads(record.pii_summary) if record.pii_summary else {}
    has_pii = any(summary.get("counts", {}).values())
    return "none" if not has_pii else ("acknowledged" if record.pii_acknowledged_at is not None else "needs_ack")


def _cheap_checks(db: Session, user: CurrentUser, req: KeyJoinIn) -> tuple[list[InputPlan], list[Problem]]:
    """파싱 없이(file_info — 내용 해시로 기억한 시트별 개수) 권한·개인정보·시트·열을 본다. 없는 시트·열은 422."""
    inputs: list[InputPlan] = []
    problems: list[Problem] = []
    records: dict[int, SourceFile] = {}
    infos: dict[int, FileInfo] = {}
    for n, item in enumerate(req.inputs):
        fid = item.file_id
        if fid not in records:
            records[fid] = get_file_for(db, user, fid)  # 볼 수 없는 파일은 없는 것처럼 404
            infos[fid] = file_info(db, records[fid])
        record, info = records[fid], infos[fid]
        own = record.uploaded_by == user.id or (record.uploaded_by is None and user.is_manager)
        found = pii_out(record)
        ip = InputPlan(n, item, record, info, None, own, _pii_state(record), found.model_dump() if found else None)
        inputs.append(ip)
        if not own:
            text = f"{ip.where()}: 다른 사용자의 파일이라 연결할 수 없습니다 (본인 파일만 — 남의 파일은 조회만)."
            problems.append(ip.add(_block("KEYJOIN_FILE_NOT_OWN", text)))
        if ip.pii == "needs_ack":
            text = f"{ip.where()}: 개인정보 의심 값이 있어 포함 확인이 필요합니다 — 포함을 확인하거나 이 파일을 빼세요."
            problems.append(ip.add(_block("KEYJOIN_PII_ACK_REQUIRED", text)))
        if not info.sheets:
            problems.append(ip.add(_block("KEYJOIN_FILE_UNREADABLE", f"{ip.where()}: 앱이 읽은 시트가 없습니다.")))
            continue
        if ip.sheet_index >= len(info.sheets):
            raise _bad(f"입력 {n + 1}: #{fid} 파일에 없는 시트입니다: {item.sheet_id} [KEYJOIN_SHEET_UNKNOWN]")
        ip.sheet = info.sheets[ip.sheet_index]
        width = len(ip.sheet.column_names)
        for cid in [*item.key_columns, *(c.column_id for c in item.columns)]:
            if int(cid[1:]) >= width:
                raise _bad(f"{ip.where()}: 없는 열입니다: {cid} [KEYJOIN_COLUMN_UNKNOWN]")
        ip.key_ids = [int(cid[1:]) for cid in item.key_columns]
        ip.pick_ids = [int(c.column_id[1:]) for c in item.columns]
        used = [*ip.key_ids, *ip.pick_ids]
        bad = [j for j in used if has_illegal(ip.sheet.column_names[j])]
        if bad or has_illegal(info.name) or has_illegal(ip.sheet.name):
            what = "열 이름" if bad else "파일·시트 이름"
            text = (
                f"{ip.where(bad[0]) if bad else ip.where()}: {what}에 엑셀·기록에 쓸 수 없는 글자(NUL·제어 문자 등)가 "
                "있습니다 — 원본에서 고쳐 다시 올리세요."
            )
            column = ip.letter(bad[0]) if bad else None
            problems.append(ip.add(_block("KEYJOIN_NAME_ILLEGAL_SOURCE", text, column=column)))
    return inputs, problems


def _header(req: KeyJoinIn, inputs: list[InputPlan], base: InputPlan, problems: list[Problem]) -> list[str]:
    """결과 열 이름: key 이름(요청 또는 기준 입력의 key 열 이름) + 입력 순 → columns 순의 결과 이름.
    사용자가 적은 이름끼리(또는 key 이름과) 같으면 차단, 서버 기본 이름이 이미 쓴 이름과 같으면 '(입력 n · #fid)'를
    붙인다 (사용자가 적은 이름은 바꾸지 않는다)."""
    if req.key_names is not None:
        key_names = list(req.key_names)
    elif base.sheet is not None:
        key_names = [base.sheet.column_names[j] for j in base.key_ids]
    else:
        key_names = list(base.req.key_columns)
    chosen: dict[str, str] = {}  # 사용자가 적은 이름·key 이름 → 위치 문구
    for i, name in enumerate(key_names):
        if name in chosen:
            text = f"key 결과 이름 {i + 1}이 다른 key 결과 이름과 같습니다 — 다른 이름을 적으세요."
            problems.append(_block("KEYJOIN_OUT_NAME_DUPLICATE", text))
        chosen.setdefault(name, f"key 결과 이름 {i + 1}")
    for ip in inputs:
        for pos, c in enumerate(ip.req.columns):
            if c.out_name is None:
                continue
            where = ip.where(ip.pick_ids[pos] if ip.pick_ids else None)
            if c.out_name in chosen:
                text = f"{where}: 적은 결과 이름이 {chosen[c.out_name]}과(와) 같습니다 — 다른 이름을 적으세요."
                problems.append(ip.add(_block("KEYJOIN_OUT_NAME_DUPLICATE", text)))
            chosen.setdefault(c.out_name, f"{where}의 결과 이름")
    used = set(chosen)
    header = list(key_names)
    for ip in inputs:
        ip.out_names, ip.default_name = [], []
        for pos, c in enumerate(ip.req.columns):
            if c.out_name is not None:
                name = c.out_name
            else:
                name = ip.sheet.column_names[ip.pick_ids[pos]] if ip.sheet is not None else c.column_id
                if name in used:  # 같은 이름의 열은 출처를 붙여 구분한다 (뒤 것에)
                    name = f"{name} (입력 {ip.index + 1} · #{ip.file_id})"
                    k = 2
                    while name in used:
                        name = f"{name} ({k})"
                        k += 1
                used.add(name)
            ip.out_names.append(name)
            ip.default_name.append(c.out_name is None)
            header.append(name)
    return header


def _budget_checks(
    req: KeyJoinIn, inputs: list[InputPlan], base: InputPlan, header: list[str]
) -> tuple[list[Problem], dict[str, int]]:
    """파싱 전 예산 (docs/MEMORY.md '공통 key 연결'): 행 합계, 엑셀 행, 결과 칸 + 인덱스 칸, 글자 추정."""
    problems: list[Problem] = []
    ready = [ip for ip in inputs if ip.sheet is not None]
    n, k = len(inputs), len(base.req.key_columns)
    rows = {ip.index: ip.sheet.kinds["data"] for ip in ready}
    rows_total = sum(rows.values())
    drive_rows = rows_total if req.scope == "all" else rows.get(base.index, 0)
    picked = sum(len(ip.pick_ids) for ip in inputs)
    indexed = [ip for ip in ready if req.scope == "all" or ip is not base]
    index_cells = sum(rows[ip.index] * (k + len(ip.pick_ids)) for ip in indexed)
    out_cells = drive_rows * (k + picked + 1 + n) + SUMMARY_LINES + SUMMARY_PER_INPUT * n
    weighted = out_cells + math.ceil(INDEX_WEIGHT * index_cells)
    text_est = sum(len(h) for h in header)
    for ip in ready:
        chars = ip.sheet.column_chars
        text_est += sum(chars[j] for j in [*ip.key_ids, *ip.pick_ids] if j < len(chars))
    budget = {
        "output_cells": out_cells,
        "index_cells": index_cells,
        "text_chars": text_est,
        "cells_weighted": weighted,
        "text_weighted": TEXT_WEIGHT * text_est,
    }
    if base.sheet is not None and rows.get(base.index, 0) == 0:
        problems.append(base.add(_block("KEYJOIN_NO_ROWS", f"{base.where()}: 기준 입력에 데이터 행이 없습니다.")))
    if rows_total > MAX_ROWS:
        text = (
            f"연결할 데이터 행이 {rows_total:,}행으로 한 번에 연결할 수 있는 {MAX_ROWS:,}행을 넘습니다 — "
            "파일·열을 나눠 만드세요."
        )
        problems.append(_block("KEYJOIN_TOO_MANY_ROWS", text))
    if drive_rows + 1 > EXCEL_MAX_ROWS:
        text = f"결과 {drive_rows:,}행은 엑셀 한 시트에 넣을 수 있는 행 수({EXCEL_MAX_ROWS:,}, 머리글 포함)를 넘습니다."
        problems.append(_block("KEYJOIN_EXCEL_ROWS", text))
    if weighted > MAX_OUTPUT_CELLS:
        text = (
            f"결과에 넣을 칸이 {out_cells:,}개(출처 시트 포함)이고 key 인덱스 {index_cells:,}칸을 더한 "
            f"{weighted:,}칸이 이 서버에서 한 번에 만들 수 있는 {MAX_OUTPUT_CELLS:,}칸을 넘습니다 — 파일·열을 나눠 "
            "만드세요."
        )
        problems.append(_block("KEYJOIN_TOO_MANY_CELLS", text))
    if TEXT_WEIGHT * text_est > MAX_OUTPUT_TEXT_CHARS:
        text = (
            f"결과에 넣을 글자가 약 {text_est:,}자(여유 {TEXT_WEIGHT}배 {TEXT_WEIGHT * text_est:,}자)로 이 서버에서 "
            f"한 번에 만들 수 있는 {MAX_OUTPUT_TEXT_CHARS:,}자를 넘습니다 — 긴 글자 열을 줄이거나 파일·열을 나눠 "
            "만드세요."
        )
        problems.append(_block("KEYJOIN_TOO_MANY_CHARS", text))
    return problems, budget


def _file_warnings(inputs: list[InputPlan]) -> list[Problem]:
    warnings: list[Problem] = []
    if any(ip.pii == "acknowledged" for ip in inputs):
        text = (
            "개인정보 포함을 확인한 파일이 있습니다 — 결과 파일에는 화면과 달리 원본 값이 가리지 않고 들어갑니다 "
            "(결과는 이 실행 기록을 볼 수 있는 사람만 내려받을 수 있음)."
        )
        warnings.append(_warn("KEYJOIN_PII_IN_RESULT", text))
    sha: dict[str, list[int]] = {}
    for fid, info in {ip.file_id: ip.info for ip in inputs}.items():
        sha.setdefault(info.sha256, []).append(fid)
    for same in sha.values():
        if len(same) > 1:
            shown = ", ".join(f"#{f}" for f in same)
            text = f"같은 내용의 파일을 함께 골랐습니다 ({shown}) — 같은 key가 서로 연결됩니다."
            warnings.append(_warn("KEYJOIN_SAME_CONTENT", text))
    return warnings


def build_plan(db: Session, user: CurrentUser, req: KeyJoinIn, keep_values: bool = False) -> Plan:
    """요청한 구성으로 입력·결과 열·한도·key 문제·미리보기를 판정한다 (파일·DB를 쓰지 않음). 실행도 같은 함수로 다시
    판정한다 (keep_values=True — 결과를 만들 값까지 둔다). 싼 검사(권한·개인정보·시트·열·예산·이름)에서 하나라도
    차단되면 파일을 읽지 않는다 (keys: null)."""
    check_request(req)
    inputs, problems = _cheap_checks(db, user, req)
    base = next(ip for ip in inputs if (ip.file_id, ip.req.sheet_id) == (req.base.file_id, req.base.sheet_id))
    header = _header(req, inputs, base, problems)
    budget_problems, budget = _budget_checks(req, inputs, base, header)
    problems.extend(budget_problems)
    k = len(base.req.key_columns)
    plan = Plan(req, user.id, inputs, base, header[:k], header, problems + _file_warnings(inputs), budget)
    plan.result = {"rows": 0, "key_columns": k, "columns": len(header), "union": 0, "intersection": 0, "excluded": 0}
    plan.preview = {"rows": [], "shown": 0, "total": 0, "pool": PREVIEW_POOL, "note": PREVIEW_NOTE}
    if not plan.blocks:
        _analyze(db, plan, keep_values)
    plan.basis = _basis(plan)
    return plan


def _analyze(db: Session, plan: Plan, keep_values: bool) -> None:
    """기준 입력부터 파일 단위로 하나씩 읽어 인덱스를 만들고(읽은 값은 다음 파일 전에 놓는다), 범위에 따라 결과 행을
    정한다."""
    req, inputs, base = plan.req, plan.inputs, plan.base
    ws_map: dict[Key, set[Key]] = {}
    pool: set[Key] = set()
    keep_first = PREVIEW_POOL if req.scope == "all" else 0
    for fid in [base.file_id, *(f for f in plan.file_ids if f != base.file_id)]:
        record, result = load_parsed(db, fid)
        formulas = formula_cells(_file_path(record))
        for ip in sorted((x for x in inputs if x.file_id == fid), key=lambda x: (x is not base, x.index)):
            sheet = result.sheets[ip.sheet_index]
            first = PREVIEW_POOL if ip is base else keep_first
            ip.ix = analyze_input(sheet, ip.key_ids, ip.pick_ids, formulas.get(sheet.name, set()), ws_map,
                                  keep_all=keep_values, keep_keys=pool, keep_first=first)  # fmt: skip
            if not keep_values:
                pool.update(ip.ix.order[:first])
        del result
    plan.analyzed = True
    others = [ip for ip in inputs if ip is not base]
    bix = base.ix
    assert bix is not None
    # 교집합: 가장 작은 입력의 key에서 줄여 간다 (입력 행을 전수 대조하지 않음)
    smallest = min(inputs, key=lambda ip: len(ip.ix.rows))
    cand = list(smallest.ix.order)
    for ip in inputs:
        if ip is not smallest:
            cand = [k for k in cand if k in ip.ix.rows]
    inter = set(cand)
    extra: set[Key] = set()  # 기준 입력에 없는 key (입력 순·행 순서로 처음 나온 차례)
    extra_order: list[Key] = []
    for ip in others:
        for k in ip.ix.order:
            if k not in bix.rows and k not in extra:
                extra.add(k)
                extra_order.append(k)
    union = len(bix.rows) + len(extra)
    if req.scope == "base":
        out_keys = list(bix.order)
    elif req.scope == "common":
        out_keys = [k for k in bix.order if k in inter]
    else:
        out_keys = list(bix.order) + extra_order
    for ip in inputs:
        if req.scope == "all":
            ip.linked = len(ip.ix.rows)
        elif req.scope == "common":
            ip.linked = len(out_keys)
        else:
            ip.linked = len(bix.rows) if ip is base else sum(1 for k in ip.ix.rows if k in bix.rows)
        ip.unlinked = len(out_keys) - ip.linked
        ip.ws_variants = ip.ix.whitespace_variants(ws_map)
    excluded = union - len(inter) if req.scope == "common" else 0
    plan.result = {
        "rows": len(out_keys),
        "key_columns": len(plan.key_names),
        "columns": len(plan.header),
        "union": union,
        "intersection": len(inter),
        "excluded": excluded,
    }
    plan.out_keys = out_keys
    _input_problems(plan, ws_map)
    _link_problems(plan, excluded)
    # 인덱스 뒤 정확한 칸 수 (줄어들 수만 있다 — 글자는 추정치 그대로)
    n, k = len(inputs), len(plan.key_names)
    picked = sum(len(ip.pick_ids) for ip in inputs)
    indexed = inputs if req.scope == "all" else others
    index_cells = sum(len(ip.ix.rows) * (k + len(ip.pick_ids)) for ip in indexed)
    out_cells = len(out_keys) * (k + picked + 1 + n) + SUMMARY_LINES + SUMMARY_PER_INPUT * n
    weighted = out_cells + math.ceil(INDEX_WEIGHT * index_cells)
    plan.budget.update(output_cells=out_cells, index_cells=index_cells, cells_weighted=weighted)
    plan.preview = _preview(plan, _masked_numbers(plan, extra))


def _input_problems(plan: Plan, ws_map: dict[Key, set[Key]]) -> None:
    """입력마다 key 문제(차단)와 연결하지 않은 행·섞인 자료형·미연결(경고)."""
    for ip in plan.inputs:
        ix = ip.ix
        assert ix is not None
        if ix.dups:
            rows = sorted(r for v in ix.dups.values() for r in v)[:PROBLEM_ROWS]
            total = sum(len(v) for v in ix.dups.values())
            j = ip.key_ids[0]
            text = (
                f"{ip.where(j)}: key가 같은 데이터 행이 {total}행({len(ix.dups)}묶음, {_rows_text(rows)}) 있습니다 — "
                "중복을 합치거나 고르지 않습니다. 원본에서 정리하거나 다른 key 열을 고르세요."
            )
            problem = _block(KEY_DUPLICATE, text, column=ip.letter(j), rows=rows, count=total, groups=len(ix.dups))
            plan.problems.append(ip.add(problem))
        for code in (*CODE_PRIORITY, KEY_PARTIAL):
            r = ix.rejects.get(code)
            if r is None:
                continue
            j = r.column if r.column is not None else ip.key_ids[0]
            text = f"{ip.where(j)}: " + REJECT_TEXT[code].format(n=r.count, rows=_rows_text(r.first))
            problem = _block(code, text, column=ip.letter(j), rows=list(r.first), count=r.count)
            plan.problems.append(ip.add(problem))
        if ix.value_illegal.count:
            cells = " · ".join(cell_place(ip.place, r, j) for r, j in ix.value_illegal_cells)
            j = ix.value_illegal.column if ix.value_illegal.column is not None else ip.pick_ids[0]
            text = (
                f"{ip.where()}: 엑셀에 쓸 수 없는 글자(제어 문자 등)가 든 값이 {ix.value_illegal.count}칸({cells}) "
                "있습니다 — 원본에서 지운 뒤 다시 올리세요."
            )
            problem = _block(VALUE_ILLEGAL, text, column=ip.letter(j), rows=list(ix.value_illegal.first),
                             count=ix.value_illegal.count)  # fmt: skip
            plan.problems.append(ip.add(problem))
        if ix.skipped:
            parts = []
            for kind, label in (("subtotal", "소계·합계"), ("note", "주석")):
                r = ix.skipped.get(kind)
                if r is not None:
                    parts.append(f"{label} 행 {r.count}행({_rows_text(r.first)})")
            rows = sorted(x for r in ix.skipped.values() for x in r.first)[:PROBLEM_ROWS]
            text = (
                f"{ip.where()}: {' · '.join(parts)}은 데이터 행이 아니라 연결하지 않습니다 — '계'·'합계'처럼 보이는 "
                "값의 행도 여기에 들어가니 행 번호를 확인하세요."
            )
            count = sum(r.count for r in ix.skipped.values())
            plan.problems.append(ip.add(_warn("KEYJOIN_NONDATA_ROWS_SKIPPED", text, rows=rows, count=count)))
        for pos, kinds in enumerate(ix.kinds):
            if kinds["int"] and kinds["text"]:
                j = ip.key_ids[pos]
                text = (
                    f"{ip.where(j)}: key 열에 정수 칸 {kinds['int']}개와 글자 칸 {kinds['text']}개가 섞여 있습니다 — "
                    "정수 123은 글자 '123'과 같은 key로, '00123'과는 다른 key로 봅니다."
                )
                plan.problems.append(ip.add(_warn("KEYJOIN_MIXED_KEY_KINDS", text, column=ip.letter(j))))
        if ip.unlinked:
            text = (
                f"{ip.where()}: 결과 {plan.result['rows']:,}행 중 {ip.unlinked:,}행은 이 입력에 같은 key가 없어 "
                "빈 칸으로 둡니다 (0으로 채우지 않음)."
            )
            plan.problems.append(ip.add(_warn("KEYJOIN_UNLINKED", text, count=ip.unlinked)))
    variants = sum(1 for s, v in ws_map.items() if len(v) + any(s in ip.ix.rows for ip in plan.inputs) >= 2)
    if variants:
        text = (
            f"앞뒤 공백만 다른 key가 {variants}묶음 있습니다 — 서로 다른 사용자로 둡니다 (공백을 지워 연결하지 않음)."
        )
        plan.problems.append(_warn("KEYJOIN_WHITESPACE_VARIANTS", text, count=variants))


def _link_problems(plan: Plan, excluded: int) -> None:
    inputs = plan.inputs
    if plan.req.scope == "common" and not plan.out_keys:
        text = (
            "모든 파일에 함께 있는 사용자가 없습니다 (교집합 0) — 만들 결과가 없어 생성하지 않습니다. "
            "범위를 바꾸거나 key 열을 확인하세요."
        )
        plan.problems.append(_block("KEYJOIN_NO_COMMON", text))
    if excluded:
        text = f"모든 파일의 공통 사용자만 넣습니다 — 어느 한 입력에만 있는 사용자 {excluded:,}명은 결과에서 빠집니다."
        plan.problems.append(_warn("KEYJOIN_EXCLUDED", text, count=excluded))
    int_union: set[Key] = set()
    for ip in inputs:
        int_union |= ip.ix.int_keys
    if int_union:
        linked = {k for ip in inputs for k in ip.ix.rows if k in int_union and k not in ip.ix.int_keys}
        if linked:
            text = (
                f"정수 칸 key와 글자 key가 같은 값으로 연결된 key가 {len(linked):,}개 있습니다 "
                "(예: 숫자 123과 글자 '123' — '00123'과는 다름)."
            )
            plan.problems.append(_warn("KEYJOIN_INT_TEXT_LINK", text, count=len(linked)))
    long_numbers = sum(ip.ix.long_numbers for ip in inputs)
    if long_numbers:
        text = f"선택한 열의 {EXCEL_DIGITS + 1}자리 이상 정수 {long_numbers:,}개는 자릿수를 지키려고 글자로 씁니다."
        plan.problems.append(_warn("KEYJOIN_LONG_NUMBERS", text, count=long_numbers))
    missing = sum(ip.ix.formula_missing for ip in inputs)
    if missing:
        text = (
            f"수식 결과값이 저장돼 있지 않은 칸 {missing:,}개는 비워 두고 메모로 표시합니다 (0으로 채우지 않음) — "
            "원본을 엑셀에서 열어 저장하면 결과값이 저장됩니다."
        )
        plan.problems.append(_warn("KEYJOIN_FORMULA_MISSING", text, count=missing))


def _masked_key(key: Key) -> Key:
    return tuple(mask_text(p) for p in key)


def _masked_numbers(plan: Plan, extra: set[Key]) -> dict[Key, int]:
    """가린 표시가 같은 서로 다른 key → 처음 나온 순서의 번호 (미리보기에만 붙인다). 원문·해시는 두지 않는다."""
    bix = plan.base.ix
    assert bix is not None and plan.out_keys is not None
    counts: dict[Key, int] = {}
    for k in (*bix.rows, *extra):
        m = _masked_key(k)
        if m != k:
            counts[m] = counts.get(m, 0) + 1
    colliding = {m for m, n in counts.items() if n + (m in bix.rows or m in extra) >= 2}
    numbers: dict[Key, int] = {}
    if not colliding:
        return numbers
    text = (
        f"가린 표시가 같은 서로 다른 key가 {len(colliding)}묶음 있습니다 — 다른 사람으로 두고 미리보기에 "
        f"'({MASKED_NO} n)'을 붙입니다 (결과 파일에는 원래 key)."
    )
    plan.problems.append(_warn("KEYJOIN_MASKED_COLLISION", text, count=len(colliding)))
    seen: dict[Key, int] = {}
    for k in plan.out_keys:
        m = _masked_key(k)
        if m in colliding:
            seen[m] = seen.get(m, 0) + 1
            numbers[k] = seen[m]
    return numbers


def _shown(value: Any, mask: Any) -> Any:
    if value is GAP or value is None:
        return None
    if isinstance(value, datetime | date):
        return value.isoformat()
    return mask(value)


def _preview(plan: Plan, numbers: dict[Key, int]) -> dict[str, Any]:
    """앞 PREVIEW_ROWS행 (값을 둔 key만 — all에서 뒤 입력에만 있는 key는 그 입력들의 값이 모두 있을 때만)."""
    assert plan.out_keys is not None
    inputs = plan.inputs
    rows: list[dict[str, Any]] = []
    for k in plan.out_keys[: PREVIEW_POOL * (len(inputs) + 1)]:
        if any(k in ip.ix.rows and k not in ip.ix.values for ip in inputs):
            continue
        key = [mask_text(p) for p in k]
        if k in numbers:
            key[-1] = f"{key[-1]} ({MASKED_NO} {numbers[k]})"
        cells: list[Any] = []
        sources = []
        for ip in inputs:
            ix = ip.ix
            mask = mask_rrn_number if ix.rrn else mask_cell
            vals = ix.values.get(k)
            cells.extend([None] * len(ip.pick_ids) if vals is None else [_shown(v, mask) for v in vals])
            if k in ix.rows:
                sources.append({"input": ip.index, "row": ix.rows[k]})
        rows.append({"key": key, "cells": cells, "sources": sources})
        if len(rows) >= PREVIEW_ROWS:
            break
    return {"rows": rows, "shown": len(rows), "total": len(plan.out_keys), "pool": PREVIEW_POOL, "note": PREVIEW_NOTE}


def _basis(plan: Plan) -> str:
    """연결 확인 뒤 바뀌었는지 보는 값: 소유권·파일 해시·개인정보 확인 상태·선택 시트/열·key 규칙·범위·출력 이름·정책·
    한도. key 값·행 수는 넣지 않는다 (내용은 sha256이 대표)."""
    files = {ip.file_id: ip.record for ip in plan.inputs}
    picks = [[c.column_id, c.out_name] for ip in plan.inputs for c in ip.req.columns]
    inputs = [
        [ip.index, ip.file_id, ip.sheet_index, list(ip.req.key_columns), len(ip.req.columns)] for ip in plan.inputs
    ]
    limits = [MAX_INPUTS, MAX_KEY_COLUMNS, MAX_PICK_COLUMNS, MAX_RESULT_COLUMNS, MAX_ROWS, MAX_OUTPUT_CELLS]
    payload = [
        POLICY,
        MODE,
        plan.user_id,
        [[fid, r.sha256, r.uploaded_by, r.pii_acknowledged_at is not None] for fid, r in files.items()],
        inputs,
        picks,
        plan.req.scope,
        [plan.base.file_id, plan.base.sheet_index],
        plan.req.key_names,
        plan.header,
        [*limits, MAX_OUTPUT_TEXT_CHARS, INDEX_WEIGHT, TEXT_WEIGHT],
    ]
    return hashlib.sha256(json.dumps(payload, ensure_ascii=False, default=str).encode()).hexdigest()


# ---------------------------------------------------------------- 결과 파일


def _output_path(run_id: int, now: datetime) -> Path:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    return OUTPUT_DIR / f"key연결_{run_id:04d}_{now.strftime('%Y%m%d_%H%M%S')}.xlsx"


def _write_result(ws: Any, plan: Plan) -> None:
    """시트 「연결 결과」: key 칸은 항상 글자, 미연결·원본 빈 칸은 None (0 아님), 결과값 없는 수식 칸은 빈 칸 + 메모."""
    assert plan.out_keys is not None
    for j, name in enumerate(plan.header, start=1):
        c = _put(ws, 1, j, name)
        c.font, c.fill = Font(bold=True), HEAD_FILL
    ws.freeze_panes = "A2"
    k = len(plan.key_names)
    for r, key in enumerate(plan.out_keys, start=2):
        for j, part in enumerate(key, start=1):
            _put(ws, r, j, part).data_type = "s"  # 선행 0·자릿수 보존 — 숫자로 바꾸지 않는다
        col = k + 1
        for ip in plan.inputs:
            vals = ip.ix.values.get(key)
            for pos in range(len(ip.pick_ids)):
                v = None if vals is None else vals[pos]
                if v is GAP:
                    c = _put(ws, r, col, None)
                    c.fill = GAP_FILL
                    c.comment = Comment("원본의 수식 결과값이 저장돼 있지 않아 비워 둠 (0으로 채우지 않음)", "Moajwo")
                else:
                    _put(ws, r, col, v)
                col += 1
    for j, name in enumerate(plan.header, start=1):
        ws.column_dimensions[get_column_letter(j)].width = min(40, max(8, len(str(name)) * 2 + 2))


def _source_head(ip: InputPlan) -> str:
    return f"입력 {ip.index + 1} · #{ip.file_id} · {sheet_place(ip.place)} 원본 행"


def _write_sources(ws: Any, plan: Plan) -> None:
    """시트 「출처」: 결과 행 번호(연결 결과 시트의 2행 = 1)와 입력별 원본 행 번호 (미연결은 빈 칸). key는 반복하지
    않는다."""
    assert plan.out_keys is not None
    for j, name in enumerate(["결과 행", *(_source_head(ip) for ip in plan.inputs)], start=1):
        c = _put(ws, 1, j, name)
        c.font, c.fill = Font(bold=True), HEAD_FILL
    ws.freeze_panes = "A2"
    for r, key in enumerate(plan.out_keys, start=2):
        _put(ws, r, 1, r - 1)
        for j, ip in enumerate(plan.inputs, start=2):
            _put(ws, r, j, ip.ix.rows.get(key))
    ws.column_dimensions["A"].width = 10
    for j in range(2, len(plan.inputs) + 2):
        ws.column_dimensions[get_column_letter(j)].width = 30


def _write_summary(ws: Any, plan: Plan, run_id: int, processed_at: str) -> None:
    """시트 「연결 요약」 — 설정·개수·경고만. key 원문·셀 값은 넣지 않는다 (가린 이름은 printable로)."""
    row = 1

    def line(*values: Any, bold: bool = False, size: int | None = None) -> None:
        nonlocal row
        for j, v in enumerate(values, start=1):
            c = _put(ws, row, j, printable(v))
            if bold or size:
                c.font = Font(bold=bold, size=size or 11)
        row += 1

    base = plan.base
    line(SUMMARY_SHEET, bold=True, size=14)
    line("실행 번호", run_id)
    line("처리 시각", processed_at)
    line("정책", POLICY)
    line("확인 기준(basis)", plan.basis)
    line("범위", SCOPE_LABEL[plan.req.scope])
    line("기준 입력", f"입력 {base.index + 1} · #{base.file_id} · {sheet_place(base.place)}")
    line("key 열 이름", *(mask_text(n) for n in plan.key_names))
    row += 1
    line("입력", "파일 #", "파일 이름", "시트 이름", "시트 순번", "key 열", "가져온 열", "데이터 행", "연결 안 한 행",
         "key 수", "연결", "미연결", bold=True)  # fmt: skip
    for ip in plan.inputs:
        ix, s = ip.ix, ip.sheet
        keys = ", ".join(f"{mask_text(s.column_names[j])}({ip.letter(j)})" for j in ip.key_ids)
        picked = ", ".join(
            f"{mask_text(s.column_names[j])} → {mask_text(ip.out_names[p])}" for p, j in enumerate(ip.pick_ids)
        )
        skipped = sum(r.count for r in ix.skipped.values())
        line(f"입력 {ip.index + 1}", f"#{ip.file_id}", mask_text(ip.info.name), mask_text(s.name),
             sheet_place(ip.place), keys, picked or "-", ix.data_rows, skipped, len(ix.rows), ip.linked,
             ip.unlinked)  # fmt: skip
    row += 1
    res = plan.result
    line("결과 행", res["rows"])
    line("결과 열", f"{res['columns']}개 (key {res['key_columns']} + 가져온 열 {res['columns'] - res['key_columns']})")
    line("모든 입력의 사용자(합집합)", res["union"])
    line("모든 입력에 있는 사용자(교집합)", res["intersection"])
    line("범위에서 빠진 사용자", res["excluded"])
    for w in plan.warnings:
        line("주의", w.text)
    row += 1
    for note in (
        "key 열은 글자 형식입니다 (Excel이 '숫자가 텍스트로 저장됨'을 표시할 수 있음 — 선행 0·자릿수 보존).",
        "미연결·원본 빈 칸은 빈 칸입니다 (0 아님). 실제 0은 0입니다.",
        f"{EXCEL_DIGITS + 1}자리 이상 정수 값은 자릿수를 지키려고 글자로 씁니다.",
        "선택하지 않은 열은 어느 시트에도 넣지 않습니다.",
        "key 원문·셀 값은 이 시트에 없습니다 (연결 결과 시트에만).",
    ):
        line("안내", note)
    ws.column_dimensions["A"].width = 24
    for col in "BCDEFGHIJKL":
        ws.column_dimensions[col].width = 16


def _start_summary(plan: Plan, request_id: str) -> dict[str, Any]:
    """시작 요약 (실패 기록에도 남는 값) — 결과 내용·key 값은 없다."""
    return {
        "mode": MODE,
        "mode_label": MODE_LABEL,
        "policy": POLICY,
        "request_id": request_id,
        "basis": plan.basis,
        "scope": plan.req.scope,
        "inputs": len(plan.inputs),
        "input_count": len(plan.inputs),
    }


def _done_summary(plan: Plan, run: IntegrationRun, request_id: str, processed_at: str) -> str:
    base = plan.base
    inputs = []
    for ip in plan.inputs:
        ix, s = ip.ix, ip.sheet
        inputs.append(
            {
                "index": ip.index,
                "file_id": ip.file_id,
                "file_name": mask_text(ip.info.name),
                "sheet_id": ip.req.sheet_id,
                "sheet": mask_text(s.name),
                "tab": s.tab,
                "key_column_ids": list(ip.req.key_columns),
                "picked": [
                    {"column_id": c.column_id, "out_name": mask_text(ip.out_names[p])}
                    for p, c in enumerate(ip.req.columns)
                ],
                "data_rows": ix.data_rows,
                "skipped_rows": sum(r.count for r in ix.skipped.values()),
                "indexed": len(ix.rows),
                "linked": ip.linked,
                "unlinked": ip.unlinked,
            }
        )
    res = plan.result
    summary = {
        **_start_summary(plan, request_id),
        "processed_at": processed_at,
        "run_id": run.id,
        "base": {"file_id": base.file_id, "sheet_id": base.req.sheet_id, "index": base.index},
        "key_columns": [
            {"name": mask_text(n), "letter": base.letter(j)} for n, j in zip(plan.key_names, base.key_ids, strict=True)
        ],
        "inputs": inputs,
        "result": {"rows": res["rows"], "columns": res["columns"], "key_columns": res["key_columns"]},
        "counts": {k: res[k] for k in ("union", "intersection", "excluded")},
        "header": [mask_text(h) for h in plan.header],
        "sheet_order": [RESULT_SHEET, SOURCE_SHEET, SUMMARY_SHEET],
        "warnings": [w.text for w in plan.warnings],
        "long_numbers": sum(ip.ix.long_numbers for ip in plan.inputs),
        "formula_missing": sum(ip.ix.formula_missing for ip in plan.inputs),
        "row_count": run.row_count,
    }
    return json.dumps(summary, ensure_ascii=False, default=str)


def make(db: Session, user: CurrentUser, plan: Plan, request_id: str) -> IntegrationRun:
    """판정이 끝난(ready, keep_values=True) 구성으로 결과 파일을 만든다. 다시 파싱하지 않는다 (검증한 입력 스냅숏).
    결과 파일을 확정하는 순서·오류 처리는 app/runlease.py (이름을 바꾼 결과 파일은 지우지 않는다)."""
    if not plan.ready or plan.out_keys is None or any(ip.ix is None for ip in plan.inputs):
        raise KeyJoinError("연결 확인이 끝나지 않은 구성입니다 — 다시 확인하세요.")
    started = datetime.now(KST)
    # 원본 기록을 먼저 읽어 잡아 둔다 — 만드는 동안 DB 질의·열린 트랜잭션이 없게 (bundle.make와 같음)
    _held = [db.get(SourceFile, fid) for fid in plan.file_ids]
    start_summary = json.dumps(_start_summary(plan, request_id), ensure_ascii=False)
    run = IntegrationRun(
        trigger="manual",
        user_id=user.id,
        started_at=started,
        file_ids=json.dumps(plan.file_ids),
        input_file_count=len(plan.file_ids),
        row_count=len(plan.out_keys),
        warning_count=len(plan.warnings),
        summary=start_summary,
    )
    now = datetime.now(KST)
    processed_at = now.strftime("%Y-%m-%d %H:%M:%S (KST)")
    target = runlease.begin(db, run, lambda run_id: _output_path(run_id, now))
    tmp_path: Path | None = None
    try:
        wb = Workbook()
        wb.remove(wb.active)
        _write_result(wb.create_sheet(RESULT_SHEET), plan)
        _write_sources(wb.create_sheet(SOURCE_SHEET), plan)
        _write_summary(wb.create_sheet(SUMMARY_SHEET), plan, run.id, processed_at)
        tmp_path = runlease.new_temp(target)
        wb.save(tmp_path)
        del wb
        summary = _done_summary(plan, run, request_id, processed_at)
        for ip in plan.inputs:  # 해시·확정 동안 인덱스·값을 쥐지 않는다
            ip.ix = None
        plan.out_keys = None
        runlease.prepare(db, run, tmp_path, summary)
        runlease.finish(db, run, tmp_path, target)
    except runlease.RunDone:  # 정리가 이 실행을 이미 완료로 확정했다 (결과 파일 해시 일치)
        db.refresh(run)
    except runlease.RunSettled as exc:  # 다른 곳이 실패로 정했다 — 그 이유로 알린다
        raise _failed(str(exc), run.id) from exc
    except runlease.RunUncertain:  # 결과 파일은 있고 완료 기록만 확인하지 못함 — 지우지 않고 503으로 알린다
        raise
    except Exception as exc:
        # 응답·실행 기록·서버 기록 모두 원문 없이 — 예외 종류와 실행 번호만 (app/failures.py)
        log_failure(logger, "keyjoin", run.id, exc)
        reason = failure_reason(exc, (KeyJoinError, runlease.RunError, HTTPException))
        try:
            runlease.fail(db, run, tmp_path, f"{reason} [{KEYJOIN_FAILED}]", start_summary)
        except runlease.RunDone:
            db.refresh(run)
            return run
        except runlease.RunSettled as settled:
            raise _failed(str(settled), run.id) from exc
        raise _failed(f"결과 파일을 만들지 못했습니다 — {reason} (실행 #{run.id}). [{KEYJOIN_FAILED}]", run.id) from exc
    finally:
        runlease.untrack(run.id)
    logger.info("keyjoin run=%s inputs=%d rows=%d scope=%s", run.id, len(plan.inputs), run.row_count, plan.req.scope)
    return run
