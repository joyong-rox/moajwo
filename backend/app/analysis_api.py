"""데이터 분석 API: 자료 고르기 → 대상 확인(계산 없음) → 실행(저장) → 결과 페이지·Excel.

- 고르기 목록은 DB의 확정 매핑만 읽는다 (파일을 다시 파싱하지 않음).
  화면을 열어도 분석·매핑·학습·모델 호출이 일어나지 않는다.
- 대상 확인과 실행은 같은 판정(analysis.evaluate)을 쓰고,
  실행은 확인 때의 입력 지문(요청 조건+파일 해시+확정 매핑)이 그대로일 때만 한다.
- 파싱·계산은 쓰기 트랜잭션 밖에서 하고, 결과 저장만 짧게 쓴다.
  결과 페이지·Excel은 저장한 결과로 만든다 (다시 계산하지 않음).
- 권한: 분석은 본인 파일로만 만든다 (최고 관리자도 남의 파일로 새 결과를 만들지 않음 — 403).
  결과는 본인 것만, 최고 관리자는 남의 결과를 조회·다운로드만 한다. 요청마다 서버가 확인한다.
"""

import io
import json
from datetime import datetime
from typing import Annotated, Any
from urllib.parse import quote

from fastapi import APIRouter, Depends, HTTPException, Query, Response, status
from openpyxl import Workbook
from openpyxl.styles import Font
from pydantic import BaseModel
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.analysis import (
    AGG_LABEL,
    MEASURE_FIELDS,
    NO_VALUE,
    REASONS,
    AnalysisIn,
    RunIn,
    compute,
    conditions,
    evaluate,
    inputs,
    label_of,
    preview_out,
)
from app.auth import CurrentUserDep, ensure_view
from app.db import get_session
from app.models import AnalysisRun, MappingDecision, SourceFile, User
from app.output_slot import output_slot
from app.pii import mask_text
from app.timefmt import utc_to_kst

router = APIRouter(prefix="/analysis", tags=["analysis"])
PAGE_MAX = 200
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class FieldOut(BaseModel):
    key: str
    label: str
    files: int
    sheets: int


@router.get("/fields", response_model=list[FieldOut])
def analysis_fields(db: Annotated[Session, Depends(get_session)], user: CurrentUserDep) -> list[FieldOut]:
    """내 파일에서 사용자가 확정한 수량 항목 (DB만 읽음)."""
    rows = db.execute(
        select(MappingDecision.field_key, MappingDecision.source_file_id, MappingDecision.sheet_name)
        .join(SourceFile, SourceFile.id == MappingDecision.source_file_id)
        .where(SourceFile.uploaded_by == user.id, MappingDecision.field_key.in_(MEASURE_FIELDS))
    ).all()
    files: dict[str, set[int]] = {}
    sheets: dict[str, set[tuple[int, str]]] = {}
    for key, fid, sheet in rows:
        files.setdefault(key, set()).add(fid)
        sheets.setdefault(key, set()).add((fid, sheet))
    out = [FieldOut(key=k, label=label_of(k), files=len(files[k]), sheets=len(sheets[k])) for k in files]
    return sorted(out, key=lambda f: f.label)


class SourceOut(BaseModel):
    file_id: int
    file_name: str
    uploaded_at: datetime | None
    sheets: list[str]  # 이 항목이 확정된 시트 (가린 이름)
    columns: int


@router.get("/sources", response_model=list[SourceOut])
def analysis_sources(
    field_key: Annotated[str, Query(max_length=64)],
    db: Annotated[Session, Depends(get_session)],
    user: CurrentUserDep,
) -> list[SourceOut]:
    """이 항목으로 확정된 열이 있는 내 파일 (DB만 읽음). 시트·열은 대상 확인 때 파일을 읽어 정한다."""
    if field_key not in MEASURE_FIELDS:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_CONTENT, "분석할 수 있는 수량 항목이 아닙니다.")
    rows = db.execute(
        select(SourceFile.id, SourceFile.original_name, SourceFile.uploaded_at, MappingDecision.sheet_name)
        .join(MappingDecision, MappingDecision.source_file_id == SourceFile.id)
        .where(SourceFile.uploaded_by == user.id, MappingDecision.field_key == field_key)
        .order_by(SourceFile.id.desc())
    ).all()
    out: dict[int, SourceOut] = {}
    for fid, name, uploaded_at, sheet in rows:
        s = out.setdefault(
            fid,
            SourceOut(
                file_id=fid, file_name=mask_text(name), uploaded_at=utc_to_kst(uploaded_at), sheets=[], columns=0
            ),
        )
        s.columns += 1
        if mask_text(sheet) not in s.sheets:
            s.sheets.append(mask_text(sheet))
    return list(out.values())


@router.post("/preview")
def analysis_preview(body: AnalysisIn, db: Annotated[Session, Depends(get_session)], user: CurrentUserDep) -> dict:
    """대상·단위·기간·제외 사유·중복 의심을 확인한다 (계산·저장 없음). 평가가 메모리를 크게 써서 결과 만들기 차례를
    쥔다 (없으면 OutputBusy — app/output_slot.py)."""
    with output_slot():
        ev = evaluate(db, user, body)
        db.rollback()
        return preview_out(ev)


@router.post("/run", status_code=status.HTTP_201_CREATED)
def analysis_run(body: RunIn, db: Annotated[Session, Depends(get_session)], user: CurrentUserDep) -> dict:
    """확인한 상태 그대로일 때만 계산하고 결과를 저장한다 (결과 만들기 차례를 쥔 채)."""
    with output_slot():
        return _analysis_run(body, db, user)


def _analysis_run(body: RunIn, db: Session, user: CurrentUserDep) -> dict:
    ev = evaluate(db, user, body)
    db.rollback()  # 읽기만 한 트랜잭션을 닫는다 (계산 중 잠금을 잡지 않음)
    if ev.fingerprint != body.expected_fingerprint:
        raise HTTPException(
            status.HTTP_409_CONFLICT,
            {"message": "대상 확인 뒤 조건·파일·매핑이 바뀌었습니다 — 다시 확인하세요.", "problems": []},
        )
    result = compute(ev)
    cond = conditions(ev)
    summary = {
        "field_label": cond["field_label"],
        "unit": cond["unit"],
        "variant": cond["variant"],
        "periods": ev.periods,
        "partial": result["partial"],
        "groups": len(result["groups"]),
        "rows_included": result["rows_included"],
        "files": len({t.file_id for t in ev.ok}),
        "group_labels": result["group_labels"],
    }
    run = AnalysisRun(
        created_by=user.id,
        field_key=body.field_key,
        partial=result["partial"],
        summary=json.dumps(summary, ensure_ascii=False),
        conditions=json.dumps(cond, ensure_ascii=False),
        inputs=json.dumps(inputs(ev), ensure_ascii=False),
        result=json.dumps(result, ensure_ascii=False),
    )
    db.add(run)
    db.commit()
    return {"id": run.id}


def _get(db: Session, user: CurrentUserDep, run_id: int) -> AnalysisRun:
    run = db.get(AnalysisRun, run_id)
    if run is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "분석 결과를 찾을 수 없습니다.")
    ensure_view(user, run.created_by, "분석 결과")
    return run


# 저장한 결과의 글자는 저장 당시 규칙으로 가린 것이라, 내보낼 때마다 지금 규칙(mask_text)으로 다시 가린다.
# 저장값은 바꾸지 않고 내보낼 사본만 가린다 (다시 가려도 그대로). 숫자·참거짓·None은 형 그대로 둔다.
# 식별자·해시·코드(조회 키로 쓰는 값)는 가리지 않는다 — sha256 같은 16진수 안의 숫자열이 전화번호 형태로 바뀔 수 있음.
_KEEP = frozenset({"sha256", "sheet_id", "column_id", "field_key", "policy", "status", "reason", "region_level"})
_KEEP |= {"group_by", "group_fields", "aggregates", "on"}


def shown(value: Any, key: str | None = None) -> Any:
    if isinstance(value, dict):
        return {k: v if k in _KEEP else shown(v, k) for k, v in value.items()}
    if isinstance(value, list):
        return [shown(v, key) for v in value]
    return mask_text(value) if isinstance(value, str) else value


@router.get("/runs")
def analysis_runs(db: Annotated[Session, Depends(get_session)], user: CurrentUserDep) -> list[dict]:
    """내 분석 결과 (최근 순)."""
    runs = db.scalars(
        select(AnalysisRun).where(AnalysisRun.created_by == user.id).order_by(AnalysisRun.id.desc()).limit(100)
    )
    return [{"id": r.id, "created_at": utc_to_kst(r.created_at), **shown(json.loads(r.summary))} for r in runs]


@router.get("/runs/{run_id}")
def analysis_detail(run_id: int, db: Annotated[Session, Depends(get_session)], user: CurrentUserDep) -> dict:
    run = _get(db, user, run_id)
    res = json.loads(run.result)
    owner = db.get(User, run.created_by) if run.created_by else None
    meta = {k: v for k, v in res.items() if k not in ("groups", "detail", "excluded")}
    return {
        "id": run.id,
        "created_at": utc_to_kst(run.created_at),  # DB 기본값(UTC) → KST
        "created_by": mask_text(owner.display_name) if owner else None,
        "read_only": run.created_by != user.id,
        "conditions": shown(json.loads(run.conditions)),
        "inputs": shown(json.loads(run.inputs)),
        "result": shown({**meta, "groups_total": len(res["groups"])}),
    }


def _page(items: list[Any], offset: int, limit: int) -> dict:
    return {"total": len(items), "offset": offset, "items": shown(items[offset : offset + limit])}


@router.get("/runs/{run_id}/groups")
def analysis_groups(
    run_id: int,
    db: Annotated[Session, Depends(get_session)],
    user: CurrentUserDep,
    offset: Annotated[int, Query(ge=0)] = 0,
    limit: Annotated[int, Query(ge=1, le=PAGE_MAX)] = 50,
) -> dict:
    return _page(json.loads(_get(db, user, run_id).result)["groups"], offset, limit)


@router.get("/runs/{run_id}/excluded")
def analysis_excluded(
    run_id: int,
    db: Annotated[Session, Depends(get_session)],
    user: CurrentUserDep,
    offset: Annotated[int, Query(ge=0)] = 0,
    limit: Annotated[int, Query(ge=1, le=PAGE_MAX)] = 50,
) -> dict:
    res = json.loads(_get(db, user, run_id).result)
    return {**_page(res["excluded"], offset, limit), "all": res["excluded_total"]}


@router.get("/runs/{run_id}/detail")
def analysis_rows(
    run_id: int,
    db: Annotated[Session, Depends(get_session)],
    user: CurrentUserDep,
    offset: Annotated[int, Query(ge=0)] = 0,
    limit: Annotated[int, Query(ge=1, le=PAGE_MAX)] = 50,
) -> dict:
    res = json.loads(_get(db, user, run_id).result)
    return {**_page(res["detail"], offset, limit), "all": res["detail_total"]}


# ---------------------------------------------------------------- Excel (저장한 결과로만 만든다)

_FORMULA_START = ("=", "+", "-", "@", "\t", "\r")


def _put(ws: Any, row: int, col: int, value: Any, bold: bool = False) -> None:
    """셀에 값을 쓴다. 글자는 수식으로 저장되지 않게 항상 글자 형식으로 둔다 ('=1+1' → 글자 그대로)."""
    cell = ws.cell(row=row, column=col, value=value)
    if isinstance(value, str) and value.startswith(_FORMULA_START):
        cell.data_type = "s"
    if bold:
        cell.font = Font(bold=True)


def _rows(ws: Any, start: int, rows: list[list[Any]], header: bool = False) -> int:
    for i, values in enumerate(rows):
        for j, v in enumerate(values):
            _put(ws, start + i, j + 1, v, bold=header and i == 0)
    return start + len(rows)


def _value(v: Any) -> Any:
    return NO_VALUE if v is None else v


def _numtext(v: float | None) -> str:
    if v is None:
        return "제한 없음"
    return f"{int(v):,}" if float(v).is_integer() else f"{v:,}"


LEVEL_TEXT = {
    None: "자동 — 하위 지역 합계가 상위 지역 값과 같을 때만 상위 행을 중복으로 뺌",
    "lower": "하위 지역만 (상위 지역 행 제외 — 사용자 선택)",
    "upper": "상위 지역만 (하위 지역 행 제외 — 사용자 선택)",
}


def _amount(v: Any) -> str:
    if v is None:
        return "확인 불가"
    return _numtext(v) if isinstance(v, int | float) else str(v)


def level_conflict_text(x: dict[str, Any]) -> str:
    up = x["upper"]
    missing = f" (값 없는 하위 지역 {x['lower_missing']}곳)" if x.get("lower_missing") else ""
    return (
        f"{x.get('period') or '-'} {up['target']}: 상위 값 {_amount(up['value'])} (#{up['file_id']} {up['row']}행) · "
        f"하위 지역 {x['lower_count']}곳 합계 {_amount(x.get('lower_sum'))}{missing} · "
        f"차이 {_amount(x.get('difference'))}"
    )


def _state_text(res: dict[str, Any]) -> str:
    if res["partial"]:
        return "부분 집계 (자료 문제로 뺀 값이 있음)"
    if res.get("level_mismatch"):
        return "선택한 지역 수준 집계 (하위 지역 합계가 상위 지역 값과 다른 곳 있음 — 조건의 주의 참고)"
    return "전체 집계 (자료 문제로 뺀 값 없음)"


def workbook_bytes(run: AnalysisRun) -> bytes:
    cond, inp, res = (shown(json.loads(x)) for x in (run.conditions, run.inputs, run.result))
    wb = Workbook()
    ws = wb.active
    ws.title = "결과"
    aggs = res["aggregates"]
    header = [*res["group_labels"], *[AGG_LABEL[a] for a in aggs]]
    if cond["bins"]:
        header.append(f"구간 ({cond['bins']['on_label']} 기준)")
    header += ["부분 집계", "자료 문제로 뺀 행", "출처 파일 수"]
    table = [header]
    for g in res["groups"]:
        row = [*(k or "(전체)" for k in g["keys"]), *[_value(g[a]) for a in aggs]]
        if cond["bins"]:
            row.append(g["bin"])
        row += ["예" if g["partial"] else "아니오", g["problem_rows"], g["files"]]
        table.append(row)
    if not res["groups"]:
        table.append(["(포함된 행 없음)"])
    _rows(ws, 1, table, header=True)

    wc = wb.create_sheet("조건")
    f = cond["filters"]
    dims = "; ".join(f"{label_of(k)}: {', '.join(v)}" for k, v in f["dims"].items() if v) or "없음"
    lines: list[list[Any]] = [
        ["항목", "값"],
        ["분석 번호", run.id],
        ["생성 시각 (KST)", utc_to_kst(run.created_at).strftime("%Y-%m-%d %H:%M:%S") if run.created_at else ""],
        ["분석 항목", cond["field_label"]],
        ["열 이름의 구분", cond["variant"] or "구분 없음"],
        ["단위", cond["unit"] or "확인 안 됨"],
        ["결과 상태", _state_text(res)],
        ["포함 행 / 뺀 행", f"{res['rows_included']} / {res['rows_excluded']} (자료 문제 {res['problem_rows']})"],
        ["기간", ", ".join(cond["periods"]) or "없음"],
        ["기간별로 나눔", "예 (기간이 여러 개라 기간끼리 더하지 않음)" if res["period_split"] else "아니오"],
        ["그룹 기준", ", ".join(res["group_labels"]) or "없음 (전체 하나)"],
        ["필터 — 지역·기관·구분", dims],
        ["필터 — 기간", f"{f['period_from'] or '처음'} ~ {f['period_to'] or '끝'}"],
        [
            "필터 — 숫자 범위",
            "없음"
            if f["value_min"] is None and f["value_max"] is None
            else f"{_numtext(f['value_min'])} ~ {_numtext(f['value_max'])}",
        ],
        ["계산 방식", ", ".join(AGG_LABEL[a] for a in aggs)],
        ["평균의 분모", cond["mean_denominator"]],
        ["값이 없을 때", f"'{NO_VALUE}' (0으로 쓰지 않음)"],
        ["지역 집계 수준", LEVEL_TEXT[cond.get("region_level")]],
    ]
    for x in cond.get("level_conflicts") or []:  # 하위 합계가 상위 값과 다르거나 확인할 수 없는 곳 (선택의 근거)
        lines.append(["지역 수준 차이", level_conflict_text(x)])
    if cond.get("base_dates"):
        lines.append(["기준일", ", ".join(f"{p}: {d}" for p, d in sorted(cond["base_dates"].items()))])
    if cond.get("records"):
        lines.append(["반복 행", "서로 다른 기록으로 확인해 모두 더함 — " + ", ".join(cond["records"])])
    if cond["bins"]:
        b = cond["bins"]
        lines.append(["구간 분류 적용 열", b["on_label"]])
        for r in b["ranges"]:
            lo = (
                "제한 없음"
                if r["min"] is None
                else f"{_numtext(r['min'])} ({'포함' if r['min_inclusive'] else '제외'})"
            )
            hi = (
                "제한 없음"
                if r["max"] is None
                else f"{_numtext(r['max'])} ({'포함' if r['max_inclusive'] else '제외'})"
            )
            lines.append([f"구간 '{r['label']}'", f"아래 경계 {lo} · 위 경계 {hi}"])
        lines.append(["구간 밖 값", "'미분류' · 값이 없는 그룹은 '계산할 값 없음'"])
    lines.append(["판정 규칙", cond["policy"]])
    for w in cond["warnings"]:
        lines.append(["주의", w])
    end = _rows(wc, 1, lines, header=True)
    files = [
        [
            "파일 번호",
            "파일",
            "시트",
            "열",
            "구분",
            "단위 (출처)",
            "기간 (출처)",
            "상태",
            "사유",
            "sha256",
            "그 시트의 확정 매핑",
        ]
    ]
    for t in inp:
        unit = f"{t['unit']} ({t['unit_source']})" if t["unit"] else "-"
        period = (
            f"{t['period']} ({t['period_source']})"
            if t["period"]
            else ("행의 기간 열: " + ", ".join(t["period_columns"]) if t["period_columns"] else "-")
        )
        state = {"ok": "사용", "excluded": "제외", "needs_input": "확인 필요"}[t["status"]]
        mapping = "; ".join(f"{c} → {label_of(k) if k else '해당 없음'}" for c, k in t["decisions"])
        files.append(
            [
                t["file_id"],
                t["file_name"],
                t["sheet"] or "-",
                t["column"] or "-",
                t["variant"] or "구분 없음",
                unit,
                period,
                state,
                "; ".join(t["reasons"]),
                t["sha256"],
                mapping,
            ]
        )
    _rows(wc, end + 1, files, header=True)

    we = wb.create_sheet("제외")
    counts = [["사유", "행 수", "자료 문제(부분 집계 원인)"]]
    for key, n in sorted(res["reason_counts"].items(), key=lambda kv: -kv[1]):
        counts.append([REASONS[key][0], n, "예" if REASONS[key][1] else "아니오"])
    end = _rows(we, 1, counts, header=True)
    if res["excluded_targets"]:
        whole = [["자료 제외 (파일·시트 단위)", "파일", "시트", "사유"]]
        for x in res["excluded_targets"]:
            whole.append([x["file_id"], x["file_name"], x["sheet"] or "-", "; ".join(x["reasons"])])
        end = _rows(we, end + 1, whole, header=True)
    listed = [["파일 번호", "파일", "시트", "원본 행", "기간", "사유", "값"]]
    for r in res["excluded"]:
        listed.append(
            [r["file_id"], r["file_name"], r["sheet"], r["row"], r["period"], REASONS[r["reason"]][0], r["value"]]
        )
    if res["excluded_total"] > len(res["excluded"]):
        listed.append([f"처음 {len(res['excluded'])}행만 적음 (전체 {res['excluded_total']}행 — 사유별 행 수는 위 표)"])
    _rows(we, end + 1, listed, header=True)

    wd = wb.create_sheet("근거 행")
    detail = [["파일 번호", "파일", "시트", "원본 행", "기간", *res["group_labels"], cond["field_label"]]]
    for r in res["detail"]:
        detail.append([r["file_id"], r["file_name"], r["sheet"], r["row"], r["period"], *r["keys"], r["value"]])
    if res["detail_total"] > len(res["detail"]):
        detail.append([f"처음 {len(res['detail'])}행만 적음 (계산은 전체 {res['detail_total']}행으로 함)"])
    _rows(wd, 1, detail, header=True)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@router.get("/runs/{run_id}/download")
def analysis_download(run_id: int, db: Annotated[Session, Depends(get_session)], user: CurrentUserDep) -> Response:
    run = _get(db, user, run_id)
    name = f"분석결과_{run.id:04d}_{label_of(run.field_key)}.xlsx"
    with output_slot():  # 결과 Excel을 메모리에서 통째로 만든다 — 결과 만들기 차례를 쥔 채 (app/output_slot.py)
        body = workbook_bytes(run)
    return Response(
        body,
        media_type=XLSX,
        headers={
            "Content-Disposition": (
                f"attachment; filename=\"analysis_{run.id:04d}.xlsx\"; filename*=UTF-8''{quote(name)}"
            )
        },
    )
